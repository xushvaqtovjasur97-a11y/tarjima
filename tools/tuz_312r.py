"""3.1.2-rasm: ball qiymatlari va misol banklar; (3.1) formula ikki qatorda; k = 1–4; matndagi misol (qizil)."""
import copy, re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
from docx.oxml.ns import qn
import kamlib as K
import wordart as W
import namuna_312r as N
from pardoz import pardoz

F = N.F
D = K.Doc(F)
cap = [p for p in D.d.paragraphs if p.text.startswith('3.1.2-rasm. Birinchi')][0]
ch = cap._p.getprevious()
rid = [e.get(qn('r:id')) for e in ch.iter('{http://schemas.openxmlformats.org/drawingml/2006/chart}chart')][0]
part = D.d.part.related_parts[rid]
xrel = [r for r in part.rels.values() if r.reltype.endswith('/package')][0]
xml, rows = N.chart(part.blob.decode('utf-8'))
part._blob = re.sub(r'(<c:externalData r:id=")[^"]+(")', r'\g<1>' + xrel.rId + r'\g<2>', xml).encode('utf-8')
xrel.target_part._blob = W._xlsx(rows)
for ext in ch.iter('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent'):
    ext.set('cy', str(int(8.5 * 360000)))
for ext in ch.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}ext'):
    if ext.getparent().tag.endswith('}xfrm'):
        ext.set('cy', str(int(8.5 * 360000)))
# (3.1) formula: birinchi shartdan keyin qator ko‘chirish, markazda, xatboshi chekinishisiz
fp = [p for p in D.d.paragraphs if p.text.startswith('si = k + (xi')][0]
runs = fp._p.findall(qn('w:r'))
txt = [''.join(t.text or '' for t in r.iter(qn('w:t'))) for r in runs]
i = next(k for k in range(1, len(runs)) if txt[k] == '; ' and txt[k - 1] == 'k+1')
br = copy.deepcopy(runs[i])
for e in list(br):
    if e.tag != qn('w:rPr'):
        br.remove(e)
br.append(br.makeelement(qn('w:br'), {}))
runs[i].addnext(br)
for t in runs[i].iter(qn('w:t')):
    t.text = ';'
ppr = fp._p.pPr
ppr.find(qn('w:ind')).set(qn('w:firstLine'), '0')
ppr.find(qn('w:jc')).set(qn('w:val'), 'center')
# izoh va misol
kp = [p for p in D.d.paragraphs if p.text.startswith('bu yerda si – i-ko‘rsatkichning bali')][0]
D.insert(kp, '(k = 1, 2, 3', ', 4', tag='312r')
mp = [p for p in D.d.paragraphs if N.OLD_M in p.text and 'Pog‘onali shkalada' in p.text][0]
D.insert(mp, N.OLD_M, ', uzluksiz shkalada esa mos ravishda 3,96 va 3,04 ball oladi', tag='312r')
D.save(F)
print(part.partname, pardoz(F))

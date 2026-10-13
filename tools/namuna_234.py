"""2.3.4-rasm: SQB da MHXS 9 bo‘yicha 3-bosqich kreditlarining milliy tasnif toifalari bo‘yicha tarkibi, 2021–2025 (namuna).
Manba: SQB MHXS hisobotlari – 2023 y. (2021–2023, S. 53–55), 2025 y. (2024–2025, P. 37–38). Mlrd so‘m."""
import re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_2.3.4-rasm.docx'
Y = ['2021', '2022', '2023', '2024', '2025']
CAT = [('Standart', [192.4, 238.3, 848.8, 1013.3, 746.4], '70AD47', 'FFFFFF'),
       ('Substandart', [797.7, 860.1, 518.6, 577.8, 2290.9], 'FFC000', '000000'),
       ('Qoniqarsiz', [953.2, 490.5, 581.2, 1326.1, 243.9], 'ED7D31', 'FFFFFF'),
       ('Shubhali', [225.3, 988.9, 910.8, 628.2, 642.1], 'C00000', 'FFFFFF'),
       ('Umidsiz', [609.4, 1206.3, 45.0, 540.2, 1386.9], '7F1F1F', 'FFFFFF')]
TOT = [round(sum(c[1][i] for c in CAT), 1) for i in range(5)]
SHARE = [round(100 * (CAT[0][1][i] + CAT[1][1][i]) / TOT[i], 1) for i in range(5)]
LINE = ('Milliy tasnifda muammoli hisoblanmagan qismi (standart va substandart), % (o‘ng o‘q)', SHARE, '1F4E79')
CATS = [f'{y}\n' + f'{t:,.1f}'.replace(',', ' ').replace('.', ',') for y, t in zip(Y, TOT)]


PCT = [(c[0], [round(100 * c[1][i] / TOT[i], 1) for i in range(5)], c[2], c[3]) for c in CAT]


def chart():
    xml, rows = W.bar_chart(CATS, PCT, ytitle='3-bosqich kreditlariga nisbatan, %', stacked=True, ymin=0, ymax=100, major=20, label_fmt='0.0', gap=45)
    sers = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)
    for s_, c in zip(sers, PCT):
        dl = ''.join(f'<c:dLbl><c:idx val="{i}"/><c:delete val="1"/></c:dLbl>' for i, v in enumerate(c[1]) if v < 4)
        if dl:
            xml = xml.replace(s_, s_.replace('<c:dLbls>', '<c:dLbls>' + dl, 1), 1)
    xml = xml.replace('<c:legendPos val="b"/>', '<c:legendPos val="r"/>')
    return xml, rows


OLD = '2024-yilda bu ulush 38,9 foizni tashkil etgan edi, ya’ni bir yil ichida ikki tasnif o‘rtasidagi nomuvofiqlik sezilarli kuchaygan.'
NEW = (f'Bu ulush 2021-yilda {SHARE[0]:.1f} foizni, 2022-yilda {SHARE[1]:.1f} foizni, 2023-yilda {SHARE[2]:.1f} foizni, 2024-yilda {SHARE[3]:.1f} foizni '
       'tashkil etgan, ya’ni ikki tasnif o‘rtasidagi nomuvofiqlik besh yil davomida muntazam saqlanib, 2025-yilda eng yuqori darajaga yetgan.').replace('.', ',', 0)
SRC23 = ('Акционерный коммерческий банк “Узбекский промышленно-строительный банк” и его дочерние организации. Консолидированная финансовая отчетность '
         'и аудиторское заключение за годы, закончившиеся 31 декабря 2023, 2022 и 2021 гг. – Ташкент, 2024. – С. 53–55')

if __name__ == '__main__':
    NEW = NEW.replace('35.6', '35,6')
    for a in SHARE:
        NEW = NEW.replace(f'{a:.1f}', f'{a:.1f}'.replace('.', ','))
    print(TOT, SHARE)
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, color=None, strike=False, indent=False):
        q = dd.add_paragraph(); q.alignment = align
        if indent:
            q.paragraph_format.first_line_indent = Cm(1.25)
        r = q.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.strike = strike
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return q

    par('2.3.4-rasm – 2021–2025-yillar (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, chart())
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.5, 10.5, '2.3.4-rasm')))
    par('2.3.4-rasm. “O‘zsanoatqurilishbank” ATB da MHXS 9 bo‘yicha 3-bosqich kreditlarining milliy tasnif toifalari bo‘yicha tarkibi, 2021–2025-yillar, %',
        bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    par('Izoh: yillar ostida – 3-bosqich kreditlarining jami summasi, mlrd so‘m.', italic=True, size=12, align=WD_ALIGN_PARAGRAPH.LEFT)
    par('Rasm nomida “mlrd so‘m” o‘rniga “2021–2025-yillar, %” (qizil); yillar ostida 3-bosqich kreditlarining jami summasi berilgan. Rasm manbasi (snoska) ga qo‘shiladi:', size=12)
    par(SRC23 + '.', size=12, color='FF0000')
    par('Rasmdan keyingi xatboshida:', size=12)
    par('Hozir: ' + OLD, size=12, color='7F7F7F', strike=True)
    par('Yangi: ' + NEW, size=12, color='FF0000')
    dd.save(DST)
    print(pardoz(DST), DST)

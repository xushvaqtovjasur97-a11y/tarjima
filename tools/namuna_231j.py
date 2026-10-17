"""2.3.1-jadval tuzatishlari namunasi (o‘zgarishlar qizil)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pardoz import pardoz

SRC = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_2.3.1-jadval.docx'
s = docx.Document(SRC)
t = [tb for tb in s.tables if tb.rows[0].cells[0].text.strip() == 'Belgi' and 'MHXS bo‘yicha hisobot' in tb.rows[0].cells[2].text][0]
tit = [p for p in s.paragraphs if p.text.startswith('Milliy prudensial hisobot va MHXS bo‘yicha hisobot o‘rtasidagi asosiy farqlar')][0]
fp = [r.target_part for r in s.part.rels.values() if r.reltype.endswith('/footnotes')][0]
from lxml import etree
x = etree.fromstring(fp.blob)
fid = [r.get(qn('w:id')) for r in tit._p.iter(qn('w:footnoteReference'))][0]
fn = [f for f in x.iter(qn('w:footnote')) if f.get(qn('w:id')) == fid][0]


def runs(el):
    out = []
    for r in el.iter(qn('w:r')):
        tx = ''.join(e.text or '' for e in r.iter(qn('w:t')))
        if tx:
            c = r.find(qn('w:rPr') + '/' + qn('w:color'))
            out.append((tx, c is not None and c.get(qn('w:val')) == 'FF0000'))
    return out


dd = docx.Document()
sec = dd.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)


def par(parts, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, ls=None):
    q = dd.add_paragraph(); q.alignment = align
    if ls:
        q.paragraph_format.line_spacing = ls
    for tx, red in parts:
        r = q.add_run(tx); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if red:
            r.font.color.rgb = RGBColor(0xFF, 0, 0)
    return q


par([('2.3.1-jadval bo‘yicha tuzatishlar (namuna, o‘zgarishlar qizil rangda)', False)], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
par([('2.3.1-jadval', False)], italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
par([('Milliy prudensial hisobot va MHXS bo‘yicha hisobot o‘rtasidagi asosiy farqlar [1]', False)], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.0)
tb = dd.add_table(rows=len(t.rows), cols=4); tb.style = 'Table Grid'; tb.autofit = False
Wd = [Cm(3.0), Cm(4.4), Cm(5.0), Cm(4.1)]
for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), Wd):
    gc.set(qn('w:w'), str(int(w.twips)))
for ri, row in enumerate(t.rows):
    for ci, c in enumerate(row.cells):
        cl = tb.cell(ri, ci); cl.width = Wd[ci]
        q = cl.paragraphs[0]; q.paragraph_format.line_spacing = 1.0
        for tx, red in runs(c._tc):
            r = q.add_run(tx); r.font.size = Pt(11); r.bold = ri == 0
            if red:
                r.font.color.rgb = RGBColor(0xFF, 0, 0)
par([('', False)])
par([('[1] ', False)] + runs(fn), size=11, ls=1.0)
dd.save(DST)
print(pardoz(DST), DST)

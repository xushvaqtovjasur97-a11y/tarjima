"""NSFR va LCR qayta hisob-kitobini tekshirish natijalari va dissertatsiyaga tuzatishlar (Word hujjat)."""
import json, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import openpyxl
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

DST = '/home/user/tarjima/NSFR_LCR_tekshiruv_va_tuzatishlar.docx'
N = json.load(open('/home/user/tarjima/manbalar/hisob_kitob/nsfr_tekshiruv.json'))
B = '/home/user/tarjima/manbalar/bank_hisobotlari/'
LF = {'SQB': B + 'Ўзсаноатқурилишбанк-31.12.2020-31.12.2025 — копия/2025/crs003L.XLSX',
      'MKB': B + 'Микрокредитбанк-31.12.2016-31.12.2025 — копия/2025/crs005L.XLSX',
      'ALOQA': B + 'Aloqabank-31.12.2020-31.12.2025/2025/crs031L.XLSX'}
BANK_ROWS = {'SQB': ('3.1', '3.2', '3.3.1.'), 'MKB': ('3.1', '3.2'), 'ALOQA': ('3.1', '3.2')}


def lcr(b):
    wb = openpyxl.load_workbook(LF[b], data_only=True)
    ws = wb['Bank Assets Analysis']
    liq = wb[[s for s in wb.sheetnames if s.startswith('Liquidity Analysis')][0]]
    L = {str(r[0]).strip(): r[2] for r in liq.iter_rows(values_only=True) if r[0]}
    v = {}
    for r in ws.iter_rows(min_row=6, values_only=True):
        c = str(r[0]).strip() if r[0] else ''
        if c == '21':
            break
        if len(r) > 23 and isinstance(r[23], (int, float)):
            v[c] = r[23]
    hq, out = L['2.3.1'], L['2.3.2']
    bank = sum(v.get(k, 0) for k in BANK_ROWS[b])
    return dict(HQLA=hq / 1e6, bank=bank / 1e6, ovn=v.get('3.3.1.', 0) / 1e6, ulush=bank / hq * 100,
                LCR=hq / out * 100, LCR_b=(hq - bank) / out * 100)


LC = {b: lcr(b) for b in LF}


def c(x, k=1):
    return f'{x:,.{k}f}'.replace(',', ' ').replace('.', ',').replace('-', '−')


d = docx.Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(1.5), Cm(2), Cm(2)
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(13)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.15
BLUE, GREY = RGBColor(0x00, 0x70, 0xC0), RGBColor(0x59, 0x59, 0x59)


def par(text='', bold=False, center=False, size=None, color=None, italic=False, indent=False):
    p = d.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color is not None: r.font.color.rgb = color
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent: p.paragraph_format.first_line_indent = Cm(1.0)
    return p


def labeled(label, text, color):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(label + ' '); r.bold = True; r.font.size = Pt(12)
    r = p.add_run(text); r.font.size = Pt(12); r.font.color.rgb = color
    return p


def table(rows, widths, hdr=1, size=11):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    lay = t._tbl.tblPr.makeelement(qn('w:tblLayout'), {qn('w:type'): 'fixed'}); t._tbl.tblPr.append(lay)
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(Cm(w).twips)))
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.cell(i, j); cell.width = Cm(widths[j]); p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i < hdr or j > 0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(v)); r.font.size = Pt(size); r.bold = i < hdr
    d.add_paragraph()
    return t


S, M, A = N['SQB'], N['MKB'], N['ALOQA']
par('Sof barqaror moliyalashtirish va likvidlikni qoplash me’yorlarining qayta hisob-kitobini tekshirish natijalari '
    'hamda dissertatsiyaga tuzatishlar (namuna)', bold=True, center=True, size=14)
par('Manba: “O‘zsanoatqurilishbank”, “Mikrokreditbank” va “Aloqabank” ATBlarining 2025-yil oxiri holatiga Markaziy bankka taqdim etilgan '
    'prudensial hisobot shakllari (crs003, crs005, crs031): regulyativ kapital hisob-kitobi (5-jadval), likvidlilik tahlili (16-jadval), '
    'bank aktivlari va majburiyatlari tahlili (17–18-jadvallar). Hisob-kitob: manbalar/hisob_kitob/nsfr_tekshiruv.py.', italic=True, size=11)

# ---------------- 1. NSFR
par('1. Sof barqaror moliyalashtirish me’yori: regulyativ kapital', bold=True)
par('Bazel III ga ko‘ra (BCBS, 2014, 21-band) mavjud barqaror moliyalashtirish tarkibiga regulyativ kapital 100 foizlik koeffitsiyent bilan '
    'kiradi, Markaziy bankning 2709-son nizomining 2025-yil 23-maygacha amal qilgan tahririda (54-band) ham regulyativ kapital belgilangan edi. '
    'Banklarning hisobot shaklida esa xususiy kapital olinadi. Demak, taklifning normativ asosi bor. Biroq hisobot shaklida kapitalga kiritilgan '
    'muddatsiz obligatsiyalar (qo‘shimcha birinchi darajali kapital) va subordinar qarz (ikkinchi darajali kapital) “Bir yildan ortiq muddatli '
    'majburiyatlar” qatorida (2.5.2) 100 foizlik koeffitsiyent bilan allaqachon hisobga olingan. Xususiy kapitalni regulyativ kapitalga shunchaki '
    'almashtirish bu instrumentlarni ikki marta hisoblaydi. Bazel III da har bir instrument bir marta hisobga olinadi.', indent=True)
rows = [['Ko‘rsatkich (2025-yil oxiri)', '“O‘zsanoat-qurilishbank”', '“Mikrokredit-bank”', '“Aloqabank”'],
        ['NSFR, hisobot shakli bo‘yicha, %', c(S['NSFR_shakl']), c(M['NSFR_shakl']), c(A['NSFR_shakl'])],
        ['Xususiy kapital, mlrd so‘m', c(S['xususiy']), c(M['xususiy']), c(A['xususiy'])],
        ['Regulyativ kapital, mlrd so‘m', c(S['regulyativ']), c(M['regulyativ']), c(A['regulyativ'])],
        ['Farq, mlrd so‘m', c(S['farq']), c(M['farq']), c(A['farq'])],
        ['shundan muddatsiz obligatsiyalar (AT1)', c(S['AT1_instr']), c(M['AT1_instr']), c(A['AT1_instr'])],
        ['shundan subordinar qarz (T2)', c(S['T2_subdebt']), c(M['T2_subdebt']), c(A['T2_subdebt'])],
        ['“Bir yildan ortiq majburiyatlar”da allaqachon hisobga olingan, mlrd so‘m', c(S['ikki_marta']), c(M['ikki_marta']), c(A['ikki_marta'])],
        ['NSFR – oddiy almashtirish (ikki marta hisoblash), %', c(S['NSFR_oddiy_almashtirish']) + ' *', c(M['NSFR_oddiy_almashtirish']), c(A['NSFR_oddiy_almashtirish']) + ' *'],
        ['NSFR – tuzatilgan (har bir instrument bir marta), %', c(S['NSFR_tuzatilgan']), c(M['NSFR_tuzatilgan']), c(A['NSFR_tuzatilgan'])],
        ['NSFR – tuzatilgan, regulyativ kapital chegirmalargacha (Bazel III), %', c(S['NSFR_tuzatilgan_chegirmagacha']),
         c(M['NSFR_tuzatilgan_chegirmagacha']), c(A['NSFR_tuzatilgan_chegirmagacha'])],
        ['Hisobot shakli bilan farq, foiz punkt',
         f'{c(S["NSFR_tuzatilgan"] - S["NSFR_shakl"])} … +{c(S["NSFR_tuzatilgan_chegirmagacha"] - S["NSFR_shakl"])}',
         f'+{c(M["NSFR_tuzatilgan"] - M["NSFR_shakl"])}', f'+{c(A["NSFR_tuzatilgan"] - A["NSFR_shakl"])}']]
table(rows, [7.6, 3.2, 3.2, 3.0])
par('* – dissertatsiyada keltirilgan qiymat (129,4 va 108,7 foiz). Xulosa: tuzatilgan hisob-kitobda ikki yondashuv o‘rtasidagi farq uchala bankda '
    '1,3 foiz punktdan oshmaydi, ya’ni “hisobot shakli me’yorni 7,1 foiz punktga past ko‘rsatadi” degan xulosa tasdiqlanmaydi. Taklifning '
    'mazmuni esa saqlanadi, ammo u “regulyativ kapitalni olish” emas, balki “kapital instrumentlarini bir marta – kapital sifatida – hisobga olish '
    'tartibini belgilash va me’yor elementlarini oshkor qilish” ko‘rinishida ifodalanishi kerak.', size=12)

# ---------------- 2. LCR
par('2. Likvidlikni qoplash koeffitsiyenti: boshqa banklardagi mablag‘lar', bold=True)
rows = [['Ko‘rsatkich (2025-yil oxiri)', '“O‘zsanoat-qurilishbank”', '“Mikrokredit-bank”', '“Aloqabank”'],
        ['Yuqori likvid aktivlar (milliy metodika), mlrd so‘m', c(LC['SQB']['HQLA']), c(LC['MKB']['HQLA']), c(LC['ALOQA']['HQLA'])],
        ['shundan boshqa banklardagi mablag‘lar (vakillik hisobvaraqlari, overnayt), mlrd so‘m', c(LC['SQB']['bank']), c(LC['MKB']['bank']), c(LC['ALOQA']['bank'])],
        ['ularning ulushi, %', c(LC['SQB']['ulush']), c(LC['MKB']['ulush']), c(LC['ALOQA']['ulush'])],
        ['LCR – milliy metodika, %', c(LC['SQB']['LCR']), c(LC['MKB']['LCR']), c(LC['ALOQA']['LCR'])],
        ['LCR – ular chiqarilganda (sof chiqim o‘zgarishsiz), %', c(LC['SQB']['LCR_b']), c(LC['MKB']['LCR_b']), c(LC['ALOQA']['LCR_b'])],
        ['Dissertatsiyada', '30,9 %; 154,1 %', 'keltirilmagan', '27,3 %; 103,2 %']]
table(rows, [7.6, 3.2, 3.2, 3.0])
par(f'Xulosa: dissertatsiyadagi “O‘zsanoatqurilishbank” va “Aloqabank” ATBlari bo‘yicha raqamlar hisobot shakllariga to‘liq mos keladi. Ikki aniqlik '
    f'kiritish lozim. Birinchidan, “O‘zsanoatqurilishbank” ATB da chiqarilgan mablag‘larning asosiy qismi ({c(LC["SQB"]["ovn"])} mlrd so‘m) boshqa '
    'banklarga overnayt joylashtirilgan mablag‘lar bo‘lib, Bazel III da ular yuqori likvid aktiv hisoblanmaydi, ammo 30 kunlik kirim sifatida '
    '100 foizlik koeffitsiyent bilan hisobga olinadi. Shu sababli 154,1 foiz koeffitsiyentning quyi chegarasi (ehtiyotkor bahosi) hisoblanadi. '
    'Vakillik hisobvaraqlari esa operatsion mablag‘ sifatida kirim bermaydi, shu sababli “Aloqabank” ATB bo‘yicha 103,2 foiz Bazel III ga mos. '
    'Ikkinchidan, “Aloqabank” ATB dagi 1 246,9 mlrd so‘mning 1 209,7 mlrd so‘mi mahalliy banklardagi xorijiy valyutadagi vakillik '
    'hisobvaraqlaridir, xorijiy banklarda esa 37,2 mlrd so‘m. Dissertatsiyada keltirilmagan “Mikrokreditbank” ATB da bu ulush eng yuqori '
    f'({c(LC["MKB"]["ulush"])} foiz) bo‘lib, koeffitsiyent {c(LC["MKB"]["LCR"])} foizdan {c(LC["MKB"]["LCR_b"])} foizgacha pasayadi. '
    'Uni 2.3-paragrafga qo‘shimcha dalil sifatida kiritish tavsiya etiladi.', size=12)

# ---------------- 3. Tuzatishlar
par('3. Dissertatsiyaga tuzatishlar', bold=True)
FIX = [
    ('3.1. Kirish – “Tadqiqotning amaliy natijalari”, 3-band',
     'sof barqaror moliyalashtirish me’yori hisob-kitobida Bazel III standartida va Markaziy bankning 2709-son nizomining 2025-yil 23-maygacha amal '
     'qilgan tahririda barqaror moliyalashtirish manbai sifatida regulyativ kapital belgilangan bo‘lsa-da, banklarning prudensial hisobot shaklida '
     'xususiy kapital olinishi me’yorni taxminan 7,1 foiz punktga pasaytirishi (122,2 foiz; regulyativ kapital olinganda – 129,4 foiz) aniqlanib, '
     'Markaziy bank tomonidan o‘rnatiladigan me’yor elementlarida barqaror moliyalashtirish manbai sifatida regulyativ kapitalni belgilash va bu '
     'elementlar hamda koeffitsiyentlarni oshkor qilish taklif etilgan;',
     'sof barqaror moliyalashtirish me’yori hisob-kitobida Bazel III standartida va Markaziy bankning 2709-son nizomining 2025-yil 23-maygacha amal '
     'qilgan tahririda barqaror moliyalashtirish manbai sifatida regulyativ kapital belgilangani, banklarning prudensial hisobot shaklida esa xususiy '
     'kapital olinib, kapitalga kiritiladigan muddatsiz obligatsiyalar va subordinar qarz bir yildan ortiq muddatli majburiyatlar tarkibida hisobga '
     'olinishi aniqlanib (kapital instrumentlari ikki marta hisoblanmaganda ikki yondashuv o‘rtasidagi farq tahlil qilingan uchta bankda 1,3 foiz '
     'punktdan oshmaydi), Markaziy bank tomonidan o‘rnatiladigan me’yor elementlarida kapital instrumentlarini bir marta – kapital sifatida – hisobga '
     'olish tartibini belgilash hamda bu elementlar va koeffitsiyentlarni oshkor qilish taklif etilgan;'),
    ('3.2. 2.3-paragraf – “O‘zsanoatqurilishbank” ATB, sof barqaror moliyalashtirish me’yori (2.3.4-jadvaldan oldingi xatboshi)',
     '“O‘zsanoatqurilishbank” ATB da regulyativ kapital (16 199,1 mlrd so‘m) xususiy kapitaldan (11 875,2 mlrd so‘m) 4 323,9 mlrd so‘mga ko‘p bo‘lgani '
     'uchun regulyativ kapital olinganda me’yor 122,2 foiz emas, 129,4 foizni tashkil etadi, ya’ni amaldagi hisobot shakli me’yorni 7,1 foiz punktga '
     'past, ehtiyotkorroq ko‘rsatadi. […] Shu sababli ularda barqaror moliyalashtirish manbai sifatida regulyativ kapitalni belgilash va elementlar '
     'hamda koeffitsiyentlarni oshkor qilish maqsadga muvofiq.',
     '“O‘zsanoatqurilishbank” ATB da regulyativ kapital (16 199,1 mlrd so‘m) xususiy kapitaldan (11 875,2 mlrd so‘m) 4 323,9 mlrd so‘mga ko‘p, biroq bu '
     'farqning asosiy qismini qo‘shimcha birinchi darajali kapitalga kiritilgan muddatsiz obligatsiyalar (3 614,6 mlrd so‘m) va ikkinchi darajali '
     'kapitalga kiritilgan subordinar qarz (1 524,8 mlrd so‘m) tashkil etadi. Ular hisobot shaklida bir yildan ortiq muddatli majburiyatlar '
     'tarkibida 100 foizlik koeffitsiyent bilan allaqachon hisobga olingan. Shu sababli kapital instrumentlari ikki marta hisoblanmagan holda '
     'regulyativ kapital olinganda me’yor 120,9 foizni (Bazel III dagidek chegirmalargacha olinganda – 123,4 foizni) tashkil etadi, ya’ni hisobot '
     'shaklidagi qiymatdan (122,2 foiz) 1,3 foiz punktdan ko‘p farq qilmaydi. […] Shu sababli ularda kapital instrumentlarini barqaror '
     'moliyalashtirish tarkibida bir marta – kapital sifatida – hisobga olish tartibini aniq belgilash hamda elementlar va koeffitsiyentlarni '
     'oshkor qilish maqsadga muvofiq.'),
    ('3.3. 2.3-paragraf – shu xatboshiga izoh (361-izohning oxirgi qismi)',
     '… (“Likvidlilik tahlili”, 16-jadval; “Prudensial me’yorlar hisob-kitobi”, 22-jadval). Regulyativ kapital regulyativ kapital monandligi '
     '(17,02 foiz) va risk bo‘yicha tortilgan aktivlar (95 179,7 mlrd so‘m) asosida muallif tomonidan hisoblandi.',
     '… (“Regulyativ kapital hisob-kitobi”, 5-jadval; “Likvidlilik tahlili”, 16-jadval; “Bank majburiyatlari tahlili”, 18-jadval; “Prudensial '
     'me’yorlar hisob-kitobi”, 22-jadval). Hisob-kitob muallif tomonidan bajarildi.'),
    ('3.4. 2.3-paragraf – “O‘zsanoatqurilishbank” ATB, likvidlikni qoplash koeffitsiyenti (2.3.4-jadvaldan keyingi xatboshi)',
     '… likvidlikni qoplash koeffitsiyenti 222,9 foizdan 154,1 foizgacha, ya’ni 68,8 foiz punktga pasayadi. Koeffitsiyent minimal talabdan yuqori '
     'qolsa-da, likvidlik buferining qariyb uchdan bir qismi xalqaro standart bo‘yicha bufer hisoblanmaydi, demak, bu blokda ham yashirin risk mavjud.',
     '… likvidlikni qoplash koeffitsiyenti 222,9 foizdan 154,1 foizgacha, ya’ni 68,8 foiz punktga pasayadi. Chiqarilgan mablag‘larning asosiy qismi '
     '(4 882,6 mlrd so‘m) boshqa banklarga overnayt joylashtirilgan mablag‘lar bo‘lib, Bazel III da ular 30 kunlik kirim sifatida hisobga olinadi, '
     'shu sababli 154,1 foiz koeffitsiyentning quyi chegarasi hisoblanadi. Koeffitsiyent minimal talabdan ancha yuqori qolsa-da, likvidlik '
     'buferining qariyb uchdan bir qismi xalqaro standart bo‘yicha yuqori likvid aktiv hisoblanmaydi, ya’ni buferning tarkibi milliy va xalqaro '
     'yondashuvda sezilarli farq qiladi.'),
    ('3.5. 2.3-paragraf – “Aloqabank” ATB (likvidlik xatboshisi)',
     '… biroq yuqori likvid aktivlarning 27,3 foizi (1 246,9 mlrd so‘m) xorijiy banklardagi vakillik hisobvaraqlaridagi mablag‘lardan iborat. […] '
     'Sof barqaror moliyalashtirish me’yori esa milliy hisobot shakli bo‘yicha 103,0 foizni tashkil etib, Bazel III dagidek xususiy kapital '
     '(3 846,9 mlrd so‘m) o‘rniga regulyativ kapital (5 026,8 mlrd so‘m) olinganda 108,7 foizga oshadi.',
     '… biroq yuqori likvid aktivlarning 27,3 foizi (1 246,9 mlrd so‘m) boshqa banklardagi xorijiy valyutadagi vakillik hisobvaraqlaridan iborat, '
     'shundan 1 209,7 mlrd so‘mi mahalliy banklarda joylashgan. […] “Mikrokreditbank” ATB da bu holat yanada yaqqol namoyon bo‘ladi, ya’ni yuqori '
     f'likvid aktivlarning {c(LC["MKB"]["ulush"])} foizi ({c(LC["MKB"]["bank"])} mlrd so‘m) boshqa banklardagi vakillik hisobvaraqlaridan iborat '
     f'bo‘lib, ular chiqarilganda koeffitsiyent {c(LC["MKB"]["LCR"])} foizdan {c(LC["MKB"]["LCR_b"])} foizgacha pasayadi. Sof barqaror '
     'moliyalashtirish me’yori esa “Aloqabank” ATB da milliy hisobot shakli bo‘yicha 103,0 foizni tashkil etadi, regulyativ kapital olinganda ham '
     '(subordinar qarz ikki marta hisoblanmagan holda) 103,5 foizdan oshmaydi, ya’ni bank barqaror moliyalashtirish bo‘yicha minimal talab '
     'chegarasida turadi.'),
    ('3.6. 2.3-paragraf – “Aloqabank” ATB xatboshisiga izoh (371-izohning oxirgi jumlasi)',
     '16-jadval, 1.4, 1.12, 2.3–2.6-qatorlar; hisob-kitob muallif tomonidan bajarildi.',
     '5-jadval; 16-jadval, 1.4, 1.12, 2.3–2.6-qatorlar; 17- va 18-jadvallar; “Mikrokreditbank” ATB bo‘yicha – crs005L shaklining 16- va '
     '17-jadvallari; hisob-kitob muallif tomonidan bajarildi.'),
    ('3.7. 3.1-paragraf – “Aniqlangan muammo – dalil – metodika elementi” jadvali, “Likvidlik me’yorlari …” qatori',
     '2.3-§: Bazel III bo‘yicha qayta hisoblashda yuqori likvid aktivlarning 27,3–30,9 foizi chiqariladi, NSFR metodika farqi 5,4–5,7 foiz punkt',
     f'2.3-§: Bazel III bo‘yicha qayta hisoblashda yuqori likvid aktivlarning 27,3–{c(LC["MKB"]["ulush"])} foizi chiqariladi, NSFR metodika '
     'farqi 5,4 foiz punkt'),
    ('3.8. Kirish – “Tadqiqot natijalarining joriy qilinishi” (yangi 8 bandlik loyiha)',
     '2-band: “sof barqaror moliyalashtirish manbai sifatida xususiy kapital o‘rniga regulyativ kapitalni olish … me’yor 122,2 foiz o‘rniga '
     '129,4 foizni tashkil etishi …”',
     'Bu band olib tashlanadi (7 band qoladi). 3-banddagi likvidlikni qoplash koeffitsiyenti bo‘yicha natija “Bazel III talablari bo‘yicha qayta '
     'hisoblanganda ham koeffitsiyent kamida 154,1 foiz bilan minimal talabdan 1,5 barobar yuqori” ko‘rinishida qoldiriladi.')]
for title, old, new in FIX:
    par(title, bold=True, size=12)
    labeled('Hozirgi matn:', old, GREY)
    labeled('Taklif:', new, BLUE)

par('4. O‘zgarmaydigan joylar', bold=True)
par('Quyidagilar tekshirildi va to‘g‘ri: 3-ilovadagi to‘rt variantli hisob-kitob (unda Bazel III variantida ham xususiy kapital olingan, ikki marta '
    'hisoblash yo‘q); 2.3.4-jadvaldagi sof barqaror moliyalashtirish me’yori qatori (122,2 → 124,8 foiz; metodika +5,4, MHXS −2,8 foiz punkt); '
    'Kirishdagi ilmiy yangilik (2-band) va joriy qilinish bandlaridagi 5,4 va 2,7–2,8 foiz punkt; II bob bo‘yicha xulosalar va umumiy xulosadagi '
    'tegishli jumlalar; II bob xulosasidagi “Aloqabank” ATB bo‘yicha 142,0 → 103,2 foiz.', size=12)
d.save(DST)
from pardoz import pardoz
print(pardoz(DST), DST)
print({b: {k: round(v, 1) for k, v in x.items()} for b, x in LC.items()})

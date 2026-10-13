"""2.1.3-rasm namunasi (asosiy fayldagi joriy holat): matn, diagramma, rasm nomi va snoskalar."""
import sys, re
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W
import namuna_135 as N135
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_2.1.3-rasm.docx'
src = docx.Document(F)
fx = etree.fromstring([r.target_part for r in src.part.rels.values() if r.reltype.endswith('/footnotes')][0].blob)
FT = {f.get(qn('w:id')): ''.join(x.text or '' for x in f.iter(qn('w:t'))).strip() for f in fx.findall(qn('w:footnote'))}
ps = src.paragraphs
i = [k for k, p in enumerate(ps) if p.text.startswith('2.1.3-rasm')][0]
before = [p for p in ps[i - 3:i] if p.text.strip()][-1]
after = ps[i + 1]
rid = [c.get(qn('r:id')) for c in ps[i - 1]._p.iter('{http://schemas.openxmlformats.org/drawingml/2006/chart}chart')][0]
xml = src.part.related_parts[rid].blob.decode('utf-8')

dd = docx.Document()
sec = dd.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)
FN = []


def par(p_or_text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, indent=True):
    q = dd.add_paragraph(); q.alignment = align
    q.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    if isinstance(p_or_text, str):
        r = q.add_run(p_or_text); r.bold = bold; r.font.size = Pt(size)
        return q
    for el in p_or_text._p:
        if el.tag != qn('w:r'):
            continue
        ref = el.find(qn('w:footnoteReference'))
        if ref is not None:
            FN.append(FT[ref.get(qn('w:id'))])
            r = q.add_run(f'[{len(FN)}]')
        else:
            t = ''.join(x.text or '' for x in el.iter(qn('w:t')))
            if not t:
                continue
            r = q.add_run(t)
            c = el.find('.//' + qn('w:color'))
            if c is not None and c.get(qn('w:val')) not in (None, 'auto', '000000'):
                r.font.color.rgb = RGBColor.from_string(c.get(qn('w:val')))
        r.bold = bold; r.font.size = Pt(size)
    return q


par(before)
q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
rid2 = W.add_chart(dd, (re.sub(r'(<c:externalData r:id=")[^"]+(")', r'\g<1>rId1\g<2>', xml), N135.rows_from(xml)))
q._p.append(etree.fromstring(W.chart_run_xml(rid2, 16.0, 8.0, '2.1.3-rasm')))
par(ps[i], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
par(after)
par('Sahifa osti snoskalari', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
for k, t in enumerate(FN, 1):
    par(f'[{k}] {t}', size=11, indent=False)
dd.save(DST)
print(pardoz(DST), DST)

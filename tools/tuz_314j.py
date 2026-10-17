"""3.1.4-jadval: “Tavsifi” ustuni ballar orqali (qizil)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
import kamlib as K
import namuna_314j as N
from pardoz import pardoz

D = K.Doc(N.F)
cap = [p for p in D.d.paragraphs if p.text.startswith('Moliyaviy barqarorlik sinflari va ularga mos choralar')][0]
t = docx.table.Table(cap._p.getnext(), D.d)
for i, new in N.YANGI.items():
    D.cell_replace(t, i, 2, t.cell(i, 2).text, new, tag='314j')
D.save(N.F)
print([t.cell(i, 2).text for i in N.YANGI], pardoz(N.F))

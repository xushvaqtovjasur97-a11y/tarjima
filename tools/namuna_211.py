"""2.1.1-rasm tuzatish namunasi: likvidlik qutichasi (NSFR – monitoring), hujjat raqamlari va to‘liq manba."""
import copy, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pardoz import pardoz

SRC = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_2.1.1-rasm_yangi.docx'
REP = {'(LCR, NSFR, milliy': '(№ 2709: LCR;',
       'normativlar)': 'NSFR – monitoring)',
       'va zaxiralar': 'va zaxiralar (№ 2696)',
       'bufer (2026)': 'bufer (№ 3697, 2026)'}
OLD_FN = 'Muallif tomonidan O‘zbekiston Respublikasi Markaziy bankining me’yoriy hujjatlari va quyidagi manba asosida tuzilgan: The Central Bank of the Republic of Uzbekistan. “Roadmap” aimed at further improving the bank regulation and supervision system in 2025–2028. – Tashkent, 17.11.2025.'
NEW_FN = ('Tadqiqot natijasida muallif tomonidan O‘zbekiston Respublikasining qonun hujjatlari, Markaziy bankning me’yoriy hujjatlari va Markaziy bankning 2025–2028-yillarga mo‘ljallangan “yo‘l xaritasi” asosida ishlab chiqildi.')


def drawing_el(doc):
    cap = [p for p in doc.paragraphs if p.text.startswith('2.1.1-rasm')][0]
    return cap._p.getprevious()


def fix(el):
    n = 0
    # “Risklarga asoslangan nazorat (42/29, 2023)” qutichasi sig‘ishi uchun shrift 1 punktga kichraytiriladi
    for sp in el.iter('{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}wsp'):
        if 'Risklarga' in ''.join(x.text or '' for x in sp.iter() if x.tag.endswith('}t')):
            for sz in sp.iter():
                if sz.tag in (qn('w:sz'), qn('w:szCs')):
                    sz.set(qn('w:val'), str(int(sz.get(qn('w:val'))) - 2))
    for t in el.iter():
        if t.tag.endswith('}t') and t.text in REP:
            t.text = REP[t.text]; n += 1
    return n


if __name__ == '__main__':
    src = docx.Document(SRC)
    el = copy.deepcopy(drawing_el(src))
    print('almashtirildi:', fix(el))
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, color=None, strike=False):
        q = dd.add_paragraph(); q.alignment = align
        r = q.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.strike = strike
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return q

    par('2.1.1-rasm bo‘yicha tuzatishlar (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    par('1. Rasm (oddiy rangda): “Likvidlik” qutichasida NSFR monitoring sifatida ko‘rsatildi (2025-yil 23-maydan talab kuchini yo‘qotgan, “yo‘l xaritasi”ga ko‘ra 2026-yildan monitoring), likvidlik, aktivlarni tasniflash va kontrsiklik bufer qutichalariga hujjat raqamlari qo‘shildi.')
    p = dd.add_paragraph(); p._p.addnext(el); p._p.getparent().remove(p._p)
    par('2.1.1-rasm. O‘zbekistonda banklar moliyaviy barqarorligini baholashning me’yoriy-huquqiy asoslari', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    par('2. Rasm manbasi (snoska):')
    par('Hozir: ' + OLD_FN, color='7F7F7F', strike=True)
    par('Yangi: ' + NEW_FN, color='C00000')
    dd.save(DST)
    print(pardoz(DST), DST)

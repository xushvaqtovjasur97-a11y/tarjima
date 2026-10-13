"""3.1.2-jadval: chap tomonda raqamlash (1, 1.1 ...) va chegaralarni asoslovchi yangi xatboshi (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_3.1.2-jadval.docx'
FN = ('O‘zbekiston Respublikasi Markaziy banki. Statistik byulleten: 2026-yil mart. – Toshkent, 2026. – 1.2-jadval. – '
      'https://cbu.uz/uz/statistics/buleten/')
ASOS = [
    'Me’yoriy yoki xalqaro tayanch qiymat mavjud bo‘lmagan chegaralar esa ikki tamoyil asosida belgilandi: '
    'birinchidan, ular mavjud tayanch qiymatdan teng qadam bilan uzoqlashtirildi; ikkinchidan, shkalaning tahlil '
    'davridagi haqiqiy qiymatlarni qamrab olishi ikkinchi bob ma’lumotlari bilan tekshirildi. Xususan, kapitalning '
    'aktivlarga nisbati uchun 8 foizlik 4 ball chegarasi milliy leverage talabidan (6 foiz) 2 foiz punkt yuqori bo‘lib, '
    '4 va 6 foiz o‘rtasidagi qadamni takrorlaydi. 3-bosqich kreditlari ulushining 15 foizlik 1 ball chegarasi 5 va 10 '
    'foizlik chegaralar qadamini davom ettiradi, 3 foizlik 4 ball chegarasi esa bank tizimida rasmiy muammoli kreditlar '
    'ulushining tahlil davridagi eng past darajasiga mos keladi. Zaxira bilan qoplanmagan muammoli kreditlarning '
    'kapitalga nisbati uchun 0 foiz muammoli kreditlar to‘liq zaxira bilan qoplangan holatni bildiradi, 10, 20 va 30 '
    'foizlik chegaralar esa teng qadamli bo‘lib, bank tizimida bu ko‘rsatkich 2021–2025-yillarda 5,7–13,1 foiz '
    'oralig‘ida bo‘lgani, “Mikrokreditbank” ATB da esa MHXS bo‘yicha 2025-yilda 29,1 foizga yetgani hisobga olindi. ',
    'Yuqori likvid aktivlar ulushi uchun chegaralar Markaziy bankning 10 foizlik minimal talabidan 5 foiz punktlik '
    'qadam bilan belgilandi, bank tizimida bu ko‘rsatkich tahlil davrida 14,9–21,1 foiz oralig‘ida bo‘lgan. '
    'Kreditlarning depozitlarga nisbati uchun 150 va 200 foizlik chegaralar 100 foizdan keyingi 50 foiz punktlik '
    'qadamni ifodalaydi: 200 foizdan yuqori qiymat kreditlarning yarmidan ko‘pi depozit bo‘lmagan manbalar hisobiga '
    'moliyalashtirilishini bildiradi, bunday holat “O‘zsanoatqurilishbank” va “Mikrokreditbank” ATBlarida kuzatilgan '
    '(206,5–414,0 foiz). Aktivlar rentabelligining 2 foizlik 4 ball chegarasi Rossiya Banki shkalasidagi qadamni '
    'taxminan davom ettiradi. Kapital rentabelligining 10 foizlik 3 ball chegarasi tahlil davridagi yillik inflyatsiya '
    'darajasiga (7,3–12,3 foiz) mos keladi',
    '[[FN]]',
    ', 15 foizlik 4 ball chegarasi esa real ijobiy rentabellikni bildiradi. Xarajatlarning daromadlarga nisbati uchun '
    '100 foiz bank zararsizlik nuqtasida turgan holatni bildiradi, 80 va 45 foizlik chegaralar esa Rossiya Banki '
    'amaliyotidagi 60 foizlik chegaradan ikki tomonga qo‘yilgan oraliq qadamlardir. Valyuta riskiga sezgirlik bloki '
    'uchun me’yoriy chegara mavjud emas, shu sababli chegaralar 15 foiz punktlik teng qadam bilan belgilanib, bank '
    'tizimidagi qiymatlar bilan bog‘landi: xorijiy valyutadagi kreditlar ulushi 2021-yilda 49,2 foiz, 2025-yilda 38,4 '
    'foiz bo‘lgan, ya’ni 50 foizlik 2 ball chegarasi tahlil davri boshidagi, 35 foizlik 3 ball chegarasi esa tizimning '
    'hozirgi darajasiga yaqin, 20 foizdan past ulush esa past dollarlashuvni bildiradi. Xorijiy valyutadagi '
    'majburiyatlar bo‘yicha chegaralar kreditlarnikidan 5 foiz punkt yuqori belgilandi, chunki bu majburiyatlar '
    'nafaqat kreditlarni, balki xorijiy valyutadagi likvid aktivlarni ham moliyalashtiradi.']
NUM = {1: '1.', 2: '1.1', 3: '1.2', 4: '1.3', 5: '2.', 6: '2.1', 7: '2.2', 8: '2.3', 9: '3.', 10: '3.1', 11: '3.2',
       12: '4.', 13: '4.1', 14: '4.2', 15: '4.3', 16: '5.', 17: '5.1', 18: '5.2'}

if __name__ == '__main__':
    src = docx.Document('/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx')
    p = [p for p in src.paragraphs if p.text.strip() == '3.1.2-jadval'][0]
    e = p._p.getnext()
    while not e.tag.endswith('}tbl'):
        e = e.getnext()
    rows = []
    for r in docx.table.Table(e, src).rows:
        seen, out = [], []
        for c in r.cells:
            if c._tc in seen: continue
            seen.append(c._tc); out.append(c.text)
        rows.append(out)
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)

    def run(q, t, size=14, bold=False, red=False):
        r = q.add_run(t); r.font.size = Pt(size); r.bold = bold
        if red: r.font.color.rgb = RGBColor(0xFF, 0, 0)
        return r
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.RIGHT; run(q, '3.1.2-jadval', bold=True)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(q, 'Moliyaviy barqarorlikning integral indeksi ko‘rsatkichlari va ularning chegaraviy qiymatlari', bold=True)
    t = dd.add_table(rows=len(rows), cols=7); t.style = 'Table Grid'
    for i, row in enumerate(rows):
        c0 = t.cell(i, 0).paragraphs[0]; c0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if i == 0:
            run(c0, '№', 11, bold=True, red=True)
        else:
            run(c0, NUM[i], 11, bold=row[1] == '', red=True)
        if len(row) == 1 or all(x == '' for x in row[1:]):
            m = t.cell(i, 1).merge(t.cell(i, 6))
            run(m.paragraphs[0], row[0], 11, bold=True)
            continue
        for j, txt in enumerate(row):
            q = t.cell(i, j + 1).paragraphs[0]
            if j > 0 or i == 0: q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(q, txt, 11, bold=(i == 0))
    t.autofit = False
    W = [1.1, 6.0, 2.4, 1.6, 1.6, 1.6, 1.6]
    for g, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), W):
        g.set(qn('w:w'), str(int(w * 567)))
    for r in t.rows:
        for j, c in enumerate(r.cells):
            c.width = Cm(W[j])
    q = dd.add_paragraph(); run(q, 'Izoh: oraliq qiymatlar uchun ball chiziqli interpolyatsiya orqali aniqlanadi; 4 ball – eng yuqori barqarorlik.', 12)
    dd.add_paragraph()
    q = dd.add_paragraph(); run(q, 'Jadvaldan keyingi matnga qo‘shiladigan yangi xatboshi (3.1-§, “Likvidlik blokida…” '
                                   'xatboshisidan keyin):', 12, bold=True)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    q.paragraph_format.first_line_indent = Cm(1.25); q.paragraph_format.line_spacing = 1.5
    for part in ASOS:
        if part == '[[FN]]':
            r = run(q, '¹', red=True); r.font.superscript = True
        else:
            run(q, part, red=True)
    q = dd.add_paragraph(); run(q, '¹ ' + FN, 10, red=True)
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

"""1.3.5-rasm (LCR talablari) tuzatish namunasi: Rossiya “o‘z mablag‘lari hisobidan” qatoridan tasdiqlanmagan
2026-yil 80 foiz nuqtasini olib tashlash, matn va manbalarni aniqlashtirish.
Manbalar: Bank Rossii press-relizlari 23.11.2023 va 16.12.2024; Interfaks 23.11.2023."""
import re, sys, io
sys.path.insert(0, '/home/user/tarjima/tools')
import docx, openpyxl
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W
from pardoz import pardoz

SRC = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_1.3.5-rasm_yangi.docx'
PR2024 = 'Банк России принял ряд решений по банковскому регулированию: пресс-релиз Банка России от 16 декабря 2024 г. – https://www.cbr.ru/press/pr/?file=638699637849513580BANK_SECTOR.htm'
IFX_OLD = 'ЦБ скорректировал условия кредитным линиям для банковского норматива краткосрочной ликвидности // Интерфакс. – 2025.'
IFX_NEW = 'ЦБ скорректировал условия кредитным линиям для банковского норматива краткосрочной ликвидности // Интерфакс. – 2023. – 23 ноября.'
OLD_S = 'Banklarning o‘z mablag‘lari hisobidan ta’minlanishi lozim bo‘lgan qism esa 2024-yil martida 40 foiz, iyulida 50 foiz, 2025-yil iyulida 60 foiz, 2026-yil yanvarida 80 foiz qilib belgilandi.'
NEW_S = 'Banklarning o‘z mablag‘lari hisobidan ta’minlanishi lozim bo‘lgan qism esa 2024-yil martida 40 foiz, iyulida 50 foiz, 2025-yil iyulida 60 foiz qilib belgilandi; 2024-yil dekabrida jadval o‘zgartirilib, 60 foizlik bosqich yarim yilga surildi'


def chart_xml(doc):
    """Asosiy fayldagi 1.3.5-rasm diagrammasi XML i (Rossiya o‘z mablag‘lari qatoridan 2026 nuqtasi olib tashlangan)."""
    cap = [p for p in doc.paragraphs if p.text.startswith('1.3.5-rasm')][0]
    rid = [c.get(qn('r:id')) for c in cap._p.getprevious().iter('{http://schemas.openxmlformats.org/drawingml/2006/chart}chart')][0]
    part = doc.part.related_parts[rid]
    x = part.blob.decode('utf-8')
    sers = re.findall(r'<c:ser>.*?</c:ser>', x, re.S)
    s = sers[-1]
    assert 'o‘z mablag‘lari hisobidan' in s and '<c:v>2026</c:v></c:pt></c:numCache>' in s
    s2 = s.replace('$I$2:$I$5', '$I$2:$I$4').replace('$J$2:$J$5', '$J$2:$J$4').replace('<c:ptCount val="4"/>', '<c:ptCount val="3"/>')
    s2 = s2.replace('<c:pt idx="3"><c:v>2026</c:v></c:pt>', '').replace('<c:pt idx="3"><c:v>80</c:v></c:pt>', '')
    return x.replace(s, s2), part


def rows_from(x):
    """Diagramma keshidan ichki Excel jadvali uchun qatorlar."""
    cells = {}
    for ref, body in re.findall(r'<c:f>([^<]+)</c:f>(.*?)</c:(?:strRef|numRef)>', x, re.S):
        m = re.match(r'Sheet1!\$([A-Z]+)\$(\d+)(?::\$[A-Z]+\$\d+)?', ref)
        col, r0 = m.group(1), int(m.group(2))
        for i, v in re.findall(r'<c:pt idx="(\d+)"><c:v>([^<]*)</c:v>', body):
            try:
                v = float(v)
            except ValueError:
                pass
            cells[(r0 + int(i), ord(col) - 64)] = v
    nr = max(r for r, _ in cells); nc = max(c for _, c in cells)
    return [[cells.get((r, c)) for c in range(1, nc + 1)] for r in range(1, nr + 1)]


if __name__ == '__main__':
    xml, _ = chart_xml(docx.Document(SRC))
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=14, color=None, strike=False):
        q = dd.add_paragraph(); q.alignment = align
        q.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
        r = q.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.strike = strike
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return q

    par('1.3.5-rasm bo‘yicha tuzatishlar (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    par('1. Diagramma: “Rossiya: o‘z mablag‘lari hisobidan” chizig‘idan 2026-yil yanvaridagi 80 foiz nuqtasi olib tashlandi (2024-yil dekabrida jadval o‘zgartirilgan, 2025-yil oktabridan esa tizimli banklar milliy normativga o‘tgan). Rasm nomi va qolgan chiziqlar o‘zgarmaydi.', size=12, indent=False)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, (xml, rows_from(xml)))
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.0, '1.3.5-rasm')))
    par('1.3.5-rasm. Likvidlikni qoplash koeffitsiyentiga qo‘yiladigan talablar: islohotning boshlanishi va davomi, %', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    par('2. Matn (“Islohotning keyingi bosqichi …” xatboshisi):', size=12, indent=False)
    par('Hozir: ' + OLD_S, size=12, indent=False, color='7F7F7F', strike=True)
    par('Yangi: ' + NEW_S + '[1].', size=12, indent=False, color='C00000')
    par('3. Manbalar:', size=12, indent=False)
    par('Snoska va adabiyotlar ro‘yxatida: ' + IFX_OLD, size=12, indent=False, color='7F7F7F', strike=True)
    par('Yangi: ' + IFX_NEW, size=12, indent=False, color='C00000')
    par('[1] ' + PR2024 + ' (yangi snoska, adabiyotlar ro‘yxatiga ham qo‘shiladi).', size=12, indent=False, color='C00000')
    dd.save(DST)
    print(pardoz(DST), DST)

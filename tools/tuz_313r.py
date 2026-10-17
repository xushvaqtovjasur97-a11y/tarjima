"""Yangi 3.1.3-rasm (uch bank MBII 2021–2025, milliy va MHXS) + sonlar jadvali, izoh, tahlil; eski 3.1.3 → 3.1.4 (qizil)."""
import copy, sys
sys.path.insert(0, '/home/user/tarjima/tools')
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import kamlib as K
import wordart as W
import namuna_313r as N
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
W._counter['chart'] = 300; W._counter['docpr'] = 9600
D = K.Doc(F)
RED = 'FF0000'

# 1) eski 3.1.3-rasm → 3.1.4-rasm
oldcap = [p for p in D.d.paragraphs if p.text.startswith('3.1.3-rasm. Moliyaviy barqarorlikning integral indeksi qiymatlarining')][0]
D.replace(oldcap, '3.1.3', '3.1.4', tag='313r')
met = D.para('Metodika natijalarining ishonchliligini tekshirish uchun indeks uch xil vaznlar')
D.replace(met, '(3.1.3-rasm)', '(3.1.4-rasm)', tag='313r')
oldchart = oldcap._p.getprevious()

# 2) yangi blok 3.1.6-jadval tahlilidan keyin
anchor = D.para('3.1.6-jadval ma’lumotlariga ko‘ra, MHXS ma’lumotlari asosida esa uchala')
lead = D.new_para_after(anchor, N.LEAD, tag='313r')

# diagramma paragrafi
cp = D.d.add_paragraph()
rid = W.add_chart(D.d, N.chart())
from lxml import etree
cp._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.0, '3.1.3-rasm')))
ppr = oldchart.find(qn('w:pPr'))
if ppr is not None:
    cp._p.insert(0, copy.deepcopy(ppr))
lead._p.addnext(cp._p)

# sonlar jadvali
t = D.d.add_table(rows=1 + len(N.SER), cols=6)
t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [Cm(6.5)] + [Cm(1.8)] * 5
t.autofit = False
tblPr = t._tbl.tblPr
lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
for gc, wd in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
    gc.set(qn('w:w'), str(int(wd / 635)))
hdr = ['Bank va hisobot turi'] + N.Y
data = [[s[0]] + [f'{v:.2f}'.replace('.', ',') for v in s[1]] for s in N.SER]
for i, r in enumerate([hdr] + data):
    for j, v in enumerate(r):
        c = t.cell(i, j); c.width = widths[j]; p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format; pf.space_before = pf.space_after = 0; pf.first_line_indent = Cm(0); pf.line_spacing = 1.0
        if j == 0 and i > 0:
            s = N.SER[i - 1]
            k = p.add_run('━━ ' if s[3] is None else '╍╍ '); k.font.size = Pt(10); k.bold = True
            k.font.color.rgb = RGBColor.from_string(s[2]); k.font.name = 'Times New Roman'
        run = p.add_run(v); run.font.size = Pt(11); run.bold = (i == 0); run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor.from_string(RED)
cp._p.addnext(t._tbl)

# izoh, rasm nomi (izohi bilan), tahlil
izsrc = D.para('Izoh: SQB – “O‘zsanoatqurilishbank” ATB, MKB – “Mikrokreditbank” ATB; qavs ichida – ball.')
iz = D.new_para_after(izsrc, N.IZOH, tag='313r')
t._tbl.addnext(iz._p)
cap = D.new_para_after(oldcap, N.CAP + '[[FN:' + N.FN + ']]', tag='313r')
iz._p.addnext(cap._p)
an = D.new_para_after(anchor, N.ANAL, tag='313r')
cap._p.addnext(an._p)
D.save(F)
print(pardoz(F))

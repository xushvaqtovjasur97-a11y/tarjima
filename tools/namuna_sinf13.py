"""3.2-§: sinf chegaralari qanday belgilangani (teng oraliqlar) va MBII-13 bo‘yicha qo‘shimcha tekshiruv (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_sinf_chegaralari.docx'
OLD1 = 'Ular ham 3.1-paragrafda me’yoriy talablar asosida belgilangan edi.'
NEW1 = 'Ular 3.1-paragrafda 1–4 ballik shkalani teng to‘rt oraliqqa bo‘lish yo‘li bilan oldindan belgilangan edi.'
ANC2 = 'shu sababli erta ogohlantirish tizimida II sinf ikki zonaga ajratilgan.'
INS2 = (' Sinflar chegaralari MBII-13 va MBII-8 uchun bir xil bo‘lib, ular MBII-8 bo‘yicha tekshirildi, chunki 25 ta bank panelida '
        'kuzatuvlar soni yetarli. MBII-13 bo‘yicha beshta bankning 2021–2024-yillardagi 20 ta kuzatuvida ham natija shu yo‘nalishda: '
        'keyingi yilda muammo hodisasi I sinfdagi 2 ta kuzatuvda kuzatilmagan, II sinfdagi 12 ta kuzatuvning 2 tasida (17 foiz), '
        'III sinfdagi 6 ta kuzatuvning 2 tasida (33 foiz) yuz bergan. Masalan, 2023-yilda III sinfda bo‘lgan “Mikrokreditbank” va '
        '“Biznesni rivojlantirish banki” ATBlarida 2024-yilda yirik zarar qayd etilgan. Kuzatuvlar soni kamligi sababli bu natija '
        'statistik dalil emas, balki MBII-8 bo‘yicha olingan xulosani tasdiqlovchi qo‘shimcha tekshiruv hisoblanadi.')

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(4)
    G, R = '7F7F7F', 'FF0000'

    def par(parts, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        q = dd.add_paragraph(); q.alignment = align
        for t, col, s in parts:
            r = q.add_run(t); r.bold = bold; r.font.strike = s
            if col: r.font.color.rgb = RGBColor.from_string(col)
    par([('Sinf chegaralari: 3.2-§ (namuna)', None, False)], True, WD_ALIGN_PARAGRAPH.CENTER)
    par([('1) 3.2.4-jadvaldan oldingi xatboshida:', None, False)], True)
    par([('Vaznlardan tashqari, integral indeksning ikkinchi metodik elementi – barqarorlik sinflari chegaralari, ya’ni 3,25; 2,5 va '
          '1,75 ball hisoblanadi. ', None, False), (OLD1, G, True), (' ', None, False), (NEW1, R, False), (' Ularni empirik tekshirish uchun …', None, False)])
    par([('2) 3.2.4-jadval tahlili oxirida:', None, False)], True)
    par([('… 4.1-paragrafda ko‘rsatiladiganidek, 2,75 balldan past bo‘lgan banklarda muammo ehtimoli keskin oshadi, ', None, False),
         (ANC2, None, False), (INS2, R, False)])
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST))

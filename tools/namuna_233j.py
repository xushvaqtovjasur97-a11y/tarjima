"""Yangi 2.3.3-jadval (namuna): SQB kredit portfeli MHXS 9 bosqichlari va milliy tasnif toifalari bo‘yicha, 2021–2025."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_2.3.3-jadval_bosqichlar.docx'
RED = 'FF0000'
# mln so‘m: segmentlar (korporativ, davlat, jismoniy) yig‘indisi
R = {2021: dict(s1=[14556470, 14246280, 3877782], s2std=[6984900, 0, 106616], s2sub=[1899935, 0, 79871],
                s3ss=192372 + 797738, s3npl=953153 + 225270 + 609407),
     2022: dict(s1=[22016653, 12615317, 5339939], s2std=[4294785, 1369382, 68124], s2sub=[1377965, 371837, 60345],
                s3ss=238290 + 860080, s3npl=490463 + 988919 + 1206289),
     2023: dict(s1=[26130681, 14170248, 7666423], s2std=[2440295, 396473, 127981], s2sub=[6426393, 5775, 139141],
                s3ss=848827 + 518587, s3npl=581177 + 910817 + 44957),
     2024: dict(s1=[26178961, 11624671, 8975818], s2std=[5228960, 2629883, 205108], s2sub=[10326319, 4318, 512366],
                s3ss=1013257 + 577835, s3npl=1326159 + 628156 + 540210),
     2025: dict(s1=[21682263, 9921826, 8774085], s2std=[12535154, 39527, 381492], s2sub=[8637105, 3249504, 477656],
                s3ss=746404 + 2290907, s3npl=243870 + 642117 + 1386888)}


def f(x):
    return f'{x / 1000:,.1f}'.replace(',', ' ').replace('.', ',')


def pc(a, b):
    return f'{100 * a / b:.1f}'.replace('.', ',')


ROWS = []
for y, v in R.items():
    s1 = sum(v['s1']); a = sum(v['s2std']); b = sum(v['s2sub']); s2 = a + b; s3 = v['s3ss'] + v['s3npl']; t = s1 + s2 + s3
    ROWS.append([str(y), f(t), f'{f(s1)}\n({pc(s1, t)})', f'{f(s2)}\n({pc(s2, t)})', f'{f(a)}\n({pc(a, s2)})', f'{f(b)}\n({pc(b, s2)})',
                 f'{f(s3)}\n({pc(s3, t)})', f'{f(v["s3ss"])}\n({pc(v["s3ss"], s3)})', f'{f(v["s3npl"])}\n({pc(v["s3npl"], s3)})'])
for r in ROWS:
    print(r)

TITLE = '“O‘zsanoatqurilishbank” ATB kredit portfelining MHXS 9 bosqichlari va milliy tasnif toifalari bo‘yicha tarkibi, 2021–2025-yillar, mlrd so‘m'
SRC = ('Tadqiqot natijasida muallif tomonidan quyidagi manbalar ma’lumotlari asosida hisoblandi: Акционерный коммерческий банк “Узбекский промышленно-строительный банк” '
       'и его дочерние организации. Консолидированная финансовая отчетность и аудиторское заключение за годы, закончившиеся 31 декабря 2023, 2022 и 2021 гг. – '
       'Ташкент, 2024. – С. 53–55; Joint Stock Commercial Bank “Uzbek Industrial and Construction Bank” and its subsidiaries. Consolidated Financial Statements '
       'for the year ended 31 December 2025 and Independent Auditor’s Report. – Tashkent, 2026. – P. 37–38.')
IZOH = ('Izoh: amortizatsiyalangan qiymatda baholanadigan mijozlarga kreditlarning yalpi qiymati (korporativ, davlat va munitsipal tashkilotlar, jismoniy shaxslar); '
        'qavs ichida – bosqichlar uchun jami kreditlardagi, toifalar uchun tegishli bosqichdagi ulushi, %; 1-bosqich kreditlarining barchasi milliy tasnifda “standart” toifasida; '
        'muammoli – “qoniqarsiz”, “shubhali” va “umidsiz” toifalari.')
TXT = ('Ikki tasnif o‘rtasidagi bog‘liqlikni bankning butun kredit portfeli bo‘yicha ko‘rib chiqish manzarani to‘ldiradi (2.3.3-jadval). Jadvaldan ko‘rinadiki, '
       '1-bosqich kreditlarining ulushi 2023-yildagi 79,4 foizdan 2025-yilda 56,9 foizgacha qisqargan, kredit riski sezilarli oshgan 2-bosqich kreditlarining ulushi esa '
       '15,8 foizdan 35,7 foizgacha, ya’ni 2,3 barobar oshgan. Milliy tasnifda bu o‘zgarish muammoli kreditlar ko‘rsatkichida umuman aks etmagan: 2-bosqich '
       'kreditlarining barchasi “standart” yoki “substandart” toifalarida bo‘lib, 2025-yilda ularning 51,2 foizi “standart” deb tasniflangan. Milliy tasnif bo‘yicha '
       'muammoli deb tan olingan kreditlarning barchasi 3-bosqichda bo‘lsa-da, ular 3-bosqich kreditlarining faqat 42,8–71,0 foizini tashkil etgan.')

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, color=RED, ls=None, indent=False):
        q = dd.add_paragraph(); q.alignment = align
        if indent:
            q.paragraph_format.first_line_indent = Cm(1.25)
        if ls:
            q.paragraph_format.line_spacing = ls
        r = q.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return q

    par('Yangi 2.3.3-jadval (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=None)
    par('Joyi: 2.3.4-rasmdan keyingi “Rasm ma’lumotlari shuni ko‘rsatadiki …” xatboshisidan so‘ng. Hozirgi 2.3.3, 2.3.4, 2.3.5-jadvallar 2.3.4, 2.3.5, 2.3.6 bo‘ladi (havolalar ham).',
        italic=True, size=12, color='595959')
    par(TXT, indent=True)
    par('2.3.3-jadval', italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    par(TITLE + ' [1]', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.0)
    t = dd.add_table(rows=2 + len(ROWS), cols=9); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    Wd = [Cm(1.2), Cm(1.8), Cm(1.9), Cm(1.9), Cm(1.8), Cm(1.8), Cm(1.8), Cm(1.9), Cm(1.9)]
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), Wd):
        gc.set(qn('w:w'), str(int(w.twips)))

    def put(c, text, bold=False, fill=None, size=10):
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.0
        for k, line in enumerate(text.split('\n')):
            r = p.add_run(line); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(RED)
            if k < len(text.split('\n')) - 1:
                r.add_break()
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if fill:
            sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), fill); c._tc.get_or_add_tcPr().append(sh)

    H = 'DEEBF7'
    t.cell(0, 0).merge(t.cell(1, 0)); put(t.cell(0, 0), 'Yil', True, H)
    t.cell(0, 1).merge(t.cell(1, 1)); put(t.cell(0, 1), 'Jami kreditlar', True, H)
    t.cell(0, 2).merge(t.cell(1, 2)); put(t.cell(0, 2), '1-bosqich (standart)', True, H)
    t.cell(0, 3).merge(t.cell(0, 5)); put(t.cell(0, 3), '2-bosqich', True, H)
    t.cell(0, 6).merge(t.cell(0, 8)); put(t.cell(0, 6), '3-bosqich', True, H)
    for j, h in zip(range(3, 9), ['jami', 'standart', 'substan-dart', 'jami', 'standart va substan-dart', 'muammoli']):
        put(t.cell(1, j), h, True, H, 9)
    for i, row in enumerate(ROWS, 2):
        for j, v in enumerate(row):
            put(t.cell(i, j), v, j == 0)
    for row in t.rows:
        for c, w in zip(row.cells, Wd):
            c.width = w
    par(IZOH, italic=True, size=11, ls=1.0)
    par('')
    par('[1] ' + SRC, size=11, ls=1.0)
    dd.save(DST)
    print(pardoz(DST), DST)

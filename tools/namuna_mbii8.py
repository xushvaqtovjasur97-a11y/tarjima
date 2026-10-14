"""MBII-13 / MBII-8 belgilarini 3.2-§ da birinchi qo‘llangan joyda kiritish (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_MBII8_tarifi.docx'
# 1) 3.2-§: logit paneli tilga olingan joy
ANC1 = 'muammo hodisasining ta’rifi esa 4.1-paragrafda batafsil yoritiladi)'
INS1 = ('. Bu panelda kapital monandligi koeffitsiyentlari, MHXS 9 bo‘yicha kreditlar bosqichlari va valyuta tarkibi bank kesimida '
        'e’lon qilinmagani sababli indeks 13 ta ko‘rsatkichning barcha banklar uchun mavjud bo‘lgan sakkiztasi asosida hisoblanadi. '
        'Ikki indeksni farqlash uchun bundan keyin 13 ko‘rsatkichli to‘liq indeks MBII-13, 8 ko‘rsatkichli panel indeksi esa MBII-8 '
        'deb belgilanadi. MBII-8 da')
OLD1 = '), unda kapital bloki'          # → ). MBII-8 da kapital bloki
# 2) 3.3-§: takroriy ta’rif o‘rniga havola
OLD2 = 'Ikki indeksni farqlash uchun keyingi o‘rinlarda 13 ko‘rsatkichli to‘liq indeks MBII-13, 8 ko‘rsatkichli panel indeksi esa MBII-8 deb belgilanadi.'
NEW2 = 'Bu panel indeksi 3.2-paragrafda kiritilgan MBII-8 belgisi bilan ifodalanadi.'
# 3) sinf chegaralari xatboshisi
OLD3 = 'Ularni empirik tekshirish uchun 25 ta bank panelida'
NEW3 = 'Ularni empirik tekshirish uchun 25 ta bank panelida MBII-8 bo‘yicha'
# 4) 3.2.4-jadval izohi
OLD4 = 'milliy hisobot asosidagi 8 ko‘rsatkichli choraklik indeks;'
NEW4 = 'milliy hisobot asosidagi 8 ko‘rsatkichli choraklik indeks (MBII-8);'

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(4)
    G, R = '7F7F7F', 'FF0000'

    def par(parts, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        q = dd.add_paragraph(); q.alignment = align
        for t, col, s in parts:
            r = q.add_run(t); r.bold = bold; r.font.strike = s
            if col: r.font.color.rgb = RGBColor.from_string(col)
    par([('MBII-13 va MBII-8 belgilari (namuna)', None, False)], True, WD_ALIGN_PARAGRAPH.CENTER)
    par([('1) 3.2-§, 3.2.1-jadvaldan keyingi xatboshi (panel birinchi tilga olingan joy):', None, False)], True)
    par([('… Logit usuli uchun 3.3-paragrafda shakllantirilgan 25 ta bank paneli qo‘llandi (panelning tarkibi 3.3-paragrafda, ', None, False),
         (ANC1, None, False), (INS1, R, False), (' kapital bloki barcha banklarda bir xil ball olgani va valyuta riski bo‘yicha ma’lumotlar '
                                                  'mavjud bo‘lmagani sababli bu ikki blok uchun tayanch vaznlar saqlandi. …', None, False)])
    par([('(“, unda” so‘zi olib tashlanadi, gap “MBII-8 da …” bilan davom etadi.)', G, False)])
    par([('2) 3.3-§ da (ta’rif takrorlanmaydi):', None, False)], True)
    par([(OLD2, G, True), (' ', None, False), (NEW2, R, False)])
    par([('3) Sinf chegaralari xatboshisida:', None, False)], True)
    par([('… ', None, False), (OLD3, None, False), (' MBII-8 bo‘yicha', R, False), (' har bir sinfdagi banklarda … (3.2.4-jadval).', None, False)])
    par([('4) 3.2.4-jadval izohida:', None, False)], True)
    par([('Izoh: 25 ta bank paneli, milliy hisobot asosidagi 8 ko‘rsatkichli choraklik indeks', None, False), (' (MBII-8)', R, False), ('; …', None, False)])
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST))

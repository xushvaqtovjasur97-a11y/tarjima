import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/2.2-paragraf.md'; DST='/home/user/tarjima/paragraflar/2.2-paragraf.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p22_raw.docx'],check=True)
d=docx.Document(S+'/p22_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^2\.2\.\d-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^2\.2\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^2\.2\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'

for t in d.tables:
  n=len(t.columns)
  if n==6:
    first=6.0 if t.rows[0].cells[1].text.strip()=='2021' else 3.6
    rest=(16.4-first)/5
    for row in t.rows:
      for k,c in enumerate(row.cells): c.width=Cm(first if k==0 else rest)
    grid=t._tbl.tblGrid
    for k,gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int((first if k==0 else rest)*567)))
    t.autofit=False
Y=['2021','2022','2023','2024','2025']
rid=W.add_chart(d,W.bar_chart(Y,[('Aktivlardagi ulushi',[81.5,78.2,67.7,65.4,63.5],BLUE),('Kreditlardagi ulushi',[85.8,83.2,70.7,68.8,67.0],ORANGE),('Depozitlardagi ulushi',[66.7,62.5,51.5,50.1,51.0],GREEN)],
  ytitle='Bank tizimi jami ko‘rsatkichiga nisbatan, %',ymax=100,major=20,label_fmt='0.0',gap=60))
W.replace_placeholder(d,'{{FIG:1}}',W.chart_run_xml(rid,16.0,8.0,'2.2.1-rasm'))
rid=W.add_chart(d,W.bar_chart(Y,[('Regulyativ kapital monandligi',[17.5,17.8,17.5,17.4,18.3],BLUE),('Asosiy kapital (CET1) monandligi',[14.6,14.4,14.1,14.3,14.7],ORANGE)],
  ytitle='%',ymax=22,major=2,label_fmt='0.0',gap=80))
W.replace_placeholder(d,'{{FIG:2}}',W.chart_run_xml(rid,16.0,8.0,'2.2.2-rasm'))
rid=W.add_chart(d,W.bar_chart(Y,[('Muammoli kreditlar, jami',[5.2,3.6,3.5,4.0,3.0],BLUE),('Muammoli kreditlar, davlat banklari',[5.4,3.9,4.0,3.9,3.2],LBLUE),
  ('Muammoli kreditlar, boshqa banklar',[4.1,2.1,2.4,4.1,2.6],GOLD),('Substandart va muammoli kreditlar',[12.5,15.9,20.0,21.0,15.0],ORANGE)],
  ytitle='Jami kreditlarga nisbatan, %',ymax=24,major=4,label_fmt='0.0',gap=50))
W.replace_placeholder(d,'{{FIG:3}}',W.chart_run_xml(rid,16.0,8.0,'2.2.3-rasm'))
rid=W.add_chart(d,W.scatter_chart([
  ('ROA',[2021,2022,2023,2024,2025],[1.4,2.5,2.6,1.4,2.2],BLUE,28575,None,'circle',('t',[0,1,2,3,4])),
  ('ROE',[2021,2022,2023,2024,2025],[6.1,13.3,14.2,6.6,12.4],ORANGE,28575,None,'square',('t',[0,1,2,3,4]))],
  'Yillar','%',2020.5,2025.5,1,0,16,2,xfmt='0',legend=True))
W.replace_placeholder(d,'{{FIG:4}}',W.chart_run_xml(rid,16.0,8.0,'2.2.4-rasm'))
rid=W.add_chart(d,W.bar_chart(Y,[('«O‘zsanoatqurilishbank»',[15.8,15.3,16.1,15.6,17.0],BLUE),('«Mikrokreditbank»',[13.8,20.7,18.6,15.4,18.5],GREEN),('Regulyativ minimum',[13,13,13,13,13],GREY)],
  ytitle='%',ymax=24,major=4,label_fmt='0.0',gap=70))
W.replace_placeholder(d,'{{FIG:5}}',W.chart_run_xml(rid,16.0,8.0,'2.2.5-rasm'))
rid=W.add_chart(d,W.bar_chart(Y,[('SQB: muammoli kreditlar',[3.8,2.8,2.2,2.8,2.5],DBLUE),('SQB: substandart kreditlar',[6.2,5.3,11.0,14.8,18.4],LBLUE),
  ('MK: muammoli kreditlar',[5.9,4.8,6.0,4.6,4.4],ORANGE),('MK: substandart kreditlar',[9.1,19.3,24.5,21.3,9.0],GOLD)],
  ytitle='Jami kreditlarga nisbatan, %',ymax=28,major=4,label_fmt='0.0',gap=50))
W.replace_placeholder(d,'{{FIG:6}}',W.chart_run_xml(rid,16.0,8.0,'2.2.6-rasm'))
rid=W.add_chart(d,W.bar_chart(Y,[('«O‘zsanoatqurilishbank»',[16.0,7.6,11.0,13.9,17.2],BLUE),('«Mikrokreditbank»',[2.2,0.7,1.1,-45.5,3.5],ORANGE)],
  ytitle='%',ymax=20,ymin=-50,major=10,label_fmt='0.0',gap=70))
W.replace_placeholder(d,'{{FIG:7}}',W.chart_run_xml(rid,16.0,8.0,'2.2.7-rasm'))
d.save(DST); print('saved')

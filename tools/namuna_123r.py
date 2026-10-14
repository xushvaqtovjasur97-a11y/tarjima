"""1.2.3-rasm: o‘q nomi qisqartirildi, legenda bir xillashtirildi, pog‘onalarga foiz yozildi,
x o‘qi qiymatlari vergul bilan; matnga kontrsiklik bufer izohi (namuna)."""
import re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_1.2.3-rasm.docx'
CAP = '1.2.3-rasm. Konservatsiya buferi'
NAME0 = 'Qoldiriladigan foyda ulushi'
YT = 'Qoldiriladigan foyda ulushi, %'
XS = [4.8125, 5.4375, 6.0625, 6.6875, 7.3125]; YS = [100, 80, 60, 40, 0]
LS = ['100 %', '80 %', '60 %', '40 %', '0 %']
XB = [4.5, 5.125, 5.75, 6.375, 7.0]; LB = ['4,5', '5,125', '5,75', '6,375', '7,0']
INS_AFTER = 'esa 40 foizi bankda qoldirilishi kerak.'
INS = (' Asosiy kapital 7 foizdan yuqori bo‘lganda esa foydani taqsimlashga cheklov qo‘llanmaydi. '
       'Shuni ta’kidlash joizki, ushbu shkala kontrsiklik bufer nolga teng bo‘lgan holat uchun keltirilgan bo‘lib, '
       'kontrsiklik bufer joriy etilganda har bir oraliq bufer miqdoriga mutanosib ravishda kengayadi.')


def _num(tag, ref, vals):
    pts = ''.join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals))
    return (f'<c:{tag}><c:numRef><c:f>{ref}</c:f><c:numCache><c:formatCode>General</c:formatCode>'
            f'<c:ptCount val="{len(vals)}"/>{pts}</c:numCache></c:numRef></c:{tag}>')


def _dlbls(texts, pos, col, sz, bold):
    one = ''.join(
        f'<c:dLbl><c:idx val="{i}"/><c:tx><c:rich><a:bodyPr/><a:lstStyle/><a:p><a:pPr>{W._rpr(sz, bold, col)}</a:pPr>'
        f'<a:r><a:rPr lang="uz-Latn-UZ" sz="{sz}" b="{int(bold)}"><a:solidFill><a:srgbClr val="{col}"/></a:solidFill>'
        f'<a:latin typeface="Times New Roman"/><a:cs typeface="Times New Roman"/></a:rPr><a:t>{t}</a:t></a:r></a:p></c:rich></c:tx>'
        f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr><c:dLblPos val="{pos}"/><c:showLegendKey val="0"/><c:showVal val="1"/>'
        f'<c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>'
        for i, t in enumerate(texts))
    return ('<c:dLbls>' + one + '<c:showLegendKey val="0"/><c:showVal val="0"/><c:showCatName val="0"/>'
            '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


def _ser(idx, name, col, xs, ys, texts, pos, lcol, sz, bold):
    c = chr(ord('A') + 2 * idx)
    c2 = chr(ord('A') + 2 * idx + 1)
    return (f'<c:ser><c:idx val="{idx}"/><c:order val="{idx}"/><c:tx><c:strRef><c:f>Sheet1!${c2}$1</c:f><c:strCache>'
            f'<c:ptCount val="1"/><c:pt idx="0"><c:v>{name}</c:v></c:pt></c:strCache></c:strRef></c:tx>'
            '<c:spPr><a:ln><a:noFill/></a:ln></c:spPr><c:marker><c:symbol val="none"/></c:marker>' +
            _dlbls(texts, pos, lcol, sz, bold) +
            _num('xVal', f'Sheet1!${c}$2:${c}${len(xs)+1}', xs) + _num('yVal', f'Sheet1!${c2}$2:${c2}${len(ys)+1}', ys) +
            '<c:smooth val="0"/></c:ser>')


def chart(xml):
    xml = xml.replace('<c:v>Taqsimlanmaydigan foyda ulushi</c:v>', f'<c:v>{NAME0}</c:v>')
    xml = xml.replace('<a:t>Taqsimlanmasdan qoldiriladigan foyda ulushi, %</a:t>', f'<a:t>{YT}</a:t>')
    sers = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)
    add = (_ser(2, 'Pog‘ona darajasi', None, XS, YS, LS, 't', '1F4E79', 1050, True) +
           _ser(3, 'X qiymatlari', None, XB, [0] * 5, LB, 'b', '000000', 1100, False))
    xml = xml.replace(sers[1], sers[1] + add, 1)
    # x o‘qi: raqamlar yashiriladi (vergulli yorliqlar bilan almashtirildi)
    xax = re.search(r'<c:valAx><c:axId val="333"/>.*?</c:valAx>', xml, re.S).group(0)
    xml = xml.replace(xax, xax.replace('<c:tickLblPos val="low"/>', '<c:tickLblPos val="none"/>'), 1)
    xml = re.sub(r'(<c:legend><c:legendPos val="[a-z]+"/>)',
                 r'\1<c:legendEntry><c:idx val="2"/><c:delete val="1"/></c:legendEntry>'
                 r'<c:legendEntry><c:idx val="3"/><c:delete val="1"/></c:legendEntry>', xml, count=1)
    X1 = [4.5, 5.125, 5.125, 5.75, 5.75, 6.375, 6.375, 7, 7, 7.625]
    Y1 = [100, 100, 80, 80, 60, 60, 40, 40, 0, 0]
    rows = [['x', NAME0, 'x (chegara)', 'Pog‘ona chegaralari', 'x (pog‘ona)', 'Pog‘ona darajasi', 'x', 'X qiymatlari']]
    for i in range(10):
        rows.append([X1[i], Y1[i]] + ([XB[i], YS[i], XS[i], YS[i], XB[i], 0] if i < 5 else [None] * 6))
    return xml, rows


def asl():
    d = docx.Document(F)
    cap = [p for p in d.paragraphs if p.text.startswith(CAP)][0]
    rid = [e.get(qn('r:id')) for e in cap._p.getprevious().iter('{http://schemas.openxmlformats.org/drawingml/2006/chart}chart')][0]
    return d.part.related_parts[rid].blob.decode('utf-8')


if __name__ == '__main__':
    xml, rows = chart(asl())
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)

    def par(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, bold=False):
        q = dd.add_paragraph(); q.alignment = align
        for t, col in parts:
            r = q.add_run(t); r.font.size = Pt(size); r.bold = bold
            if col: r.font.color.rgb = RGBColor.from_string(col)
        return q
    par([('1.2.3-rasm (namuna)', None)], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, (xml, rows))
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.5, '1.2.3-rasm')))
    par([('1.2.3-rasm. Konservatsiya buferi bo‘yicha foydani taqsimlash cheklovlari shkalasi', None)],
        WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    dd.add_paragraph()
    par([('Rasmdagi o‘zgarishlar:', None)], size=12, bold=True)
    for t in ['1) chap o‘q nomi qisqartirildi va to‘liq ko‘rinadi: “' + YT + '”;',
              '2) legendadagi nom o‘q nomi bilan bir xillashtirildi: “' + NAME0 + '”;',
              '3) har bir pog‘ona ustiga qoldiriladigan foyda ulushi yozildi (100 %, 80 %, 60 %, 40 %, 0 %);',
              '4) gorizontal o‘qdagi qiymatlar vergul bilan berildi (4,5; 5,125; 5,75; 6,375; 7,0), ortiqcha 7,625 olib tashlandi.']:
        par([(t, None)], size=12)
    dd.add_paragraph()
    par([('Matndagi qo‘shimcha (rasmdan keyingi xatboshi):', None)], size=12, bold=True)
    par([('Rasmdan ko‘rinadiki, asosiy kapital 4,5–5,125 foiz oralig‘ida bo‘lsa foydaning 100 foizi, 5,125–5,75 foiz oralig‘ida 80 foizi, '
          '5,75–6,375 foiz oralig‘ida 60 foizi, 6,375–7 foiz oralig‘ida ' + INS_AFTER, None), (INS, 'FF0000'), (' Bu shkala …', None)])
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

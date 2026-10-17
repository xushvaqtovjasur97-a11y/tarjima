"""3.2.4-jadval: nomi va izohi (davr, indeks turi); keyingi xatboshida 'tayanch vaznli indeks' (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_3.2.4-jadval.docx'
CAP_OLD = '2021–2025-yillar'
CAP_NEW = '2021–2024-yillar kuzatuvlari'
CAP_FULL = 'Integral indeks sinflari bo‘yicha muammo hodisalarining yuz berish chastotasi, '
IZOH = ('Izoh: 25 ta bank paneli, milliy hisobot asosidagi 8 ko‘rsatkichli choraklik indeks; muammo hodisalari har bir kuzatuvdan '
        'keyingi to‘rt chorak, ya’ni 2025-yil oxirigacha bo‘lgan davrda kuzatilgan. 2025-yil kuzatuvlari keyingi to‘rt chorak '
        'ma’lumotlari to‘liq bo‘lmagani sababli kiritilmagan.')
OLD3 = 'Shundan so‘ng empirik asoslangan indeks'
NEW3 = 'Shundan so‘ng empirik tekshiruvdan o‘tgan, tayanch vaznli indeks'

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(4)
    G, R = '7F7F7F', 'FF0000'

    def par(parts, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        q = dd.add_paragraph(); q.alignment = align
        for t, col, s in parts:
            r = q.add_run(t); r.bold = bold; r.font.strike = s
            if col: r.font.color.rgb = RGBColor.from_string(col)
    par([('3.2.4-jadval (namuna)', None, False)], True, WD_ALIGN_PARAGRAPH.CENTER)
    par([('1) Jadval nomi:', None, False)], True)
    par([(CAP_FULL, None, False), (CAP_OLD, G, True), (' ', None, False), (CAP_NEW, R, False)], True, WD_ALIGN_PARAGRAPH.CENTER)
    par([('2) Jadval ostida yangi izoh:', None, False)], True)
    par([(IZOH, R, False)])
    par([('3) Jadvaldan keyingi xatboshida:', None, False)], True)
    par([(OLD3, G, True), (' ', None, False), (NEW3, R, False),
         (' beshta bankning 2021–2025-yillardagi moliyaviy barqarorligi dinamikasini baholash uchun qo‘llandi (3.2.2-rasm).', None, False)])
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST))

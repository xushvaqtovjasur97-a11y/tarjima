import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/3.1-paragraf.md'; DST='/home/user/tarjima/paragraflar/3.1-paragraf.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p31_raw.docx'],check=True)
d=docx.Document(S+'/p31_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^3\.1\.\d-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^3\.1\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^3\.1\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'

for t in d.tables:
  n=len(t.columns)
  if n==6:
    first=6.0 if t.rows[0].cells[1].text.strip()=='2021' else 3.6
    rest=(16.4-first)/5
    for row in t.rows:
      for k,c in enumerate(row.cells): c.width=Cm(first if k==0 else rest)
    grid=t._tbl.tblGrid
    for k,gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int((first if k==0 else rest)*567)))
    t.autofit=False

def setw(t,ws):
  for row in t.rows:
    for k,c in enumerate(row.cells): c.width=Cm(ws[k])
  for k,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(ws[k]*567)))
  t.autofit=False

import sys
sys.path.insert(0,'/home/user/tarjima/manbalar/hisob_kitob')
import mbii as M
def shade(row,fill='F2F2F2'):
  for c in row.cells:
    tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),fill); tcPr.append(sh)
for t in d.tables:
  h=[c.text.strip() for c in t.rows[0].cells]
  if h[0]=='Aniqlangan muammo': setw(t,[5.4,5.2,5.8])
  elif h[0]=='Blok va ko‘rsatkich':
    setw(t,[6.2,2.4,1.95,1.95,1.95,1.95])
    for row in t.rows[1:]:
      if row.cells[1].text.strip()=='': shade(row)
  elif h[0]=='Blok': setw(t,[3.2,2.2,2.0,2.4,6.6])
  elif h[0]=='Sinf': setw(t,[2.6,2.2,5.2,6.4])
  elif h[0]=='Ko‘rsatkich': setw(t,[5.6,2.7,2.7,2.7,2.7])
  elif h[0]=='Belgi': setw(t,[3.0,3.2,3.4,3.2,3.6])
  for row in t.rows[1:]:
    for c in row.cells[1:]:
      for p in c.paragraphs:
        if h[0] in ('Blok va ko‘rsatkich','Ko‘rsatkich') or (h[0]=='Blok' and c is not row.cells[-1]) or (h[0]=='Sinf' and c is row.cells[1]): p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
# 3.1.1 algoritm
PURPLE='7030A0'; RED='C00000'
sh=[]
st=[('1','Ma’lumotlar:\nmilliy hisobot /\nMHXS hisoboti',DBLUE),('2','MHXS yo‘q bo‘lsa:\ntuzatish\nkoeffitsiyentlari',ORANGE),('3','14 ta ko‘rsatkichni\nhisoblash',BLUE),('4','Uzluksiz\nshkalada ballash\n(1–4)',BLUE),
    ('5','Bloklar\nballari',GREEN),('6','Vaznlash:\n0,25; 0,30; 0,20;\n0,15; 0,10',GREEN),('7','MBII va\nbarqarorlik\nsinfi',PURPLE),('8','Nazorat choralari,\nΔMBII, erta\nogohlantirish',RED)]
bw,bh,gx=3.6,2.0,0.6
for i,(n,t,c) in enumerate(st):
  r,cix=divmod(i,4)
  if r==1: cix=3-cix
  x=0.1+cix*(bw+gx); y=0.2+r*3.0
  sh.append(W.box(x,y,bw,bh,n+'. '+t,fill=c,line=c,sz=20,bold=True,color='FFFFFF'))
for i in range(7):
  r,cix=divmod(i,4); r2,c2=divmod(i+1,4)
  if r==0 and r2==0: x=0.1+cix*(bw+gx)+bw; sh.append(W.line(x+0.02,1.2,x+gx-0.02,1.2,color='404040',lw=19050))
  elif r==0 and r2==1: x=0.1+3*(bw+gx)+bw/2; sh.append(W.line(x,2.22,x,3.18,color='404040',lw=19050))
  else:
    cix=3-cix; x=0.1+cix*(bw+gx); sh.append(W.line(x-0.02,4.2,x-gx+0.02,4.2,color='404040',lw=19050))
W.replace_placeholder(d,'{{FIG:1}}',W.group_run_xml(sh,16.5,5.4,'3.1.1-rasm'))
# 3.1.2 ballash funksiyasi
rid=W.add_chart(d,W.scatter_chart([
  ('Uzluksiz shkala (taklif)',[3,4.5,5.75,7,8,9.25,10.5,13],[1,1,1.5,2,3,3.5,4,4],BLUE,28575,None,'circle',None),
  ('Pog‘onali shkala',[3,7,7,8,8,10.5,10.5,13],[1,1,2,2,3,3,4,4],ORANGE,19050,'dash',None,None)],
  'Birinchi darajali asosiy kapital monandligi, %','Ball',3,13,1,0.5,4.5,0.5,xfmt='0',yfmt='0.0',legend=True))
W.replace_placeholder(d,'{{FIG:2}}',W.chart_run_xml(rid,16.0,8.0,'3.1.2-rasm'))
B=['K','A','L','D','V']; BN=['Kapital','Aktivlar sifati','Likvidlik','Daromadlilik','Valyuta riski']
ser=[]
for k,c in [('SQB milliy',LBLUE),('SQB MHXS',DBLUE),('MKB milliy','A9D18E'),('MKB MHXS','38761D')]:
  m,bs=M.mbii(M.DATA[k],M.W['taklif']); ser.append((k.replace('SQB','«O‘zsanoatqurilishbank»:').replace('MKB','«Mikrokreditbank»:'),[round(bs[b],2) for b in B],c))
rid=W.add_chart(d,W.bar_chart(BN,ser,ytitle='Ball',ymax=4,major=1,label_fmt='0.00',gap=50))
W.replace_placeholder(d,'{{FIG:3}}',W.chart_run_xml(rid,16.0,8.0,'3.1.3-rasm'))
V=['SQB milliy','SQB MHXS','MKB milliy','MKB MHXS']; VN=['SQB: milliy','SQB: MHXS','MKB: milliy','MKB: MHXS']
ser=[(n,[round(M.mbii(M.DATA[v],M.W[w])[0],2) for v in V],c) for w,n,c in [('taklif','Taklif etilgan vaznlar',BLUE),('teng','Teng vaznlar',GREY),('CAMEL','CAMEL nisbati',ORANGE)]]
rid=W.add_chart(d,W.bar_chart(VN,ser,ytitle='MBII',ymax=4,major=0.5,label_fmt='0.00',gap=60))
W.replace_placeholder(d,'{{FIG:4}}',W.chart_run_xml(rid,16.0,8.0,'3.1.4-rasm'))
d.save(DST); print('saved')

"""2.2.2-jadvalni uch bank bo‘yicha uchta jadvalga ajratish va har biri bo‘yicha to‘liq tahlil (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import namuna_222j as J
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_2.2.2-jadval_3_bank_tahlil.docx'
RED, BLUE, GREY = 'FF0000', '0070C0', '595959'
F_SQB_IFRS = 'Joint Stock Commercial Bank “Uzbek Industrial and Construction Bank” and its subsidiaries. Consolidated Financial Statements for the year ended 31 December 2025 and Independent Auditor’s Report. – Tashkent, 2026. – P. 28–31.'
F_SQB = '“O‘zsanoatqurilishbank” ATB ' + J.FORM.format('crs003', 30)
F_MKB = '“Mikrokreditbank” ATB ' + J.FORM.format('crs005', 30)
F_ALQ = '“Aloqabank” ATB ' + J.FORM.format('crs031', 31)
F_HEN = 'Hendrickson J.M. Regulation and Instability in U.S. Commercial Banking: A History of Crises. – Basingstoke: Palgrave Macmillan, 2011. – P. 225–227.'

# matn: [(qism, rang, snoska yoki None)]
SQB1 = [('2.2.2-jadval ma’lumotlari “O‘zsanoatqurilishbank” ATB moliyaviy barqarorligi ikki xil yo‘nalishda rivojlanganini ko‘rsatadi. Bankning aktivlari 2021–2025-yillarda '
         '56,5 trln so‘mdan 102,4 trln so‘mga, ya’ni 1,8 barobar o‘sgan. Regulyativ kapital monandligi butun davr mobaynida 15,3–17,0 foiz oralig‘ida bo‘lib, 13 foizlik '
         'minimumdan atigi 2,3–4,0 foiz punkt yuqori turgan, asosiy kapital monandligi esa 2022-yildagi 12,1 foizdan 2025-yilda 10,2 foizgacha izchil pasayib, 8 foizlik '
         'minimumga yaqinlashgan. 2025-yilda birinchi darajali kapital monandligining 14 foizgacha oshishi asosan xalqaro bozorda joylashtirilgan qo‘shimcha birinchi darajali '
         'kapital obligatsiyalari hisobiga ta’minlangan, sof asosiy kapital esa o‘smagan.', RED, None)]
SQB2 = [('Aktivlar sifati bo‘yicha rasmiy muammoli kreditlar ulushi 2,2–3,8 foiz oralig‘ida past darajada saqlangan, biroq substandart kreditlar ulushi 2021-yildagi 6,2 foizdan '
         '2025-yilda 18,4 foizgacha, ya’ni qariyb uch barobar oshgan. Bu kredit riskining yashirin to‘planishidan dalolat beradi va MHXS hisoboti bilan tasdiqlanadi: 2025-yil '
         'yakunida bankning MHXS 9 bo‘yicha 3-bosqich kreditlari ulushi 7,48 foizni, 2-bosqich kreditlari ulushi esa 35,7 foizni tashkil etgan', RED, F_SQB_IFRS),
        ('. Daromadlilik 2022-yildagi pasayishdan (ROE 7,6 foiz) keyin tiklanib, 2025-yilda ROE 17,2 foizga yetgan, ya’ni tizim o‘rtacha darajasidan yuqori. Likvidlik '
         'koeffitsiyentlari Bazel III talablaridan yuqori bo‘lsa-da, 2023-yilda NSFR 104,2 foizgacha pasaygan. Kreditlarning depozitlarga nisbati esa 206,5–414,0 foiz bo‘lib, '
         'bank kreditlarni asosan xalqaro moliya institutlari, xorijiy banklar va davlat fondlarining resurslari hisobiga moliyalashtiradi: 2025-yil yakunida bankning olingan '
         'kreditlari 35,5 trln so‘mni tashkil etib, depozitlardan qariyb 1,1 barobar ko‘p bo‘lgan', RED, F_SQB), ('.', RED, None)]
MKB1 = [('“Mikrokreditbank” ATB ning profili tubdan farq qiladi (2.2.3-jadval). Bankning aktivlari 13,5 trln so‘mdan 30,7 trln so‘mga, ya’ni 2,3 barobar o‘sgan. Kapital '
         'monandligi yuqoriroq bo‘lsa-da, keskin tebranadi: 2021-yilda 13,8 foiz bilan minimum chegarasida bo‘lgan ko‘rsatkich 2022-yilda 20,7 foizga ko‘tarilgan, 2024-yilda '
         '15,4 foizgacha tushib, 2025-yilda yana 18,5 foizga yetgan. Bu tebranishlarning sababi kapital manbalarida: bankning xususiy kapitali 2022-yilda 2,3 trln so‘mdan '
         '3,6 trln so‘mga oshgan, holbuki shu yili bank atigi 21,3 mlrd so‘m sof foyda olgan; 2025-yilda ham kapital 3,6 trln so‘mdan 5,7 trln so‘mga ko‘paygan, sof foyda esa '
         '160,8 mlrd so‘mni tashkil etgan. Demak, bankning kapital yetarliligi o‘z foydasi hisobiga emas, balki davlat tomonidan amalga oshirilgan qayta kapitallashtirish '
         'hisobiga saqlab turilgan. Bu holat 1.3-paragrafda keltirilgan AQSh tajribasini eslatadi: J.M. Hendrickson ijtimoiy kreditlash majburiyati bank kapitaliga qo‘yiladigan '
         'talablarga zid kelishini va bunday kreditlash CAMELS reytingini yomonlashtirishini ko‘rsatgan edi', RED, F_HEN), ('.', RED, None)]
MKB2 = [('Aktivlar sifati ham zaif: rasmiy muammoli kreditlar ulushining o‘zi 4,4–6,0 foiz bo‘lib, 2021 va 2023-yillarda 5 foizlik erta ogohlantirish chegarasidan oshgan, '
         'substandart kreditlar ulushi esa 2023-yilda 24,5 foizga yetgan. Bu muammolar daromadlilikka bevosita ta’sir ko‘rsatgan: ROE 2021–2023-yillarda 0,7–1,9 foiz bilan '
         'inflyatsiya darajasidan ancha past bo‘lgan, 2024-yilda esa bank 1 703,8 mlrd so‘m zarar ko‘rgan. Bu zarar aktivlar bo‘yicha ehtimoliy yo‘qotishlarga qarshi '
         'zaxiralarning keskin oshirilishi natijasida yuzaga kelgan: zaxiralar ajratilgandan keyingi sof foizli daromad 1 167,3 mlrd so‘m zararni tashkil etgan', RED, F_MKB),
        ('. Boshqacha aytganda, ilgari kamsitib ko‘rsatilgan kredit riski 2024-yilda bir yo‘la tan olingan va kapitalni yemirgan, keyin esa yana davlat kapitali hisobiga '
         'to‘ldirilgan. Likvidlik bo‘yicha bank Bazel III talablarini bajargan (LCR 131,7–231,9 foiz, NSFR 108,3–117,1 foiz), ammo kreditlarning depozitlarga nisbati '
         '208,5–343,0 foiz bo‘lib, u ham kreditlarni asosan qarz mablag‘lari hisobiga moliyalashtiradi.', RED, None)]
ALQ1 = [('“Aloqabank” ATB tizimli ahamiyatga molik banklar qatoriga kirmaydi va tahlilga natijalarning umumiyligini tekshirish uchun kiritildi (2.2.4-jadval). Bankning aktivlari '
         '12,6 trln so‘mdan 31,5 trln so‘mga, ya’ni 2,5 barobar o‘sgan. Kapital bo‘yicha zaxira uch bank ichida eng kichik: regulyativ kapital monandligi 2023-yildan tashqari '
         'barcha yillarda 12,5–14,1 foiz bo‘lib, 13 foizlik minimum atrofida bo‘lgan, 2022-yilda esa hisoblangan qiymat minimumdan 0,5 foiz punkt past bo‘lgan; asosiy kapital '
         'monandligi 9,4–11,8 foiz oralig‘ida o‘zgargan. 2023-yildagi kapital o‘sishi (847,8 mlrd so‘m) asosan o‘z foydasi (607,7 mlrd so‘m) hisobiga ta’minlangan, '
         '2024–2025-yillarda esa kapital o‘sishining yarmidan ko‘pi tashqi manbalarga to‘g‘ri kelgan', BLUE, F_ALQ), ('.', BLUE, None)]
ALQ2 = [('Rasmiy muammoli kreditlar ulushi 1,5–4,1 foiz bilan uch bank ichida eng past, ammo substandart kreditlar ulushi 2022-yildagi 1,1 foizdan 2023-yilda 11,8 foizga keskin '
         'ko‘tarilgan va 2025-yilda 6,4 foizni tashkil etgan. Daromadlilik 2023-yilda eng yuqori bo‘lgan (ROA 3,8 foiz, ROE 27,4 foiz), keyingi yillarda esa ROE 6,6–7,3 foizgacha '
         'pasaygan. Bankning yuqoridagi ikki bankdan asosiy farqi moliyalashtirish tuzilmasida: kreditlarning depozitlarga nisbati 93,7–121,1 foiz bo‘lib, u kreditlarni asosan '
         'mijozlar depozitlari hisobiga moliyalashtiradi. Biroq likvidlik zaxirasi tor: LCR 2023-yilda 103,1 foizgacha, NSFR esa 2022-yilda 100,9 foizgacha tushib, Bazel III '
         'minimumiga yaqinlashgan.', BLUE, None)]
QIY = [('Uch bankni o‘zaro solishtirish ularning moliyaviy barqarorlik profili turlicha bo‘lsa-da, umumiy qonuniyatlar mavjudligini ko‘rsatadi. Birinchidan, uchala bankda ham '
        'kapital bo‘yicha zaxira kichik yoki beqaror bo‘lib, kapital yetarliligi ko‘p hollarda bankning o‘z foydasi emas, balki aksiyador – davlat tomonidan qo‘shimcha kapital '
        'kiritilishi hisobiga saqlab turilgan. Ikkinchidan, uchala bankda ham rasmiy muammoli kreditlar ulushi past yoki o‘rtacha bo‘lgani holda substandart kreditlar ulushi '
        'tahlil davrining ayrim yillarida 11,8–24,5 foizgacha ko‘tarilgan, ya’ni kredit riski rasmiy ko‘rsatkichdan tashqarida to‘planmoqda. Uchinchidan, moliyalashtirish '
        'tuzilmasi bo‘yicha banklar ikki guruhga bo‘linadi: “O‘zsanoatqurilishbank” ATB va “Mikrokreditbank” ATB kreditlarni asosan tashqi qarzlar hisobiga, “Aloqabank” ATB '
        'esa depozitlar hisobiga moliyalashtiradi. Rasmiy ko‘rsatkichlarning bu manzarasini uchala bankning auditdan o‘tgan MHXS bo‘yicha moliyaviy hisobotlari bilan '
        'solishtirish kredit riskining haqiqiy ko‘lamini aniqroq ko‘rsatadi (2.2.5-rasm).', RED, None)]

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)
    FN = []

    def par(parts, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=14, ls=None):
        q = dd.add_paragraph(); q.alignment = align
        q.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
        if ls:
            q.paragraph_format.line_spacing = ls
        for text, col, fn in parts:
            r = q.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
            if col:
                r.font.color.rgb = RGBColor.from_string(col)
            if fn:
                FN.append(fn)
                s = q.add_run(str(len(FN))); s.font.superscript = True; s.font.size = Pt(size)
                if col:
                    s.font.color.rgb = RGBColor.from_string(col)
        return q

    def table(num, title, data, color, src):
        par([(num + '-jadval', color, None)], italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
        par([(title, color, src)], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, ls=1.0)
        t = dd.add_table(rows=len(J.ROWS) + 1, cols=6); t.style = 'Table Grid'; t.autofit = False
        Wd = [Cm(6.0)] + [Cm(2.1)] * 5
        for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), Wd):
            gc.set(qn('w:w'), str(int(w.twips)))
        for ri, row in enumerate([['Ko‘rsatkich'] + J.Y if hasattr(J, 'Y') else ['Ko‘rsatkich', '2021', '2022', '2023', '2024', '2025']] + [[nm] + v for nm, v in zip(J.ROWS, data)]):
            for ci, v in enumerate(row):
                c = t.cell(ri, ci); c.width = Wd[ci]
                q = c.paragraphs[0]; q.paragraph_format.line_spacing = 1.0
                q.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 and ri else WD_ALIGN_PARAGRAPH.CENTER
                s = v if isinstance(v, str) else f'{v:.1f}'.replace('.', ',').replace('-', '−')
                r = q.add_run(s); r.font.size = Pt(12); r.bold = ri == 0
                if ri:
                    r.font.color.rgb = RGBColor.from_string(color)
        par([(J.IZOH, color, None)], italic=True, size=12, indent=False, ls=1.0)

    par([('2.2.2-jadval – uch bank bo‘yicha alohida jadvallar va tahlil (namuna)', None, None)], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    par([('Qizil – yangi/qayta yozilgan matn; ko‘k – “Aloqabank” ATB. Bu qism hozirgi “Tizim darajasidagi agregat ko‘rsatkichlar …” xatboshisidan boshlab 2.2.7-rasmgacha bo‘lgan '
          'qismning o‘rnini egallaydi (2.2.5, 2.2.6, 2.2.8-rasmlar olib tashlanadi, ularning tahlili jadvallardan keyingi matnga ko‘chirildi).', GREY, None)], italic=True, size=12, indent=False)
    par([(J.INTRO, RED, None)])
    table('2.2.2', '“O‘zsanoatqurilishbank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari', J.SQB, RED, J.SRC + F_SQB)
    par(SQB1); par(SQB2)
    table('2.2.3', '“Mikrokreditbank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari', J.MKB, RED, J.SRC + F_MKB)
    par(MKB1); par(MKB2)
    table('2.2.4', '“Aloqabank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari', J.ALQ, BLUE, J.SRC + F_ALQ)
    par(ALQ1); par(ALQ2)
    par(QIY)
    par([('Shundan keyin 2.2.5-rasm (hozirgi 2.2.7-rasm: MHXS 9 bo‘yicha 3-bosqich va rasmiy muammoli kreditlar) va uning tahlili (“ikkala bankda” → “uchala bankda”), '
          'so‘ngra umumlashtiruvchi 2.2.5-jadval (“Aloqabank” ustuni bilan) keladi.', GREY, None)], italic=True, size=12, indent=False)
    par([('Sahifa osti snoskalari', None, None)], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12)
    for i, s in enumerate(FN, 1):
        par([(f'{i} {s}', None, None)], size=11, indent=False, ls=1.0)
    dd.save(DST)
    print(pardoz(DST), DST)

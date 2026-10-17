"""4.2-§ ga ikki yangi rasm (ko‘k): 4.2.1-rasm – muammoli kreditlar ulushi va 4 chorak oldingi YaIM o‘sishi (nuqtalar va
regressiya chizig‘i), 4.2.2-rasm – 25 ta bankda muammoli kreditlar ulushi va ROA (nuqtalar buluti va regressiya chizig‘i).
Mavjud 4.2.1–4.2.3-rasmlar 4.2.3–4.2.5 ga suriladi. Diagrammalar: tools/namuna_42modellar.py (chart_a, chart_b)."""
import copy, re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
from docx.oxml.ns import qn
from lxml import etree
import kamlib as K
import wordart as W
import namuna_42modellar as N
from pardoz import pardoz

K.RED = '0070C0'
F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
W._counter['chart'] = 310; W._counter['docpr'] = 9700
D = K.Doc(F)


def tx(el):
    return ''.join(x.text or '' for x in el.iter(qn('w:t')))


# 1) mavjud rasmlar raqami: 4.2.3→4.2.5, 4.2.2→4.2.4, 4.2.1→4.2.3 (kamayish tartibida)
n = 0
for old, new in (('4.2.3-rasm', '4.2.5-rasm'), ('4.2.2-rasm', '4.2.4-rasm'), ('4.2.1-rasm', '4.2.3-rasm')):
    for p in D.d.paragraphs:
        for _ in range(tx(p._p).count(old)):
            D.replace(p, old, new); n += 1
assert n == 6, n

(A, ia), (B, ib) = N.chart_a(), N.chart_b()
oldcap = D.para('4.2.3-rasm. Bank tizimida muammoli kreditlar ulushining haqiqiy dinamikasi')
chart_ppr = oldcap._p.getprevious().find(qn('w:pPr'))
iz_src = D.para('Izoh: qavs ichida standart xatolar (1-modelda')
FN = 'Tadqiqotlar jarayonida muallif tomonidan tuzilgan.'


def f2(x, k=2):
    return f'{x:.{k}f}'.replace('.', ',').replace('-', '−')


def block(anchor, xr, name, izoh, cap, anal):
    cp = D.d.add_paragraph()
    rid = W.add_chart(D.d, xr)
    cp._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 9.0, name)))
    cp._p.insert(0, copy.deepcopy(chart_ppr))
    anchor._p.addnext(cp._p)
    iz = D.new_para_after(iz_src, izoh); cp._p.addnext(iz._p)
    c = D.new_para_after(oldcap, cap + '[[FN:' + FN + ']]'); iz._p.addnext(c._p)
    a = D.new_para_after(anchor, anal); c._p.addnext(a._p)
    return a


# 2) 4.2.1-rasm: tizim darajasi, YaIM (t−4) va muammoli kreditlar ulushi
p692 = D.para('Prognozning asosini YaIM o‘sishining muammoli kreditlar ulushiga')
D.replace(p692, 'Tanlangan model va bank darajasidagi satellit modellar natijalari 4.2.2-jadvalda keltirilgan.',
          'Muammoli kreditlar ulushi va to‘rt chorak oldingi YaIM o‘sishi o‘rtasidagi bog‘liqlik 4.2.1-rasmda keltirilgan.')
block(p692, A, '4.2.1-rasm',
      f'Izoh: har bir nuqta – bir chorak; regressiya chizig‘i y = {f2(ia["a"])} − {f2(-ia["b"], 3)}x (R² = {f2(ia["r2"])}, n = {ia["n"]}); '
      'oldingi chorak darajasi hisobga olingan baho (−0,151) 4.2.2-jadvalda.',
      '4.2.1-rasm. Muammoli kreditlar ulushi va to‘rt chorak oldingi YaIM o‘sish sur’ati o‘rtasidagi bog‘liqlik, 2021–2025-yillar',
      'Rasmdan ko‘rinib turibdiki, nuqtalar manfiy qiyalikka ega regressiya chizig‘i atrofida joylashgan, ya’ni YaIM o‘sishi '
      'sekinlashgan davrdan to‘rt chorak o‘tib muammoli kreditlar ulushi oshadi. Ikki ko‘rsatkich o‘rtasidagi korrelyatsiya koeffitsiyenti '
      f'{f2(ia["r"])} ga teng bo‘lib, YaIM o‘sishining 1 foiz punktga sekinlashuvi muammoli kreditlar ulushini o‘rtacha {f2(-ia["b"])} foiz punktga '
      'oshiradi. Eng yuqori qiymatlar 2020-yildagi pandemiya shokidan keyingi 2021-yil choraklariga to‘g‘ri keladi. Shu bilan birga, '
      'nuqtalarning chiziqdan sezilarli og‘ishi muammoli kreditlar darajasiga YaIM o‘sishidan tashqari boshqa omillar ham ta’sir qilishini '
      'ko‘rsatadi, shu sababli tizim darajasidagi modelga muammoli kreditlar ulushining oldingi chorakdagi qiymati ham kiritildi. '
      'Tanlangan model va bank darajasidagi satellit modellar natijalari 4.2.2-jadvalda keltirilgan.')

# 3) 4.2.2-rasm: bank paneli, muammoli kreditlar ulushi va ROA
p697 = D.para('Jadvaldan ko‘rinib turibdiki, ikki mustaqil yondashuv YaIM o‘sishining kredit riskiga')
D.insert(p697, '', ' Kredit riski va daromadlilik o‘rtasidagi bog‘liqlik bank darajasida 4.2.2-rasmda keltirilgan.')
block(p697, B, '4.2.2-rasm',
      'Izoh: har bir nuqta – bir bankning bir choragi; ko‘rsatkichlar 1 va 99 foizlik chegaralarda vinzorlangan; regressiya chizig‘i '
      f'barcha kuzatuvlar bo‘yicha y = {f2(ib["a"])} − {f2(-ib["b"])}x (R² = {f2(ib["r2"])}); bank effektlari hisobga olingan baho (−0,201) 4.2.2-jadvalda.',
      '4.2.2-rasm. Tanlangan 25 ta bankda muammoli kreditlar ulushi va aktivlar rentabelligi o‘rtasidagi bog‘liqlik, 2022–2025-yillar',
      'Rasmdan ko‘rinib turibdiki, nuqtalar keng sochilgan bo‘lsa-da, muammoli kreditlar ulushi yuqori bo‘lgan kuzatuvlarda aktivlar '
      f'rentabelligi past (korrelyatsiya koeffitsiyenti {f2(ib["r"])}). Xususan, muammoli kreditlar ulushi 6 foizdan yuqori bo‘lgan 46 ta kuzatuvda '
      'aktivlar rentabelligi o‘rtacha −0,05 foizni, 3 foizdan past bo‘lgan 210 ta kuzatuvda esa 3,78 foizni tashkil etgan. Davlat ulushi '
      'mavjud banklarning nuqtalari asosan rasmning o‘ng-pastki qismida joylashgan, ya’ni ularda muammoli kreditlar ulushi o‘rtacha 4,59 foiz, '
      'aktivlar rentabelligi esa 0,81 foiz bo‘lgan, boshqa banklarda bu ko‘rsatkichlar mos ravishda 2,83 va 3,76 foizni tashkil etgan. '
      f'Barcha kuzatuvlar bo‘yicha regressiya chizig‘ining qiyaligi ({f2(ib["b"])}) 3-modeldagi bahodan (−0,20) yuqori, chunki u banklar o‘rtasidagi '
      'doimiy farqlarni ham o‘z ichiga oladi. Shu sababli prognozda har bir bankning o‘ziga xos xususiyatlari hisobga olingan 3-model bahosi qo‘llandi.')
D.save(F)
print(n, pardoz(F))

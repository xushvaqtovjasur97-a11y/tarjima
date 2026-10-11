import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/4.2-paragraf.md'; DST='/home/user/tarjima/paragraflar/4.2-paragraf.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p42_raw.docx'],check=True)
d=docx.Document(S+'/p42_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^4\.2\.\d-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^4\.2\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^4\.2\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'
PURPLE='7030A0'

def setw(t,ws):
  for row in t.rows:
    for k,c in enumerate(row.cells): c.width=Cm(ws[k])
  for k,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(ws[k]*567)))
  t.autofit=False
for t in d.tables:
  h=[c.text.strip() for c in t.rows[0].cells]; n=len(h)
  if h[0]=='Ssenariy': setw(t,[2.6,1.8,1.8,1.8,8.4])
  elif h[0]=='Ko‘rsatkich' and n==10: setw(t,[3.8]+[1.4]*9)
  elif h[0]=='Ko‘rsatkich' and h[1]=='Guruh': setw(t,[3.8,2.6,2.0,2.0,2.0,2.0,2.0])
  elif h[0]=='Bank va hisobot': setw(t,[3.6,2.6,2.6,2.6,2.5,2.5])
  elif h[0]=='Prognoz qilingan yil': setw(t,[1.8,3.6,3.6,2.4,2.4,2.6])
  elif n==4: setw(t,[4.0,4.2,4.2,4.0])
  for row in t.rows[1:]:
    for c in row.cells[1:]:
      if h[0]!='Ssenariy' or True:
        for p in c.paragraphs: p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
  if h[0]=='Ko‘rsatkich' and h[1]!='Guruh':
    for row in t.rows:
      for c in row.cells:
        for p in c.paragraphs:
          for r in p.runs: r.font.size=Pt(10)
# "Izoh:" paragraflari – 12 pt, xatboshisiz
for p in d.paragraphs:
  if p.text.startswith('Izoh:'):
    p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(6)
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.font.size=Pt(12)
  if p.text.startswith('MBII') and '=' in p.text:
    p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0)
    for r in p.runs: r.italic=True; r.font.size=Pt(13)
import json, numpy as np
O=json.load(open('/home/user/tarjima/manbalar/hisob_kitob/fc42.json'))
import pandas as pd
D=pd.read_csv('/home/user/tarjima/manbalar/hisob_kitob/tizim42.csv',index_col=0).loc['2021_I':'2025_IV']
def xq(k): y,q=k.split('_'); return int(y)-1+{'I':0.25,'II':0.5,'III':0.75,'IV':1.0}[q]   # IV chorak – yil belgisi ostida
hx=[xq(k) for k in D.index]; hy=[round(v,2) for v in D.NPL]
fk=[f'{y}_{q}' for y in (2026,2027,2028) for q in ('I','II','III','IV')]; fx=[xq(k) for k in fk]
N=len(hx)+len(fx)
def pad(xs,ys): return xs+[None]*(N-len(xs)), ys+[None]*(N-len(ys))
S=O['tizim_NPL_yo‘li']
ser=[('Haqiqiy (2021–2025)',*pad(hx,hy),'000000',28575,None,'circle',('t',[2,7,11,15,19]))]
for nm,col,pos in (('Asosiy',BLUE,'t'),('Ijobiy',GREEN,'b'),('Salbiy',GOLD,'t'),('Keskin salbiy','C00000','t')):
  xs=[hx[-1]]+fx; ys=[hy[-1]]+S[nm]
  ser.append((nm+' ssenariy',*pad(xs,ys),col,22225,'dash','circle' if False else None,(pos,[8],col)))
xml,rows=W.scatter_chart(ser,'Yillar','Muammoli kreditlar ulushi, %',2020,2028,1,2.0,6.0,1.0,xfmt='[&lt;2020.5]&quot;&quot;;0',yfmt='0.0',legend=True)
xml=xml.replace('<c:dLbls><c:spPr><a:noFill/>','<c:dLbls><c:spPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>')
rid=W.add_chart(d,(xml,rows))
W.replace_placeholder(d,'{{FIG:1}}',W.chart_run_xml(rid,16.0,9.0,'4.2.1-rasm'))
T=O['tizim']; zc=[('2025','2025','2025 (haqiqiy)')]+[(s,'2027',s) for s in ('Asosiy','Ijobiy','Salbiy','Keskin salbiy')]
rid=W.add_chart(d,W.bar_chart([c[2] for c in zc],[('Yashil zona',[T[f'{a}|{b}']['yashil'] for a,b,_ in zc],GREEN,'FFFFFF'),('Sariq zona',[T[f'{a}|{b}']['sariq'] for a,b,_ in zc],GOLD),('Qizil zona',[T[f'{a}|{b}']['qizil'] for a,b,_ in zc],'C00000','FFFFFF')],ytitle='Banklar soni',stacked=True,ymax=25,major=5,label_fmt='0',gap=60))
W.replace_placeholder(d,'{{FIG:2}}',W.chart_run_xml(rid,16.0,8.0,'4.2.2-rasm'))
B=O['banklar']; cats=['2025 (haqiqiy)','Asosiy','Ijobiy','Salbiy','Keskin salbiy']
def v(code,var): return [B[code]['2025'][var]]+[B[code][f'{s}|2027'][var] for s in cats[1:]]
rid=W.add_chart(d,W.bar_chart(cats,[('SQB: milliy',v('SQB','milliy'),LBLUE),('SQB: MHXS',v('SQB','MHXS'),DBLUE),('MKB: milliy',v('MKB','milliy'),'F4B183'),('MKB: MHXS',v('MKB','MHXS'),ORANGE)],ytitle='MBII, ball',ymax=4,ymin=0,major=0.5,label_fmt='0.00',gap=50))
W.replace_placeholder(d,'{{FIG:3}}',W.chart_run_xml(rid,16.0,8.0,'4.2.3-rasm'))
d.save(DST); print('saved')

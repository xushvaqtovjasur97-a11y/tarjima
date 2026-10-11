import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/3.2-paragraf.md'; DST='/home/user/tarjima/paragraflar/3.2-paragraf.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p32_raw.docx'],check=True)
d=docx.Document(S+'/p32_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^3\.2\.\d-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^3\.2\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^3\.2\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'
PURPLE='7030A0'

def setw(t,ws):
  for row in t.rows:
    for k,c in enumerate(row.cells): c.width=Cm(ws[k])
  for k,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(ws[k]*567)))
  t.autofit=False
for t in d.tables:
  h=[c.text.strip() for c in t.rows[0].cells]; n=len(h)
  if h[0]=='Usul': setw(t,[3.0,5.4,3.8,4.2])
  elif h[0]=='Blok': setw(t,[3.4,2.2,2.2,2.2,2.2,2.0,2.2])
  elif h[0]=='Ko‘rsatkich': setw(t,[5.0,2.3,2.3,2.3,2.2,2.3])
  elif h[0]=='Barqarorlik sinfi': setw(t,[3.4,3.0,2.6,4.4,3.0])
  for row in t.rows[1:]:
    for c in row.cells[1:]:
      if h[0]!='Usul':
        for p in c.paragraphs: p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for p in d.paragraphs:
  if p.text.startswith('Izoh:'):
    p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(6)
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.font.size=Pt(12)
import json, pandas as pd
V=json.load(open('/home/user/tarjima/manbalar/hisob_kitob/vazn32.json'))
Wt=pd.DataFrame(V['vaznlar']); BLK=['K','A','L','D','V']; BN=['Kapital','Aktivlar sifati','Likvidlik','Daromadlilik','Valyuta riski']
cols=[('Ekspert','Ekspert',DBLUE),('Entropiya','Entropiya',ORANGE),('Bosh komponentlar','Bosh komponentlar',GREEN),('Logit (K, V – ekspert)','Logit',GOLD),('O‘rtacha (empirik)','Empirik o‘rtacha',GREY)]
rid=W.add_chart(d,W.bar_chart(BN,[(nm,[round(Wt[c][b],2) for b in BLK],col) for c,nm,col in cols],ytitle='Vazn',ymax=0.6,ymin=0,major=0.1,label_fmt='0.00',gap=40))
W.replace_placeholder(d,'{{FIG:1}}',W.chart_run_xml(rid,16.0,8.5,'3.2.1-rasm'))
NM={'SQB':'SQB','MKB':'MKB','BRB':'BRB','XALQ':'Xalq banki','IPOTEKA':'Ipoteka-bank'}
def per(nm): s=pd.Series(V[nm]['bank_yil']); return s.groupby(lambda k:k.split('|')[0]).mean()
a,b=per('mc_ekspert_atrofida'),per('mc_bir_tekis'); order=['SQB','MKB','BRB','XALQ','IPOTEKA']
rid=W.add_chart(d,W.bar_chart([NM[k] for k in order]+['O‘rtacha'],[('Vaznlar ekspert vaznlari atrofida',[round(a[k],1) for k in order]+[round(a.mean(),1)],BLUE),('Vaznlar to‘liq tasodifiy',[round(b[k],1) for k in order]+[round(b.mean(),1)],ORANGE)],ytitle='%',ymax=100,ymin=0,major=20,label_fmt='0.0',gap=60))
W.replace_placeholder(d,'{{FIG:2}}',W.chart_run_xml(rid,16.0,8.0,'3.2.2-rasm'))
R=json.load(open('/home/user/tarjima/manbalar/hisob_kitob/mbii32.json'))
B5=[('XALQ','«Xalq banki»',GREEN),('BRB','«Biznesni rivojlantirish banki»',GOLD),('IPOTEKA','«Ipoteka-bank»',PURPLE),('SQB','«O‘zsanoatqurilishbank»',BLUE),('MKB','«Mikrokreditbank»',ORANGE)]
Y=['2021','2022','2023','2024','2025']; mk=['circle','square','triangle','diamond','x']
ser=[(n,[2021,2022,2023,2024,2025],[round(R[c][y]['mbii'],2) for y in Y],col,28575,None,mk[i],('t',[2,4])) for i,(c,n,col) in enumerate(B5)]
rid=W.add_chart(d,W.scatter_chart(ser,'Yillar','MBII, ball',2020.5,2025.5,1,1.5,3.5,0.5,xfmt='0',yfmt='0.00',legend=True))
W.replace_placeholder(d,'{{FIG:3}}',W.chart_run_xml(rid,16.0,8.0,'3.2.3-rasm'))
d.save(DST); print('saved')

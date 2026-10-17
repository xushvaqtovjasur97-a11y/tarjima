"""3.2.1-rasm (Monte-Karlo): qisqartmalar izohi, vergulli yorliqlar; matnda Dirixle parametri va sezgir banklar; umumiy izoh (namuna)."""
import re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W

DST = '/home/user/tarjima/Namuna_3.2.1-rasm_MK.docx'
CATS = ['SQB', 'MKB', 'BRB', 'Xalq banki', 'Ipoteka-bank', 'O‘rtacha']
IZOH = 'Izoh: SQB – “O‘zsanoatqurilishbank” ATB, MKB – “Mikrokreditbank” ATB, BRB – “Biznesni rivojlantirish banki” ATB.'
SER = [('Vaznlar tayanch vaznlar atrofida', [91.7, 76.9, 71.5, 78.1, 76.1, 78.9], '2A78D6'),
       ('Vaznlar to‘liq tasodifiy', [57.0, 52.7, 58.6, 57.7, 53.1, 55.8], 'EB6834')]
ANC1 = 'vaznlar tayanch vaznlar atrofida tanlangan holatda'
INS1 = ' (Dirixle taqsimotining parametri α = 20·w, bu yerda w – blokning tayanch vazni)'
ANC2 = 'kuratorlarning asoslangan mulohazasiga tayanishi lozim.'
INS2 = (' Tahlilda bunday holat, masalan, “Biznesni rivojlantirish banki” ATB ning 2021-yildagi ko‘rsatkichida kuzatilgan: vaznlar '
        'tayanch vaznlar atrofida o‘zgarganda uning sinfi atigi 51,2 foiz holatda saqlanadi, vaznlar to‘liq tasodifiy olinganda esa eng '
        'sezgir kuzatuv “Xalq banki” ATB ning 2024-yildagi ko‘rsatkichi bo‘lib, uning sinfi 39,4 foiz holatda saqlanadi.')
FN_NEW = ('Tadqiqot natijasida muallif tomonidan “O‘zsanoatqurilishbank”, “Mikrokreditbank”, “Biznesni rivojlantirish banki”, '
          '“Xalq banki” va “Ipoteka-bank” ATBlarining 2021–2025-yillardagi MHXS bo‘yicha moliyaviy hisobotlari va “SNS” reyting '
          'agentligi bazasi asosida hisoblandi.')


def _lbls(vals, col):
    one = ''.join(
        f'<c:dLbl><c:idx val="{i}"/><c:tx><c:rich><a:bodyPr wrap="none"/><a:lstStyle/><a:p><a:pPr>{W._rpr(1000, True, "000000")}</a:pPr>'
        f'<a:r><a:rPr lang="uz-Latn-UZ" sz="1000" b="1"><a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
        f'<a:latin typeface="Times New Roman"/><a:cs typeface="Times New Roman"/></a:rPr><a:t>{v:.1f}</a:t></a:r></a:p></c:rich></c:tx>'
        '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr><c:dLblPos val="outEnd"/><c:showLegendKey val="0"/><c:showVal val="1"/>'
        '<c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>'.replace(f'{v:.1f}', f'{v:.1f}'.replace('.', ','))
        for i, v in enumerate(vals))
    return ('<c:dLbls>' + one + '<c:showLegendKey val="0"/><c:showVal val="0"/><c:showCatName val="0"/>'
            '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


def chart():
    xml, rows = W.bar_chart(CATS, SER, ytitle='%', ymin=0, ymax=100, major=20, gap=60)
    for s, (nm, vals, col) in zip(re.findall(r'<c:ser>.*?</c:ser>', xml, re.S), SER):
        xml = xml.replace(s, re.sub(r'<c:dLbls>.*?</c:dLbls>', _lbls(vals, col), s, flags=re.S), 1)
    return xml, rows


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(4)

    def par(parts, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12):
        q = dd.add_paragraph(); q.alignment = align
        for t, col in parts:
            r = q.add_run(t); r.bold = bold; r.font.size = Pt(size)
            if col: r.font.color.rgb = RGBColor.from_string(col)
    par([('3.2.1-rasm (namuna)', None)], True, WD_ALIGN_PARAGRAPH.CENTER, 14)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, chart())
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.5, '3.2.1-rasm')))
    par([(IZOH, 'FF0000')], size=11, align=WD_ALIGN_PARAGRAPH.LEFT)
    par([('3.2.1-rasm. Monte-Karlo tahlilida banklar barqarorlik sinfining saqlanish ulushi, foizda', None)], True, WD_ALIGN_PARAGRAPH.CENTER, 14)
    par([('Rasm izohi (uzun ro‘yxat o‘rniga): ', None), (FN_NEW, 'FF0000')], size=11)
    par([('Matndagi qo‘shimchalar:', None)], True)
    par([('1) Rasmdan oldingi xatboshida: …', None), (ANC1, None), (INS1, 'FF0000'), (' va vaznlar to‘liq tasodifiy … tanlangan holatda (3.2.1-rasm).', None)])
    par([('2) Rasmdan keyingi xatboshida (“Ikkinchidan” bandi oxirida): …', None), (ANC2, None), (INS2, 'FF0000')])
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST))

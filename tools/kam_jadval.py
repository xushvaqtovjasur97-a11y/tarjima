# Kamchiliklar jadvaliga «Tuzatish yo‘li va natijasi» ustunini (qizil) qo'shish
import sys, copy
sys.path.insert(0, '/home/user/tarjima/tools')
import kamlib as K

SRC = '/home/user/tarjima/taqriz/Kamchiliklar_jadvali.docx'
DST = '/home/user/tarjima/Kamchiliklar_jadvali_tuzatish_yollari.docx'

T = {
1: 'Tasdiqlandi: xususiy kapital olinganda me’yor 122,2%, regulyativ kapital olinganda 129,4% (−7,1 p.p.). Kirish (taqrizchi tuzatgan) saqlandi; 2.3-§ ga nizom 53–55-bandlari va bankning 16-, 22-jadval shakllariga havola bilan hisob-kitob qo‘shildi.',
2: 'Kirish va Xulosadagi «2,7–2,8» saqlandi; 2.3-§ matniga farq dekompozitsiya tartibiga bog‘liqligi (MB metodikasi bo‘yicha 2,7, Bazel III bo‘yicha 2,8) izohi qo‘shildi.',
3: 'Hal qilindi: SQB va MKB uchun barcha joylarda 2.3-§ metodikasi (CET1 8,87%, regulyativ kapital 15,32%) qo‘llandi; mbii32.py qayta hisoblandi. MBII SQB 2025: 2,83 → 2,81 (sinf o‘zgarmadi). Yangilandi: 3.1.5-jadval va izohi, 3.1.3–3.1.4-rasmlar, ΔMBII (−0,46; aktivlar sifati ulushi 83–86%), 1-ilova (1.1, 1.2), 3.2.3-jadval va rasm, 4.1.3–4.1.4-rasmlar va matn (2024-yilda SQB MHXS indeksi 2,72 – qizil zona), 4.2.6-jadval va 4.2.3-rasm, 4.3.2-jadval, III bob xulosalari va Xulosa.',
4: 'Hal qilindi: MKB MHXS CET1 hamma joyda 18,24% (regulyativ kapital 19,96%); ball va MBII o‘zgarmadi.',
5: 'Belgilar ajratildi: 3.3-§ da MBII-13 (to‘liq) va MBII-8 (panel) ta’riflandi; 4.2.6 va 4.3.2-jadvallar sarlavhalari belgilandi; 4.2-§ ga bir xil tarkibdagi (MBII-13) taqqoslash qo‘shildi: SQB 3,27 → 2,81 (−0,46), MKB 3,13 → 2,47 (−0,66), zona o‘zgarishi xulosasi saqlanadi.',
6: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.', 7: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.', 8: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.',
9: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.', 10: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.', 11: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.',
12: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.', 13: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.',
14: 'Tasdiqlandi: panelda «Ipoteka-bank» ATB 2023-yil II chorakgacha davlat banki (D=1), III chorakdan xususiy. 3.3.3-jadval izohi aniqlashtirildi («deyarli o‘zgarmaydi; faqat Ipoteka-bankda 2023-y. III chorakda o‘zgargan»).',
15: 'Hal qilindi: 2.1.1-jadval Markaziy bankning 2025-yil moliyaviy barqarorlik sharhi (1-jadval, 129-b.) va SQB ning 22-jadval shakli bo‘yicha qayta yozildi: 2025-yilgacha CET1 8%, jami 13% (3% bufer bilan), leverej 6%; 2026-yildan 4,5% / 8% / 3% + buferlar (CCoB 2,5%, CCyB 1,5%, D-SIB 1%). «HLA/aktivlar ≥ 10%» qatori qo‘shildi. 2.1-§ stress-test chegaralari aniqlashtirildi. Yangi ilmiy natija qo‘shildi (2.3-§): 2026-yildan D-SIB uchun 9,5% CET1 talabiga nisbatan SQB milliy hisobot bo‘yicha +0,71 p.p. yuqori, MHXS bo‘yicha −0,63 p.p. past.',
16: 'Hal qilindi: Garant bank (ijobiy, 2027) aniq qiymati 3,248 → «3,248 (sariq)»; chegaraga yaqin boshqa qiymatlar ham uch xonada (Davr bank 3,251; Tenge bank 3,252; Xalq banki 3,254); 7-ilova izohiga qoida qo‘shildi.',
17: 'Hal qilindi: noto‘g‘ri izoh o‘rniga «substandart va muammoli kreditlar ulushi 2024–2025-yillar uchun manbada butun songacha yaxlitlangan» izohi yozildi.',
18: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.',
19: 'To‘liq hal qilindi: 3.3.6-jadval izohiga makro koeffitsiyentlar 8 ta vaqt nuqtasidan aniqlangani va «indikativ» ekani, prognozda 4.2-§ dagi uzunroq qator bahosi qo‘llanishi yozildi (taqrizchi izohlari bilan birga).',
20: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.', 21: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.', 22: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.',
23: 'To‘liq hal qilindi (robust_kam.py): bank bo‘yicha klasterlangan bootstrap – MBII AUC 0,864 [0,65; 0,96]; MBII−Z farqi [0,19; 0,98] – ahamiyatli; MBII−NPL farqi [−0,11; +0,07] – ahamiyatsiz. Muqobil hodisa ta’riflari: 3% chegara – AUC 0,79 [0,62; 0,92], 10% – 0,90; Z-score 0,33–0,48. Qoplanmaydigan oynalar (faqat IV chorak) – 0,73 [0,50; 0,92]. 4.1-§ ga yangi paragraf, Kirish (yangilik 5) ga oraliq qo‘shildi.',
24: 'Hal qilindi: I bob xulosalariga va 1.2-§ dagi ishchi ta’rifga bog‘lovchi jumla qo‘shildi (boshqaruv sifati – kurator mulohazasi, valyuta riski – bozor riskining asosiy shakli). 1.1.7-rasm konseptual model sifatida saqlandi.',
25: 'Hal qilindi (vazn32.py): empirik o‘rtacha faqat 5 blokni baholagan 3 usuldan (entropiya, variatsiya, bosh komponentlar): K 0,08; A 0,31; L 0,08; D 0,38; V 0,16. 3.2.3-jadvalga variatsiya koeffitsiyenti ustuni qo‘shildi; AUC 0,871. Xulosa rostlandi: aktivlar sifati vazni empirik oraliqda (0,27–0,57), tarqoqlik usullarida daromadlilik birinchi, logitda aktivlar sifati yagona ahamiyatli blok. 3.2-§, III bob xulosasi, Xulosa yangilandi.',
26: 'Hal qilindi: L1 uchun 10% – SQB ning 22-jadval shaklidagi «yuqori likvid aktivlarning jami aktivlardagi ulushi» prudensial me’yori (lahzali likvidlik emas – taqrizchi fikri asossiz, manba keltirildi). A2 chegaralari EBA Risk Dashboard (Q4 2024, 33-b.) bilan tashqi tekshirildi: YeI o‘rtacha 2-bosqich ulushi 9,6–9,7%, maksimum 15,9%. Sinf oraliqlari yarim ochiq ko‘rinishda yozildi (3.1.4 va 3.2.4-jadvallar), teng oraliqlar asoslandi.',
27: 'Hal qilindi: 3.2.4-jadvalga Wilson 95% oraliqlari ustuni (0,8–5,0; 7,9–21,4; 59,6–98,2) va Fisher testi (p < 0,001) qo‘shildi; xulosa «tahlil davri ma’lumotlariga mos» deb yumshatildi (3.2-§, III bob xulosasi, Xulosa).',
28: 'To‘liq hal qilindi: 3.3.6-jadval izohiga Nickell siljishi (S. Nickell, 1981) va inersiya koeffitsiyentining haqiqiy qiymati FE va Pooled baholari oralig‘ida (0,62–0,81) ekani yozildi.',
29: 'Hal qilindi (lag8_kam.py): 8 chorak lagli modellar baholandi (N=194): Pooled −0,0007***, FE +0,0005 (p=0,05), Z-score – ahamiyatsiz. 3.3-§ matni va 3.3.7-jadval (H2) to‘ldirildi.',
30: 'Hal qilindi: 3.1.5-jadval izohiga va 3.1-§ cheklovlariga L1 milliy balansdan olingani, shu sababli likvidlik blokidagi ΔMBII faqat L2 hisobiga shakllanishi yozildi.',
31: 'Hal qilindi: mualliflik ta’rifi «amaliy natijalar»dan «ilmiy yangilik»ning birinchi bandiga ko‘chirildi.',
32: 'Hal qilindi: LCR Bazel III bo‘yicha qayta hisoblandi (16-jadval shakli): banklardagi mablag‘lar HQLA dan chiqarilganda 222,9% → 154,1% (−68,8 p.p.); garovga qo‘yilgan aktivlar shaklda yuqori likvid aktivlar jamisidan allaqachon chiqarilgan. 2.3.4-jadvalga qator va 2.3-§ ga matn qo‘shildi.',
33: 'Hal qilindi: usullar ro‘yxatida «ekspert baholash (vaznlarni muallif belgilashi va ularni entropiya, variatsiya, bosh komponentlar va logit usullari bilan tekshirish), bootstrap va Monte-Karlo» deb aniqlashtirildi.',
34: 'Muallif hal qiladi: Markaziy bank (va/yoki boshqa vakolatli tashkilot) dalolatnomasini olish tavsiya etiladi; matnga o‘zgartirish kiritilmadi.',
35: 'Hal qilindi: 2.3-§ ga ochiq/yopiq ma’lumotlar chegarasi va takrorlanuvchanlik izohi qo‘shildi («banklarning ruxsati bilan» – muallif tasdiqlashi kerak).',
36: 'Hal qilindi (kapkanal_kam.py): ssenariylarga kapital kanali qo‘shilib tekshirildi – MBII ko‘pi bilan 0,04 ballga pasayadi, keskin salbiy ssenariyda qizil banklar soni o‘zgarmaydi (4 ta), faqat salbiy 2028 da MKB sariqdan qizilga o‘tadi (2,755 → 2,747). 4.2-§ cheklovlariga yozildi va natija «sezgirlik (stress-ssenariy) tahlili» deb belgilandi. Paragraf nomi o‘zgartirilmadi (muallif xohlasa «ssenariy prognozi» deb o‘zgartirishi mumkin).',
37: 'Taqrizchi tuzatgan. Qo‘shimcha: 1.3-§ ga EBA risk dashbordi havolasi, M. Bowman (S&C, 2026) jumlasi; 4.3-§ ga S&C va EBA havolalari qo‘shildi.',
38: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.',
39: 'Hal qilindi: 2.1-§ dagi raqamlangan kichik sarlavhalar va «Paragraf bo‘yicha xulosa» sarlavhasi olib tashlandi, xulosa boshqa paragraflardagi kabi «…tadqiq etgan holda xulosa qilish mumkinki» shaklida yozildi (mundarijani F9 bilan yangilang).',
40: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.', 41: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.', 42: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.',
43: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.', 44: 'Taqrizchi tuzatgan – tekshirildi, to‘g‘ri.',
45: 'Muallif to‘ldiradi: ilmiy-tadqiqot mavzusi nomi, aprobatsiya va nashrlar soni (joylari qizil bilan belgilandi).',
46: 'O‘zgartirilmadi: Ilmiy kengash Toshkentda (Tashkent International University) joylashgani uchun «Toshkent – 2026» saqlandi; OAK talabi boshqacha bo‘lsa, muallif o‘zgartiradi.',
47: 'Hal qilindi: 169–174 (sayt manzillari) olib tashlandi; Kilde blogi o‘rniga Markaziy bank sharhi (aktivlar, 30-b.) va SNS bazasi (kreditlar 603,8 trln) keltirildi; 9 ta normativ hujjat (PF-158, PQ-2344, PQ-3270, 3254, 3302, 3427, 4/13, 34/3, 42/29) va 6 ta yangi manba qo‘shildi; V.V. Ivanov ro‘yxatdan chiqarildi; ro‘yxat qayta raqamlandi – 183 ta manba, Kirish yangilandi. Afanasyeva–Dubova va Sokolinskaya kitoblarining yili, Shatalova kitobining nashriyoti – muallif kiritadi (qizil belgi qo‘yildi).',
48: 'Hal qilindi: formulalar raqamlandi – (2.1), (2.2), (3.1), (3.2), (3.3).',
49: 'Taqrizchi tuzatgan – mundarijani yangilash (F9) kerak.',
}
ST = {i: 'Tuzatildi' for i in T}
for i in (34, 45):
    ST[i] = 'Muallif hal qiladi'
for i in (6, 7, 8, 9, 10, 11, 12, 13, 18, 20, 21, 22, 38, 40, 41, 42, 43, 44, 49):
    ST[i] = 'Taqrizchi tuzatgan'
ST[46] = 'Asoslandi'
ST[47] = 'Tuzatildi (3 ta nashr ma’lumoti – muallif)'

D = K.Doc(SRC)
from docx.oxml.ns import qn
t = [tb for tb in D.d.tables if '№' in tb.rows[0].cells[0].text or any('Kamchilik' in c.text for c in tb.rows[0].cells)][0]
vals = []
for r in t.rows:
    c0 = r.cells[0].text.strip()
    if c0 == '№':
        vals.append('Tuzatish yo‘li va natijasi (muallif, 2026-yil oktabr)')
    elif c0.isdigit():
        n = int(c0); vals.append(f'{ST[n]}. {T[n]}')
    else:
        vals.append('')
D.add_col(t, 5, vals, tag='jadval')
# ustun kengliklari va yangi ustun formati
from docx.oxml.ns import qn as _qn
grid = t._tbl.find(_qn('w:tblGrid')); gc = grid.findall(_qn('w:gridCol'))
tot = sum(int(g.get(_qn('w:w'))) for g in gc)
prop = [0.04, 0.09, 0.08, 0.21, 0.22, 0.08, 0.28]
ws = [int(tot * p) for p in prop]
for g, w in zip(gc, ws): g.set(_qn('w:w'), str(w))
for tr in t._tbl.findall(_qn('w:tr')):
    tcs = tr.findall(_qn('w:tc'))
    if len(tcs) == 7:
        for tc, w in zip(tcs, ws):
            tcw = tc.find(_qn('w:tcPr') + '/' + _qn('w:tcW'))
            if tcw is not None: tcw.set(_qn('w:w'), str(w)); tcw.set(_qn('w:type'), 'dxa')
        tc = tcs[6]
        tcpr = tc.find(_qn('w:tcPr'))
        if tcpr is not None:
            for e in tcpr.findall(_qn('w:shd')): tcpr.remove(e)
        for p in tc.iter(_qn('w:p')):
            ppr = p.find(_qn('w:pPr'))
            if ppr is not None:
                for e in ppr.findall(_qn('w:jc')): ppr.remove(e)
D.d.add_paragraph()
p = D.d.add_paragraph('Izoh: qizil rang bilan muallif tomonidan kiritilgan tuzatishlar belgilangan. Dissertatsiya fayli: Dissertatsiya_Maxmudov_O.T_tuzatilgan_v2.docx (taqrizchi tuzatishlari – yashil, muallif tuzatishlari – qizil). Qayta hisob-kitoblar: manbalar/hisob_kitob/ (mbii32.py, vazn32.py, fc42.py, robust_kam.py, lag8_kam.py, kapkanal_kam.py).')
for r in p.runs:
    r.font.color.rgb = __import__('docx').shared.RGBColor(0xFF, 0, 0)
# albom format
from docx.enum.section import WD_ORIENT
for s in D.d.sections:
    if s.page_width < s.page_height:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = s.page_height, s.page_width
D.save(DST)
print('saqlandi', DST)

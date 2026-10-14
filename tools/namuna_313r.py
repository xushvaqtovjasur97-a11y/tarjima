"""Yangi 3.1.3-rasm: uch bank MBII 2021–2025, milliy (uzluksiz) va MHXS (uzuq), ostida sonlar jadvali (namuna)."""
import json, re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from lxml import etree
import wordart as W

DST = '/home/user/tarjima/Namuna_3.1.3-rasm_dinamika.docx'
J = json.load(open('/home/user/tarjima/manbalar/hisob_kitob/mbii31_dyn3.json'))
Y = ['2021', '2022', '2023', '2024', '2025']
BANK = [('SQB', '“O‘zsanoatqurilishbank”', '2A78D6'), ('MKB', '“Mikrokreditbank”', 'EB6834'), ('ALQ', '“Aloqabank”', '548235')]
SER = []
for k, nm, col in BANK:
    for var, lab, dash in (('milliy', 'milliy hisobot', None), ('MHXS', 'MHXS', 'dash')):
        SER.append((f'{nm}: {lab}', [round(J[f'{k} {var}'][y]['mbii'], 2) for y in Y], col, dash))
GR = [('I va II sinflar chegarasi', [3.25] * 5, 'C00000', 'sysDash'), ('II va III sinflar chegarasi', [2.5] * 5, 'C00000', 'sysDash')]
CAP = '3.1.3-rasm. Uch bankda moliyaviy barqarorlikning integral indeksi: milliy hisobot va MHXS asosida, 2021–2025-yillar, ball'
LEAD = ('2025-yilda aniqlangan tafovut bir yillik holat emasligini tekshirish uchun integral indeks uchala bank bo‘yicha '
        '2021–2025-yillar uchun ham ikki variantda hisoblandi (3.1.3-rasm).')
ANAL = ('Rasmdan ko‘rinib turibdiki, MHXS ma’lumotlari asosidagi indeks uchala bankda ham tahlil davrining barcha yillarida milliy '
        'hisobot asosidagi indeksdan past bo‘lgan, farq 0,23–0,66 ballni tashkil etgan. “Mikrokreditbank” ATB milliy hisobot bo‘yicha '
        'butun davr mobaynida II sinfda, MHXS bo‘yicha esa III sinfda joylashgan. “Aloqabank” ATB milliy hisobot bo‘yicha 2021–2024-yillarda '
        'I sinfda bo‘lsa, MHXS bo‘yicha barcha yillarda II sinfda qolgan. “O‘zsanoatqurilishbank” ATB ham MHXS bo‘yicha butun davr '
        'mobaynida II sinfda bo‘lgan, milliy hisobot bo‘yicha esa 2021 va 2025-yillarda I sinfga ko‘tarilgan. Demak, ikki variant '
        'o‘rtasidagi tafovut tasodifiy emas, balki muntazam xususiyatga ega bo‘lib, u 2.3-paragrafda aniqlangan yashirin kredit riskining '
        'integral indeksdagi aksidir.')
FN = ('Tadqiqot natijasida muallif tomonidan banklarning Markaziy bankka taqdim etilgan prudensial hisobot shakllari va MHXS bo‘yicha '
      'moliyaviy hisobotlari asosida hisoblandi.')
IZOH = 'Izoh: uzluksiz chiziq – milliy hisobot, uzuq chiziq – MHXS; qizil chiziqlar – I va II sinflar (3,25 ball) hamda II va III sinflar (2,50 ball) chegaralari.'


def _lastlbl(text):
    return (f'<c:dLbls><c:dLbl><c:idx val="4"/><c:tx><c:rich><a:bodyPr/><a:lstStyle/><a:p><a:pPr>{W._rpr(1050, True, "C00000")}</a:pPr>'
            f'<a:r><a:rPr lang="uz-Latn-UZ" sz="1050" b="1"><a:solidFill><a:srgbClr val="C00000"/></a:solidFill>'
            f'<a:latin typeface="Times New Roman"/><a:cs typeface="Times New Roman"/></a:rPr><a:t>{text}</a:t></a:r></a:p></c:rich></c:tx>'
            '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr><c:dLblPos val="r"/><c:showLegendKey val="0"/><c:showVal val="1"/>'
            '<c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>'
            '<c:showLegendKey val="0"/><c:showVal val="0"/><c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


def chart():
    xml, rows = W.line_table_chart(Y, SER + GR, ytitle='ball', ymin=2.0, ymax=3.75, major=0.25, fmt='0.00', w=28575, dtable=False,
                                   layout=(0.11, 0.04, 0.70, 0.84))
    xml = xml.replace('<c:tickLblPos val="none"/>', '<c:tickLblPos val="low"/>', 1)
    xml = xml.replace('<c:numFmt formatCode="0" sourceLinked="0"/>', '<c:numFmt formatCode="0.00" sourceLinked="0"/>', 1)
    sers = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)
    for s, t in zip(sers[6:], ['3,25 (I/II)', '2,50 (II/III)']):
        ns = re.sub(r'<c:marker>.*?</c:marker>', '<c:marker><c:symbol val="none"/></c:marker>', s, flags=re.S)
        ns = ns.replace('<c:cat>', _lastlbl(t) + '<c:cat>', 1).replace('w="28575"', 'w="19050"')
        xml = xml.replace(s, ns, 1)
    return xml, rows


def add_table(dd):
    t = dd.add_table(rows=1 + len(SER), cols=6)
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(6.5)] + [Cm(1.8)] * 5
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = tblPr.makeelement(qn('w:tblLayout'), {qn('w:type'): 'fixed'}); tblPr.append(lay)
    for gc, wd in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(wd.twips if hasattr(wd, 'twips') else wd / 635)))
    hdr = ['Bank va hisobot turi'] + Y
    data = [[s[0]] + [f'{v:.2f}'.replace('.', ',') for v in s[1]] for s in SER]
    for i, r in enumerate([hdr] + data):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.width = widths[j]; p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            if j == 0 and i > 0:
                s = SER[i - 1]
                k = p.add_run('━━ ' if s[3] is None else '╍╍ '); k.font.size = Pt(10); k.bold = True
                k.font.color.rgb = RGBColor.from_string(s[2])
            run = p.add_run(v); run.font.size = Pt(11); run.bold = (i == 0)
    return t


def build(dd, red=None):
    def par(text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, size=14, color=red, indent=True):
        q = dd.add_paragraph(); q.alignment = align
        if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
            q.paragraph_format.first_line_indent = Cm(1.25)
        r = q.add_run(text); r.bold = bold; r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        return q
    par(LEAD)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, chart())
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.0, '3.1.3-rasm')))
    add_table(dd)
    par(IZOH, WD_ALIGN_PARAGRAPH.LEFT, size=11, color=None, indent=False)
    par(CAP, WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=None)
    par(ANAL)


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.5
    h = dd.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run('Yangi 3.1.3-rasm (namuna)'); r.bold = True
    n = dd.add_paragraph(); r = n.add_run('Joyi: 3.1.6-jadval tahlilidan keyin, vaznlarga sezgirlik tahlilidan oldin. Qizil – yangi matn. '
                                          'Hozirgi 3.1.3-rasm 3.1.4-rasmga aylanadi (nomi va matndagi havolalar).'); r.font.size = Pt(11)
    build(dd, red='FF0000')
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

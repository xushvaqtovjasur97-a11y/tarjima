"""Milliy tasnif va MHXS 9 bosqichlari mezonlari – qiyosiy jadval (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_tasniflash_mezonlari_jadval.docx'
RED = 'FF0000'
ROWS = [
    ('Tuzilishi', 'Besh toifa: standart, substandart, qoniqarsiz, shubhali, umidsiz',
     'Uch bosqich; bankda besh ichki guruh: 1–2-guruh – 1-bosqich, 3–4-guruh – 2-bosqich, 5-guruh – 3-bosqich'),
    ('Asosiy yondashuv', 'Kreditning joriy holati; hal qiluvchi belgi – to‘lov intizomi va muddati o‘tgan kunlar soni',
     'Kredit riskining dastlabki tan olingan sanaga nisbatan o‘zgarishi; kelajakka yo‘naltirilgan ma’lumotlar ham hisobga olinadi'),
    ('Muddati o‘tgan qarzdorlik', 'Toifa asosan muddati o‘tgan kunlar soni bo‘yicha belgilanadi',
     '16–30 kun – 1-bosqich (2-guruh); 31–90 kun – 2-bosqich; 90 kundan ortiq – 3-bosqich'),
    ('Milliy toifa bilan bog‘liqlik', '–',
     '“Substandart” – kamida 2-bosqich; “qoniqarsiz”, “shubhali”, “umidsiz” – 3-bosqich'),
    ('Qarz oluvchining moliyaviy holati', 'Toifa ta’riflarida umumiy baho: barqaror holat, noqulay tendensiyalar, aniq kamchiliklar',
     'Tushum 60 foizdan yoki marja 40 foizdan ko‘p kamaysa – 3-bosqich; auditorlik xulosasida faoliyat uzluksizligiga shubha bo‘lsa – 2-bosqich'),
    ('Garov', '“Substandart” – ishonchli ikkilamchi manba (garov, kafillik) bilan ta’minlangan; “qoniqarsiz” – qaytarish uchun garovni sotish zarur',
     'Yirik kredit bo‘yicha garov qiymati 50 foizdan ko‘p tushsa – 2-bosqich'),
    ('Restrukturizatsiya', '6 oygacha imtiyozli davr tasnifga ta’sir qilmaydi; yashovchanlik bahosi o‘tkazilmaydi',
     'Bir martadan ortiq, oxirgisi so‘nggi 6 oyda bo‘lsa – 2-bosqich; yangi shartlardagi pul oqimlarining joriy qiymati 10 foizdan ko‘p o‘zgarsa – 3-bosqich'),
    ('Tashqi kredit reytingi', 'Toifa ta’riflarida hisobga olinmaydi',
     '2 va undan ortiq pog‘ona pasaysa – 2-bosqich; “CC” va undan past bo‘lsa – 3-bosqich'),
    ('Boshqa banklardagi qarzdorlik va kross-defolt', 'Toifa ta’riflarida alohida ko‘rsatilmagan',
     'Boshqa banklarda 31–90 kunlik muddati o‘tgan qarz – 2-bosqich; 90 kundan ortiq yoki boshqa majburiyat bo‘yicha defolt – 3-bosqich'),
    ('Huquqiy holat', 'Toifa ta’riflarida alohida ko‘rsatilmagan',
     'Moliyaviy holatni yomonlashtiruvchi sud jarayoni – 2-bosqich; bankrotlik, litsenziyaning bekor qilinishi, 12 oy davomida aloqa yo‘qligi – 3-bosqich'),
    ('Yaxshilanish (qayta tasniflash)', 'Sinov muddati yo‘q: to‘lov tiklansa, kredit yuqori toifaga qaytariladi',
     '2-bosqichdan 1-bosqichga – bir chorak davomida yomonlashuv belgilari bo‘lmasa; restrukturizatsiya qilingan kreditlar 3-bosqichdan – 6 oylik barqaror to‘lovdan keyin'),
    ('Zaxira', 'Toifa bo‘yicha qat’iy stavkalar: 1, 10, 25, 50, 100 foiz',
     '1-bosqich – 12 oylik kutilgan kredit zararlari; 2 va 3-bosqich – butun muddat uchun kutilgan kredit zararlari'),
]
TITLE = 'Kreditlarni milliy tasnif va MHXS 9 bosqichlari bo‘yicha tasniflash mezonlarining qiyosiy tavsifi (“O‘zsanoatqurilishbank” ATB misolida)'
SRC = ('Tadqiqot natijasida muallif tomonidan quyidagi manbalar asosida ishlab chiqildi: Tijorat banklarida aktivlar sifatini tasniflash va aktivlar bo‘yicha ehtimoliy '
       'yo‘qotishlarni qoplash uchun zaxiralarni shakllantirish hamda ulardan foydalanish tartibi to‘g‘risidagi nizom (O‘zbekiston Respublikasi Markaziy banki Boshqaruvining '
       'qarori, ro‘yxat raqami 2696, 14.07.2015) // https://lex.uz; Joint Stock Commercial Bank “Uzbek Industrial and Construction Bank” and its subsidiaries. Consolidated '
       'Financial Statements for the year ended 31 December 2025 and Independent Auditor’s Report. – Tashkent, 2026; International Monetary Fund. Republic of Uzbekistan: '
       'Financial Sector Assessment Program – Financial System Stability Assessment. IMF Country Report No. 25/145. – Washington, DC: IMF, 2025. – P. 30.')
TXT1 = ('Ikki tasnif o‘rtasidagi tafovutning sabablarini aniqlash uchun bankning MHXS hisobotida oshkor qilingan bosqichlarni belgilash mezonlari Markaziy bank tasnifi '
        'bilan solishtirildi (2.3.x-jadval).')
TXT2 = ('Jadvaldan ko‘rinadiki, bank MHXS 9 bo‘yicha bosqichlarni belgilashda milliy toifalarni quyi chegara sifatida qo‘llaydi: “substandart” kredit kamida 2-bosqichga, '
        'muammoli toifadagi kredit esa 3-bosqichga kiritiladi. Biroq MHXS 9 bunga qo‘shimcha ravishda restrukturizatsiya, boshqa banklardagi qarzdorlik, tashqi reyting, '
        'garov qiymati, moliyaviy ko‘rsatkichlar va huquqiy holat kabi belgilarni ham hisobga oladi, yaxshilanish uchun esa sinov muddatini talab qiladi. Shu sababli '
        '3-bosqich har doim milliy tasnif bo‘yicha muammoli kreditlardan katta bo‘ladi. 2025-yilda bank hukumat ko‘rsatmalari asosida 14 129,2 mlrd so‘mlik kreditlarni '
        'restrukturizatsiya qilgani va bundan 184,7 mlrd so‘m zarar tan olgani 2-bosqich kreditlarining keskin ko‘payishini izohlaydi. Milliy tasnifda esa 6 oygacha '
        'imtiyozli davrli restrukturizatsiya tasnifga ta’sir qilmagani sababli bu kreditlar “standart” yoki “substandart” toifasida qolgan.')

# ---- 2025-yil: milliy toifalar × MHXS 9 bosqichlari (mln so‘m, manba: SQB MHXS 2025)
M25 = [('Standart', 40378174, 12956173, 746404), ('Substandart', 0, 12364265, 2290907),
       ('Qoniqarsiz', 0, 0, 243870), ('Shubhali', 0, 0, 642117), ('Umidsiz', 0, 0, 1386888)]
ECL25 = (266281 + 149064 + 120546, 1433186 + 98163 + 180057, 1742976 + 3141 + 257472)
TITLE2 = '“O‘zsanoatqurilishbank” ATB kreditlarining milliy tasnif toifalari va MHXS 9 bosqichlari bo‘yicha taqsimoti, 2025-yil 31-dekabr holatiga, mlrd so‘m'
TXT3 = ('Tafovutning miqdoriy ko‘lamini 2025-yil yakunidagi ma’lumotlar ko‘rsatadi (2.3.y-jadval). Milliy tasnif bo‘yicha muammoli kreditlar 2 272,9 mlrd so‘mni, ya’ni '
        'jami kreditlarning 3,2 foizini tashkil etgan bo‘lsa, MHXS 9 bo‘yicha 3-bosqich kreditlari 5 310,2 mlrd so‘m yoki 7,5 foizni tashkil etgan, ya’ni 2,3 barobar ko‘p. '
        'Bundan tashqari, milliy tasnifda “standart” deb ko‘rsatilgan 54 080,8 mlrd so‘mlik kreditlarning 24,0 foizi MHXS 9 bo‘yicha 2-bosqichda, 1,4 foizi esa 3-bosqichda '
        'bo‘lgan. Kutilgan kredit zararlari bo‘yicha zaxira 4 250,9 mlrd so‘mni tashkil etib, uning 47,1 foizi 3-bosqich, 40,3 foizi 2-bosqich kreditlariga to‘g‘ri kelgan.')


def m(x):
    return f'{x / 1000:,.1f}'.replace(',', ' ').replace('.', ',')


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, color=RED, ls=None, indent=False):
        q = dd.add_paragraph(); q.alignment = align
        if indent:
            q.paragraph_format.first_line_indent = Cm(1.25)
        if ls:
            q.paragraph_format.line_spacing = ls
        r = q.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return q

    par('Tasniflash mezonlari – qiyosiy jadval (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=None)
    par('Joyi: 2.3.4-rasm tahlilidagi “Bank bu farqni shunday izohlaydi …” jumlasidan keyin. Jadval raqami kiritishda aniqlanadi (2.3.3 yoki 2.3.4), keyingi jadvallar qayta raqamlanadi.',
        italic=True, size=12, color='595959')
    par(TXT1, indent=True)
    par('2.3.x-jadval', italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    par(TITLE + ' [1]', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.0)
    t = dd.add_table(rows=len(ROWS) + 1, cols=3); t.style = 'Table Grid'; t.autofit = False
    Wd = [Cm(3.9), Cm(5.6), Cm(7.0)]
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), Wd):
        gc.set(qn('w:w'), str(int(w.twips)))
    for ri, row in enumerate([('Mezon', 'Milliy tasnif (Markaziy bank)', 'MHXS 9 bosqichlari (bank siyosati)')] + ROWS):
        for ci, v in enumerate(row):
            c = t.cell(ri, ci); c.width = Wd[ci]
            q = c.paragraphs[0]; q.paragraph_format.line_spacing = 1.0
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = q.add_run(v); r.font.size = Pt(11); r.bold = ri == 0 or ci == 0; r.font.color.rgb = RGBColor.from_string(RED)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'DEEBF7'); c._tc.get_or_add_tcPr().append(sh)
    par('')
    par(TXT2, indent=True)
    par('[1] ' + SRC, size=11, ls=1.0)
    par('')
    par(TXT3, indent=True)
    par('2.3.y-jadval', italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    par(TITLE2 + ' [2]', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, ls=1.0)
    hdr = ['Milliy tasnif toifasi', '1-bosqich', '2-bosqich', '3-bosqich', 'Jami', 'Ulushi, %']
    tot = [sum(r[i] for r in M25) for i in (1, 2, 3)]
    T = sum(tot)
    body = [[r[0], m(r[1]) if r[1] else '–', m(r[2]) if r[2] else '–', m(r[3]), m(r[1] + r[2] + r[3]),
             f'{100 * (r[1] + r[2] + r[3]) / T:.1f}'.replace('.', ',')] for r in M25]
    body.append(['Jami', m(tot[0]), m(tot[1]), m(tot[2]), m(T), '100,0'])
    body.append(['Ulushi, %'] + [f'{100 * x / T:.1f}'.replace('.', ',') for x in tot] + ['100,0', ''])
    body.append(['Kutilgan kredit zararlari bo‘yicha zaxira'] + [m(x) for x in ECL25] + [m(sum(ECL25)), ''])
    body.append(['Zaxira bilan qoplanish, %'] + [f'{100 * e / x:.1f}'.replace('.', ',') for e, x in zip(ECL25, tot)] + [f'{100 * sum(ECL25) / T:.1f}'.replace('.', ','), ''])
    t2 = dd.add_table(rows=len(body) + 1, cols=6); t2.style = 'Table Grid'; t2.autofit = False
    W2 = [Cm(4.6), Cm(2.4), Cm(2.4), Cm(2.4), Cm(2.5), Cm(2.2)]
    for gc, w in zip(t2._tbl.tblGrid.findall(qn('w:gridCol')), W2):
        gc.set(qn('w:w'), str(int(w.twips)))
    for ri, row in enumerate([hdr] + body):
        for ci, v in enumerate(row):
            c = t2.cell(ri, ci); c.width = W2[ci]
            q = c.paragraphs[0]; q.paragraph_format.line_spacing = 1.0
            q.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 and ri else WD_ALIGN_PARAGRAPH.CENTER
            r = q.add_run(v); r.font.size = Pt(11); r.bold = ri == 0 or row[0] == 'Jami'; r.font.color.rgb = RGBColor.from_string(RED)
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'DEEBF7'); c._tc.get_or_add_tcPr().append(sh)
    par('Izoh: muammoli – “qoniqarsiz”, “shubhali” va “umidsiz” toifalari (2 272,9 mlrd so‘m, 3,2 foiz).', italic=True, size=11, ls=1.0)
    par('[2] Tadqiqot natijasida muallif tomonidan quyidagi manba ma’lumotlari asosida hisoblandi: Joint Stock Commercial Bank “Uzbek Industrial and Construction Bank” and its '
        'subsidiaries. Consolidated Financial Statements for the year ended 31 December 2025 and Independent Auditor’s Report. – Tashkent, 2026. – P. 37–38.', size=11, ls=1.0)
    dd.save(DST)
    print(pardoz(DST), DST)

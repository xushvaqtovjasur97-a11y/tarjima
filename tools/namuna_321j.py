"""3.2.1-jadval: entropiya va logit qatorlari aniqlashtirildi; matnda 0,31 → 0,30; 3.3 va 4.1 ga havola (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_3.2.1-jadval.docx'
HDR = ['Usul', 'Mazmuni', 'Vazn nimani aks ettiradi', 'Qo‘llangan ma’lumotlar']
# (qator, ustun, eski, yangi)
CELLS = [
    (1, 1, 'K. Shannon entropiyasi asosida har bir ko‘rsatkichning banklar o‘rtasidagi tarqoqligi o‘lchanadi',
     'K. Shannon entropiyasi asosida har bir ko‘rsatkich ballarining 25 ta bank-yil kuzatuvi o‘rtasidagi tarqoqligi o‘lchanadi'),
    (4, 2, 'Blokning muammolarni oldindan aniqlashdagi hissasi',
     'Blokning muammolarni oldindan aniqlashdagi hissasi (koeffitsiyentlar modulining ulushi; kutilganiga teskari ishorali koeffitsiyentga ega blok vazni – 0)'),
    (4, 3, '25 bank, 2021–2024-yillar choraklik ma’lumotlari, 4.1-paragrafdagi muammo hodisalari',
     '25 bank, 2021–2024-yillar choraklik ma’lumotlari (308 kuzatuv), 8 ko‘rsatkichli panel indeksining 3 bloki (aktivlar sifati, likvidlik, daromadlilik), 4.1-paragrafdagi muammo hodisalari'),
]
ROWS = [HDR,
        ['Entropiya usuli', None, 'Banklarni bir-biridan qanchalik farqlashi', '5 bank, 2021–2025-yillar, 13 ko‘rsatkich (MHXS)'],
        ['Variatsiya koeffitsiyenti usuli', 'Ko‘rsatkich ballarining standart chetlanishi o‘rtacha qiymatga nisbati',
         'Ko‘rsatkich qiymatlarining nisbiy o‘zgaruvchanligi', '5 bank, 2021–2025-yillar, 13 ko‘rsatkich (MHXS)'],
        ['Bosh komponentlar usuli', 'Bloklar ballarining korrelyatsiya tuzilmasi asosida umumiy omillar ajratiladi',
         'Blokning umumiy barqarorlik omiliga hissasi', '5 bank, 2021–2025-yillar, 5 blok (MHXS)'],
        ['Logit (bashorat) usuli', 'Muammo hodisasi ehtimoli bloklar ballariga regressiya qilinadi', None, None]]
OLD_T = 'uchta tarqoqlik usulining o‘rtachasiga (0,31) deyarli teng'
NEW_T = 'uchta tarqoqlik usulining o‘rtachasiga (0,30) deyarli teng'
ANC_33 = 'Logit usuli uchun 3.3-paragrafda shakllantirilgan 25 ta bank paneli qo‘llandi'
INS_33 = ' (panelning tarkibi 3.3-paragrafda, muammo hodisasining ta’rifi esa 4.1-paragrafda batafsil yoritiladi)'

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)

    def par(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, bold=False):
        q = dd.add_paragraph(); q.alignment = align
        for t, col, strike in parts:
            r = q.add_run(t); r.font.size = Pt(size); r.bold = bold; r.font.strike = strike
            if col: r.font.color.rgb = RGBColor.from_string(col)
        return q
    par([('3.2.1-jadval (namuna)', None, False)], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    par([('3.2.1-jadval', None, False)], WD_ALIGN_PARAGRAPH.RIGHT, size=12)
    par([('Integral indeks vaznlarini aniqlashning empirik usullari', None, False)], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    t = dd.add_table(rows=5, cols=4); t.style = 'Table Grid'
    new = {(r, c): n for r, c, o, n in CELLS}
    for i, row in enumerate(ROWS):
        for j, v in enumerate(row):
            p = t.cell(i, j).paragraphs[0]
            txt = new.get((i, j), v)
            run = p.add_run(txt); run.font.size = Pt(11); run.bold = (i == 0)
            if (i, j) in new:
                run.font.color.rgb = RGBColor.from_string('FF0000')
    dd.add_paragraph()
    par([('Matndagi o‘zgarishlar:', None, False)], size=12, bold=True)
    par([('1) Jadvaldan keyingi xatboshida: …', None, False), ('Logit usuli uchun 3.3-paragrafda shakllantirilgan 25 ta bank paneli qo‘llandi', None, False),
         (INS_33, 'FF0000', False), (', unda kapital bloki barcha banklarda bir xil ball olgani …', None, False)], size=12)
    par([('2) 3.2.2-jadval tahlilida: …', None, False), ('(0,31)', '7F7F7F', True), (' → ', None, False), ('(0,30)', 'FF0000', False),
         (' – jadvaldagi o‘rtacha bilan bir xil: (0,320 + 0,322 + 0,272) / 3 = 0,305.', None, False)], size=12)
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

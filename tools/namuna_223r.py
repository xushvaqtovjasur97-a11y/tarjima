"""2.2.3-rasm: chiziqli diagramma + ostida sonlar jadvali (raqamlar ko‘rinadigan qilib) – namuna."""
import re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from lxml import etree
import wordart as W

DST = '/home/user/tarjima/Namuna_2.2.3-rasm.docx'
Y = ['2021', '2022', '2023', '2024', '2025']
SER = [('Substandart va muammoli kreditlar', [12.5, 15.9, 19.5, 21.0, 15.0], '7030A0', None),
       ('Muammoli kreditlar, davlat ulushi mavjud banklar', [5.4, 3.9, 4.0, 3.9, 3.2], 'EB6834', None),
       ('Muammoli kreditlar, jami', [5.2, 3.6, 3.5, 4.0, 3.0], '2A78D6', None),
       ('Muammoli kreditlar, boshqa banklar', [4.1, 2.1, 2.4, 4.1, 2.6], '70AD47', None)]
GR = [('Erta ogohlantirish chegarasi (5 %)', [5.0] * 5, 'C00000', 'sysDash')]
CAP = '2.2.3-rasm. O‘zbekiston bank tizimida aktivlar sifati ko‘rsatkichlari, %'
IZOH = 'Izoh: diagramma ostidagi jadvalda ko‘rsatkichlarning yil oxiridagi qiymatlari, foizda; qizil uzuq chiziq – 5 foizlik erta ogohlantirish chegarasi.'


def _lbls(color, pos='t'):
    return (f'<c:dLbls><c:numFmt formatCode="0.0" sourceLinked="0"/><c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
            f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr>{W._rpr(1100, True, color)}</a:pPr><a:endParaRPr lang="uz-Latn-UZ"/></a:p></c:txPr>'
            f'<c:dLblPos val="{pos}"/><c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/>'
            '<c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


def chart():
    xml, rows = W.line_table_chart(Y, SER + GR, ytitle='Kreditlardagi ulushi, %', ymin=0, ymax=24, major=4, fmt='0.0', w=28575,
                                   dtable=False, layout=(0.10, 0.05, 0.86, 0.80))
    xml = xml.replace('<c:tickLblPos val="none"/>', '<c:tickLblPos val="low"/>', 1)
    sers = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)
    # yuqori chiziq (12–21 %) – raqamlar nuqtalar ustida
    xml = xml.replace(sers[0], sers[0].replace('<c:cat>', _lbls('7030A0') + '<c:cat>', 1), 1)
    # chegara chizig‘i – marker yo‘q, ingichka
    ns = re.sub(r'<c:marker>.*?</c:marker>', '<c:marker><c:symbol val="none"/></c:marker>', sers[-1], flags=re.S).replace('w="28575"', 'w="19050"')
    xml = xml.replace(sers[-1], ns, 1)
    return xml, rows


def add_table(dd):
    allser = SER + GR
    t = dd.add_table(rows=1 + len(allser), cols=6)
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(7.4)] + [Cm(1.7)] * 5
    t.autofit = False
    tblPr = t._tbl.tblPr
    tblPr.append(tblPr.makeelement(qn('w:tblLayout'), {qn('w:type'): 'fixed'}))
    for gc, wd in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(wd.twips)))
    hdr = ['Ko‘rsatkich'] + Y
    for i, r in enumerate([hdr] + [[s[0]] + [f'{v:.1f}'.replace('.', ',') for v in s[1]] for s in allser]):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.width = widths[j]; p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and i > 0) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            if j == 0 and i > 0:
                s = allser[i - 1]
                k = p.add_run('━━ ' if s[3] is None else '╍╍ '); k.font.size = Pt(11); k.bold = True
                k.font.color.rgb = RGBColor.from_string(s[2])
            run = p.add_run(v); run.font.size = Pt(12); run.bold = (i == 0)
    return t


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.5
    h = dd.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run('2.2.3-rasm – yangi ko‘rinish (namuna)'); r.bold = True
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, chart())
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.5, '2.2.3-rasm')))
    add_table(dd)
    iz = dd.add_paragraph(); rr = iz.add_run(IZOH); rr.font.size = Pt(11); iz.paragraph_format.line_spacing = 1.0
    c = dd.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; rr = c.add_run(CAP); rr.bold = True
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

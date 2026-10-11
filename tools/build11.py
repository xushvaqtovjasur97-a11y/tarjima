import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/1.1-paragraf.md'; DST='/home/user/tarjima/paragraflar/1.1-paragraf.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p11_raw.docx'],check=True)
d=docx.Document(S+'/p11_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^1\.1\.\d-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^1\.1\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^1\.1\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'

def wt(c): return '000000' if c in (GOLD,GREY,LBLUE) else 'FFFFFF'

# ---- 1: concepts
sh=[]
top=[('Stabillik',BLUE,'holatning saqlanishi;\nbarqarorlik sharti'),
     ('Ishonchlilik',GREEN,'majburiyatlarning\no‘z vaqtida\nbajarilishi'),
     ('Muvozanat',GOLD,'kapital, aktivlar,\nmajburiyatlar va\nfoydaning mutanosib\no‘sishi'),
     ('Rivojlanish',ORANGE,'resurslarni\nkengaytirilgan\ntakror ishlab\nchiqarish')]
w=3.85; g=0.25
for i,(h,c,t) in enumerate(top):
  x=0.1+i*(w+g)
  sh.append(W.box(x,0.1,w,0.8,h,fill=c,line=c,sz=22,bold=True,color=wt(c)))
  sh.append(W.box(x,1.0,w,1.9,t,fill='F2F2F2',line=c,sz=19))
  sh.append(W.line(x+w/2,2.9,8.2,3.6,color='404040',lw=15875))
sh.append(W.box(3.2,3.6,10.0,1.2,'Moliyaviy barqarorlik\n(natija, dinamik va kompleks xususiyat)',fill=DBLUE,line=DBLUE,sz=22,bold=True,color='FFFFFF'))
sh.append(W.line(8.2,4.8,8.2,5.35,color='7F7F7F',lw=12700,dash='dash'))
sh.append(W.box(3.2,5.35,10.0,0.85,'Moliyaviy holat – muayyan sanadagi statik baho',fill='FFFFFF',line='7F7F7F',sz=20,italic=True))
W.replace_placeholder(d,'{{FIG:1}}',W.group_run_xml(sh,16.4,6.35,'1.1.1-rasm'))

# ---- 2: four approaches
sh=[]
sh.append(W.box(4.7,0.1,7.0,0.9,'Moliyaviy barqarorlik tushunchasi',fill=DBLUE,line=DBLUE,sz=22,bold=True,color='FFFFFF'))
ap=[('Resurs\nyondashuvi',BLUE,'Kapital, foyda,\nzaxiralarning\nbarqaror o‘sishi\n\n(O.I. Lavrushin,\nI.D. Mamonova)'),
    ('Institutsional-\nfunksional',GREEN,'Vositachilik\nfunksiyalarini\nuzluksiz bajarish\n\n(O.I. Lavrushin,\nG.J. Schinasi)'),
    ('Ishonch\nyondashuvi',GOLD,'Tashqi yordamsiz\nyuqori ishonch\n\n(J.M. Hendrickson,\nA. Krokett)'),
    ('Risk va\nchidamlilik',ORANGE,'Kapital va likvidlik\nbuferlari, shoklarga\nchidamlilik\n\n(Bazel qo‘mitasi,\nXVJ)')]
w=3.85; g=0.25
for i,(h,c,t) in enumerate(ap):
  x=0.1+i*(w+g)
  sh.append(W.line(8.2,1.0,x+w/2,1.6,color='404040',lw=15875))
  sh.append(W.box(x,1.6,w,1.1,h,fill=c,line=c,sz=21,bold=True,color=wt(c)))
  sh.append(W.box(x,2.85,w,2.6,t,fill='F2F2F2',line=c,sz=18))
  sh.append(W.line(x+w/2,5.45,8.2,6.0,color='7F7F7F',lw=12700))
sh.append(W.box(2.7,6.0,11.0,0.9,'Integral yondashuv: daraja + dinamika + muvozanat',fill='DEEBF7',line=BLUE,sz=21,bold=True))
W.replace_placeholder(d,'{{FIG:2}}',W.group_run_xml(sh,16.4,7.05,'1.1.2-rasm'))

# ---- 3: levels
sh=[]
lv=[(0.1,16.2,'Moliya tizimi: banklar, nobank moliya tashkilotlari, moliya bozorlari, infratuzilma',DBLUE),
    (2.0,12.4,'Bank tizimi (bank sektori): Markaziy bank va tijorat banklari',BLUE),
    (3.9,8.6,'Alohida tijorat banki',LBLUE)]
y=0.1
for x,w_,t,c in lv:
  sh.append(W.box(x,y,w_,1.0,t,fill=c,line=c,sz=20,bold=True,color=wt(c)))
  y+=1.35
sh.append(W.line(14.6,1.1,14.6,3.9,color=ORANGE,lw=28575))
sh.append(W.box(13.0,4.0,3.3,1.2,'Makroshoklar,\numumiy risklar',fill='FBE5D6',line=ORANGE,sz=18))
sh.append(W.line(1.8,3.8,1.8,1.1,color=GREEN,lw=28575))
sh.append(W.box(0.1,4.0,3.4,1.2,'Zararlanish,\ntizimli risk',fill='E2EFDA',line=GREEN,sz=18))
sh.append(W.box(3.9,4.2,8.6,1.0,'Tizim ≠ banklarning oddiy yig‘indisi',fill='FFFFFF',line='7F7F7F',sz=19,italic=True))
W.replace_placeholder(d,'{{FIG:3}}',W.group_run_xml(sh,16.4,5.35,'1.1.3-rasm'))

# ---- 4: factors
sh=[]
ext='1. Makroiqtisodiy holat\n2. Davlatning maqsadli mo‘ljallari\n3. Talab va to‘lov qobiliyati\n4. Pul bozori holati\n5. Pul muomalasi va inflyatsiya\n6. Raqobat\n7. Bank tizimining holati\n8. Qonunchilik\n9. Infratuzilma\n10. Ishonch (fundamental omil)'
inn='1. Strategiya va prognozlash\n2. Miqdor va sifat ko‘rsatkichlari o‘sishi\n3. Resurslarni jalb qilish va likvidlik\n4. Risklarga qarshi tura olish\n5. Xarajatlarni tejash\n6. Ichki infratuzilma va kadrlar\n7. Marketing va boshqaruv\n8. Texnologiyalar\n9. Ish tashkiloti'
sh.append(W.box(4.7,0.1,7.0,0.9,'Moliyaviy barqarorlik omillari',fill=DBLUE,line=DBLUE,sz=22,bold=True,color='FFFFFF'))
for i,(h,c,t) in enumerate([('Tashqi omillar (10)',ORANGE,ext),('Ichki omillar (9)',GREEN,inn)]):
  x=0.1+i*8.25
  sh.append(W.line(8.2,1.0,x+3.95,1.55,color='404040',lw=15875))
  sh.append(W.box(x,1.55,7.9,0.85,h,fill=c,line=c,sz=21,bold=True,color='FFFFFF'))
  sh.append(W.box(x,2.5,7.9,4.6,t,fill='F2F2F2',line=c,sz=19,align='l'))
W.replace_placeholder(d,'{{FIG:4}}',W.group_run_xml(sh,16.4,7.25,'1.1.4-rasm'))

# ---- 5: US causes
rid=W.add_chart(d,W.bar_chart(['Aktivlar sifatining\nyomonligi','Rejalashtirish va\nboshqaruv zaifligi','Ichki auditning\nyo‘qligi','Firibgarlik'],[
  ('Muammoli banklar ulushi',[98,90,25,11],BLUE)],
  ytitle='Muammoli banklar sonidagi ulushi, %',ymax=100,major=20,label_fmt='0',legend=False,gap=60))
W.replace_placeholder(d,'{{FIG:5}}',W.chart_run_xml(rid,15.0,7.5,'1.1.5-rasm'))

# ---- 6: Z-score accuracy
rid=W.add_chart(d,W.bar_chart(['Faqat Z-score','To‘liq model','To‘liq model\n(inqiroz davri)'],[
  ('Muammoli banklar to‘g‘ri aniqlangan',[75,81,91],ORANGE),('Sog‘lom banklar to‘g‘ri aniqlangan',[81,87,83],BLUE)],
  ytitle='To‘g‘ri tasniflash ulushi, %',ymax=100,major=20,label_fmt='0',gap=70))
W.replace_placeholder(d,'{{FIG:6}}',W.chart_run_xml(rid,15.0,7.5,'1.1.6-rasm'))

# ---- 7: author's conceptual model
sh=[]
sh.append(W.box(0.1,0.1,16.2,0.75,'Makroiqtisodiy muhit va tizimli risklar',fill='FBE5D6',line=ORANGE,sz=20,bold=True))
bl=[('Kapital',BLUE),('Aktivlar\nsifati',ORANGE),('Likvidlik',GREEN),('Daromad-\nlilik',GOLD),('Boshqaruv\nsifati',GREY)]
w=3.05; g=0.24
for i,(h,c) in enumerate(bl):
  x=0.1+i*(w+g)
  sh.append(W.box(x,1.1,w,1.05,h,fill=c,line=c,sz=20,bold=True,color=wt(c)))
  sh.append(W.line(x+w/2,2.15,8.2,2.6,color='7F7F7F',lw=12700,arrow=False))
dm=[('Daraja','xalqaro va milliy\nchegaralarga nisbatan'),('Dinamika','3–5 yillik trend\nva tebranish'),('Muvozanat','bloklar o‘rtasidagi\nmutanosiblik')]
for i,(h,t) in enumerate(dm):
  x=0.1+i*5.45
  sh.append(W.box(x,2.6,5.2,1.25,h+'\n'+t,fill='DEEBF7',line=LBLUE,sz=18))
sh.append(W.line(8.2,3.85,8.2,4.3,color='404040',lw=19050))
sh.append(W.box(3.7,4.3,9.0,0.95,'Integral moliyaviy barqarorlik ko‘rsatkichi',fill=DBLUE,line=DBLUE,sz=22,bold=True,color='FFFFFF'))
sh.append(W.line(6.0,5.25,4.2,5.75,color='404040',lw=15875))
sh.append(W.line(10.4,5.25,12.2,5.75,color='404040',lw=15875))
sh.append(W.box(0.1,5.75,7.9,0.95,'Erta ogohlantirish signallari',fill='E2EFDA',line=GREEN,sz=20,bold=True))
sh.append(W.box(8.4,5.75,7.9,0.95,'Nazorat choralari',fill='E2EFDA',line=GREEN,sz=20,bold=True))
sh.append(W.box(0.1,6.9,16.2,0.75,'Kelajakka yo‘naltirilgan baho: stress-test va kutilgan kredit zararlari',fill='FFFFFF',line='7F7F7F',sz=19,italic=True))
W.replace_placeholder(d,'{{FIG:7}}',W.group_run_xml(sh,16.4,7.8,'1.1.7-rasm'))
d.save(DST); print('saved')

"""Butun matn bo‘yicha imlo, grammatika va tinish belgilari tuzatishlari (ko‘k rangda).

Tahrirlar: tools/imlo_tahrirlar/e[1-7].json  (turlari: para, cell, tx)
Paragraf va jadval indekslari tahrirlar tayyorlangan paytdagi hujjat holatiga mos (sahifalashdan oldin qo‘llanadi).
"""
import json, glob, sys, collections
sys.path.insert(0, '/home/user/tarjima/tools')
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import kamlib as K
from pardoz import pardoz

K.RED = '0070C0'
F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
E = []
for f in sorted(glob.glob('/home/user/tarjima/tools/imlo_tahrirlar/e*.json')):
    E += json.load(open(f))
D = K.Doc(F)
PS = D.d.paragraphs
TS = D.d.tables


def text_of(pe):
    return ''.join(t.text or '' for t in pe.iter(qn('w:t'))
                   if not any(a.tag == qn('w:txbxContent') for a in t.iterancestors()))


def minimal(p, old, new):
    a = 0
    while a < min(len(old), len(new)) and old[a] == new[a]:
        a += 1
    b = 0
    while b < min(len(old), len(new)) - a and old[-1 - b] == new[-1 - b]:
        b += 1
    while a > 0 and old[a - 1] not in ' (“':
        a -= 1
    while b > 0 and old[len(old) - b] not in ' ,.;)”':
        b -= 1
    mo, mn = old[a:len(old) - b], new[a:len(new) - b]
    full = text_of(p._p)
    pos = full.find(old)
    if not mo:
        D.replace(p, old, new, occ=full[:pos].count(old))
        return
    D.replace(p, mo, mn, occ=full[:pos + a].count(mo))


done = collections.Counter(); skip = []
for e in E:
    t = e['type']
    try:
        if t == 'para':
            p = PS[e['p']]
            ktext = ''.join((r.find(qn('w:t')).text or '') for r in K.Doc._runs(p))
            if text_of(p._p).count(e['old']) != 1 or ktext.count(e['old']) != 1:
                skip.append((t, e['p'], e['old'][:40])); continue
            minimal(p, e['old'], e['new'])
        elif t == 'cell':
            cell = TS[e['t']].rows[e['r']].cells[e['c']]
            hits = [p for p in cell.paragraphs if e['old'] in text_of(p._p)]
            if len(hits) != 1 or text_of(hits[0]._p).count(e['old']) != 1:
                skip.append((t, e['t'], e['old'][:40])); continue
            minimal(hits[0], e['old'], e['new'])
        elif t == 'tx':
            k = 0
            for tb in PS[e['p']]._p.iter(qn('w:txbxContent')):
                for pe in tb.iter(qn('w:p')):
                    ts = list(pe.iter(qn('w:t')))
                    s = ''.join(x.text or '' for x in ts)
                    if e['old'] in s:
                        ts[0].text = s.replace(e['old'], e['new'])
                        ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                        for x in ts[1:]:
                            x.text = ''
                        k += 1
            if not k:
                skip.append((t, e['p'], e['old'][:40])); continue
        done[t] += 1
    except Exception as ex:
        skip.append((t, e.get('p', e.get('t')), e['old'][:40], str(ex)[:60]))

# “kamsitib ko‘rsatish” (kamsitmoq – kamsitish ma’nosida) → “kamaytirib ko‘rsatish” (understate)
nk = 0
for p in D.paras():
    while 'kamsitib' in text_of(p._p):
        D.replace(p, 'kamsitib', 'kamaytirib'); nk += 1
done['kamsitib→kamaytirib'] = nk
D.save(F)
print('kiritildi:', dict(done), '| o‘tkazib yuborildi:', len(skip), pardoz(F))
for s in skip:
    print('  skip', s)

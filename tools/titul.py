import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
d=docx.Document()
s=d.sections[0]; s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(14)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
def P(t='',al='c',b=False,sz=14,sa=0,it=False,sb=0):
  p=d.add_paragraph(); p.alignment={'c':WD_ALIGN_PARAGRAPH.CENTER,'r':WD_ALIGN_PARAGRAPH.RIGHT,'l':WD_ALIGN_PARAGRAPH.LEFT}[al]
  pf=p.paragraph_format; pf.space_after=Pt(sa); pf.space_before=Pt(sb); pf.line_spacing=1.15
  if t: r=p.add_run(t); r.bold=b; r.italic=it; r.font.size=Pt(sz); r.font.name='Times New Roman'
  return p
P('TASHKENT INTERNATIONAL UNIVERSITY HUZURIDAGI ILMIY DARAJALAR BERUVCHI DSc.03/2025.27.12.I.30.01 RAQAMLI ILMIY KENGASH',b=True,sa=18)
P('DENOV TADBIRKORLIK VA PEDAGOGIKA UNIVERSITETI',b=True,sa=18)
P('Qo‘lyozma huquqida',al='r',it=True)
P('UO‘K: 336.71',al='r',it=True,sa=30)
P('MAXMUDOV OMON TO‘XTAYEVICH',b=True,sa=24)
P('TIJORAT BANKLARI MOLIYAVIY BARQARORLIGINI XALQARO STANDARTLAR ASOSIDA BAHOLASH METODOLOGIYASINI TAKOMILLASHTIRISH',b=True,sz=16,sa=24)
P('08.00.07 – Moliya, pul muomalasi va kredit',sa=24)
P('Iqtisodiyot fanlari doktori (DSc)',b=True)
P('ilmiy darajasini olish uchun yozilgan',)
P('DISSERTATSIYA',b=True,sz=16,sa=36,sb=8)
p=P('',al='r'); r=p.add_run('Ilmiy rahbar: '); r.bold=True; r.font.size=Pt(14)
r=p.add_run('Tulakov Ulug‘bek Toshmamatovich'); r.font.size=Pt(14)
P('iqtisodiyot fanlari doktori, professor',al='r',sa=48)
P('Toshkent – 2026')
p=d.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
P('MUNDARIJA',b=True,sa=12)
P('{{TOC}}',al='l')
for nm in ('Heading 1','Heading 2'):
  h=d.styles[nm]; h.font.name='Times New Roman'; h.font.size=Pt(14); h.font.bold=True; h.font.color.rgb=None
  h.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); h.element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); h.element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
  h.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
from docx.shared import RGBColor
for nm in ('Heading 1','Heading 2'): d.styles[nm].font.color.rgb=RGBColor(0,0,0)
d.save('/tmp/claude-0/merge_titul.docx'); print('ok')

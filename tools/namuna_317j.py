"""3.1.7-jadval aniqlashtirishlari (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_3.1.7-jadval.docx'
G, R = '7F7F7F', 'FF0000'
N = None
# har bir katak: [(matn, rang, o‘chirilganmi)]
ROWS = [
    [[('Belgi', G, True), (' ', N, False), ('Taqqoslash mezoni', R, False)],
     [('CAMELS (AQSh)', N, False)], [('Rossiya Banki metodikasi', N, False)],
     [('XVJ barqarorlik ko‘rsatkichlari', N, False)], [('Taklif etilayotgan MBII', N, False)]],
    [[('Natija', N, False)], [('1–5 ballik kompozit reyting', N, False)],
     [('Bloklar bo‘yicha ball, guruhlarga ajratish', N, False)],
     [('Ko‘rsatkichlar to‘plami, yagona baho yo‘q', N, False)],
     [('1–4 oraliqdagi uzluksiz indeks va 4 sinf', N, False)]],
    [[('Ochiqlik', N, False)], [('Nazorat siri', N, False)], [('Metodika ochiq, natija yopiq', N, False)],
     [('Ochiq', N, False)], [('Metodika va ma’lumotlar ochiq', N, False)]],
    [[('Xalqaro standartlarga bog‘liqlik', N, False)], [('Qisman', N, False)],
     [('Milliy normativlarga bog‘langan', N, False)],
     [('XVJ standarti', G, True), (' ', N, False), ('Bazel III ga yaqin ta’riflar (XVJ qo‘llanmasi, 2019)', R, False)],
     [('Bazel III, MHXS 9, XVJ', N, False)]],
    [[('Aktivlar sifati', N, False)], [('Nazoratchi bahosi', N, False)], [('Milliy tasnif', N, False)],
     [('Milliy ta’rif', G, True), (' ', N, False), ('90 kunlik mezon (amalda milliy ta’rif)', R, False)],
     [('MHXS 9 bosqichlari', N, False)]],
    [[('Hisobot standarti farqi', N, False)], [('Hisobga olinmaydi', N, False)], [('Hisobga olinmaydi', N, False)],
     [('Hisobga olinmaydi', N, False)], [('ΔMBII va tuzatish koeffitsiyentlari', N, False)]],
    [[('Ball shkalasi', N, False)], [('Pog‘onali', N, False)], [('Pog‘onali', N, False)], [('–', N, False)],
     [('Uzluksiz (chiziqli interpolyatsiya)', N, False)]],
    [[('Vaznlar', R, False)], [('Qat’iy belgilanmagan, nazoratchi mulohazasi', R, False)],
     [('Har bir ko‘rsatkichga 1–3 vazn', R, False)], [('Qo‘llanmaydi', R, False)],
     [('Bloklar vaznlari 0,25/0,30/0,20/0,15/0,10, empirik tekshirilgan (3.2-§)', R, False)]],
    [[('Valyuta riski', N, False)], [('Bozor riski doirasida', N, False)], [('Alohida blok yo‘q', N, False)],
     [('Sof ochiq pozitsiya', N, False)], [('Alohida blok', N, False)]],
    [[('Boshqaruv sifati', R, False)], [('M komponenti', R, False)],
     [('Boshqaruv sifati va shaffoflik guruhlari', R, False)], [('Yo‘q', R, False)],
     [('Yo‘q (ochiq ma’lumotlarda baholanmaydi)', R, False)]],
    [[('Nazorat choralari bilan bog‘liqlik', N, False)], [('Bevosita', N, False)], [('Bevosita', N, False)],
     [('Yo‘q', N, False)],
     [('Sinflar', N, False), (' va erta ogohlantirish zonalari', R, False), (' bo‘yicha choralar', N, False)]],
]
FN_OLD = ('Оценка финансовой устойчивости кредитной организации: учебник / под ред. О.И. Лаврушина, '
          'И.Д. Мамоновой. – М.: КНОРУС, 2011. – С. 91–92, 135–154;')
FN_NEW = 'Указание Банка России от 03.04.2017 № 4336-У “Об оценке экономического положения банков”. – https://www.cbr.ru;'
TXT_OLD = 'Bu cheklovlar indeksning bashorat kuchini tekshirish orqali dissertatsiyaning to‘rtinchi bobida qisman bartaraf etiladi.'
TXT_NEW = ('Bu cheklovlar 3.2-paragrafda vaznlar va sinflar chegaralarini empirik tekshirish hamda to‘rtinchi bobda '
           'indeksning bashorat kuchini baholash orqali qisman bartaraf etiladi.')

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(4)

    def runs(q, parts, size=None, bold=False):
        for t, col, s in parts:
            r = q.add_run(t); r.font.strike = s; r.bold = bold
            if size: r.font.size = Pt(size)
            if col: r.font.color.rgb = RGBColor.from_string(col)

    def par(parts, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        q = dd.add_paragraph(); q.alignment = align; runs(q, parts, bold=bold); return q

    par([('3.1.7-jadval', N, False)], align=WD_ALIGN_PARAGRAPH.RIGHT)
    par([('Moliyaviy barqarorlikning integral indeksi va mavjud baholash metodikalarining qiyosiy tavsifi', N, False),
         ('¹', R, False)], True, WD_ALIGN_PARAGRAPH.CENTER)
    t = dd.add_table(rows=len(ROWS), cols=5); t.style = 'Table Grid'
    for i, row in enumerate(ROWS):
        for j, cell in enumerate(row):
            q = t.cell(i, j).paragraphs[0]
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or j > 0) else WD_ALIGN_PARAGRAPH.LEFT
            runs(q, cell, size=10, bold=(i == 0))
    for row in t.rows:
        for j, w in enumerate([3.4, 3.2, 3.4, 3.2, 3.3]):
            row.cells[j].width = Cm(w)
    dd.add_paragraph()
    q = par([('¹ Tadqiqot natijasida muallif tomonidan quyidagi manbalar asosida ishlab chiqildi: ', N, False),
             (FN_OLD, G, True), (' ', N, False), (FN_NEW, R, False),
             (' Cargill T.F. … 2017. – P. 193; International Monetary Fund. 2019 Financial Soundness Indicators '
              'Compilation Guide …; Karaliyev T.M., Sayfiddinov I.F. … 2019. – B. 208–213.', N, False)])
    for r in q.runs: r.font.size = Pt(10)
    par([('Jadvaldan keyingi xatboshi oxirida:', N, False)], True)
    par([('… Metodikaning cheklovlarini ham qayd etish lozim: … MHXS variantida esa yuqori likvid aktivlar ulushi milliy '
          'balansdan olinadi. ', N, False), (TXT_OLD, G, True), (' ', N, False), (TXT_NEW, R, False)])
    par([('Kulrang ustidan chizilgan – olib tashlanadi, qizil – qo‘shiladi.', G, False)])
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST))

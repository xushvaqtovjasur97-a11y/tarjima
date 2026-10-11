"""Native Word charts (DrawingML chart parts with embedded xlsx) and
native Word shape diagrams (wpg groups) for python-docx documents."""
import io
from xml.sax.saxutils import escape
import openpyxl
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

EMU_CM = 360000
FONT = 'Times New Roman'
NS_C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS_WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'

_counter = {'chart': 0, 'docpr': 1000, 'shape': 1}


def _rpr(sz, bold=False, color='000000'):
    return (f'<a:defRPr sz="{sz}" b="{1 if bold else 0}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="{FONT}"/><a:cs typeface="{FONT}"/></a:defRPr>')


def _txpr(sz=1100, bold=False, rot=None):
    r = f' rot="{rot}" vert="horz"' if rot is not None else ''
    return f'<c:txPr><a:bodyPr{r}/><a:lstStyle/><a:p><a:pPr>{_rpr(sz, bold)}</a:pPr><a:endParaRPr lang="uz-Latn-UZ"/></a:p></c:txPr>'


def _title(text, sz=1100, rot=None):
    r = f' rot="{rot}" vert="horz"' if rot is not None else ''
    runs = ''.join(f'<a:p><a:pPr>{_rpr(sz, False)}</a:pPr><a:r><a:rPr lang="uz-Latn-UZ" sz="{sz}" b="0"><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:latin typeface="{FONT}"/></a:rPr><a:t>{escape(line)}</a:t></a:r></a:p>' for line in text.split('\n'))
    return f'<c:title><c:tx><c:rich><a:bodyPr{r}/><a:lstStyle/>{runs}</c:rich></c:tx><c:overlay val="0"/></c:title>'


def _col(i):
    return chr(ord('A') + i)


def _strref(ref, vals):
    pts = ''.join(f'<c:pt idx="{i}"><c:v>{escape(str(v))}</c:v></c:pt>' for i, v in enumerate(vals))
    return f'<c:strRef><c:f>{ref}</c:f><c:strCache><c:ptCount val="{len(vals)}"/>{pts}</c:strCache></c:strRef>'


def _numref(ref, vals, fmt='General'):
    pts = ''.join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals) if v is not None)
    return f'<c:numRef><c:f>{ref}</c:f><c:numCache><c:formatCode>{fmt}</c:formatCode><c:ptCount val="{len(vals)}"/>{pts}</c:numCache></c:numRef>'


def _ln(color, w=12700, dash=None):
    if not w:
        return '<a:ln><a:noFill/></a:ln>'
    d = f'<a:prstDash val="{dash}"/>' if dash else ''
    if ':' in color:  # 'RRGGBB:alpha%' -> semi-transparent line
        c, al = color.split(':')
        return f'<a:ln w="{w}"><a:solidFill><a:srgbClr val="{c}"><a:alpha val="{int(al)*1000}"/></a:srgbClr></a:solidFill>{d}</a:ln>'
    return f'<a:ln w="{w}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>{d}</a:ln>'


def _dlbls(sz=1050, pos=None, color='000000', show=True, fmt=None):
    if not show:
        return '<c:dLbls><c:delete val="1"/></c:dLbls>'
    nf = f'<c:numFmt formatCode="{fmt}" sourceLinked="0"/>' if fmt else ''
    p = f'<c:dLblPos val="{pos}"/>' if pos else ''
    return (f'<c:dLbls>{nf}<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
            f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr>{_rpr(sz, True, color)}</a:pPr><a:endParaRPr lang="uz-Latn-UZ"/></a:p></c:txPr>'
            f'{p}<c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/>'
            f'<c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


GRID = '<c:majorGridlines><c:spPr><a:ln w="6350"><a:solidFill><a:srgbClr val="D9D9D9"/></a:solidFill></a:ln></c:spPr></c:majorGridlines>'
AXLN = '<c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="595959"/></a:solidFill></a:ln></c:spPr>'


def _valax(axid, cross, pos, title=None, mn=None, mx=None, major=None, grid=True, fmt='General', delete=False, crosses='autoZero'):
    sc = '<c:scaling><c:orientation val="minMax"/>' + (f'<c:max val="{mx}"/>' if mx is not None else '') + (f'<c:min val="{mn}"/>' if mn is not None else '') + '</c:scaling>'
    rot = -5400000 if pos == 'l' else None
    t = _title(title, 1100, rot) if title else ''
    mu = f'<c:majorUnit val="{major}"/>' if major else ''
    return (f'<c:valAx><c:axId val="{axid}"/>{sc}<c:delete val="{1 if delete else 0}"/><c:axPos val="{pos}"/>'
            f'{GRID if grid else ""}{t}<c:numFmt formatCode="{fmt}" sourceLinked="0"/><c:majorTickMark val="out"/>'
            f'<c:minorTickMark val="none"/><c:tickLblPos val="low"/>{AXLN}{_txpr(1100)}<c:crossAx val="{cross}"/>'
            f'<c:crosses val="{crosses}"/><c:crossBetween val="midCat"/>{mu}</c:valAx>')


def _catax(axid, cross, title=None):
    t = _title(title, 1100) if title else ''
    return (f'<c:catAx><c:axId val="{axid}"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:delete val="0"/>'
            f'<c:axPos val="b"/>{t}<c:numFmt formatCode="General" sourceLinked="0"/><c:majorTickMark val="none"/>'
            f'<c:minorTickMark val="none"/><c:tickLblPos val="low"/>{AXLN}{_txpr(1100)}<c:crossAx val="{cross}"/>'
            f'<c:crosses val="autoZero"/><c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>')


def _legend(pos='b'):
    return f'<c:legend><c:legendPos val="{pos}"/><c:overlay val="0"/>{_txpr(1050)}</c:legend>'


def _space(plot, legend=True, legend_pos='b'):
    leg = _legend(legend_pos) if legend else ''
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<c:chartSpace xmlns:c="{NS_C}" xmlns:a="{NS_A}" xmlns:r="{NS_R}"><c:date1904 val="0"/><c:lang val="uz-Latn-UZ"/>'
            '<c:roundedCorners val="0"/><c:chart><c:autoTitleDeleted val="1"/>'
            f'<c:plotArea><c:layout/>{plot}<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr></c:plotArea>{leg}'
            '<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/></c:chart>'
            '<c:spPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>'
            f'{_txpr(1100)}<c:externalData r:id="rId1"><c:autoUpdate val="0"/></c:externalData></c:chartSpace>')


def bar_chart(categories, series, ytitle=None, xtitle=None, stacked=False, ymax=None, ymin=0, major=None, fmt='General', label_fmt=None, legend=True, gap=80):
    """series: list of (name, values, color, label_color)."""
    rows = [[''] + [s[0] for s in series]] + [[c] + [s[1][i] for s in series] for i, c in enumerate(categories)]
    n = len(categories)
    sers = ''
    for k, s in enumerate(series):
        name, vals, color = s[0], s[1], s[2]
        lc = s[3] if len(s) > 3 else '000000'
        col = _col(k + 1)
        sers += (f'<c:ser><c:idx val="{k}"/><c:order val="{k}"/><c:tx>{_strref(f"Sheet1!${col}$1", [name])}</c:tx>'
                 f'<c:spPr><a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln w="6350"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:ln></c:spPr>'
                 f'<c:invertIfNegative val="0"/>{_dlbls(1050, "ctr" if stacked else "outEnd", lc, fmt=label_fmt)}'
                 f'<c:cat>{_strref(f"Sheet1!$A$2:$A${n+1}", categories)}</c:cat>'
                 f'<c:val>{_numref(f"Sheet1!${col}$2:${col}${n+1}", vals)}</c:val></c:ser>')
    grouping = 'stacked' if stacked else 'clustered'
    overlap = '<c:overlap val="100"/>' if stacked else ''
    plot = (f'<c:barChart><c:barDir val="col"/><c:grouping val="{grouping}"/><c:varyColors val="0"/>{sers}'
            f'<c:gapWidth val="{gap}"/>{overlap}<c:axId val="111"/><c:axId val="222"/></c:barChart>'
            + _catax(111, 222, xtitle) + _valax(222, 111, 'l', ytitle, ymin, ymax, major, True, fmt))
    return _space(plot, legend), rows


def _sc_lbls(lab, n, color):
    """lab: (pos, indices_to_show) -> per-point labels; others deleted."""
    if not lab:
        return ''
    pos, show = lab[:2]
    if len(lab) > 2:
        color = lab[2]
    hid = ''.join(f'<c:dLbl><c:idx val="{i}"/><c:delete val="1"/></c:dLbl>' for i in range(n) if i not in show)
    d = _dlbls(1000, pos, color)
    return d.replace('<c:dLbls>', '<c:dLbls>' + hid, 1)


def scatter_chart(series, xtitle, ytitle, xmin, xmax, xmajor, ymin, ymax, ymajor, xfmt='General', yfmt='General', legend=False):
    """series: list of (name, xs, ys, color, width, dash, marker)."""
    rows = []
    sers = ''
    maxlen = max(len(s[1]) for s in series)
    header = []
    for k, s in enumerate(series):
        header += [s[0] + ' (x)', s[0]]
    rows.append(header)
    for i in range(maxlen):
        r = []
        for s in series:
            r += [s[1][i] if i < len(s[1]) else None, s[2][i] if i < len(s[2]) else None]
        rows.append(r)
    for k, s in enumerate(series):
        name, xs, ys, color, w, dash, marker = s[:7]
        lab = s[7] if len(s) > 7 else None
        cx, cy = _col(2 * k), _col(2 * k + 1)
        n = len(xs)
        mcol = color.split(':')[0]
        mk = ('<c:marker><c:symbol val="none"/></c:marker>' if not marker else
              f'<c:marker><c:symbol val="{marker}"/><c:size val="7"/><c:spPr><a:solidFill><a:srgbClr val="{mcol}"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr></c:marker>')
        sers += (f'<c:ser><c:idx val="{k}"/><c:order val="{k}"/><c:tx>{_strref(f"Sheet1!${cy}$1", [name])}</c:tx>'
                 f'<c:spPr>{_ln(color, w, dash)}</c:spPr>{mk}{_sc_lbls(lab, len(xs), color)}'
                 f'<c:xVal>{_numref(f"Sheet1!${cx}$2:${cx}${n+1}", xs)}</c:xVal>'
                 f'<c:yVal>{_numref(f"Sheet1!${cy}$2:${cy}${n+1}", ys)}</c:yVal><c:smooth val="0"/></c:ser>')
    plot = (f'<c:scatterChart><c:scatterStyle val="lineMarker"/><c:varyColors val="0"/>{sers}'
            f'<c:axId val="333"/><c:axId val="444"/></c:scatterChart>'
            + _valax(333, 444, 'b', xtitle, xmin, xmax, xmajor, False, xfmt)
            + _valax(444, 333, 'l', ytitle, ymin, ymax, ymajor, True, yfmt))
    return _space(plot, legend), rows


def _xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def add_chart(doc, xml_and_rows):
    xml, rows = xml_and_rows
    _counter['chart'] += 1
    n = _counter['chart']
    pkg = doc.part.package
    chart = Part(PackURI(f'/word/charts/chart{n}.xml'),
                 'application/vnd.openxmlformats-officedocument.drawingml.chart+xml', b'', pkg)
    xlsx = Part(PackURI(f'/word/embeddings/Microsoft_Excel_Worksheet{n}.xlsx'),
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', _xlsx(rows), pkg)
    rid_x = chart.relate_to(xlsx, RT.PACKAGE)
    chart._blob = xml.replace('r:id="rId1"', f'r:id="{rid_x}"').encode('utf-8')
    return doc.part.relate_to(chart, RT.CHART)


def chart_run_xml(rid, w_cm, h_cm, name):
    _counter['docpr'] += 1
    cx, cy = int(w_cm * EMU_CM), int(h_cm * EMU_CM)
    return (f'<w:r xmlns:w="{NS_W}" xmlns:wp="{NS_WP}" xmlns:a="{NS_A}" xmlns:c="{NS_C}" xmlns:r="{NS_R}">'
            f'<w:rPr><w:noProof/></w:rPr><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
            f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
            f'<wp:docPr id="{_counter["docpr"]}" name="{escape(name)}"/><wp:cNvGraphicFramePr/>'
            f'<a:graphic><a:graphicData uri="{NS_C}"><c:chart r:id="{rid}"/></a:graphicData></a:graphic>'
            f'</wp:inline></w:drawing></w:r>')


# ---------------- shapes ----------------
NS_WPG = 'http://schemas.microsoft.com/office/word/2010/wordprocessingGroup'
NS_WPS = 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'
NS_MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'


def _e(v):
    return int(round(v * EMU_CM))


def _para(text, sz, bold, color, align='center', italic=False):
    align = {'l': 'left', 'r': 'right', 'c': 'center'}.get(align, align)
    out = ''
    for line in text.split('\n'):
        out += (f'<w:p><w:pPr><w:spacing w:before="0" w:after="0" w:line="228" w:lineRule="auto"/><w:jc w:val="{align}"/>'
                f'<w:ind w:firstLine="0"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:cs="{FONT}"/>'
                + ('<w:b/>' if bold else '') + ('<w:i/>' if italic else '') +
                f'<w:color w:val="{color}"/><w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr><w:t xml:space="preserve">{escape(line)}</w:t></w:r></w:p>')
    return out


def box(x, y, w, h, text='', fill='FFFFFF', line='404040', sz=22, bold=False, color='000000', geom='roundRect', lw=12700, italic=False, align='center', anchor='ctr'):
    _counter['shape'] += 1
    i = _counter['shape']
    fill_x = f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>' if fill else '<a:noFill/>'
    line_x = f'<a:ln w="{lw}"><a:solidFill><a:srgbClr val="{line}"/></a:solidFill></a:ln>' if line else '<a:ln><a:noFill/></a:ln>'
    av = '<a:avLst><a:gd name="adj" fmla="val 10000"/></a:avLst>' if geom == 'roundRect' else '<a:avLst/>'
    txb = (f'<wps:txbx><w:txbxContent>{_para(text, sz, bold, color, align, italic)}</w:txbxContent></wps:txbx>' if text else '')
    return (f'<wps:wsp><wps:cNvPr id="{i}" name="Shape {i}"/><wps:cNvSpPr/><wps:spPr><a:xfrm><a:off x="{_e(x)}" y="{_e(y)}"/>'
            f'<a:ext cx="{_e(w)}" cy="{_e(h)}"/></a:xfrm><a:prstGeom prst="{geom}">{av}</a:prstGeom>{fill_x}{line_x}</wps:spPr>'
            f'{txb}<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="36000" tIns="18000" rIns="36000" bIns="18000" anchor="{anchor}" anchorCtr="0"><a:noAutofit/></wps:bodyPr></wps:wsp>')


def line(x1, y1, x2, y2, color='404040', lw=15875, arrow=True, dash=None):
    _counter['shape'] += 1
    i = _counter['shape']
    fh = ' flipH="1"' if x2 < x1 else ''
    fv = ' flipV="1"' if y2 < y1 else ''
    d = f'<a:prstDash val="{dash}"/>' if dash else ''
    tail = '<a:tailEnd type="triangle" w="med" len="med"/>' if arrow else ''
    return (f'<wps:wsp><wps:cNvPr id="{i}" name="Connector {i}"/><wps:cNvCnPr/><wps:spPr><a:xfrm{fh}{fv}>'
            f'<a:off x="{_e(min(x1,x2))}" y="{_e(min(y1,y2))}"/><a:ext cx="{_e(abs(x2-x1))}" cy="{_e(abs(y2-y1))}"/></a:xfrm>'
            f'<a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom><a:ln w="{lw}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>{d}{tail}</a:ln>'
            f'</wps:spPr><wps:bodyPr/></wps:wsp>')


def group_run_xml(shapes, w_cm, h_cm, name):
    _counter['docpr'] += 1
    cx, cy = _e(w_cm), _e(h_cm)
    return (f'<w:r xmlns:w="{NS_W}" xmlns:wp="{NS_WP}" xmlns:a="{NS_A}" xmlns:mc="{NS_MC}" xmlns:wpg="{NS_WPG}" xmlns:wps="{NS_WPS}">'
            f'<w:rPr><w:noProof/></w:rPr><mc:AlternateContent><mc:Choice Requires="wpg"><w:drawing>'
            f'<wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="{cx}" cy="{cy}"/>'
            f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:docPr id="{_counter["docpr"]}" name="{escape(name)}"/><wp:cNvGraphicFramePr/>'
            f'<a:graphic><a:graphicData uri="{NS_WPG}"><wpg:wgp><wpg:cNvGrpSpPr/><wpg:grpSpPr><a:xfrm><a:off x="0" y="0"/>'
            f'<a:ext cx="{cx}" cy="{cy}"/><a:chOff x="0" y="0"/><a:chExt cx="{cx}" cy="{cy}"/></a:xfrm></wpg:grpSpPr>'
            + ''.join(shapes) +
            '</wpg:wgp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice><mc:Fallback/></mc:AlternateContent></w:r>')


def replace_placeholder(doc, key, run_xml):
    from docx.shared import Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for p in doc.paragraphs:
        if p.text.strip() == key:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            for el in list(p._element):
                if el.tag != f'{{{NS_W}}}pPr':
                    p._element.remove(el)
            p._element.append(etree.fromstring(run_xml))
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = Cm(0)
            pf.line_spacing = 1.0
            pf.space_before = 0
            pf.space_after = 0
            pf.keep_with_next = True
            return True
    raise KeyError(key)

"""3.2.2-jadval: izohga logit likvidlik vazni sababi, 428-izohga 2022–2024 hisobotlari; 3.2.1-rasm olib tashlanadi,
3.2.2→3.2.1, 3.2.3→3.2.2; tahlil matni jadvalga bog‘lanadi (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_3.2.2-jadval.docx'
IZ_ANC = 'tayanch vazni saqlangan.'
IZ_INS = (' Logit usulida likvidlik blokining koeffitsiyenti kutilganiga teskari ishorali va statistik ahamiyatsiz (p = 0,79) '
          'bo‘lgani sababli uning vazni nolga teng olindi.')
FN_INS = (' 2021–2023-yillar ko‘rsatkichlari esa shu banklarning 2022, 2023 va 2024-yillar uchun auditdan o‘tgan konsolidatsiyalangan '
          'moliyaviy hisobotlaridagi qiyosiy (qayta hisoblangan) ma’lumotlardan olindi.')
# (eski, yangi) – matndagi o‘zgarishlar
TXT = [
    ('Eng katta farq likvidlik va daromadlilik bloklarida kuzatiladi:', 'Likvidlik va daromadlilik bloklarida ham sezilarli farq kuzatiladi:'),
    (' Usullar natijalarining farqi 3.2.1-rasmda yaqqol ko‘rinadi.', ''),
    ('Rasmdan ko‘rinib turibdiki, usullar o‘rtasidagi farq asosan likvidlik va daromadlilik bloklariga tegishli.',
     'Jadval ma’lumotlaridan ko‘rinib turibdiki, usullar o‘rtasidagi farq asosan kapital, likvidlik va daromadlilik bloklariga tegishli.'),
    ('(3.2.2-rasm)', '(3.2.1-rasm)'), ('3.2.2-rasm. Monte-Karlo', '3.2.1-rasm. Monte-Karlo'),
    ('(3.2.3-rasm)', '(3.2.2-rasm)'), ('3.2.3-rasm. Tanlangan banklarda', '3.2.2-rasm. Tanlangan banklarda'),
    ('(3.2.3-rasmda uzuq chiziq', '(3.2.2-rasmda uzuq chiziq'),
]

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(4)

    def par(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False):
        q = dd.add_paragraph(); q.alignment = align
        for t, col, strike in parts:
            r = q.add_run(t); r.bold = bold; r.font.strike = strike
            if col: r.font.color.rgb = RGBColor.from_string(col)
        return q
    G, R = '7F7F7F', 'FF0000'
    par([('3.2.2-jadval va 3.2.1-rasm (namuna)', None, False)], WD_ALIGN_PARAGRAPH.CENTER, True)
    par([('1) 3.2.2-jadval izohida:', None, False)], bold=True)
    par([('Izoh: * – logit usulida baholanmagan, ', None, False), (IZ_ANC, None, False), (IZ_INS, R, False), (' Bosh komponentlar usulida …', None, False)])
    par([('2) Jadval nomidagi 428-izoh oxirida (2025-yil hisobotlari va “SNS” bazasidan keyin):', None, False)], bold=True)
    par([('… https://cbu.uz/uz/statistics/bankstats/.', None, False), (FN_INS, R, False)])
    par([('3) Jadval tahlilida:', None, False)], bold=True)
    par([('… ', None, False), (TXT[0][0], G, True), (' ', None, False), (TXT[0][1], R, False), (' empirik usullar likvidlik blokiga …', None, False)])
    par([('… qoplanmagan muammoli kreditlarning kapitalga nisbatiga (0,125) bergan.', None, False), (TXT[1][0], G, True)])
    par([('4) 3.2.1-rasm (diagramma, nomi va 429-izoh) to‘liq olib tashlanadi – u 3.2.2-jadval raqamlarini takrorlaydi.', None, False)], bold=True)
    par([('5) Rasmdan keyingi xatboshi saqlanadi, boshi jadvalga bog‘lanadi:', None, False)], bold=True)
    par([(TXT[2][0], G, True), (' ', None, False), (TXT[2][1], R, False), (' Bu farqni to‘g‘ri talqin qilish muhim. …', None, False)])
    par([('6) Keyingi rasmlar raqami bittaga kamayadi:', None, False)], bold=True)
    for o, n in TXT[3:]:
        par([('– ', None, False), (o, G, True), (' → ', None, False), (n, R, False)])
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

# 2.2.7-rasm: «Aloqabank» ATB qo'shilgan variant – chiziqli diagramma va ostida unga moslangan Word jadvali
import sys, json
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W
A = json.load(open('/home/user/tarjima/manbalar/hisob_kitob/aloqa23.json'))
Y = ['2021', '2022', '2023', '2024', '2025']
SER = [
    ('SQB: rasmiy muammoli', [3.8, 2.8, 2.2, 2.8, 2.5], '5B9BD5', 'dash'), ('SQB: 3-bosqich (MHXS)', [6.2, 7.4, 4.8, 5.9, 7.5], '1F4E79'),
    ('MKB: rasmiy muammoli', [5.9, 4.8, 6.0, 4.6, 4.4], '70AD47', 'dash'), ('MKB: 3-bosqich (MHXS)', [17.4, 13.5, 16.6, 15.2, 12.3], '375623'),
    ('Aloqabank: rasmiy muammoli', [round(A[y]['NPL_n'], 1) for y in Y], 'ED7D31', 'dash'), ('Aloqabank: 3-bosqich (MHXS)', [round(A[y]['S3_i'], 1) for y in Y], 'C00000')]
WCM, LEFT = 16.5, 0.38          # diagramma kengligi va chap ustun ulushi (jadval bilan bir xil)
FW = (1 - LEFT - 0.01) / len(Y)


def build(d, after_para=None):
    """Diagramma paragrafi va jadvalni qaytaradi."""
    p = d.add_paragraph('{{FIG227}}')
    rid = W.add_chart(d, W.line_table_chart(Y, SER, ymin=0, ymax=20, major=4, layout=(LEFT, 0.04, 1 - LEFT - 0.01, 0.92), dtable=False))
    W.replace_placeholder(d, '{{FIG227}}', W.chart_run_xml(rid, WCM, 6.5, '2.2.7-rasm'))
    t = d.add_table(rows=len(SER) + 1, cols=len(Y) + 1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    widths = [Cm(WCM * LEFT)] + [Cm(WCM * FW)] * len(Y)
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(w.twips if hasattr(w, 'twips') else w / 635 for w in widths)))); tw.set(qn('w:type'), 'dxa')
    for e in tblPr.findall(qn('w:tblW')):
        tblPr.remove(e)
    tblPr.append(tw)
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w / 635)))
    for col, w in zip(t.columns, widths):
        col.width = w
    for r_i, row in enumerate(t.rows):
        for c_i, cell in enumerate(row.cells):
            cell.width = widths[c_i]
            par = cell.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i else WD_ALIGN_PARAGRAPH.LEFT
            pf = par.paragraph_format; pf.space_before = pf.space_after = 0; pf.first_line_indent = Cm(0); pf.left_indent = Cm(0); pf.right_indent = Cm(0); pf.line_spacing = 1.0
            if r_i == 0:
                txt = '' if c_i == 0 else Y[c_i - 1]
                rr = par.add_run(txt); rr.bold = True
            elif c_i == 0:
                nm, vals, col = SER[r_i - 1][:3]; dash = len(SER[r_i - 1]) > 3
                k = par.add_run('━ ━ ' if dash else '━━━ '); k.font.color.rgb = RGBColor.from_string(col); k.bold = True
                k.font.size = Pt(13)
                rr = par.add_run(nm); rr.bold = True
            else:
                rr = par.add_run(f'{SER[r_i - 1][1][c_i - 1]:.1f}'.replace('.', ',')); rr.bold = True
            for j, run in enumerate(par.runs):
                if not (r_i and c_i == 0 and j == 0):
                    run.font.name = 'Times New Roman'; run.font.size = Pt(9.5)
    return p, t


if __name__ == '__main__':
    d = docx.Document()
    s = d.sections[0]; s.page_width, s.page_height = Cm(21), Cm(29.7); s.left_margin = Cm(3); s.right_margin = Cm(1.5)
    build(d)
    c = d.add_paragraph('2.2.7-rasm. «O‘zsanoatqurilishbank» ATB, «Mikrokreditbank» ATB va «Aloqabank» ATB da MHXS 9 bo‘yicha 3-bosqich kreditlari va rasmiy muammoli kreditlar ulushi, %')
    c.alignment = 1
    for r in c.runs:
        r.font.name = 'Times New Roman'; r.font.size = Pt(12); r.bold = True
    n = d.add_paragraph('Namuna (ko‘rish uchun): uzuq chiziq – rasmiy muammoli kreditlar, to‘liq chiziq – MHXS 9 bo‘yicha 3-bosqich kreditlari. Aloqabank ma’lumotlari: muallif hisob-kitobi (prudensial shakllar va MHXS hisobotlari).')
    for r in n.runs:
        r.font.name = 'Times New Roman'; r.font.size = Pt(11); r.italic = True
    d.save('/home/user/tarjima/Namuna_2.2.7-rasm_Aloqabank.docx'); print('saqlandi')

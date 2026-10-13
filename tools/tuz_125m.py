"""1.2.5-rasm o‘rniga “standart × blok” matritsa-jadval (1.2.5-jadval) kiritiladi; matn havolalari tuzatiladi (qizil)."""
import copy, sys
sys.path.insert(0, '/home/user/tarjima/tools')
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table
import kamlib as K
import namuna_125m as N
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
RED = RGBColor(0xFF, 0, 0)
D = K.Doc(F)
p_before = D.para('Jadval ma’lumotlari shuni ko‘rsatadiki, ko‘rib chiqilgan standartlarning birortasi')
cap = D.para('1.2.5-rasm. Xalqaro standartlarning moliyaviy barqarorlikni integral')
p_after = D.para('Rasmdan ko‘rinadiki, kapital yetarliligi va likvidlik bloklari')
lab_src = [q for q in D.d.paragraphs if q.text.strip() == '1.2.4-jadval'][0]
assert lab_src.text.strip() == '1.2.4-jadval'
t_src = lab_src._p.getnext().getnext()
assert t_src.tag == qn('w:tbl')
title_src = lab_src._p.getnext()

# 1) rasmni olib tashlash
drw = cap._p.getprevious()
assert drw.xpath('.//w:drawing') or drw.xpath('.//mc:AlternateContent')
drw.getparent().remove(drw)

# 2) jadval raqami (o‘ng tomonda, kursiv, qizil)
lab = copy.deepcopy(lab_src._p)
for r in lab.findall(qn('w:r')):
    lab.remove(r)
cap._p.addprevious(lab)
r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr'); i_ = OxmlElement('w:i'); col = OxmlElement('w:color'); col.set(qn('w:val'), 'FF0000')
rpr.append(i_); rpr.append(col); r.append(rpr); t = OxmlElement('w:t'); t.text = '1.2.5-jadval'; r.append(t); lab.append(r)

# 3) rasm nomi -> jadval nomi (snoska saqlanadi)
cap._p.remove(cap._p.find(qn('w:pPr')))
cap._p.insert(0, copy.deepcopy(title_src.find(qn('w:pPr'))))
D.replace(cap, '1.2.5-rasm. ', '', tag='125M')
for rr in cap.runs:
    rr.font.color.rgb = RED

# 4) jadval
rows = 1 + len(N.M) + 1
tb = D.d.add_table(rows=rows, cols=6)
tbl = tb._tbl
tbl.remove(tbl.tblPr); tbl.insert(0, copy.deepcopy(t_src.find(qn('w:tblPr'))))
W_ = [1700] + [1531] * 5
for gc, w in zip(tbl.tblGrid.findall(qn('w:gridCol')), W_):
    gc.set(qn('w:w'), str(w))


def shade(cell, color):
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(sh)


def put(cell, text, bold=False, size=11, left=False, fill=None):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format; pf.line_spacing = 1.0; pf.first_line_indent = Cm(0); pf.space_before = Pt(0); pf.space_after = Pt(0)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RED
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade(cell, fill)


for row in tb.rows:
    for i, cl in enumerate(row.cells):
        cl.width = Cm(W_[i] / 567)
put(tb.cell(0, 0), 'Standart / blok', bold=True, fill='DEEBF7')
for j, b in enumerate(N.BLOK, 1):
    put(tb.cell(0, j), b, bold=True, size=10.5, fill='DEEBF7')
for i, (s, cells) in enumerate(N.M, 1):
    put(tb.cell(i, 0), s, bold=True, left=True, fill='F2F2F2')
    for j, v in enumerate(cells, 1):
        put(tb.cell(i, j), v, size=10)
last = rows - 1
roles = ['Bazel III mezonlari', 'MHXS 9 bo‘yicha baholash', 'XVJ va CAMELS ko‘rsatkichlari', 'Bazel III mezonlari', 'XVJ ko‘rsatkichi, dollarlashuvni hisobga olish']
put(tb.cell(last, 0), 'Integral bahoga hissasi', bold=True, left=True, fill='DEEBF7')
for j, v in enumerate(roles, 1):
    put(tb.cell(last, j), v, bold=True, size=10, fill='DEEBF7')
cap._p.addnext(tbl)

# 5) izoh
iz = copy.deepcopy(title_src)
for e in list(iz):
    if e.tag != qn('w:pPr'):
        iz.remove(e)
ppr = iz.find(qn('w:pPr'))
for e in list(ppr):
    if e.tag in (qn('w:keepNext'), qn('w:jc')):
        ppr.remove(e)
jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'both'); ppr.append(jc)
sp = ppr.find(qn('w:spacing')); sp.set(qn('w:after'), '240')
tbl.addnext(iz)
r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
for tag, val in (('w:i', None), ('w:color', 'FF0000'), ('w:sz', '22'), ('w:szCs', '22')):
    e = OxmlElement(tag)
    if val:
        e.set(qn('w:val'), val)
    rpr.append(e)
r.append(rpr); t = OxmlElement('w:t'); t.text = 'Izoh: ● – blok uchun asosiy mezon yoki ko‘rsatkich manbai; ○ – bilvosita ta’sir; – – blokni qamramaydi.'
r.append(t); iz.append(r)

# 6) matn havolalari
D.replace(p_before, '(1.2.5-rasm)', '(1.2.5-jadval)', tag='125M')
D.replace(p_after, 'Rasmdan ko‘rinadiki', 'Jadvaldan ko‘rinadiki', tag='125M')
D.save(F)
print(pardoz(F))

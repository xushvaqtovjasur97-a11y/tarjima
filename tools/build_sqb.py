import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/SQB_milliy_MHXS.md'; DST='/home/user/tarjima/paragraflar/SQB_milliy_MHXS.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/sqb_raw.docx'],check=True)
d=docx.Document(S+'/sqb_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^\d+-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^\d+-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^\d+-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'
rid=W.add_chart(d,W.bar_chart(['Muammoli kreditlar\n(milliy) / 3-bosqich (MHXS)','Substandart kreditlar\n(milliy) / 2-bosqich (MHXS)'],[
  ('Milliy hisobot',[2.53,18.44],BLUE),('MHXS',[7.48,35.66],ORANGE)],
  ytitle='Jami kreditlardagi ulushi, %',ymax=40,major=10,label_fmt='0.00',gap=80))
W.replace_placeholder(d,'{{FIG:1}}',W.chart_run_xml(rid,15.0,7.5,'1-rasm'))
rid=W.add_chart(d,W.bar_chart(['Xususiy kapital','Sof foyda','Kreditlar bo‘yicha\nzaxiralar'],[
  ('Milliy hisobot',[11.88,1.88,2.95],BLUE),('MHXS',[10.27,1.54,4.25],ORANGE)],
  ytitle='trln so‘m',ymax=14,major=2,label_fmt='0.00',gap=80))
W.replace_placeholder(d,'{{FIG:2}}',W.chart_run_xml(rid,15.0,7.5,'2-rasm'))
rid=W.add_chart(d,W.bar_chart(['Markaziy bank\nmetodikasi','Bazel III\nmetodikasi'],[
  ('Milliy hisobot',[122.2,127.6],BLUE),('MHXS tuzatishlari bilan',[119.5,124.8],ORANGE),('Minimal talab',[100,100],GREY)],
  ytitle='%',ymax=140,ymin=0,major=20,label_fmt='0.0',gap=80))
W.replace_placeholder(d,'{{FIG:3}}',W.chart_run_xml(rid,15.0,7.5,'3-rasm'))
d.save(DST); print('saved')

import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/2.1-paragraf.md'; DST='/home/user/tarjima/paragraflar/2.1-paragraf.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p21_raw.docx'],check=True)
d=docx.Document(S+'/p21_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^2\.1\.\d-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^2\.1\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^2\.1\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'


# ---- 2.1.1 normative framework
sh=[]
lv=[('Qonunlar',DBLUE,['«O‘zbekiston Respublikasining\nMarkaziy banki to‘g‘risida»gi Qonun','«Banklar va bank faoliyati\nto‘g‘risida»gi Qonun']),
    ('Prudensial\ntalablar',BLUE,['Kapital monandligi\n(№ 3697, 2025)','Likvidlik\n(LCR, NSFR, milliy\nnormativlar)','Aktivlarni tasniflash\nva zaxiralar']),
    ('Ichki boshqaruv\ntalablari',GREEN,['Korporativ boshqaruv\n(№ 3254, 2020)','Risklarni boshqarish\ntizimi (№ 3427, 2023)','Ichki audit\n(№ 3302, 2021)']),
    ('Nazorat\njarayonlari',ORANGE,['Risklarga asoslangan\nnazorat (42/29, 2023)','Asoslangan mulohaza\n(34/3, 2023)','Tizimli banklar\n(4/13, 2023)','Kontrsiklik\nbufer (2026)'])]
y=0.1; H=1.35; G=0.45
for k,(lab,c,items) in enumerate(lv):
  sh.append(W.box(0.1,y,2.8,H,lab,fill=c,line=c,sz=22,bold=True,color='FFFFFF'))
  n=len(items); W_=13.3; gap=0.2; w=(W_-gap*(n-1))/n
  for i,it in enumerate(items):
    sh.append(W.box(3.1+i*(w+gap),y,w,H,it,fill='F2F2F2',line=c,sz=20))
  if k<3: sh.append(W.line(1.5,y+H,1.5,y+H+G,color='404040',lw=19050))
  y+=H+G
W.replace_placeholder(d,'{{FIG:2.1.1}}',W.group_run_xml(sh,16.4,y-G+0.15,'2.1.1-rasm'))

# ---- 2.1.2 assessment system
sh=[]
sh.append(W.box(5.7,0.1,5.0,0.9,'O‘zbekiston Respublikasi\nMarkaziy banki',fill=DBLUE,line=DBLUE,sz=22,bold=True,color='FFFFFF'))
cols=[('Prudensial normativlar',BLUE,'Kapital monandligi;\nleverej;\nLCR, NSFR va milliy\nlikvidlik normativlari;\nhisobotlar asosida\nmuntazam nazorat'),
      ('Risklarga asoslangan\nnazorat',GREEN,'Bank risk profili va\nichki nazoratining\navtomatik reytingi;\nkuratorlar tuzatishi;\nasoslangan mulohaza;\njoyida tekshiruv'),
      ('Makroprudensial\nmonitoring',ORANGE,'Moliyaviy barqarorlik\nsharhi (yiliga 2 marta);\nto‘lov qobiliyati va\nlikvidlik stress-testlari;\nkontrsiklik bufer;\ntizimli banklar')]
for i,(h,c,t) in enumerate(cols):
  x=0.1+i*5.5
  sh.append(W.line(8.2,1.0,x+2.55,1.6,color='404040',lw=15875))
  sh.append(W.box(x,1.6,5.1,1.0,h,fill=c,line=c,sz=22,bold=True,color='FFFFFF'))
  sh.append(W.box(x,2.75,5.1,2.6,t,fill='F2F2F2',line=c,sz=20))
sh.append(W.box(0.1,5.55,10.6,0.65,'Mikro daraja – alohida bank',fill='DEEBF7',line='9DC3E6',sz=20,bold=True))
sh.append(W.box(11.1,5.55,5.1,0.65,'Makro daraja – bank tizimi',fill='FBE5D6',line='F4B183',sz=20,bold=True))
sh.append(W.box(0.1,6.4,16.1,0.75,'Tashqi baholash: xalqaro reyting agentliklari, auditorlar (MHXS bo‘yicha hisobot), XVJ va Jahon banki (FSAP)',fill='FFFFFF',line='7F7F7F',sz=20,italic=True))
W.replace_placeholder(d,'{{FIG:2.1.2}}',W.group_run_xml(sh,16.4,7.3,'2.1.2-rasm'))

# ---- 2.1.3 CAR dynamics
rid=W.add_chart(d,W.bar_chart(['2023-y. oxiri','2024-y. oxiri','2025-y. I yarim','2025-y. oxiri','2026-y. 1-iyul'],[
  ('Jami kapital monandligi',[17.5,17.0,17.4,18.3,18.5],BLUE),('Asosiy kapital (CET1) monandligi',[None,None,14.6,14.7,15.0],ORANGE)],
  ytitle='Risk bo‘yicha tortilgan aktivlarga nisbatan, %',ymax=22,major=2,label_fmt='0.0',gap=80))
W.replace_placeholder(d,'{{FIG:2.1.3}}',W.chart_run_xml(rid,15.5,8.0,'2.1.3-rasm'))

# ---- 2.1.4 stress test
rid=W.add_chart(d,W.bar_chart(['Asosiy kapital (CET1)','Jami kapital'],[
  ('Amaldagi daraja (2025-y. oxiri)',[14.7,18.3],BLUE),
  ('Salbiy ssenariy: 2028-y. I yarmi (2025-y. I yarim sharhi)',[6.1,8.4],GOLD),
  ('Salbiy ssenariy: 2028-y. oxiri (2025-y. sharhi)',[5.6,6.8],ORANGE),
  ('Regulyativ minimum',[8.0,13.0],GREY)],
  ytitle='Risk bo‘yicha tortilgan aktivlarga nisbatan, %',ymax=20,major=4,label_fmt='0.0',gap=60))
W.replace_placeholder(d,'{{FIG:2.1.4}}',W.chart_run_xml(rid,15.5,8.5,'2.1.4-rasm'))

# ---- 2.1.5 problems
sh=[]
sh.append(W.box(4.7,0.1,7.0,1.0,'Xalqaro standartlarga moslashtirish muammolari',fill=DBLUE,line=DBLUE,sz=22,bold=True,color='FFFFFF'))
grp=[('Kapitalni o‘lchash',BLUE,'subordinar qarzlar Bazel\nmezonlariga mos emas;\nRWA metodikasida\nchetlanishlar'),
     ('Aktivlar sifati',ORANGE,'NPL va restrukturizatsiya\nta’riflari tor;\nkamsitib ko‘rsatish;\nMHXS 9 to‘liq emas'),
     ('Axborot va\nshaffoflik',GREEN,'oshkoralik bo‘yicha\nminimal talablar yo‘q;\nmilliy hisobot va MHXS\nfarqi'),
     ('Institutsional',GREY,'davlat ulushi yuqori;\nkonsolidatsiyalangan\nnazorat va sanatsiya\nrejimi shakllanmoqda'),
     ('Baholash vositalari',GOLD,'erta ogohlantirish va\nPCA yo‘q; 2-ustun yo‘q;\nreytinglar yopiq; integral\nko‘rsatkich yo‘q')]
w=3.05; gap=0.22
for i,(h,c,t) in enumerate(grp):
  x=0.1+i*(w+gap)
  sh.append(W.line(8.2,1.1,x+w/2,1.7,color='404040',lw=15875))
  sh.append(W.box(x,1.7,w,1.0,h,fill=c,line=c,sz=20,bold=True,color='000000' if c in (GOLD,GREY) else 'FFFFFF'))
  sh.append(W.box(x,2.85,w,2.3,t,fill='F2F2F2',line=c,sz=18))
  sh.append(W.line(x+w/2,5.15,8.2,5.75,color='7F7F7F',lw=12700))
sh.append(W.box(2.7,5.75,11.0,1.0,'Ochiq ma’lumotlarga asoslangan integral baholash va erta ogohlantirish vositasiga ehtiyoj',fill='DEEBF7',line=BLUE,sz=22,bold=True))
W.replace_placeholder(d,'{{FIG:2.1.5}}',W.group_run_xml(sh,16.4,6.9,'2.1.5-rasm'))
d.save(DST); print('saved')

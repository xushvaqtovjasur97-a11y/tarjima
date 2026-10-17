"""3.2.1-rasm: variatsiya koeffitsiyenti usuli qo‘shildi, empirik o‘rtacha jadvaldagi kabi, logitda baholanmagan vaznlar shtrixli (*),
vergulli vertikal yorliqlar; matnda farqlar talqini aniqlashtirildi (namuna)."""
import re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W

DST = '/home/user/tarjima/Namuna_3.2.1-rasm.docx'
CATS = ['Kapital', 'Aktivlar sifati', 'Likvidlik', 'Daromadlilik', 'Valyuta riski']
SER = [('Tayanch', [0.25, 0.30, 0.20, 0.15, 0.10], '1F4E79'),
       ('Entropiya', [0.02, 0.32, 0.03, 0.46, 0.17], 'ED7D31'),
       ('Variatsiya koeffitsiyenti', [0.07, 0.32, 0.06, 0.37, 0.18], '5B9BD5'),
       ('Bosh komponentlar', [0.14, 0.27, 0.16, 0.30, 0.13], '70AD47'),
       ('Logit', [0.25, 0.57, 0.00, 0.08, 0.10], 'FFC000'),
       ('Empirik o‘rtacha', [0.08, 0.30, 0.08, 0.38, 0.16], '7F7F7F')]
STAR = {4: {0, 4}}          # logit: kapital va valyuta – baholanmagan (tayanch vazn)
IZOH = ('Izoh: * – logit usulida baholanmagan, tayanch vazni saqlangan (shtrixli ustunlar); empirik o‘rtacha entropiya, '
        'variatsiya koeffitsiyenti va bosh komponentlar usullari bo‘yicha hisoblangan.')
OLD1 = 'Eng katta farq likvidlik va daromadlilik bloklarida kuzatiladi:'
NEW1 = 'Likvidlik va daromadlilik bloklarida ham sezilarli farq kuzatiladi:'
OLD2 = 'usullar o‘rtasidagi farq asosan likvidlik va daromadlilik bloklariga tegishli.'
NEW2 = 'usullar o‘rtasidagi farq asosan kapital, likvidlik va daromadlilik bloklariga tegishli.'


def _f(v):
    return f'{v:.2f}'.replace('.', ',')


def _lbls(k, vals, col):
    one = ''
    for i, v in enumerate(vals):
        t = _f(v) + ('*' if i in STAR.get(k, set()) else '')
        one += (f'<c:dLbl><c:idx val="{i}"/><c:tx><c:rich><a:bodyPr rot="-5400000" vert="horz" wrap="none"/><a:lstStyle/><a:p><a:pPr>{W._rpr(850, True, "000000")}</a:pPr>'
                f'<a:r><a:rPr lang="uz-Latn-UZ" sz="850" b="1"><a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
                f'<a:latin typeface="Times New Roman"/><a:cs typeface="Times New Roman"/></a:rPr><a:t>{t}</a:t></a:r></a:p></c:rich></c:tx>'
                '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr><c:dLblPos val="outEnd"/><c:showLegendKey val="0"/><c:showVal val="1"/>'
                '<c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>')
    return ('<c:dLbls>' + one + '<c:showLegendKey val="0"/><c:showVal val="0"/><c:showCatName val="0"/>'
            '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


def chart():
    xml, rows = W.bar_chart(CATS, SER, ytitle='Vazn', ymin=0, ymax=0.7, major=0.1, fmt='0.0', gap=60)
    sers = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)
    for k, (s, (nm, vals, col)) in enumerate(zip(sers, SER)):
        ns = re.sub(r'<c:dLbls>.*?</c:dLbls>', _lbls(k, vals, col), s, flags=re.S)
        if k in STAR:
            dpts = ''.join(f'<c:dPt><c:idx val="{i}"/><c:invertIfNegative val="0"/><c:bubble3D val="0"/><c:spPr>'
                           f'<a:pattFill prst="wdUpDiag"><a:fgClr><a:srgbClr val="{col}"/></a:fgClr><a:bgClr><a:srgbClr val="FFFFFF"/></a:bgClr></a:pattFill>'
                           f'<a:ln w="9525"><a:solidFill><a:srgbClr val="{col}"/></a:solidFill></a:ln></c:spPr></c:dPt>' for i in sorted(STAR[k]))
            ns = ns.replace('<c:invertIfNegative val="0"/>', '<c:invertIfNegative val="0"/>' + dpts, 1)
        xml = xml.replace(s, ns, 1)
    return xml, rows


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)

    def par(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, bold=False):
        q = dd.add_paragraph(); q.alignment = align
        for t, col, strike in parts:
            r = q.add_run(t); r.font.size = Pt(size); r.bold = bold; r.font.strike = strike
            if col: r.font.color.rgb = RGBColor.from_string(col)
        return q
    par([('3.2.1-rasm (namuna)', None, False)], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, chart())
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 9.5, '3.2.1-rasm')))
    par([(IZOH, 'FF0000', False)], WD_ALIGN_PARAGRAPH.LEFT, size=11)
    par([('3.2.1-rasm. Integral indeks bloklarining tayanch va empirik vaznlari', None, False)], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    dd.add_paragraph()
    par([('Matndagi o‘zgarishlar:', None, False)], size=12, bold=True)
    par([('1) Rasmdan oldingi xatboshida: ', None, False), (OLD1, '7F7F7F', True), (' → ', None, False), (NEW1, 'FF0000', False)], size=12)
    par([('2) Rasmdan keyingi xatboshida: Rasmdan ko‘rinib turibdiki, ', None, False), (OLD2, '7F7F7F', True), (' → ', None, False), (NEW2, 'FF0000', False)], size=12)
    par([('Asos: empirik o‘rtachaning tayanch vazndan farqi – daromadlilik 0,23, kapital 0,17, likvidlik 0,12, valyuta riski 0,06, aktivlar sifati 0,00.', None, False)], size=12)
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

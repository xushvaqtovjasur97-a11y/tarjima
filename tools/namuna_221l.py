"""2.2.1-rasm chiziqli diagramma ko‘rinishida (namuna)."""
import sys, re
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_2.2.1-rasm_chiziqli.docx'
Y = ['2021', '2022', '2023', '2024', '2025']
SER = [('Kreditlardagi ulushi', [85.8, 83.2, 70.7, 68.8, 67.0], 'EB6834'),
       ('Aktivlardagi ulushi', [81.5, 78.2, 67.7, 65.4, 63.5], '2A78D6'),
       ('Depozitlardagi ulushi', [66.7, 62.5, 51.5, 50.1, 51.0], '70AD47')]
POS = ('t', 'b', 'b')


def chart():
    xml, rows = W.line_table_chart(Y, SER, ytitle='Bank tizimidagi ulushi, %', ymin=40, ymax=90, major=10, fmt='0.0', w=31750, dtable=False)
    xml = xml.replace('<c:tickLblPos val="none"/>', '<c:tickLblPos val="low"/>', 1)
    sers = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)
    for s, (nm, v, col), pos in zip(sers, SER, POS):
        xml = xml.replace(s, s.replace('<c:cat>', W._dlbls(1050, pos, col, fmt='0.0') + '<c:cat>', 1), 1)
    xml = xml.replace('<c:size val="5"/>', '<c:size val="7"/>')
    xml = xml.replace('</c:plotArea>', '</c:plotArea>' + W._legend('b'), 1)
    return xml, rows


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, chart())
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.5, '2.2.1-rasm')))
    c = dd.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run('2.2.1-rasm. Davlat ulushi mavjud banklarning bank tizimidagi ulushi, %'); r.bold = True
    dd.save(DST)
    print(pardoz(DST), DST)

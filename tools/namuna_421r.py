"""4.2.1-rasm: haqiqiy qiymatlar (yil oxiri va 2021-yil cho‘qqisi) chiziq bilan ustma-ust tushmaydigan joyda,
2027- va 2028-yil prognoz qiymatlari barcha ssenariylar uchun ko‘rsatiladi – namuna.
Diagramma ma’lumotlari o‘zgarmaydi, faqat yorliqlar, marker o‘lchami va X o‘qi chegarasi (2028-yil yorliqlari uchun joy)."""
import io, re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import openpyxl
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import kamlib as K
import wordart as W

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_4.2.1-rasm.docx'
CAP0 = '4.2.1-rasm. Bank tizimida muammoli kreditlar ulushining'
W_CM, H_CM = 16.0, 9.0
# seriya nomi → {nuqta indeksi: yorliq joyi}; haqiqiy: 2021-y. III chorak cho‘qqisi va har yil oxiri (IV chorak)
LBL = {'Haqiqiy (2021–2025)': {2: 't', 3: 'l', 7: 'l', 11: 'b', 15: 'b', 19: 'b'},
       'Asosiy ssenariy': {8: 't', 12: 'r'},
       'Ijobiy ssenariy': {8: 'b', 12: 'r'},
       'Salbiy ssenariy': {8: 't', 12: 'r'},
       'Keskin salbiy ssenariy': {8: 't', 12: 't'}}


def _lbl(i, pos, color):
    return (f'<c:dLbl><c:idx val="{i}"/><c:numFmt formatCode="0.00" sourceLinked="0"/><c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
            f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr>{W._rpr(1000, True, color)}</a:pPr><a:endParaRPr lang="uz-Latn-UZ"/></a:p></c:txPr>'
            f'<c:dLblPos val="{pos}"/><c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/>'
            '<c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>')


def fix(xml):
    out = xml
    for s in re.findall(r'<c:ser>.*?</c:ser>', xml, re.S):
        name = re.search(r'<c:tx>.*?<c:v>([^<]*)</c:v>', s, re.S).group(1)
        n = int(re.search(r'<c:yVal>.*?<c:ptCount val="(\d+)"/>', s, re.S).group(1))
        color = re.search(r'<c:spPr><a:ln[^>]*><a:solidFill><a:srgbClr val="(\w+)"/>', s).group(1)
        show = LBL[name]
        body = ''.join(_lbl(i, show[i], color) if i in show else f'<c:dLbl><c:idx val="{i}"/><c:delete val="1"/></c:dLbl>' for i in range(n))
        d = (f'<c:dLbls>{body}<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
             '<c:showLegendKey val="0"/><c:showVal val="0"/><c:showCatName val="0"/><c:showSerName val="0"/>'
             '<c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')
        ns = re.sub(r'<c:dLbls>.*?</c:dLbls>', d, s, count=1, flags=re.S)
        ns = ns.replace('<c:symbol val="circle"/><c:size val="7"/>', '<c:symbol val="circle"/><c:size val="5"/>')
        out = out.replace(s, ns, 1)
    # X o‘qi 2028,6 gacha – 2028-yil yorliqlari chizish maydoni ichida qoladi (belgilar 2020–2028)
    out, k = re.subn(r'<c:max val="2028"/><c:min val="2020"/>', '<c:max val="2028.6"/><c:min val="2020"/>', out)
    assert k == 1
    return out


def source():
    D = K.Doc(F)
    cap = D.para(CAP0)
    chp = cap._p.getprevious()
    rid = [e.get(qn('r:id')) for e in chp.iter('{http://schemas.openxmlformats.org/drawingml/2006/chart}chart')][0]
    part = D.d.part.related_parts[rid]
    return D, cap, chp, part


if __name__ == '__main__':
    D, cap, chp, part = source()
    xml = fix(part.blob.decode('utf-8'))
    xrel = [r for r in part.rels.values() if r.reltype.endswith('/package')][0]
    ws = openpyxl.load_workbook(io.BytesIO(xrel.target_part.blob)).active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    xml = re.sub(r'(<c:externalData r:id=")[^"]+(")', r'\g<1>rId1\g<2>', xml)
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.5
    h = dd.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run('4.2.1-rasm – tuzatilgan ko‘rinish (namuna)'); r.bold = True
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, (xml, rows))
    q._p.append(etree.fromstring(W.chart_run_xml(rid, W_CM, H_CM, '4.2.1-rasm')))
    c = dd.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; rr = c.add_run(cap.text.rstrip()); rr.bold = True
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

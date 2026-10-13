"""2.2-§ qayta tuzilishi: uch bank bo‘yicha alohida jadvallar (2.2.2–2.2.4) va tahlil, takroriy rasmlar (2.2.5, 2.2.6, 2.2.8) olib tashlandi,
2.2.7 → 2.2.5-rasm, 2.2.3-jadval → 2.2.5-jadval (“Aloqabank” ustuni), 2.2.3 va 2.2.4-rasmlar chiziqli + mezon chiziqlari.
Qizil – tuzatishlar, ko‘k – “Aloqabank”."""
import copy, re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
from lxml import etree
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
import kamlib as K
import wordart as W
import namuna_222j as J
import namuna_222t as T
import namuna_22v2 as V
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
RED, BLUE = 'FF0000', '0070C0'
D = K.Doc(F)
body = D.d.element.body
CHART = '{http://schemas.openxmlformats.org/drawingml/2006/chart}chart'


def P(start):
    r = [p for p in D.d.paragraphs if p.text.startswith(start)]
    assert len(r) == 1, (start, len(r))
    return r[0]


def color(c):
    K.RED = c


def text_of(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


# ---------- 0) qidiruvlar (tahrirdan oldin)
p_open = P('Tahlil ikki darajada amalga oshirildi.')
p_intro = P('Tizim darajasidagi agregat ko‘rsatkichlar alohida banklar')
lab222 = [p for p in D.d.paragraphs if p.text.strip() == '2.2.2-jadval'][0]
tit222 = P('“O‘zsanoatqurilishbank” ATB va “Mikrokreditbank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari')
tbl = tit222._p.getnext(); assert tbl.tag == qn('w:tbl')
izoh = Paragraph(tbl.getnext(), tit222._parent); assert izoh.text.startswith('Izoh: muammoli kreditlar')
p_old1 = P('Jadval ma’lumotlari ikki bankning moliyaviy barqarorlik profili')
cap5 = P('2.2.5-rasm.'); cap6 = P('2.2.6-rasm.'); cap7 = P('2.2.7-rasm.'); cap8 = P('2.2.8-rasm.')
p_teb = P('Rasmda aks etgan tebranishlarning sababi')
p_akt = P('Aktivlar sifati bo‘yicha ham ikki bank turlicha')
p_akt2 = P('Rasmdan ko‘rinib turibdiki, “O‘zsanoatqurilishbank” ATB da rasmiy muammoli')
p_mhxs = P('Rasmdan ko‘rinib turibdiki, ikkala bankda ham MHXS 9')
p_roe0 = P('Aktivlar sifatidagi bu muammolar “Mikrokreditbank”')
p_roe = P('Rasm ma’lumotlari shuni ko‘rsatadiki, “O‘zsanoatqurilishbank” ATB ning kapital rentabelligi')
p_lik = P('Likvidlik bo‘yicha esa ikkala bank ham Bazel III')
p_sum = P('Tahlil natijalarini umumlashtirish uchun bank tizimi va ikki bankning')
lab223 = [p for p in D.d.paragraphs if p.text.strip() == '2.2.3-jadval'][0]
tit_sum = P('Bank tizimi va tanlangan banklarning moliyaviy barqarorlik ko‘rsatkichlari: xalqaro mezonlar')
tbl_sum = Table(tit_sum._p.getnext(), tit_sum._parent)
p_xul = P('O‘zbekiston bank tizimi va alohida banklar moliyaviy barqarorlik ko‘rsatkichlarini tadqiq etgan holda')
cap3 = P('2.2.3-rasm.'); cap4 = P('2.2.4-rasm.')
p_after3 = Paragraph(cap3._p.getnext(), cap3._parent)
p_after4 = Paragraph(cap4._p.getnext(), cap4._parent)
assert p_after3.text.startswith('Rasmiy ko‘rsatkich va MHXS 9') and p_after4.text.startswith('CAMELS tizimida aktivlar rentabelligi')
fn_tit = [r.get(qn('w:id')) for r in tit222._p.iter(qn('w:footnoteReference'))][0]
fn_sum = [r.get(qn('w:id')) for r in tit_sum._p.iter(qn('w:footnoteReference'))][0]
TEMPL = copy.deepcopy(p_old1._p)        # oddiy matn xatboshisi namunasi


def new_par(anchor_el, parts):
    """anchor_el dan keyin oddiy xatboshi; parts – [(matn, rang, snoska|None)]."""
    np_ = copy.deepcopy(TEMPL)
    runs = [r for r in np_ if r.tag == qn('w:r')]
    tmpl = [r for r in runs if r.find(qn('w:t')) is not None][0]
    for e in list(np_):
        if e.tag != qn('w:pPr'):
            np_.remove(e)
    for text, col, fn in parts:
        color(col)
        s = text + (f'[[FN:{fn}]]' if fn else '')
        for e in D._build(tmpl, s):
            np_.append(e)
    color(RED)
    anchor_el.addnext(np_)
    return np_


# ---------- 1) 2.2-§ boshidagi xatboshi (tizimli bank ta’rifi + “Aloqabank”)
D.replace(p_open, 'tizimli ahamiyatga ega ikki davlat banki – “O‘zsanoatqurilishbank” ATB va “Mikrokreditbank” ATB ning',
          'uchta davlat banki – tizimli ahamiyatga molik “O‘zsanoatqurilishbank” ATB, “Mikrokreditbank” ATB va ', tag='22')
color(BLUE)
D.insert(p_open, 'Mikrokreditbank” ATB va ', '“Aloqabank” ATB ning', tag='22')
color(RED)
D.insert(p_open, 'Mikrokreditbank” ATB va “Aloqabank” ATB ning', ' Markaziy bankka taqdim etilgan prudensial hisobot shakllari', tag='22')
D.replace(p_open, ' Markaziy bankka taqdim etilgan prudensial hisobot shakllari Markaziy bankka', ' Markaziy bankka', tag='22') if False else None
D.replace(p_open, 'ikkinchisi davlat dasturlari doirasida mikrokreditlash va ijtimoiy yo‘naltirilgan kreditlash vazifasini bajaradi. Shu sababli ular davlat banklari faoliyatining ikki xil modelini',
          'ikkinchisi davlat dasturlari doirasida mikrokreditlash va ijtimoiy yo‘naltirilgan kreditlash vazifasini bajaradi', tag='22')
color(BLUE)
D.insert(p_open, 'ijtimoiy yo‘naltirilgan kreditlash vazifasini bajaradi',
         ', uchinchisi esa tizimli ahamiyatga molik banklar qatoriga kirmaydigan, kreditlarni asosan mijozlar depozitlari hisobiga moliyalashtiruvchi o‘rta hajmdagi bank', tag='22')
color(RED)
D.insert(p_open, 'kreditlarni asosan mijozlar depozitlari hisobiga moliyalashtiruvchi o‘rta hajmdagi bank',
         '. Shu sababli ular davlat banklari faoliyatining uch xil modelini', tag='22')

# ---------- 2) kirish xatboshi
D.replace(p_intro, p_intro.text, J.INTRO, tag='22')

# ---------- 3) jadvallar
tbl_mk = copy.deepcopy(tbl); tbl_al = copy.deepcopy(tbl)
trs = tbl.findall(qn('w:tr'))
hdr_idx = [i for i, tr in enumerate(trs) if 'Mikrokreditbank' in text_of(tr)][0]
for tr in trs[hdr_idx:] + [trs[1]]:
    tbl.remove(tr)
for t_ in (tbl_mk, tbl_al):
    trs_ = t_.findall(qn('w:tr'))
    for tr in trs_[1:hdr_idx + 1]:
        t_.remove(tr)
D.replace(tit222, '“O‘zsanoatqurilishbank” ATB va “Mikrokreditbank” ATB ning', '“O‘zsanoatqurilishbank” ATB ning', tag='22')
fnp = D.fn_para(fid=fn_tit)
D.replace(fnp, '; “Mikrokreditbank” ATB ning Markaziy bankka taqdim etilgan prudensial hisobot shakllari (crs005C, crs005L, crs005O): bank balansi, daromadlar, regulyativ kapital hisob-kitobi, aktivlarni tasniflash va prudensial me’yorlar jadvallari, 2020–2025-yillar 30-dekabr holatiga', '', tag='22')
D.replace(fnp, 'quyidagi manbalar ma’lumotlari', 'quyidagi manba ma’lumotlari', tag='22')

F_MKB_SRC = J.SRC + T.F_MKB
F_ALQ_FORM = '“Aloqabank” ATB ning Markaziy bankka taqdim etilgan prudensial hisobot shakllari (crs031C, crs031L, crs031O): bank balansi, daromadlar, regulyativ kapital hisob-kitobi, aktivlarni tasniflash, likvidlilik tahlili va prudensial me’yorlar jadvallari, 2020–2025-yillar 31-dekabr holatiga.'


def block(after_el, num, title, tbl_el, col, src, data):
    """jadval raqami, nomi (snoska bilan), jadval, izoh – after_el dan keyin; oxirgi elementni qaytaradi."""
    lab = copy.deepcopy(lab222._p)
    for r in [r for r in lab if r.tag == qn('w:r')]:
        lab.remove(r)
    rr = [r for r in lab222._p if r.tag == qn('w:r') and r.find(qn('w:t')) is not None][0]
    color(col)
    for e in D._build(rr, num + '-jadval'):
        lab.append(e)
    after_el.addnext(lab)
    tit = copy.deepcopy(tit222._p)
    tr0 = [r for r in tit if r.tag == qn('w:r') and r.find(qn('w:t')) is not None][0]
    for r in [r for r in tit if r.tag == qn('w:r')]:
        tit.remove(r)
    for e in D._build(tr0, title + f'[[FN:{src}]]'):
        tit.append(e)
    lab.addnext(tit)
    tit.addnext(tbl_el)
    t = Table(tbl_el, tit222._parent)
    for i, (nm, vals) in enumerate(zip(J.ROWS, data), 1):
        if col == BLUE:
            D.set_cell(t, i, 0, nm, tag='22')
        for j, v in enumerate(vals, 1):
            s = f'{v:.1f}'.replace('.', ',').replace('-', '−')
            if col == BLUE or t.rows[i].cells[j].text.strip() != s:
                D.set_cell(t, i, j, s, tag='22')
    iz = copy.deepcopy(izoh._p)
    if col == BLUE:
        ir = [r for r in iz if r.tag == qn('w:r') and r.find(qn('w:t')) is not None][0]
        for r in [r for r in iz if r.tag == qn('w:r')]:
            iz.remove(r)
        for e in D._build(ir, J.IZOH):
            iz.append(e)
    tbl_el.addnext(iz)
    color(RED)
    return iz


a = new_par(izoh._p, T.SQB1)
a = new_par(a, T.SQB2)
a = block(a, '2.2.3', '“Mikrokreditbank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari', tbl_mk, RED, F_MKB_SRC, J.MKB)
a = new_par(a, T.MKB1)
a = new_par(a, T.MKB2)
a = block(a, '2.2.4', '“Aloqabank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari', tbl_al, BLUE, J.SRC + F_ALQ_FORM, J.ALQ)
alq = [(x[0], x[1], F_ALQ_FORM if x[2] == T.F_ALQ else x[2]) for x in T.ALQ1]
a = new_par(a, alq)
a = new_par(a, T.ALQ2)
a = new_par(a, T.QIY)

# ---------- 4) takroriy rasmlar va ularning matni olib tashlanadi
for cap in (cap5, cap6, cap8):
    ch = cap._p.getprevious(); assert list(ch.iter(CHART))
    ch.getparent().remove(ch)
    cap._p.getparent().remove(cap._p)
for p in (p_old1, p_teb, p_akt, p_akt2, p_roe0, p_roe, p_lik):
    p._p.getparent().remove(p._p)

# ---------- 5) 2.2.7 → 2.2.5-rasm va uning tahlili
D.replace(cap7, '2.2.7-rasm', '2.2.5-rasm', tag='22')
D.replace(p_mhxs, 'ikkala bankda ham MHXS 9', 'uchala bankda ham MHXS 9', tag='22')
D.replace(p_mhxs, 'Rasmda taqqoslash uchun keltirilgan uchinchi davlat banki – “Aloqabank” ATB da ham', '“Aloqabank” ATB da ham', tag='22')
D.replace(p_mhxs, 'har ikki bankda ham', 'uchala bankda ham', tag='22')

# ---------- 6) umumlashtiruvchi jadval → 2.2.5-jadval, “Aloqabank” ustuni
D.replace(p_sum, 'bank tizimi va ikki bankning', 'bank tizimi va uch bankning', tag='22')
D.replace(p_sum, '(2.2.3-jadval)', '(2.2.5-jadval)', tag='22')
D.replace(lab223, '2.2.3-jadval', '2.2.5-jadval', tag='22')
color(BLUE)
D.add_col(tbl_sum, 3, ['“Aloqabank”'] + [r[4] for r in V.SUMROWS], tag='22')
color(RED)
for i, r in enumerate(V.SUMROWS, 1):
    if tbl_sum.rows[i].cells[6].text.strip() != r[6]:
        D.set_cell(tbl_sum, i, 6, r[6], tag='22')
color(BLUE)
D.insert(D.fn_para(fid=fn_sum), '', '; ' + F_ALQ_FORM[:-1] + '; Joint Stock Commercial “Aloqabank” and its subsidiaries. Consolidated Financial Statements and Independent Auditor’s Report for the year ended 31 December 2025. – Tashkent: PricewaterhouseCoopers, 2026.', tag='22')
D.insert(p_xul, 'kapital yetarliligining davlat qayta kapitallashtirishiga bog‘liqligi aniqlandi.',
         ' Tizimli ahamiyatga molik bo‘lmagan “Aloqabank” ATB misolida esa bu qonuniyatlar faqat yirik davlat banklariga xos emasligi tasdiqlandi.', tag='22')
color(RED)

# ---------- 7) 2.2.3 va 2.2.4-rasmlar: chiziqli, mezon chiziqlari bilan
for cap, series, yt, mn, mx, mj, pos in ((cap3, V.C223, 'Kreditlardagi ulushi, %', 0, 25, 5, ('t', '', 'b', '', None)),
                                         (cap4, V.C224, 'Rentabellik, %', 0, 16, 4, ('t', 't', None))):
    ch = cap._p.getprevious()
    rid = [e.get(qn('r:id')) for e in ch.iter(CHART)][0]
    part = D.d.part.related_parts[rid]
    xrel = [r for r in part.rels.values() if r.reltype.endswith('/package')][0]
    xml, rows = V.line_chart(series, yt, mn, mx, mj, pos)
    part._blob = re.sub(r'(<c:externalData r:id=")[^"]+(")', r'\g<1>' + xrel.rId + r'\g<2>', xml).encode('utf-8')
    xrel.target_part._blob = W._xlsx(rows)
    for ext in ch.iter('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent'):
        ext.set('cy', str(int(8.5 * 360000)))
D.replace(p_after3, 'Rasmiy ko‘rsatkich va MHXS 9',
          'Rasmda ko‘rinadiki, rasmiy muammoli kreditlar ulushi 5 foizlik erta ogohlantirish chegarasidan past bo‘lsa-da, substandart kreditlar bilan birgalikda bu ko‘rsatkich chegaradan 2,5–4,2 barobar yuqori. Rasmiy ko‘rsatkich va MHXS 9', tag='22')
D.insert(p_after4, 'bu mezonga mos keladi.', ' Rasmda bu mezon uzuq chiziq bilan ko‘rsatilgan: aktivlar rentabelligi 2021 va 2024-yillarda (1,3 va 1,4 foiz) unga yaqinlashgan.', tag='22')

# ---------- 8) havolasi qolmagan snoskalarni o‘chirish
refs = {r.get(qn('w:id')) for r in body.iter(qn('w:footnoteReference'))}
for f in list(D.fx.findall(qn('w:footnote'))):
    if f.get(qn('w:type')) is None and f.get(qn('w:id')) not in refs:
        D.fx.remove(f)
D.save(F)
print(pardoz(F))

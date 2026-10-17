"""1.3.3-rasmning yangi varianti (namuna): Yevropa Ittifoqi yirik banklarida muammolarni bashorat qilish
modellarining aniqligi, tanlanmadan tashqari davr 2022-yil III chorak – 2023-yil III chorak.
Manba: Malikkidou D., Strohbach W. EBA Staff Paper No. 21 (2025), Table 7."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_1.3.3-rasm_yangi.docx'
CATS = ['Logit modeli', 'Tasodifiy o‘rmon\n(random forest)', 'Qaror daraxti', 'Neyron tarmoq']
SENS = [71, 86, 71, 79]       # Table 7: Sensitivity 0,7143; 0,8571; 0,7143; 0,7857
SPEC = [84, 89, 89, 79]       # Table 7: Specificity 0,8438; 0,8894; 0,8894; 0,7885
SER = [('Zaif banklar to‘g‘ri aniqlangan', SENS, 'EB6834'), ('Barqaror banklar to‘g‘ri aniqlangan', SPEC, '2A78D6')]
xml, rows = W.bar_chart(CATS, SER, ytitle='To‘g‘ri tasniflash ulushi, %', ymax=100, ymin=0, major=20, label_fmt='0', gap=80)

SRC = 'Malikkidou D., Strohbach W. Predicting Bank Distress in Europe: Using Machine Learning and a Novel Definition of Distress: EBA Staff Paper Series No. 21. – Paris: European Banking Authority, January 2025.'


def texts(fn):
    s_fims = ('O.I. Lavrushin va I.D. Mamonova keltirgan ma’lumotlarga ko‘ra, FIMS modeli avvalgi UBSS tizimiga nisbatan to‘g‘ri bashorat ulushini 50 foizdan 74,6 foizga oshirgan. Bugungi kunda bunday modellar mashinali o‘qitish usullari asosida rivojlanmoqda va ularning imkoniyatlarini Yevropa bank nazorati organi (EBA) mutaxassislarining 2025-yilgi tadqiqoti yaqqol ko‘rsatadi (1.3.3-rasm).')
    caption = 'Yevropa Ittifoqi yirik banklarida zaif banklarni bashorat qilish modellarining aniqligi, 2022-yil III chorak – 2023-yil III chorak, %'
    cap_fn = ('Tadqiqot natijasida muallif tomonidan quyidagi manba ma’lumotlari asosida ishlab chiqildi: ' + SRC[:-1] + ' – P. 27 (Table 7). Izoh: modellar 2017-yil I chorak – 2022-yil II chorak ma’lumotlarida baholangan va undan keyingi davrda, ya’ni baholashda foydalanilmagan ma’lumotlarda tekshirilgan.')
    a = ('Tadqiqotda Yevropa iqtisodiy hududining 27 ta mamlakatidagi 176 ta yirik bankning 2017–2023-yillardagi nazorat hisobotlari asosida asosiy kapital (CET1), leverej koeffitsiyenti yoki likvidlikni qoplash me’yorining erta ogohlantirish chegarasidan pastga tushishi “zaif bank” holati sifatida qabul qilingan'
         + fn(SRC[:-1] + ' – P. 5–6.')
         + '. Rasm ma’lumotlari shuni ko‘rsatadiki, modellarni baholashda foydalanilmagan keyingi davrda eng yaxshi natijani tasodifiy o‘rmon usuli bergan: u zaif banklarning 86 foizini, barqaror banklarning 89 foizini to‘g‘ri aniqlagan (AUC 0,93). An’anaviy logit modeli zaif banklarning 71 foizini, barqaror banklarning 84 foizini to‘g‘ri aniqlagan (AUC 0,91), ya’ni oddiy statistik model ham yuqori bashorat kuchiga ega. Mualliflar xulosasiga ko‘ra, bashoratda eng muhim omillar rentabellik va to‘lov qobiliyatiga oid ko‘rsatkichlar bo‘lib, tanlanmani muvozanatlash usullari zaif banklarni aniqlash ulushini oshirsa-da, barqaror banklarni noto‘g‘ri belgilash ehtimolini ham oshiradi'
         + fn(SRC[:-1] + ' – P. 1, 23–27.')
         + '. Bu natija integral baholash metodikasida chegaraviy qiymatlarni tanlash xatolar narxini hisobga olgan holda amalga oshirilishi kerakligini ko‘rsatadi.')
    return dict(s_fims=s_fims, caption=caption, cap_fn=cap_fn, a=a)


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=14, color=None):
        q = dd.add_paragraph(); q.alignment = align
        q.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
        r = q.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)

    FN = []

    def fn(t):
        FN.append(t); return f'[{len(FN)}]'

    FN.append(texts(lambda x: '')['cap_fn'])
    T = texts(fn)
    par('1.3.3-rasmning yangi varianti (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    par('Ma’qul bo‘lsa: 1.3-§ dagi “O.I. Lavrushin va I.D. Mamonova keltirgan ma’lumotlarga ko‘ra, 48 306 kuzatuv asosida … (1.3.3-rasm).” jumlasi quyidagi jumlalar bilan, 1.3.3-rasm va undan keyingi “Rasm ma’lumotlari FIMS modeli …” xatboshi esa quyidagi rasm va tahlil bilan almashtiriladi.',
        italic=True, size=12, indent=False, color='595959')
    par('… ' + T['s_fims'])
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, (xml, rows))
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.5, '1.3.3-rasm')))
    par('1.3.3-rasm. ' + T['caption'] + '[1]', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    par(T['a'])
    par('Sahifa osti snoskalari', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    for i, t in enumerate(FN, 1):
        par(f'[{i}] {t}', size=11, indent=False)
    dd.save(DST)
    print(pardoz(DST), DST)

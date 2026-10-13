"""3.1.2-jadval: chap tomonda raqamlash ustuni; chegaralarni asoslovchi yangi xatboshi (qizil)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.oxml.ns import qn
import kamlib as K
import namuna_312j as N
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
D = K.Doc(F)
p = [p for p in D.d.paragraphs if p.text.strip() == '3.1.2-jadval'][0]
e = p._p.getnext()
while not e.tag.endswith('}tbl'):
    e = e.getnext()
t = docx.table.Table(e, D.d)
n = len(t.rows)
vals = ['№'] + [N.NUM[i] for i in range(1, n)]
D.add_col(t, 0, vals, tag='312j')
# yangi ustunni (1) chapga ko‘chirish
for tr in t._tbl.findall(qn('w:tr')):
    tcs = tr.findall(qn('w:tc'))
    tcs[0].addprevious(tcs[1])
grid = t._tbl.find(qn('w:tblGrid'))
gc = grid.findall(qn('w:gridCol'))
gc[0].addprevious(gc[1])
gc = grid.findall(qn('w:gridCol'))
w0, w1 = int(gc[0].get(qn('w:w'))), int(gc[1].get(qn('w:w')))
NW = 600
gc[1].set(qn('w:w'), str(w0 + w1 - NW)); gc[0].set(qn('w:w'), str(NW))
ws = [int(g.get(qn('w:w'))) for g in gc]
for tr in t._tbl.findall(qn('w:tr')):
    for tc, w in zip(tr.findall(qn('w:tc')), ws):
        tcw = tc.find(qn('w:tcPr') + '/' + qn('w:tcW'))
        if tcw is not None:
            tcw.set(qn('w:w'), str(w)); tcw.set(qn('w:type'), 'dxa')
# raqam kataklari: markazda; blok qatorlarida qalin
for i, r in enumerate(t.rows):
    c0 = r.cells[0].paragraphs[0]
    jc = c0._p.get_or_add_pPr().find(qn('w:jc'))
    if jc is None:
        jc = c0._p.get_or_add_pPr().makeelement(qn('w:jc'), {}); c0._p.get_or_add_pPr().append(jc)
    jc.set(qn('w:val'), 'center')
    blok = i > 0 and N.NUM[i].endswith('.')
    for rr in c0.runs:
        rr.bold = True if (i == 0 or blok) else False
# yangi xatboshi
anchor = [q for q in D.d.paragraphs if q.text.startswith('Aktivlar sifati blokida 3-bosqich kreditlari ulushi uchun 5 foiz')][0]
txt = ''.join('[[FN:' + N.FN + ']]' if s == '[[FN]]' else s for s in N.ASOS)
D.new_para_after(anchor, txt, tag='312j')
D.save(F)
print(pardoz(F))
# Keyin: ustun kengliklari [600, qolgan, 1450, 800, 800, 800, 800] dxa qilib to‘g‘rilandi (Standart ustuni sig‘ishi uchun).

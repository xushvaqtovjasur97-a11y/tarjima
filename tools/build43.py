import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/4.3-paragraf.md'; DST='/home/user/tarjima/paragraflar/4.3-paragraf.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p43_raw.docx'],check=True)
d=docx.Document(S+'/p43_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^4\.3\.\d-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^4\.3\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^4\.3\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'
PURPLE='7030A0'

def setw(t,ws):
  for row in t.rows:
    for k,c in enumerate(row.cells): c.width=Cm(ws[k])
  for k,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(ws[k]*567)))
  t.autofit=False
for t in d.tables:
  h=[c.text.strip() for c in t.rows[0].cells]; n=len(h)
  if h[0]=='Yo‘nalish': setw(t,[2.6,4.6,4.4,4.8])
  elif h[0]=='Bosqich': setw(t,[2.2,5.6,2.4,2.4,3.8])
  elif h[0]=='Bank': setw(t,[3.4,1.8,1.8,1.9,1.8,1.8,1.9,2.0])
  elif h[0]=='Jihat': setw(t,[3.0,6.6,6.8])
  for row in t.rows:
    for c in row.cells:
      for p in c.paragraphs:
        for r in p.runs: r.font.size=Pt(10)
# "Izoh:" paragraflari – 12 pt, xatboshisiz
for p in d.paragraphs:
  if p.text.startswith('Izoh:'):
    p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(6)
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.font.size=Pt(12)
  if p.text.startswith('MBII') and '=' in p.text:
    p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0)
    for r in p.runs: r.italic=True; r.font.size=Pt(13)
sh=[]
cols=[('1. Ma’lumotlar',['Milliy prudensial hisobot (choraklik)','MHXS hisoboti (yillik)','Yashirin kredit riski koeffitsiyenti bilan tuzatish'],DBLUE),
      ('2. Baholash',['Integral indeks: 5 blok, 13 ko‘rsatkich','Davlat ko‘magisiz indeks (davlat banklari)','MHXS va milliy indekslar farqi (ΔMBII)'],BLUE),
      ('3. Signal',['Erta ogohlantirish zonalari: yashil ≥ 3,25; sariq 2,75–3,25; qizil < 2,75','Stress-test natijalari zonalarda'],'C55A11'),
      ('4. Choralar',['Yashil: odatiy monitoring','Sariq: kuchaytirilgan monitoring, tuzatish rejasi','Qizil: maqsadli tekshiruv, kapitalni tiklash, dividend cheklovi'],'7F6000')]
bw,gx=3.75,0.42; x0=0.1
for ci,(t,items,c) in enumerate(cols):
  x=x0+ci*(bw+gx)
  sh.append(W.box(x,0.1,bw,0.75,t,fill=c,line=c,sz=20,bold=True,color='FFFFFF'))
  y=1.05
  for it in items:
    hgt=1.25 if len(it)>40 else 1.0
    sh.append(W.box(x,y,bw,hgt,it,fill='F2F2F2',line=c,sz=17))
    y+=hgt+0.15
  if ci<3: sh.append(W.line(x+bw+0.03,0.47,x+bw+gx-0.03,0.47,color='404040',lw=19050))
yb=5.55
sh.append(W.box(x0,yb,4*bw+3*gx,0.7,'Keyingi chorakda indeks qayta hisoblanadi: choralar samaradorligi baholanadi (yopiq sikl)',fill='FFFFFF',line='404040',sz=17,italic=True,geom='rect'))
xr=x0+3*(bw+gx)+bw/2; xl=x0+bw/2
sh.append(W.line(xr,5.25,xr,yb-0.02,color='404040',lw=15875))
sh.append(W.line(xl,yb-0.02,xl,5.25,color='404040',lw=15875))
W.replace_placeholder(d,'{{FIG:1}}',W.group_run_xml(sh,16.6,6.4,'4.3.1-rasm'))
d.save(DST); print('saved')

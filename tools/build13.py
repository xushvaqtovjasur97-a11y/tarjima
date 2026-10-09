import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/1.3-paragraf.md'; DST='/home/user/tarjima/paragraflar/1.3-paragraf.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p13_raw.docx'],check=True)
d=docx.Document(S+'/p13_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^1\.3\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^1\.3\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if t.startswith(('CAMELS tizimida','XVJ moliyaviy barqarorlik ko‘rsatkichlarining asosiy','Banklar moliyaviy barqarorligini baholash bo‘yicha xalqaro')):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'

# ---- 1.3.1 scheme
sh=[]
heads=[('Bazel I\n(1988)',BLUE),('Bozor riski tuzatishi\n(1996)',LBLUE),('Bazel II\n(2004)',ORANGE),('Bazel III\n(2010)',GREEN)]
dets=['Faqat kredit riski;\nTier 1 ≥ 4%;\njami kapital ≥ 8%;\nrisk vaznlari\n0, 20, 50, 100%',
      'Bozor riski uchun\nkapital;\nVaR modeli;\n10 kunlik gorizont;\n×3 koeffitsiyent',
      'Uch ustun;\nkredit, bozor va\noperatsion risklar;\nIRB yondashuvi\n(PD, LGD, EAD, M)',
      'CET1 ≥ 4,5%;\nkonservatsiya va\nkontrsiklik buferlar;\nleverej ≥ 3%;\nLCR, NSFR ≥ 100%']
for i,((h,c),dt) in enumerate(zip(heads,dets)):
  x=0.1+i*4.25
  sh.append(W.box(x,0.2,3.5,1.3,h,fill=c,line=c,sz=22,bold=True,color='FFFFFF' if c!=LBLUE else '000000'))
  sh.append(W.line(x+1.75,1.5,x+1.75,2.0,arrow=False,color='7F7F7F'))
  sh.append(W.box(x,2.0,3.5,4.4,dt,fill='F2F2F2',line='A6A6A6',sz=22))
  if i<3: sh.append(W.line(x+3.52,0.85,x+4.23,0.85,color='404040',lw=19050))
W.replace_placeholder(d,'{{FIG:1.3.1}}',W.group_run_xml(sh,16.4,6.6,'1.3.1-rasm'))

# ---- 1.3.2 stacked bar
rid=W.add_chart(d,W.bar_chart(['Bazel II','Bazel III'],[
  ('Asosiy kapital (CET1)',[2,4.5],DBLUE,'FFFFFF'),('Qo‘shimcha birinchi darajali kapital',[2,1.5],LBLUE),
  ('Ikkinchi darajali kapital',[4,2],GREY),('Konservatsiya buferi',[None,2.5],ORANGE),('Kontrsiklik bufer (0–2,5%)',[None,2.5],GOLD)],
  ytitle='Risk bo‘yicha tortilgan aktivlarga nisbatan, %',stacked=True,ymax=14,major=2,gap=90))
W.replace_placeholder(d,'{{FIG:1.3.2}}',W.chart_run_xml(rid,15.5,8.5,'1.3.2-rasm'))

# ---- 1.3.3 step
xs=[4.5,5.125,5.125,5.75,5.75,6.375,6.375,7.0,7.0,7.625]; ys=[100,100,80,80,60,60,40,40,0,0]
rid=W.add_chart(d,W.scatter_chart([('Taqsimlanmaydigan foyda ulushi',xs,ys,BLUE,34925,None,None),
   ('Pog‘ona chegaralari',[4.5,5.125,5.75,6.375,7.0],[100,80,60,40,0],ORANGE,0,None,'circle')],
   'Asosiy kapital (CET1) yetarliligi, %','Taqsimlanmasdan qoldiriladigan foyda ulushi, %',4.5,7.625,0.625,0,110,20,xfmt='0.0##',legend=True))
W.replace_placeholder(d,'{{FIG:1.3.3}}',W.chart_run_xml(rid,15.5,7.5,'1.3.3-rasm'))

# ---- 1.3.4 CCyB
rid=W.add_chart(d,W.scatter_chart([('Kontrsiklik bufer',[-4,2,10,14],[0,0,2.5,2.5],ORANGE,34925,None,None),
   ('Chegaraviy nuqtalar (L = 2; H = 10)',[2,10],[0,2.5],BLUE,0,None,'circle')],
   'Kredit gepi (kredit/YaIM nisbatining trenddan og‘ishi), foiz punkti','Kontrsiklik bufer, RWA ga nisbatan %',-4,14,2,0,3,0.5,yfmt='0.0',legend=True))
W.replace_placeholder(d,'{{FIG:1.3.4}}',W.chart_run_xml(rid,15.5,7.5,'1.3.4-rasm'))

# ---- 1.3.5 IFRS9 scheme
sh=[]
sh.append(W.box(0.1,0.0,16.2,0.55,'Dastlabki tan olingandan keyin kredit sifatining yomonlashuvi',fill=None,line=None,sz=22,italic=True))
sh.append(W.line(0.2,0.75,16.2,0.75,color='7F7F7F',lw=31750))
cols=[('1-bosqich',GREEN,'Kredit riski sezilarli\noshmagan aktivlar','12 oylik kutilgan\nkredit zararlari','Foiz daromadi yalpi\nbalans qiymatiga\nhisoblanadi'),
      ('2-bosqich',GOLD,'Kredit riski sezilarli\noshgan aktivlar','Butun muddat uchun\nkutilgan kredit zararlari','Foiz daromadi yalpi\nbalans qiymatiga\nhisoblanadi'),
      ('3-bosqich',ORANGE,'Qadrsizlangan\n(defolt) aktivlar','Butun muddat uchun\nkutilgan kredit zararlari','Foiz daromadi sof\n(zaxiradan keyingi)\nqiymatga hisoblanadi')]
for i,(h,c,a,b_,e) in enumerate(cols):
  x=0.1+i*5.5
  sh.append(W.box(x,1.05,5.1,0.85,h,fill=c,line=c,sz=24,bold=True,color='000000' if c==GOLD else 'FFFFFF'))
  sh.append(W.box(x,2.1,5.1,1.35,a,fill='F2F2F2',line='A6A6A6',sz=22))
  sh.append(W.box(x,3.65,5.1,1.35,b_,fill='DEEBF7',line='9DC3E6',sz=22,bold=True))
  sh.append(W.box(x,5.2,5.1,1.5,e,fill='F2F2F2',line='A6A6A6',sz=22))
W.replace_placeholder(d,'{{FIG:1.3.5}}',W.group_run_xml(sh,16.4,6.8,'1.3.5-rasm'))

# ---- 1.3.6 bar + exp trend
xml,rows=W.bar_chart(['AAA','AA','A','BBB','BB','B','CCC/C'],[('5 yillik kumulyativ defolt darajasi, %',[0.34,0.28,0.39,1.36,5.75,15.60,46.53],BLUE)],
   ytitle='Defolt darajasi, %',xtitle='S&P reyting toifasi',ymax=50,major=10,label_fmt='0.00',legend=True,gap=60)
trend=('<c:trendline><c:name>Eksponensial trend</c:name><c:spPr><a:ln w="22225"><a:solidFill><a:srgbClr val="EB6834"/></a:solidFill>'
       '<a:prstDash val="dash"/></a:ln></c:spPr><c:trendlineType val="exp"/><c:dispRSqr val="0"/><c:dispEq val="0"/></c:trendline>')
xml=xml.replace('<c:cat>',trend+'<c:cat>',1)
rid=W.add_chart(d,(xml,rows))
W.replace_placeholder(d,'{{FIG:1.3.6}}',W.chart_run_xml(rid,15.5,7.5,'1.3.6-rasm'))

# ---- 1.3.7 mapping scheme
sh=[]
sh.append(W.box(0.1,0.0,3.6,0.6,'Xalqaro standartlar',fill=None,line=None,sz=22,bold=True))
sh.append(W.box(6.0,0.0,4.8,0.6,'Baholash bloklari',fill=None,line=None,sz=22,bold=True))
src=[('Bazel III',GREEN,'FFFFFF'),('MHXS 9',ORANGE,'FFFFFF'),('XVJ FSI',LBLUE,'000000'),('CAMELS',BLUE,'FFFFFF'),('Reyting agentliklari',GREY,'000000')]
blk=['Kapital yetarliligi','Aktivlar sifati','Daromadlilik','Likvidlik','Bozor (valyuta) riskiga\nsezgirlik']
links={0:[0,3],1:[1],2:[0,1,2,3,4],3:[0,1,2,3,4],4:[0,1,2,3]}
ys=[0.8+i*1.65 for i in range(5)]; hh=1.1
cl=[GREEN,ORANGE,'5B9BD5',BLUE,'7F7F7F']
for i,js in links.items():
  for j in js: sh.append(W.line(3.7,ys[i]+hh/2,6.0,ys[j]+hh/2,color=cl[i],lw=12700,arrow=False))
for i,(t,c,tc) in enumerate(src): sh.append(W.box(0.1,ys[i],3.6,hh,t,fill=c,line=c,sz=22,bold=True,color=tc))
for j,t in enumerate(blk): sh.append(W.box(6.0,ys[j],4.8,hh,t,fill='DEEBF7',line='5B9BD5',sz=22))
for j in range(5): sh.append(W.line(10.8,ys[j]+hh/2,12.6,4.55,color='404040',lw=15875))
sh.append(W.box(12.6,3.15,3.8,2.8,'Moliyaviy\nbarqarorlikning\nintegral bahosi',fill=DBLUE,line=DBLUE,sz=24,bold=True,color='FFFFFF'))
W.replace_placeholder(d,'{{FIG:1.3.7}}',W.group_run_xml(sh,16.5,8.5,'1.3.7-rasm'))

d.save(DST); print('saved')

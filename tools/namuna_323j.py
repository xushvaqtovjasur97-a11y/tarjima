"""3.2.3-jadval: AUC oralig‘i, 'deyarli bir xil', logit izohi, umumiy izoh, ustun sarlavhasi (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_3.2.3-jadval.docx'
HDR_OLD, HDR_NEW = 'Ko‘rsatkich', 'Bank / ko‘rsatkich'
R1 = ('biroz ustun (AUC 0,864–0,907)', 'biroz ustun (AUC 0,865–0,907, tayanch vaznlarda – 0,864)')
R2 = ('bosh komponentlar usuli bilan bir xil', 'bosh komponentlar usuli bilan deyarli bir xil')
ANC = 'ya’ni vaznlarni empirik usulda almashtirish bashorat kuchini sezilarli oshirmaydi.'
INS = (' Logit usulining nisbatan yuqori AUC qiymati (0,907) esa qisman uning vaznlari aynan shu panel ma’lumotlari va muammo '
       'hodisalari asosida baholangani bilan izohlanadi, shu sababli bu ustunlik tanlanmadan tashqari tekshiruvni talab qiladi.')
FN_NEW = ('Tadqiqot natijasida muallif tomonidan “O‘zsanoatqurilishbank”, “Mikrokreditbank”, “Biznesni rivojlantirish banki”, '
          '“Xalq banki” va “Ipoteka-bank” ATBlarining 2021–2025-yillardagi MHXS bo‘yicha moliyaviy hisobotlari, “SNS” reyting '
          'agentligi bazasi va Markaziy bank statistikasi asosida hisoblandi.')

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(4)
    G, R = '7F7F7F', 'FF0000'

    def par(parts, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        q = dd.add_paragraph(); q.alignment = align
        for t, col, s in parts:
            r = q.add_run(t); r.bold = bold; r.font.strike = s
            if col: r.font.color.rgb = RGBColor.from_string(col)
    par([('3.2.3-jadval (namuna)', None, False)], True, WD_ALIGN_PARAGRAPH.CENTER)
    par([('1) Jadval birinchi ustuni sarlavhasi: ', None, False), (HDR_OLD, G, True), (' → ', None, False), (HDR_NEW, R, False)])
    par([('2) Jadval nomidagi izoh (uzun ro‘yxat o‘rniga):', None, False)], True)
    par([(FN_NEW, R, False)])
    par([('3) Jadval tahlilida:', None, False)], True)
    par([('… Muammo hodisalarini bashorat qilish kuchi bo‘yicha empirik vaznlar tayanch vaznlardan ', None, False),
         (R1[0][:13], None, False), (R1[0][13:], G, True), (' ', None, False), (R1[1][13:], R, False),
         (', biroq farq katta emas, tayanch vaznlarning bashorat kuchi esa ', None, False), (R2[0], G, True), (' ', None, False), (R2[1], R, False), ('. …', None, False)])
    par([('… atigi 0,007 ga yuqori, ', None, False), (ANC, None, False), (INS, R, False)])
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST))

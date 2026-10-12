# Titul + mundarija + barcha yozilgan bo'limlarni bitta docx ga yig'ish (docxcompose; diagrammalar Excel ma'lumotlari bilan saqlanadi)
import docx, copy, re, subprocess, os
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docxcompose.composer import Composer
P='/home/user/tarjima/paragraflar/'
OUT='/home/user/tarjima/Dissertatsiya_Maxmudov_O.T.docx'
TMP='/tmp/claude-0/birlashgan_tmp.docx'
ITEMS=[('D',P+'Kirish.docx'),
('B','I BOB. TIJORAT BANKLARI MOLIYAVIY BARQARORLIGINI BAHOLASHNING NAZARIY ASOSLARI VA XALQARO STANDARTLARI'),
('D',P+'1.1-paragraf.docx'),('D',P+'1.2-paragraf.docx'),('D',P+'1.3-xorijiy-tajriba.docx'),('D',P+'I_bob_xulosalar.docx'),
('B','II BOB. O‘ZBEKISTON BANK TIZIMIDA MOLIYAVIY BARQARORLIKNI BAHOLASH AMALIYOTINING XALQARO STANDARTLARGA MUVOFIQLIGI TAHLILI'),
('D',P+'2.1-paragraf.docx'),('D',P+'2.2-paragraf.docx'),('D',P+'2.3-paragraf.docx'),('D',P+'II_bob_xulosalar.docx'),
('B','III BOB. XALQARO STANDARTLARGA ASOSLANGAN INTEGRAL BAHOLASH METODIKASI VA UNING EKONOMETRIK ASOSLANISHI'),
('D',P+'3.1-paragraf.docx'),('D',P+'3.2-paragraf.docx'),('D',P+'3.3-paragraf.docx'),('D',P+'III_bob_xulosalar.docx'),
('B','IV BOB. TIJORAT BANKLARI MOLIYAVIY BARQARORLIGINI BAHOLASH METODOLOGIYASINI TAKOMILLASHTIRISH YO‘NALISHLARI'),
('D',P+'4.1-paragraf.docx'),('D',P+'4.2-paragraf.docx'),('D',P+'4.3-paragraf.docx'),('D',P+'IV_bob_xulosalar.docx'),('D',P+'Xulosa.docx'),('D',P+'Adabiyotlar.docx')]
def heads(d):
  return [p for p in d.paragraphs if p.style.name.startswith('Heading')]
def build(toc):
  master=docx.Document('/tmp/claude-0/merge_titul.docx')
  # mundarija o'rni
  ph=[p for p in master.paragraphs if p.text=='{{TOC}}'][0]
  for lvl,text,page in toc:
    q=ph.insert_paragraph_before(); pf=q.paragraph_format
    pf.line_spacing=1.0; pf.space_after=Pt(4); pf.left_indent=Cm(0 if lvl==1 else 0.8)
    pf.tab_stops.add_tab_stop(Cm(16.5),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    r=q.add_run(f'{text}\t{page}'); r.font.size=Pt(13); r.font.name='Times New Roman'; r.bold=(lvl==1)
  ph._element.getparent().remove(ph._element)
  # sahifa raqami (titulda yo'q)
  sec=master.sections[0]; sec.different_first_page_header_footer=True
  fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
  r=fp.add_run()
  for t,v in (('begin',None),('instr',' PAGE '),('end',None)):
    if t=='instr': e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=v
    else: e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),t)
    r._r.append(e)
  r.font.name='Times New Roman'; r.font.size=Pt(12)
  comp=Composer(master); prevB=False
  for kind,x in ITEMS:
    if kind=='B':
      h=comp.doc.add_paragraph(x,style='Heading 1'); h.paragraph_format.page_break_before=True
      h.alignment=WD_ALIGN_PARAGRAPH.CENTER; prevB=True; continue
    d=docx.Document(x)
    if not prevB: d.paragraphs[0].paragraph_format.page_break_before=True
    hs=heads(d)
    comp.append(d); prevB=False
  for p in comp.doc.paragraphs:
    if p.text.strip() in ('KIRISH','XULOSA','FOYDALANILGAN ADABIYOTLAR RO‘YXATI') and p.style.name.startswith('Heading'): p.style=comp.doc.styles['Heading 1']
  comp.save(TMP)
def pages():
  subprocess.run(['soffice','--headless','--convert-to','pdf',TMP,'--outdir','/tmp/claude-0/'],check=True,capture_output=True)
  pdf=TMP.replace('.docx','.pdf'); n=int(re.search(r'Pages:\s+(\d+)',subprocess.run(['pdfinfo',pdf],capture_output=True,text=True).stdout).group(1))
  txt=[subprocess.run(['pdftotext','-f',str(i),'-l',str(i),pdf,'-'],capture_output=True,text=True).stdout for i in range(1,n+1)]
  return txt,n
def norm(s): return re.sub(r'\s+',' ',s).replace('ʻ','‘').strip()
# 1-o'tish: mundarija bandlarini yig'ish (sahifasiz), 2-o'tish: sahifa raqamlari
entries=[]
for kind,x in ITEMS:
  if kind=='B': entries.append((1,x)); continue
  hs=heads(docx.Document(x)); t=hs[0].text.strip()
  entries.append((1 if t in ('KIRISH','XULOSA','FOYDALANILGAN ADABIYOTLAR RO‘YXATI') else 2,t))
toc=[(l,t,'00') for l,t in entries]
for it in range(2):
  build(toc); txt,n=pages()
  res=[]; start=2
  for l,t in entries:
    key=norm(t)[:45]; pg=None
    for i in range(start,n):
      if key in norm(txt[i]): pg=i+1; break
    if pg is None: pg='?'
    else: start=pg-1
    res.append((l,t,pg))
  toc=res
build(toc); os.replace(TMP,OUT)
print(toc)

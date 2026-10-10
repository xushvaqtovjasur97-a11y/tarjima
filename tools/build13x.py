import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/1.3-xorijiy-tajriba.md'; DST='/home/user/tarjima/paragraflar/1.3-xorijiy-tajriba.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p13x_raw.docx'],check=True)
d=docx.Document(S+'/p13x_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^1\.3\.\d-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^1\.3\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^1\.3\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'

def wt(c): return '000000' if c in (GOLD,GREY,LBLUE) else 'FFFFFF'
# ---- 1: classification
sh=[]
sh.append(W.box(3.7,0.1,9.0,1.0,'Banklar moliyaviy barqarorligini baholash tizimlari\n(Xalqaro hisob-kitoblar banki tasnifi)',fill=DBLUE,line=DBLUE,sz=22,bold=True,color='FFFFFF'))
cols=[('Nazorat reyting\ntizimlari',BLUE,'CAMELS (AQSh)\nORAP (Fransiya)\nPATROL (Italiya)\nRossiya banki\nreytingi'),
      ('Koeffitsiyentlar va\nguruhlab tahlil',GREEN,'BAKIS (Germaniya)\nEBA risk\nindikatorlari\n(Yevropa Ittifoqi)'),
      ('Risklarni kompleks\nbaholash',ORANGE,'BAITE, RATE\n(Buyuk Britaniya)\nRAST\n(Niderlandiya)'),
      ('Statistik\nmodellar',GOLD,'FIMS, SEER,\nSCOR (AQSh)\nSAABA ekspert\ntizimi (Fransiya)')]
w=3.85; g=0.25
for i,(h,c,t) in enumerate(cols):
  x=0.1+i*(w+g)
  sh.append(W.line(8.2,1.1,x+w/2,1.75,color='404040',lw=15875))
  sh.append(W.box(x,1.75,w,1.1,h,fill=c,line=c,sz=21,bold=True,color=wt(c)))
  sh.append(W.box(x,3.0,w,2.3,t,fill='F2F2F2',line=c,sz=20))
sh.append(W.box(0.1,5.5,16.1,0.75,'Joriy holatni baholash  ←————————————————→  Kelgusi holatni bashorat qilish',fill='DEEBF7',line=LBLUE,sz=20,italic=True))
W.replace_placeholder(d,'{{FIG:1}}',W.group_run_xml(sh,16.4,6.4,'1.3.1-rasm'))

# ---- 2: PCA ladder
sh=[]
lv=[('Yaxshi kapitallashgan','Jami kapital ≥ 10%; 1-darajali ≥ 6%; leverej ≥ 5%','Cheklovlar yo‘q',GREEN),
    ('Yetarli kapitallashgan','Jami kapital ≥ 8%; 1-darajali ≥ 4%; leverej ≥ 4%','Broker depozitlarini jalb qilish cheklanadi',LBLUE),
    ('Kapitali yetarli emas','Jami kapital < 8%; 1-darajali < 4%; leverej < 4%','Kapitalni tiklash rejasi; aktivlar o‘sishi va dividendlar cheklanadi',GOLD),
    ('Kapitali sezilarli yetarli emas','Jami kapital < 6%; 1-darajali < 3%; leverej < 3%','Rahbariyat almashtirilishi, faoliyat cheklanishi mumkin',ORANGE),
    ('Kapitali keskin yetarli emas','Moddiy kapital ≤ 2%','Vaqtinchalik boshqaruv yoki tugatish','C00000')]
y=0.1; H=1.1
for i,(n,cr,ac,c) in enumerate(lv):
  x=0.1+i*0.4
  sh.append(W.box(x,y,4.6,H,n,fill=c,line=c,sz=20,bold=True,color=wt(c) if c!='C00000' else 'FFFFFF'))
  sh.append(W.box(x+4.75,y,5.4,H,cr,fill='F2F2F2',line=c,sz=19))
  sh.append(W.box(x+10.3,y,5.8-i*0.4,H,ac,fill='FFFFFF',line=c,sz=19,italic=True))
  y+=H+0.2
sh.append(W.line(16.45,0.2,16.45,y-0.3,color='C00000',lw=28575))
W.replace_placeholder(d,'{{FIG:2}}',W.group_run_xml(sh,16.6,y,'1.3.2-rasm'))

# ---- 3: FIMS vs UBSS
rid=W.add_chart(d,W.bar_chart(['To‘g‘ri bashorat','Birinchi tur xatosi','Ikkinchi tur xatosi'],[
  ('UBSS tizimi',[50,32.7,12.2],GREY),('FIMS modeli',[74.6,17,7.4],BLUE)],
  ytitle='Kuzatuvlar ulushi, %',ymax=80,major=20,label_fmt='0.0',gap=70))
W.replace_placeholder(d,'{{FIG:3}}',W.chart_run_xml(rid,15.5,8.0,'1.3.3-rasm'))

# ---- 4: BAITE matrix
sh=[]
sh.append(W.box(0.1,0.95,2.3,0.55,'Risk darajasi',fill=DBLUE,line=DBLUE,sz=19,bold=True,color='FFFFFF'))
sh.append(W.box(2.5,0.1,13.7,0.7,'Ichki nazorat sifati',fill=DBLUE,line=DBLUE,sz=20,bold=True,color='FFFFFF'))
sh.append(W.box(2.5,0.95,6.8,0.55,'Yuqori',fill='DEEBF7',line=LBLUE,sz=20,bold=True))
sh.append(W.box(9.4,0.95,6.8,0.55,'Past',fill='DEEBF7',line=LBLUE,sz=20,bold=True))
cells=[('Past',[('A','Tekshiruv davri:\n18–24 oy',GREEN),('B','Tekshiruv davri:\n12–18 oy',GOLD)]),
       ('Yuqori',[('C','Tekshiruv davri:\n6–12 oy',GOLD),('D','Tekshiruv davri:\ntaxminan 1 yil va\nqo‘shimcha choralar',ORANGE)])]
for r,(lab,row) in enumerate(cells):
  y=1.6+r*2.55
  for c,(k,t,col) in enumerate(row):
    x=2.5+c*6.9
    sh.append(W.box(x,y,6.8,2.45,k+'\n'+t,fill=col,line=col,sz=22,bold=False,color=wt(col)))
  sh.append(W.box(0.1,y,2.3,2.45,lab,fill='DEEBF7',line=LBLUE,sz=19,bold=True))
W.replace_placeholder(d,'{{FIG:4}}',W.group_run_xml(sh,16.4,6.8,'1.3.4-rasm'))

# ---- 5: LCR phase-in
rid=W.add_chart(d,W.scatter_chart([
  ('Bazel qo‘mitasi',[2015,2016,2017,2018,2019],[60,70,80,90,100],BLUE,28575,None,'circle'),
  ('Rossiya (tizimli banklar)',[2016,2017,2018,2019],[70,80,90,100],ORANGE,28575,None,'square'),
  ('O‘zbekiston',[2016,2017,2018],[80,90,100],GREEN,28575,None,'triangle')],
  xtitle='Yillar',ytitle='Minimal talab, %',xmin=2014,xmax=2020,xfmt='0',xmajor=1,ymin=50,ymax=105,ymajor=10,legend=True))
W.replace_placeholder(d,'{{FIG:5}}',W.chart_run_xml(rid,15.5,8.0,'1.3.5-rasm'))

# ---- 6: stress test architecture
sh=[]
st=[('Ssenariylar','Bazaviy, salbiy,\nkeskin salbiy',DBLUE),
    ('Makroiqtisodiy\nmodel','YaIM, valyuta kursi,\ninflyatsiya, daromadlar,\ninvestitsiyalar',BLUE),
    ('Bank\nko‘rsatkichlari','Muammoli kreditlar\nulushi, foiz stavkalari,\nzararlar',GREEN),
    ('Balans imitatsion\nmodeli','Kapital va likvidlik\nprognozi (9 chorak)',ORANGE),
    ('Natija','Kapital va likvidlik\ntanqisligi; nazorat\nchoralari',GOLD)]
w=2.95; g=0.35
for i,(h,t,c) in enumerate(st):
  x=0.1+i*(w+g)
  sh.append(W.box(x,0.1,w,1.1,h,fill=c,line=c,sz=20,bold=True,color=wt(c)))
  sh.append(W.box(x,1.35,w,2.0,t,fill='F2F2F2',line=c,sz=18))
  if i<4: sh.append(W.line(x+w,0.65,x+w+g,0.65,color='404040',lw=19050))
sh.append(W.line(14.3,3.35,14.3,3.9,color='7F7F7F',lw=12700,arrow=False))
sh.append(W.line(14.3,3.9,1.6,3.9,color='7F7F7F',lw=12700,arrow=False,dash='dash'))
sh.append(W.line(1.6,3.9,1.6,3.35,color='7F7F7F',lw=12700,dash='dash'))
sh.append(W.box(4.0,3.6,8.4,0.6,'Teskari aloqa: kredit taklifining iqtisodiyotga ta’siri',fill='FFFFFF',line='FFFFFF',sz=18,italic=True))
W.replace_placeholder(d,'{{FIG:6}}',W.group_run_xml(sh,16.4,4.3,'1.3.6-rasm'))

# ---- 7: China NPL
rid=W.add_chart(d,W.bar_chart(['2023-y. oxiri','2024-y. oxiri','2025-y. III chorak','Ilmiy baho:\nquyi chegara','Ilmiy baho:\nyuqori chegara','XVJ: risk ostidagi\nkreditlar (2016)'],[
  ('Muammoli kreditlar',[1.59,1.50,1.52,None,None,None],BLUE,'FFFFFF'),
  ('Alohida e’tibor talab qiluvchi kreditlar',[2.20,2.22,2.20,None,None,None],GOLD),
  ('Muqobil baholar',[None,None,None,3.0,4.6,15.5],ORANGE)],
  ytitle='Jami kreditlarga nisbatan, %',stacked=True,ymax=16,major=4,label_fmt='0.00',gap=50))
W.replace_placeholder(d,'{{FIG:7}}',W.chart_run_xml(rid,15.5,7.2,'1.3.7-rasm'))
d.save(DST); print('saved')

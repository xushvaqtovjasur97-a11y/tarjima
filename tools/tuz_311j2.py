"""3.1.1-jadval: dalillar tahlil qilingan banklar bo‘yicha umumlashtirildi; kirish gapi (qizil)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
import kamlib as K
import namuna_311j2 as N
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
D = K.Doc(F)
intro = [p for p in D.d.paragraphs if N.INTRO_OLD in p.text and '3.1.1-jadval' in p.text][0]
D.insert(intro, 'ikkinchi bobda ', 'bank tizimi darajasida va tahlil qilingan uchta davlat banki misolida ', tag='311j2')
cap = [p for p in D.d.paragraphs if p.text.startswith('Amaldagi baholash tizimi muammolari')][0]
t = docx.table.Table(cap._p.getnext(), D.d)
for i, new in N.FIX.items():
    old = t.cell(i, 1).text
    pre = old[:7]                      # '2.3-§: ' yoki '2.2-§: '
    assert pre == new[:7], (i, old)
    if i == 2:
        D.insert(t.cell(i, 1).paragraphs[0], pre, 'tahlil qilingan banklarda ', tag='311j2')
    else:
        D.cell_replace(t, i, 1, old[7:], new[7:], tag='311j2')
D.save(F)
print(pardoz(F))

"""3.1.1-jadval: dalillar tahlil qilingan uch bank bo‘yicha umumlashtirilgan (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_3.1.1-jadval_umumiy.docx'
FIX = {2: '2.3-§: tahlil qilingan banklarda yashirin kredit riski koeffitsiyenti 1,65–4,90',
       3: '2.3-§: tahlil qilingan banklarda MHXS bo‘yicha xususiy kapital milliy hisobotdagidan 13,5 foizgacha kam, '
          'birinchi darajali asosiy kapital monandligi 1,4 foiz punktgacha past',
       4: '2.3-§: Bazel III bo‘yicha qayta hisoblashda yuqori likvid aktivlarning 27,3–30,9 foizi chiqariladi, '
          'NSFR metodika farqi 5,4–5,7 foiz punkt',
       7: '2.2-§: tahlil qilingan banklarda kapital yetarliligi ko‘p hollarda davlat tomonidan qo‘shimcha kapital '
          'kiritilishi hisobiga saqlangan'}
INTRO_OLD = 'ikkinchi bobda aniqlangan har bir muammo'
INTRO_NEW = 'ikkinchi bobda bank tizimi darajasida va tahlil qilingan uchta davlat banki misolida aniqlangan har bir muammo'

if __name__ == '__main__':
    src = docx.Document('/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx')
    cap = [p for p in src.paragraphs if p.text.startswith('Amaldagi baholash tizimi muammolari')][0]
    intro = [p for p in src.paragraphs if INTRO_OLD in p.text and '3.1.1-jadval' in p.text][0].text
    rows = [[c.text for c in r.cells] for r in docx.table.Table(cap._p.getnext(), src).rows]
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)

    def run(q, t, size=14, bold=False, red=False, strike=False):
        r = q.add_run(t); r.font.size = Pt(size); r.bold = bold; r.font.strike = strike
        if red: r.font.color.rgb = RGBColor(0xFF, 0, 0)
        if strike: r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)

    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    q.paragraph_format.first_line_indent = Cm(1.25); q.paragraph_format.line_spacing = 1.5
    a, b = intro.split(INTRO_OLD)
    run(q, a); run(q, 'ikkinchi bobda '); run(q, 'bank tizimi darajasida va tahlil qilingan uchta davlat banki misolida ', red=True)
    run(q, 'aniqlangan har bir muammo'); run(q, b)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.RIGHT; run(q, '3.1.1-jadval', bold=True)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(q, cap.text, bold=True)
    t = dd.add_table(rows=len(rows), cols=3); t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            c = t.cell(i, j); q = c.paragraphs[0]
            if i == 0: q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if j == 1 and i in FIX:
                run(q, txt, 11, strike=True)
                run(c.add_paragraph(), FIX[i], 11, red=True)
            else:
                run(q, txt, 11, bold=(i == 0))
    for j, w in enumerate([5.0, 5.5, 6.0]):
        for c in t.columns[j].cells: c.width = Cm(w)
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

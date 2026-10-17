"""3.3-§ jadval va diagrammalarini tekshirish natijalari bo‘yicha tuzatishlar."""
import re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree
import kamlib as K
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
D = K.Doc(F)
C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'


def R(old, new, n=1):
    ps = [p for p in D.d.paragraphs if old in p.text]
    cnt = sum(p.text.count(old) for p in ps)
    assert cnt == n, (old, cnt)
    for p in ps:
        for _ in range(p.text.count(old)):
            D.replace(p, old, new)


def CELL(old, new, n=1, exact=False):
    k = 0
    for t in D.d.tables:
        seen = set()
        for row in t.rows:
            for c in row.cells:
                if c._tc in seen:
                    continue
                seen.add(c._tc)
                for p in c.paragraphs:
                    if exact and p.text.strip() != old:
                        continue
                    for _ in range(p.text.count(old)):
                        D.replace(p, old, new); k += 1
    assert k == n, (old, k)


def FN_tail_cap(cap_start, new_tail, pre='Tadqiqot natijasida muallif tomonidan '):
    ps = [p for p in D.d.paragraphs if p.text.startswith(cap_start)]
    assert len(ps) == 1, (cap_start, len(ps))
    ids = [r.get(qn('w:id')) for r in ps[0]._p.iter(qn('w:footnoteReference'))]
    assert len(ids) == 1, (cap_start, ids)
    f = [f for f in D.fx.findall(qn('w:footnote')) if f.get(qn('w:id')) == ids[0]][0]
    pe = f.find(qn('w:p'))
    old = ''.join(t.text or '' for t in pe.iter(qn('w:t')))
    i = old.index(pre) + len(pre)
    D.replace(Paragraph(pe, None), old[i:], new_tail)


# ---------------- matn
R('8 chorakli siljuvchi oyna bo‘yicha', '8 chorakli (tanlanma boshida kamida 4 chorakli) siljuvchi oyna bo‘yicha')
R('2024-yilda davlat banklari indeksining 2,88 ballgacha pasayishiga ularning aktivlar rentabelligi medianasining 0,1 foizgacha tushib '
  'ketishi sabab bo‘lgan, bu II bobda “Mikrokreditbank” ATB misolida aniqlangan yashirin kredit riskining bir yo‘la tan olinishi bilan '
  'bog‘liq. 2025-yilda farq 0,46 ballgacha qisqargan, biroq bunga davlat banklari indeksining oshishi emas, asosan boshqa banklar '
  'indeksining 3,89 dan 3,61 ballgacha izchil pasayishi sabab bo‘lgan.',
  '2024-yilda davlat banklari indeksining 2,88 ballgacha pasayishiga ularning aktivlar rentabelligi medianasining 0,1 foizgacha tushib '
  'ketishi sabab bo‘lgan. Bu yilda to‘qqizta davlat bankidan beshtasida ROA 0,13 foizdan oshmagan, “Mikrokreditbank” va “Biznesni '
  'rivojlantirish banki” ATBlari esa zarar bilan yakunlagan, “Mikrokreditbank” ATB zarari esa II bobda aniqlangan yashirin kredit '
  'riskining bir yo‘la tan olinishi bilan bog‘liq. 2025-yilda farq 0,47 ballgacha qisqargan. Bunga davlat banklari indeksining 2,88 dan '
  '3,15 ballgacha tiklanishi bilan birga boshqa banklar indeksining 2022-yildagi 3,89 dan 3,61 ballgacha izchil pasayishi sabab bo‘lgan.')
R('1,7–3,3 barobar', '1,65–4,90 barobar', 2)
R('kapitalning aktivlarga nisbati esa davlat tomonidan kiritilgan mablag‘lar hisobiga yuqori darajada saqlanadi.',
  'kapitalning aktivlarga nisbati esa boshqa banklarnikidan past bo‘lsa-da (2022–2025-yillarda medianasi 13,3–14,0 foizga nisbatan '
  '15,8–16,6 foiz), aktivlar rentabelligining standart chetlanishi ancha kichik (2025-yilda 1,16 foiz punktga nisbatan 0,71 foiz punkt).')
p543 = [p for p in D.d.paragraphs if p.text.startswith('MBIIit = α')]
assert len(p543) == 1
D.replace(p543[0], 'MBII', 'MBII-8')
R('(qayd etilgan effektli modelda 5 foiz, qolgan modellarda 1 foiz darajada)', '(barcha modellarda 1 foiz darajada)')
R('0,16 ballga oshiradi', 'qariyb 0,17 ballga oshiradi', 3)
R('indeks 0,05 ballga past bo‘ladi', 'indeks qariyb 0,04 ballga past bo‘ladi')
R('(qayd etilgan effektli modelda koeffitsiyent +0,0005, p = 0,05)',
  '(koeffitsiyent qayd etilgan effektli modelda +0,0005, p = 0,05, tasodifiy effektlar modelida +0,0005, p = 0,007)')
R('Demak, jadal kreditlashning Z-score ga ta’siri asosan bir yil ichida namoyon bo‘ladi.',
  'Demak, jadal kreditlashning Z-score ga ta’siri asosan bir yil ichida namoyon bo‘ladi (3.3.5-jadval).')
R('davlat banklarining integral indeksi xususiy banklarnikidan o‘rtacha 0,69', 'davlat banklarining integral indeksi boshqa banklarnikidan o‘rtacha 0,69', 2)
R('xususiy banklarnikidan 0,69 ballga past', 'boshqa banklarnikidan 0,69 ballga past', 2)
R('davlat banklarini xususiy banklardan barqarorroq', 'davlat banklarini boshqa banklardan barqarorroq')
R('jumladan qayd etilgan effektli modelda ham statistik ahamiyatli: bank ichida kreditlarning yillik o‘sishi 10 foiz punktga oshishi bir '
  'yildan keyin Z-score ni 1,7–2,7 foizga pasaytiradi.',
  'jumladan qayd etilgan effektli modelda ham statistik ahamiyatli. RE va FE modellariga ko‘ra, kreditlarning yillik o‘sishi 10 foiz '
  'punktga oshishi bir yildan keyin Z-score ni 1,7–2,7 foizga pasaytiradi, Hausman testi esa bu holatda tasodifiy effektlar modelini '
  'rad etmaydi (p = 0,145).')
R('So‘m kursining ta’siri esa', 'AQSh dollari kursining (so‘m qadrsizlanishining) ta’siri esa')
R('beshta gipotezadan to‘rttasi to‘liq, bittasi qisman tasdiqlangan',
  'beshta gipotezadan to‘rttasi (H1, H3 va H4 – integral indeks bo‘yicha, H5 – muammoli kreditlar bo‘yicha) tasdiqlangan, bittasi (H2) '
  'qisman tasdiqlangan')
R('Z-score logarifmining medianasi, har yilning IV choragi', 'Z-score logarifmining medianasi, 2022–2025-yillarning IV choragi')

# ---------------- jadvallar
CELL('YaIM – manfiy, stavka – musbat (muammoli kreditlarga)',
     'Muammoli kreditlar ulushiga: YaIM – manfiy; stavka va so‘m qadrsizlanishi – musbat; inflyatsiya – noaniq')
CELL('Davlat banki (dummy)', 'Davlat banki (fiktiv o‘zgaruvchi)', 6)
FN_tail_cap('Tadqiqot gipotezalari va ularni tekshirishda',
            'ishlab chiqildi; gipotezalarning nazariy asoslari jadvaldan keyingi matnda keltirilgan manbalarga tayanadi.')
CELL('Panel integral indeksi (MBII), ball', 'Panel integral indeksi (MBII-8), ball')
t332 = D.table_after('Tahlilda foydalanilgan o‘zgaruvchilarning tavsifiy statistikasi')
izsrc = [p for p in D.d.paragraphs if p.text.startswith('Izoh: qavs ichida bank bo‘yicha klasterlangan standart xatolar')][0]
iz = D.new_para_after(izsrc, 'Izoh: muammoli kreditlar ulushi, kreditlar o‘sishi, depozitlar/aktivlar, sof foiz marjasi va ln(Z-score) '
                             'regressiyalardagi kabi 1 va 99 foizli persentillar bo‘yicha vinzorlangan.')
t332._tbl.addnext(iz._p)
CELL('MBII (t−1)', 'MBII-8 (t−1)')
CELL('0,007** (0,003)', '0,007*** (0,003)')
CELL('0,021 (0,013)', '0,021* (0,013)')
CELL('0,275*** (0,095)', '0,274*** (0,095)')
CELL('1-model: funding nazoratisiz', '1-model: moliyalashtirish tuzilmasi nazoratisiz')
CELL('2-model: depozitlar va foiz marjasi nazorati bilan', '2-model: depozitlar ulushi va sof foiz marjasi nazorati bilan')
CELL('funding modeli orqali', 'moliyalashtirish modeli orqali')
CELL('−0,245 (0,150)', '−0,246 (0,150)')
CELL('Integral indeks (MBII)', 'Integral indeks (MBII-8)')
CELL('barcha modellarda musbat (**/***)', 'barcha modellarda musbat (***)')
CELL('musbat (RE, FE va dinamik modellarda **/***)', 'barcha modellarda musbat (Pooled modelda *, RE, FE va dinamik modellarda **/***)')
CELL('bank ichida ahamiyatsiz', 'bank ichida manfiy ta’sir aniqlanmadi')

# ---------------- 3.3.1-rasm (chart30)
x, part = D.chart_xml('/word/charts/chart30.xml')
for t in x.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'):
    if t.text == 'MBII, ball':
        t.text = 'MBII-8, ball'
ax = [a for a in x.iter('{%s}valAx' % C) if a.find('{%s}axPos' % C).get('val') == 'b'][0]
sc = ax.find('{%s}scaling' % C)
sc.find('{%s}max' % C).set('val', '2025')
sc.find('{%s}min' % C).set('val', '2021')
for dl in x.iter('{%s}dLbls' % C):
    if dl.find('{%s}numFmt' % C) is None:
        nf = etree.Element('{%s}numFmt' % C); nf.set('formatCode', '0.00'); nf.set('sourceLinked', '0')
        dl.insert(0, nf)
part._blob = etree.tostring(x, xml_declaration=True, encoding='UTF-8', standalone=True)

# ---------------- adabiyotlar: Nickell takrori


def bib_paras():
    on = False; out = []
    for p in D.d.paragraphs:
        if p.text.strip() == 'FOYDALANILGAN ADABIYOTLAR RO‘YXATI':
            on = True; out = []; continue
        if on and p.text.strip() == 'ILOVALAR':
            break
        if on:
            out.append(p)
    return out


nk = [p for p in bib_paras() if 'Nickell S. Biases in Dynamic Models with Fixed Effects' in p.text]
assert len(nk) == 2
nk[1]._p.getparent().remove(nk[1]._p)
n = 0
for p in bib_paras():
    m = re.match(r'^(\d+)\. ', p.text)
    if not m:
        continue
    n += 1
    r0 = D._runs(p)[0].find(qn('w:t'))
    if r0.text.startswith(m.group(1) + '.'):
        r0.text = str(n) + r0.text[len(m.group(1)):]
    elif r0.text == m.group(1):
        r0.text = str(n)
    else:
        raise RuntimeError(p.text[:60])
R('foydalanilgan adabiyotlar ro‘yxati 210 ta manbani', 'foydalanilgan adabiyotlar ro‘yxati %d ta manbani' % n)
D.save(F)
print('manbalar:', n, pardoz(F))

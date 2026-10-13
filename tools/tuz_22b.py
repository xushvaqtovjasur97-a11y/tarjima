"""2.2-§ ko‘rinishini tuzatish: “Aloqabank” jadvali (markazlash, izoh), 2.2.5-jadval ustun kengliklari."""
import copy, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
d = docx.Document(F)
P = lambda s: [p for p in d.paragraphs if p.text.startswith(s)][0]
izoh_sqb = Paragraph(P('“O‘zsanoatqurilishbank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari')._p.getnext().getnext(), None)
tit_al = P('“Aloqabank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari')
tal = Table(tit_al._p.getnext(), tit_al._parent)
iz_al = tal._tbl.getnext()
assert ''.join(t.text or '' for t in iz_al.iter(qn('w:t'))).startswith('Izoh:')
iz_al.addprevious(copy.deepcopy(izoh_sqb._p)); iz_al.getparent().remove(iz_al)
# markazlash (SQB jadvalidagi kabi)
tsqb = Table(P('“O‘zsanoatqurilishbank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari')._p.getnext(), None)
for ri, row in enumerate(tal.rows):
    for ci, c in enumerate(row.cells):
        src = tsqb.rows[min(ri, len(tsqb.rows) - 1)].cells[ci].paragraphs[0]._p.find(qn('w:pPr'))
        for p in c.paragraphs:
            old = p._p.find(qn('w:pPr'))
            if old is not None:
                p._p.remove(old)
            if src is not None:
                p._p.insert(0, copy.deepcopy(src))
# 2.2.5-jadval kengliklari
tit_s = P('Bank tizimi va tanlangan banklarning moliyaviy barqarorlik ko‘rsatkichlari: xalqaro')
ts = Table(tit_s._p.getnext(), None)
Wd = [1900, 900, 1080, 1080, 1000, 1500, 1895]
grid = ts._tbl.tblGrid.findall(qn('w:gridCol'))
for gc, w in zip(grid, Wd):
    gc.set(qn('w:w'), str(w))
for tr in ts._tbl.findall(qn('w:tr')):
    for tc, w in zip(tr.findall(qn('w:tc')), Wd):
        tcpr = tc.get_or_add_tcPr()
        tw = tcpr.find(qn('w:tcW'))
        if tw is None:
            tw = OxmlElement('w:tcW'); tcpr.insert(0, tw)
        tw.set(qn('w:w'), str(w)); tw.set(qn('w:type'), 'dxa')
d.save(F)
print(pardoz(F))

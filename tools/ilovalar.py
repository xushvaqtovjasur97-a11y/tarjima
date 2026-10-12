# Ilovalar: katta hisob-kitob jadvallari, albom (landshaft) formatda
import sys, json, re, numpy as np, pandas as pd
sys.path.insert(0,'/home/user/tarjima/manbalar/hisob_kitob'); sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import mbii as M
H='/home/user/tarjima/manbalar/hisob_kitob/'
d=docx.Document(); s=d.sections[0]
s.orientation=WD_ORIENT.LANDSCAPE; s.page_width,s.page_height=Cm(29.7),Cm(21.0)
s.left_margin=Cm(2); s.right_margin=Cm(2); s.top_margin=Cm(2.5); s.bottom_margin=Cm(1.5)
st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(12); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
for nm in ('Heading 1','Heading 2'):
  h=d.styles[nm]; h.font.name='Times New Roman'; h.font.size=Pt(14); h.font.bold=True; h.font.color.rgb=None
  h.element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); h.element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
from docx.shared import RGBColor
for nm in ('Heading 1','Heading 2'): d.styles[nm].font.color.rgb=RGBColor(0,0,0)
def n(x,dd=2):
  if x is None or (isinstance(x,float) and np.isnan(x)): return '–'
  return f'{x:,.{dd}f}'.replace(',',' ').replace('.',',').replace('-','−')
def para(t,al='c',b=False,i=False,sz=12,sa=4):
  p=d.add_paragraph(); p.alignment={'c':WD_ALIGN_PARAGRAPH.CENTER,'r':WD_ALIGN_PARAGRAPH.RIGHT,'l':WD_ALIGN_PARAGRAPH.JUSTIFY}[al]
  p.paragraph_format.space_after=Pt(sa); p.paragraph_format.line_spacing=1.0
  r=p.add_run(t); r.bold=b; r.italic=i; r.font.size=Pt(sz); r.font.name='Times New Roman'; return p
first=[True]
hp=d.add_paragraph('ILOVALAR',style='Heading 1'); hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
def ilova(num,title):
  if not first[0]:
    p=d.add_paragraph(); p.paragraph_format.page_break_before=True; p.paragraph_format.space_after=Pt(0)
    para(f'{num}-ilova',al='r',b=True,i=True)
  else:
    first[0]=False; para(f'{num}-ilova',al='r',b=True,i=True)
  para(title,b=True,sa=6)
def table(head,rows,widths,sz=9,bold_rows=()):
  t=d.add_table(rows=1+len(rows),cols=len(head)); t.style='Table Grid'; t.autofit=False
  for j,hh in enumerate(head):
    c=t.rows[0].cells[j]; c.text=''; r=c.paragraphs[0].add_run(hh); r.bold=True; r.font.size=Pt(sz)
    c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),'DEEBF7'); c._tc.get_or_add_tcPr().append(sh)
  for i,row in enumerate(rows):
    for j,v in enumerate(row):
      c=t.rows[i+1].cells[j]; c.text=''; r=c.paragraphs[0].add_run(str(v)); r.font.size=Pt(sz); r.bold=(i in bold_rows)
      c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
  for row in t.rows:
    trPr=row._tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    for j,c in enumerate(row.cells):
      c.width=Cm(widths[j])
      for p in c.paragraphs: p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
  hd=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(hd)
  for j,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(widths[j]*567)))
  lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); t._tbl.tblPr.append(lay)
  return t
def manba(t): para('Manba: '+t,al='l',i=True,sz=10,sa=8)
NAMES={'SQB':'«O‘zsanoatqurilishbank» ATB','MKB':'«Mikrokreditbank» ATB','BRB':'«Biznesni rivojlantirish banki» ATB','XALQ':'«Xalq banki» ATB','IPOTEKA':'«Ipoteka-bank» ATB'}
LBL={'CET1':'K1. CET1 monandligi, %','CAR':'K2. Regulyativ kapital monandligi, %','LEV':'K3. Kapital / aktivlar, %','S3':'A1. 3-bosqich kreditlari ulushi, %','S2':'A2. 2-bosqich kreditlari ulushi, %','NPLnet':'A3. Qoplanmagan NPL / kapital, %','HLA':'L1. Yuqori likvid aktivlar / aktivlar, %','LTD':'L2. Kreditlar / depozitlar, %','ROA':'D1. ROA, %','ROE':'D2. ROE, %','CI':'D3. Xarajatlar / daromadlar, %','FXL':'V1. Valyutadagi kreditlar ulushi, %','FXLI':'V2. Valyutadagi majburiyatlar ulushi, %'}
BLN={'K':'Kapital yetarliligi bloki','A':'Aktivlar sifati bloki','L':'Likvidlik bloki','D':'Daromadlilik bloki','V':'Valyuta riski bloki'}
Y=['2021','2022','2023','2024','2025']
def indtab(get):  # get(y)-> (ind dict, bloklar, mbii)
  rows=[]
  for k,b,nm,dr,a in M.IND:
    r=[LBL[nm]]
    for y in Y:
      x,_,_=get(y); v=x.get(nm)
      r.append('–' if v is None else f'{n(v,1)} ({n(M.score(v,dr,a),2)})')
    rows.append(r)
  for b in 'KALDV': rows.append([BLN[b]]+[n(get(y)[1].get(b),2) for y in Y])
  rows.append(['MBII (sinf)']+[f'{n(get(y)[2],2)} ({M.cls(get(y)[2])})' for y in Y])
  return rows
# 1-ilova: 5 bank, MHXS
R32=json.load(open(H+'mbii32.json'))
IF='Tadqiqot natijasida muallif tomonidan banklarning 2021–2025-yillar uchun auditdan o‘tgan MHXS bo‘yicha moliyaviy hisobotlari (har bir yil ko‘rsatkichlari keyingi yil hisobotidagi qayta hisoblangan qiyosiy ma’lumotlardan) va «SNS» reyting agentligi bazasi asosida hisoblandi.'
ilova(1,'Tanlangan banklarda moliyaviy barqarorlik integral indeksi ko‘rsatkichlari va ballari, MHXS hisobotlari asosida, 2021–2025-yillar')
para('Izoh: har bir katakda ko‘rsatkich qiymati va qavs ichida uning bali (1–4) keltirilgan; bloklar vaznlari: kapital – 0,25; aktivlar sifati – 0,30; likvidlik – 0,20; daromadlilik – 0,15; valyuta riski – 0,10.',al='l',i=True,sz=10)
for i,c in enumerate(['SQB','MKB','BRB','XALQ','IPOTEKA']):
  para(f'1.{i+1}. {NAMES[c]}',b=True,sz=11,sa=2)
  table(['Ko‘rsatkich']+Y,indtab(lambda y,c=c:(R32[c][y]['ind'],R32[c][y]['bloklar'],R32[c][y]['mbii'])),[8.5]+[3.6]*5,sz=9,bold_rows=(18,))
  para('',sa=2)
manba(IF)
# 2-ilova: SQB va MKB, milliy hisobot
import mbii_milliy as MM
ilova(2,'«O‘zsanoatqurilishbank» ATB va «Mikrokreditbank» ATB da moliyaviy barqarorlik integral indeksi ko‘rsatkichlari va ballari, milliy prudensial hisobot asosida, 2021–2025-yillar')
for i,(c,cn) in enumerate((('SQ','SQB'),('MK','MKB'))):
  para(f'2.{i+1}. {NAMES[cn]}',b=True,sz=11,sa=2)
  def g(y,c=c):
    x=MM.ind(c,int(y)); m,bs=M.mbii(x,M.W['taklif']); return x,bs,m
  table(['Ko‘rsatkich']+Y,indtab(g),[8.5]+[3.6]*5,sz=9,bold_rows=(18,))
  para('',sa=2)
manba('Tadqiqot natijasida muallif tomonidan banklarning Markaziy bankka taqdim etilgan prudensial hisobot shakllari va «SNS» reyting agentligi bazasi asosida hisoblandi.')
# 3-ilova: NSFR 4 variant
ilova(3,'«O‘zsanoatqurilishbank» ATB ning sof barqaror moliyalashtirish me’yori: to‘rt variant bo‘yicha hisob-kitob, 2025-yil yakuni, mlrd so‘m')
L=[l for l in open('/home/user/tarjima/paragraflar/NSFR_4_variant.md').read().split('\n')]
i0=[k for k,l in enumerate(L) if l.startswith('| Tarkibiy qismlar |')][0]
rows=[]; k=i0+2
while k<len(L) and L[k].startswith('|'):
  cells=[x.strip().replace('**','') for x in L[k].strip('|').split('|')]; rows.append(cells); k+=1
table(['Tarkibiy qismlar','1-variant: Markaziy bank metodikasi, milliy hisobot','2-variant: Bazel III metodikasi, milliy hisobot','3-variant: Markaziy bank metodikasi, MHXS','4-variant: Bazel III metodikasi, MHXS'],rows,[9.7,4.0,4.0,4.0,4.0],sz=10,bold_rows=(3,9,10))
manba('Tadqiqot natijasida muallif tomonidan «O‘zsanoatqurilishbank» ATB ning prudensial hisobot shakllari va 2025-yil uchun MHXS bo‘yicha konsolidatsiyalangan moliyaviy hisoboti asosida hisoblandi.')
# 4-ilova: 25 bank paneli MBII
p=pd.read_csv(H+'panel33.csv'); pp=p[p.tanlov]
Q=sorted(pp.davr.unique(),key=lambda x:(int(x[:4]),['I','II','III','IV'].index(x[5:])))
ilova(4,'Tanlangan 25 ta tijorat bankida moliyaviy barqarorlik panel integral indeksi (8 ko‘rsatkich, milliy hisobot asosida), 2021-yil IV chorak – 2025-yil IV chorak')
for part,qq in enumerate((Q[:9],Q[9:])):
  pv=pp.pivot_table(index='bank',columns='davr',values='MBII')
  rows=[[b+(' (D)' if pp[pp.bank==b].DAVLAT.iloc[-1]==1 else '')]+[n(pv.loc[b,q],2) if q in pv.columns and pv.loc[b,q]==pv.loc[b,q] else '–' for q in qq] for b in sorted(pv.index)]
  med=['Mediana']+[n(pv[q].median(),2) for q in qq]
  table(['Bank']+[q.replace('_',' ') for q in qq],rows+[med],[6.0]+[19.7/len(qq)]*len(qq),sz=8.5,bold_rows=(len(rows),))
  para('',sa=4)
para('Izoh: (D) – davlat ishtirokidagi bank (2025-yil holatiga); «–» – ma’lumot mavjud emas.',al='l',i=True,sz=10)
manba('Tadqiqot natijasida muallif tomonidan O‘zbekiston Respublikasi Markaziy bankining banklar statistikasi va «SNS» reyting agentligi bazasi asosida hisoblandi.')
# 5-ilova: korrelyatsiya matritsasi
ilova(5,'3.3-paragrafdagi o‘zgaruvchilarning juftlik korrelyatsiya koeffitsiyentlari, 2021-yil IV chorak – 2025-yil IV chorak')
V=[('MBII','MBII'),('lnZ','ln(Z)'),('lnTA','ln(aktivlar)'),('KR_osish_l4','Kredit o‘sishi (t−4)'),('DEP_A','Depozit/aktiv'),('NIM','NIM'),('DAVLAT','Davlat'),('YaIM_osish','YaIM o‘sishi'),('Inflyatsiya','Inflyatsiya'),('Asosiy_stavka','Asosiy stavka'),('USD_osish','USD kursi')]
C=pp[[v for v,_ in V]].corr()
table(['']+[l for _,l in V],[[l]+[n(C.loc[a,b],2) for b,_ in V] for a,l in V],[3.4]+[2.2]*len(V),sz=8.5)
manba('Tadqiqot natijasida muallif tomonidan hisoblandi.')
# 6-ilova: signallar usuli – to'liq jadval
ilova(6,'Integral indeksning turli chegaralarida erta ogohlantirish signallarining aniqligi (to‘liq jadval)')
g=p.sort_values(['bank','t']).groupby('bank'); p2=p.sort_values(['bank','t']).copy()
p2['lossq']=-p2.NPq/g.Kapital.shift(1)*100; p2['EV']=(p2.lossq>=5).astype(int)
gg=p2.groupby('bank')
for j in range(1,5): p2[f'e{j}']=gg.EV.shift(-j).where(gg.t.shift(-j)==p2.t+j)
p2['Y']=p2[[f'e{j}' for j in range(1,5)]].max(axis=1)
dd=p2[p2.tanlov&(p2.t<=2024*4+3)&(p2.EV==0)].dropna(subset=['Y','MBII'])
rows=[]
for th in np.round(np.arange(2.40,3.81,0.05),2):
  S=dd.MBII<th; y=dd.Y==1; TP=int((S&y).sum()); FN=int((~S&y).sum()); FP=int((S&~y).sum()); TN=int((~S&~y).sum())
  t1=FN/(TP+FN)*100; t2=FP/(FP+TN)*100; nsr=(t2/100)/(1-t1/100) if t1<100 else np.nan
  rows.append([n(th,2),TP,FN,FP,TN,n(t1,1),n(t2,1),n(nsr,2),n(100-t1-t2,1)])
table(['Chegara (MBII)','To‘g‘ri signal','O‘tkazib yuborilgan hodisa','Noto‘g‘ri signal','To‘g‘ri «signalsiz»','I tur xato, %','II tur xato, %','Shovqin/signal nisbati','Youden indeksi, %'],rows,[3.0,2.6,3.0,2.6,2.8,2.6,2.6,3.0,3.0],sz=9)
manba('Tadqiqot natijasida muallif tomonidan 25 ta bank paneli (308 ta kuzatuv, 26 tasi muammo hodisasidan oldingi davr) asosida hisoblandi.')
# 7-ilova: 4.2 bank darajasidagi prognoz
ilova(7,'Tanlangan 25 ta bankning integral indeksi va erta ogohlantirish zonasi: 2025-yil (haqiqiy) va 2027-yil (ssenariylar bo‘yicha prognoz)')
import fc42 as F  # noqa – faqat hisoblangan F jadvali uchun
FF=F.F; b25=F.b25
def z(m): return 'yashil' if m>=3.25 else 'sariq' if m>=2.75 else 'qizil'
rows=[]
for b in sorted(b25.index):
  r=[b+(' (D)' if b25.loc[b].DAVLAT==1 else ''),f'{n(b25.loc[b].MBII,2)} ({z(b25.loc[b].MBII)})']
  for sc in ('Asosiy','Ijobiy','Salbiy','Keskin salbiy'):
    x=FF[(FF.bank==b)&(FF.ssenariy==sc)&(FF.yil==2027)].iloc[0]; r.append(f'{n(x.MBII,2)} ({z(x.MBII)})')
  x=FF[(FF.bank==b)&(FF.ssenariy=='Keskin salbiy')&(FF.yil==2027)].iloc[0]; r.append(f'{n(b25.loc[b].S3,1)} → {n(x.S3,1)}')
  rows.append(r)
table(['Bank','2025 (haqiqiy)','Asosiy 2027','Ijobiy 2027','Salbiy 2027','Keskin salbiy 2027','Muammoli kreditlar ulushi: 2025 → keskin salbiy 2027, %'],rows,[5.6,3.3,3.3,3.3,3.3,3.5,4.3],sz=9)
para('Izoh: (D) – davlat ishtirokidagi bank; qavs ichida erta ogohlantirish zonasi.',al='l',i=True,sz=10)
manba('Tadqiqot natijasida muallif tomonidan hisoblandi.')
# 8-ilova: 3.2 – ko'rsatkichlar entropiya vaznlari va Monte-Karlo
VZ=json.load(open(H+'vazn32.json'))
ilova(8,'Integral indeks ko‘rsatkichlarining entropiya vaznlari va Monte-Karlo tahlilida banklar sinfining saqlanish ulushi')
ie=VZ['ind_entropiya']
table(['Ko‘rsatkich','Entropiya vazni','Blok'],[[LBL[k],n(v,3),BLN[[b for _,b,nm,_,_ in M.IND if nm==k][0]]] for k,v in ie.items()],[11.0,4.0,10.7],sz=9)
para('',sa=4)
a=VZ['mc_ekspert_atrofida']['bank_yil']; b=VZ['mc_bir_tekis']['bank_yil']
rows=[]
for c in ['SQB','MKB','BRB','XALQ','IPOTEKA']:
  rows.append([NAMES[c]]+[f"{n(a[f'{c}|{y}'],1)} / {n(b[f'{c}|{y}'],1)}" for y in Y])
table(['Bank']+Y,rows,[7.7]+[3.6]*5,sz=9)
para('Izoh: har bir katakda vaznlar ekspert vaznlari atrofida va to‘liq tasodifiy tanlanganda bank sinfi saqlangan holatlar ulushi (%) keltirilgan; 4000 ta tasodifiy vazn to‘plami.',al='l',i=True,sz=10)
manba('Tadqiqot natijasida muallif tomonidan hisoblandi.')
d.save('/home/user/tarjima/paragraflar/Ilovalar.docx'); print('saved')

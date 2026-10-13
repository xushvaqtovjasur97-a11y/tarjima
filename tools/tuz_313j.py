"""3.1.3-jadval: uch manba, Asoslash ustuni, vaznlar o‘zgarishi asoslanishi (qizil)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
import kamlib as K
import namuna_313j as N
from pardoz import pardoz

F = N.F
D = K.Doc(F)
ps = D.d.paragraphs
p1 = [p for p in ps if 'uch manba asosida asoslandi' in p.text][0]
D.insert(p1, N.INS1[0], N.INS1[1], tag='313j')
p2 = [p for p in ps if p.text.startswith('Jadval ma’lumotlari shuni ko‘rsatadiki, taklif etilayotgan vaznlar')][0]
D.replace(p2, N.REP2[0], N.REP2[1], tag='313j')
D.insert(p2, N.INS3[0], N.INS3[1], tag='313j')
cap = [p for p in ps if p.text.startswith('Moliyaviy barqarorlikning integral indeksi bloklari vaznlari')][0]
t = docx.table.Table(cap._p.getnext(), D.d)
r = [i for i, row in enumerate(t.rows) if row.cells[4].text == N.CELL[0]][0]
D.insert(t.cell(r, 4).paragraphs[0], N.CELL[0], N.CELL[1], tag='313j')
D.save(F)
print(pardoz(F))

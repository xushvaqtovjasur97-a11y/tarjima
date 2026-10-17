"""Yashirin kredit riski: 2.3.2-jadvalga koeffitsiyent qatori, tuzatish formulasi (2.3), 4.3-§ da havola (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_yashirin_kredit_riski.docx'
ROW = ['Yashirin kredit riski koeffitsiyenti', '1,65', '2,61', '2,14', '2,08', '2,96']
AFTER_P = 'tahlil xulosalariga sezilarli ta’sir qilmaydi.'
P1 = 'Yashirin kredit riski koeffitsiyenti asosida milliy hisobot bo‘yicha muammoli kreditlar ulushi xalqaro standartlar darajasiga quyidagicha keltiriladi:'
F3 = 'NPLtuz = (NPL / Lmilliy) × Kykr,  (2.3)'
P2 = ('bu yerda NPLtuz – MHXS 9 darajasiga keltirilgan muammoli kreditlar ulushi. MHXS hisoboti mavjud bo‘lgan yillarda tuzatilgan '
      'ko‘rsatkich 3-bosqich kreditlari ulushiga teng bo‘ladi, shu sababli tuzatish asosan MHXS hisoboti hali e’lon qilinmagan davrlar, '
      'ya’ni joriy yilning choraklari va prognoz davri uchun ahamiyatli: bunda bankning oxirgi ma’lum yillik koeffitsiyenti qo‘llanadi. '
      'Tuzatilgan ko‘rsatkich integral indeksning aktivlar sifati blokida milliy muammoli kreditlar ulushi o‘rniga ishlatiladi.')
OLD43 = '2.3-paragrafda taklif etilgan yashirin kredit riski koeffitsiyenti bilan tuzatgan holda'
NEW43 = '2.3-paragrafda taklif etilgan yashirin kredit riski koeffitsiyenti bilan (2.3) formula bo‘yicha tuzatgan holda'

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)

    def par(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, bold=False, indent=True):
        q = dd.add_paragraph(); q.alignment = align
        if indent and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
            q.paragraph_format.first_line_indent = Cm(1.25)
        for t, col, strike in parts:
            r = q.add_run(t); r.font.size = Pt(size); r.bold = bold; r.font.strike = strike
            if col: r.font.color.rgb = RGBColor.from_string(col)
        return q
    par([('Yashirin kredit riski koeffitsiyenti (namuna)', None, False)], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    dd.add_paragraph()
    par([('1) 2.3-§, (2.2) formuladan keyingi izoh oxiridan so‘ng yangi matn va formula:', None, False)], size=12, bold=True, indent=False)
    par([('… ', None, False), (AFTER_P, None, False)])
    par([(P1, 'FF0000', False)])
    par([(F3, 'FF0000', False)], WD_ALIGN_PARAGRAPH.CENTER)
    par([(P2, 'FF0000', False)], indent=False)
    dd.add_paragraph()
    par([('2) 2.3.2-jadval (“O‘zsanoatqurilishbank” ATB) – oxirgi qator qo‘shiladi, boshqa ikki bank jadvallaridagidek:', None, False)],
        size=12, bold=True, indent=False)
    t = dd.add_table(rows=3, cols=6); t.style = 'Table Grid'
    rows = [['Ko‘rsatkich', '2021', '2022', '2023', '2024', '2025'],
            ['Muammoli kreditlar (milliy) / 3-bosqich (MHXS), %', '3,8 / 6,2', '2,8 / 7,4', '2,2 / 4,8', '2,8 / 5,9', '2,5 / 7,5'],
            ROW]
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = ''
            run = c.paragraphs[0].add_run(v); run.font.size = Pt(12); run.bold = (i == 0)
            if i == 2: run.font.color.rgb = RGBColor.from_string('FF0000')
    dd.add_paragraph()
    par([('3) 4.3-§, uchinchi taklif xatboshisida:', None, False)], size=12, bold=True, indent=False)
    par([('…milliy muammoli kreditlar ko‘rsatkichini 2.3-paragrafda taklif etilgan yashirin kredit riski koeffitsiyenti bilan ', None, False),
         ('(2.3) formula bo‘yicha ', 'FF0000', False), ('tuzatgan holda integral indeksga kiritish maqsadga muvofiq.', None, False)])
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

"""3.1.5-jadval ikkiga bo‘linadi: milliy hisobot (3 bank) va MHXS (3 bank) + tahlil matnlari (namuna)."""
import json, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_3.1.5-3.1.6-jadval.docx'
O = json.load(open('/home/user/tarjima/manbalar/hisob_kitob/mbii31_3bank.json'))
B = [('SQB', '“O‘zsanoatqurilishbank” ATB'), ('MKB', '“Mikrokreditbank” ATB'), ('ALQ', '“Aloqabank” ATB')]
IND = [('K1. CET1 monandligi, %', 'CET1', 2), ('K2. Regulyativ kapital monandligi, %', 'CAR', 2), ('K3. Kapital / aktivlar, %', 'LEV', 1),
       ('A1. {} kreditlar ulushi, %', 'S3', 1), ('A2. {} kreditlar ulushi, %', 'S2', 1),
       ('A3. Qoplanmagan muammoli kreditlar / kapital, %', 'NPLnet', 1),
       ('L1. Yuqori likvid aktivlar / aktivlar, %', 'HLA', 1), ('L2. Kreditlar / depozitlar, %', 'LTD', 1),
       ('D1. ROA, %', 'ROA', 1), ('D2. ROE, %', 'ROE', 1), ('D3. Xarajatlar / daromadlar, %', 'CI', 1),
       ('V1. Xorijiy valyutadagi kreditlar, %', 'FXL', 1), ('V2. Xorijiy valyutadagi majburiyatlar, %', 'FXLI', 1)]
BL = [('Kapital yetarliligi bloki', 'K'), ('Aktivlar sifati bloki', 'A'), ('Likvidlik bloki', 'L'), ('Daromadlilik bloki', 'D'),
      ('Valyuta riskiga sezgirlik bloki', 'V')]
SINF = {'barqaror': 'I', 'qoniqarli': 'II', 'zaif': 'III', 'nobarqaror': 'IV'}


def f(v, n=1):
    s = f'{v:,.{n}f}'.replace(',', ' ').replace('.', ',')
    return s.replace('-', '−')


def jadval_rows(var):
    rows = [['Ko‘rsatkich', 'SQB', 'MKB', 'Aloqabank']]
    s3, s2 = ('Muammoli', 'Substandart') if var == 'milliy' else ('3-bosqich', '2-bosqich')
    for lab, k, n in IND:
        lab = lab.format(s3 if k == 'S3' else s2)
        r = [lab]
        for c, _ in B:
            o = O[f'{c} {var}']; v = o['ind'].get(k)
            r.append('–' if v is None else f'{f(v, n)} ({f(o["ball"][k], 2)})')
        rows.append(r)
    for lab, b in BL:
        rows.append([lab] + [f(O[f'{c} {var}']['bloklar'][b], 2) for c, _ in B])
    rows.append(['MBII'] + [f(O[f'{c} {var}']['mbii'], 3 if (c == 'ALQ' and var == 'milliy') else 2) for c, _ in B])
    rows.append(['Barqarorlik sinfi'] + [SINF[O[f'{c} {var}']['sinf']] for c, _ in B])
    return rows


INTRO_OLD = ('ikki bankning', 'uch bankning')
INTRO_OLD2 = ('Hisob-kitob ikki variantda – milliy hisobot va Markaziy bank metodikasi asosida hamda MHXS va Bazel III metodikasi asosida bajarildi (3.1.5-jadval).',
              'Hisob-kitob ikki variantda bajarildi: milliy hisobot va Markaziy bank metodikasi asosida (3.1.5-jadval) hamda MHXS va Bazel III metodikasi asosida (3.1.6-jadval).')
CAP1 = 'Uch bank bo‘yicha integral indeks ko‘rsatkichlari va ballari: milliy hisobot, 2025-yil yakuni'
CAP2 = 'Uch bank bo‘yicha integral indeks ko‘rsatkichlari va ballari: MHXS hisobotlari, 2025-yil yakuni'
IZOH1 = ('Izoh: SQB – “O‘zsanoatqurilishbank” ATB, MKB – “Mikrokreditbank” ATB; qavs ichida – ball. “Aloqabank” ATB ning indeksi I sinf chegarasidan (3,25) 0,002 ball past bo‘lgani sababli '
         'uch xona aniqlikda keltirildi.')
IZOH2 = ('Izoh: SQB – “O‘zsanoatqurilishbank” ATB, MKB – “Mikrokreditbank” ATB; qavs ichida – ball. CET1 va regulyativ kapital monandligi 2.3-paragrafda Markaziy bank metodikasi bo‘yicha MHXS '
         'ma’lumotlari asosida qayta hisoblangan qiymatlarda olindi (2.3.4, 2.3.5-jadvallar va “Mikrokreditbank” ATB bo‘yicha jadval), '
         'chunki banklar MHXS hisobotida oshkor qiladigan kapital koeffitsiyentlari milliy buxgalteriya hisobi ma’lumotlari asosida '
         'hisoblanadi. Yuqori likvid aktivlar MHXS hisobotlarida alohida oshkor qilinmagani sababli L1 ko‘rsatkichi ikkala variantda '
         'milliy balans asosida olindi. “Mikrokreditbank” va “Aloqabank” ATBlarining MHXS hisobotlarida xorijiy valyutadagi kreditlar '
         'ulushi oshkor qilinmagan, shu sababli ularda valyuta bloki bitta ko‘rsatkich bo‘yicha baholandi.')
T1 = ('Milliy hisobot ma’lumotlari bo‘yicha uchala bank yuqori baholanadi: “O‘zsanoatqurilishbank” ATB ning indeksi 3,27 ballni '
      'tashkil etib, I sinfga kiradi, “Aloqabank” ATB niki 3,248 ball bilan I sinf chegarasidan atigi 0,002 ball past, '
      '“Mikrokreditbank” ATB niki esa 3,13 ball bilan II sinfning yuqori qismida joylashadi. Uchala bankda ham kapital yetarliligi '
      '(3,76–4,00 ball) va aktivlar sifati (3,42–3,80 ball) bloklari yuqori baholangan, chunki rasmiy muammoli kreditlar ulushi '
      '2,5–4,4 foizni tashkil etadi. Zaif bo‘g‘inlar esa banklarning faoliyat modeliga qarab farqlanadi: “O‘zsanoatqurilishbank” '
      'ATB da valyuta riskiga sezgirlik (1,47 ball, kreditlarning 60,1 foizi xorijiy valyutada) va likvidlik (2,33 ball), '
      '“Mikrokreditbank” ATB da likvidlik (2,05 ball) va daromadlilik (2,09 ball, kapital rentabelligi 3,5 foiz), “Aloqabank” ATB '
      'da esa daromadlilik (2,22 ball, xarajatlar daromadlarning 82,2 foizini tashkil etadi) va likvidlik (2,56 ball) bloklari. '
      'Demak, milliy hisobot asosida uchala bank barqaror yoki qoniqarli darajada ko‘rinadi, ularning zaifliklari esa asosan '
      'likvidlik, daromadlilik va valyuta riski bilan bog‘liq bo‘lib, aktivlar sifati muammo sifatida namoyon bo‘lmaydi.')
T2 = ('MHXS ma’lumotlari asosida esa uchala bankning indeksi pasayadi: “O‘zsanoatqurilishbank” ATB niki 2,81 ballga (II sinf), '
      '“Aloqabank” ATB niki 2,78 ballga (II sinf), “Mikrokreditbank” ATB niki 2,47 ballga (III sinf) tushadi. ΔMBII mos ravishda '
      '−0,46; −0,47 va −0,66 ballni tashkil etadi: “O‘zsanoatqurilishbank” va “Mikrokreditbank” ATBlari bir sinf pastga tushadi, '
      '“Aloqabank” ATB esa I sinf chegarasidan II sinfning quyi qismiga, 4.1-paragrafdagi qizil zona chegarasiga (2,75) yaqin '
      'darajaga tushadi. Indeks pasayishining asosiy manbai uchala bankda ham aktivlar sifati blokidir: uning bali '
      '“O‘zsanoatqurilishbank” ATB da 3,55 dan 2,28 ga, “Mikrokreditbank” ATB da 3,42 dan 1,53 ga, “Aloqabank” ATB da 3,80 dan '
      '2,93 ga pasaygan, bu esa ΔMBII ning mos ravishda 83, 86 va 56 foizini tashkil etadi. Ikkinchi manba – daromadlilik bloki: '
      'u “Aloqabank” ATB da 2,22 dan 1,18 ballga pasayib, ΔMBII ning 33 foizini belgilaydi, chunki MHXS bo‘yicha bank 2025-yilda '
      '193,9 mlrd so‘m zarar ko‘rgan, “Mikrokreditbank” ATB da esa 2,09 dan 1,33 ballga tushgan (ΔMBII ning 17 foizi). Kapital bloki '
      '“O‘zsanoatqurilishbank” va “Aloqabank” ATBlarida biroz pasaygan, “Mikrokreditbank” ATB da esa maksimal darajada qolgan. '
      'Shunday qilib, tizimli ahamiyatga molik bo‘lmagan uchinchi bank misolida ham milliy hisobot asosidagi baholash barqarorlikni '
      'oshirib ko‘rsatishi, ΔMBII esa asosan aktivlar sifati va daromadlilik bloklari hisobiga shakllanishi tasdiqlandi.')
FN1 = ('Tadqiqot natijasida muallif tomonidan “O‘zsanoatqurilishbank” ATB (crs003C, crs003L, crs003O), “Mikrokreditbank” ATB '
       '(crs005C, crs005L, crs005O) va “Aloqabank” ATB (crs031C, crs031L, crs031O) ning Markaziy bankka taqdim etilgan prudensial '
       'hisobot shakllari asosida, 2025-yil 31-dekabr holatiga hisoblandi.')
FN2 = ('Tadqiqot natijasida muallif tomonidan quyidagi manbalar asosida hisoblandi: Joint Stock Commercial Bank “Uzbek Industrial and '
       'Construction Bank” and its subsidiaries. Consolidated Financial Statements for the year ended 31 December 2025 and Independent '
       'Auditor’s Report. – Tashkent, 2026. – P. 2–3, 28–31, 37–38, 74; Акционерно-коммерческий банк “Микрокредитбанк”. '
       'Консолидированная финансовая отчетность за год, закончившийся 31 декабря 2025 года, и заключение независимых аудиторов. – '
       'Ташкент, 2026. – С. 11–12, 48; Joint Stock Commercial “Aloqabank” and its subsidiaries: Consolidated Financial Statements and '
       'Independent Auditor’s Report, 31 December 2025. – Tashkent: Audit Organization “PricewaterhouseCoopers” LLC, 2026. – '
       'P. 1–2, 35, 55, 64–65.')

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)

    def run(q, t, size=14, bold=False, red=False, it=False, sup=False):
        r = q.add_run(t); r.font.size = Pt(size); r.bold = bold; r.italic = it; r.font.superscript = sup
        if red: r.font.color.rgb = RGBColor(0xFF, 0, 0)

    def body(parts):
        q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        q.paragraph_format.first_line_indent = Cm(1.25); q.paragraph_format.line_spacing = 1.5
        for t, red in parts: run(q, t, red=red)

    def table(no, cap, rows, fn):
        q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.RIGHT; run(q, no, bold=True, it=True, red=True)
        q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(q, cap, bold=True, red=True); run(q, fn, bold=True, red=True, sup=True)
        t = dd.add_table(rows=len(rows), cols=4); t.style = 'Table Grid'; t.autofit = False
        Wd = [7.2, 3.1, 3.1, 3.1]
        for g, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), Wd): g.set(qn('w:w'), str(int(w * 567)))
        for i, row in enumerate(rows):
            for j, txt in enumerate(row):
                c = t.cell(i, j); c.width = Cm(Wd[j]); q = c.paragraphs[0]
                if j or i == 0: q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run(q, txt, 11, bold=(i == 0 or i >= 14), red=True)

    def izoh(t):
        q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; run(q, t, 12, it=True, red=True)

    src = docx.Document('/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx')
    intro = [p for p in src.paragraphs if p.text.startswith('Taklif etilayotgan metodikaning amaliy qo‘llanilishini')][0].text
    a, b = intro.split(INTRO_OLD[0]); b1 = b.replace(INTRO_OLD2[0], '')
    body([(a, False), (INTRO_OLD[1], True), (b1, False), (INTRO_OLD2[1], True)])
    table('3.1.5-jadval', CAP1, jadval_rows('milliy'), '¹')
    izoh(IZOH1)
    body([(T1, True)])
    table('3.1.6-jadval', CAP2, jadval_rows('MHXS'), '²')
    izoh(IZOH2)
    body([(T2, True)])
    q = dd.add_paragraph(); run(q, '¹ ' + FN1, 10, red=True)
    q = dd.add_paragraph(); run(q, '² ' + FN2, 10, red=True)
    dd.add_paragraph()
    q = dd.add_paragraph(); run(q, 'Eslatma (dissertatsiyaga kiritilmaydi): mavjud 3.1.6-jadval va undan keyingi 3.1-§ jadvallari bir raqamga '
                                   'suriladi; eski 3.1.5-jadval, uning izohi va tahlil xatboshisi yangilari bilan almashtiriladi. '
                                   '“Aloqabank” ATB ning milliy K3 ko‘rsatkichi boshqa banklar kabi xususiy kapital / aktivlar sifatida '
                                   'qayta hisoblandi (12,2 foiz; avval prudensial leverage 9,9 foiz olingan), ball o‘zgarmaydi (4,00).', 11, bold=True)
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

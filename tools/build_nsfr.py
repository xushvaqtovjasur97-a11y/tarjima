import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/NSFR.md'; DST='/home/user/tarjima/paragraflar/NSFR.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/nsfr_raw.docx'],check=True)
d=docx.Document(S+'/nsfr_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^\d+-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^\d+-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^\d+-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'
# formula
sh=[]
sh.append(W.box(0.6,0.75,2.6,1.0,'NSFR =',fill=None,line=None,sz=28,bold=True))
sh.append(W.box(3.4,0.1,8.6,1.0,'Mavjud barqaror moliyalashtirish (ASF)',fill='E2EFDA',line=GREEN,sz=22,bold=True))
sh.append(W.line(3.4,1.25,12.0,1.25,color='000000',lw=19050,arrow=False))
sh.append(W.box(3.4,1.4,8.6,1.0,'Talab qilinadigan barqaror moliyalashtirish (RSF)',fill='FBE5D6',line=ORANGE,sz=22,bold=True))
sh.append(W.box(12.3,0.75,3.0,1.0,'≥ 100%',fill=None,line=None,sz=28,bold=True,color='C00000'))
W.replace_placeholder(d,'{{FIG:F}}',W.group_run_xml(sh,16.0,2.5,'formula'))
# 1: structure
sh=[]
asf='Kapital va 1 yildan ortiq resurslar – 100%\nBarqaror chakana depozitlar – 95%\nKamroq barqaror chakana depozitlar – 90%\nKorxonalarning 1 yilgacha resurslari – 50%\nMoliya institutlarining qisqa\nmuddatli resurslari – 0%'
rsf='Naqd pul, markaziy bankdagi zaxiralar – 0%\n1-darajali likvid aktivlar – 5%\n1 yilgacha kreditlar – 50%\nIpoteka (1 yildan ortiq) – 65%\nBoshqa uzoq muddatli kreditlar – 85%\nMuammoli kreditlar, asosiy vositalar – 100%'
sh.append(W.box(0.1,0.1,7.8,0.9,'Mavjud barqaror moliyalashtirish\n(passivlar × ASF koeffitsiyenti)',fill=GREEN,line=GREEN,sz=20,bold=True,color='FFFFFF'))
sh.append(W.box(8.5,0.1,7.8,0.9,'Talab qilinadigan barqaror moliyalashtirish\n(aktivlar × RSF koeffitsiyenti)',fill=ORANGE,line=ORANGE,sz=20,bold=True,color='FFFFFF'))
sh.append(W.box(0.1,1.1,7.8,3.3,asf,fill='F2F2F2',line=GREEN,sz=18,align='left'))
sh.append(W.box(8.5,1.1,7.8,3.3,rsf,fill='F2F2F2',line=ORANGE,sz=18,align='left'))
sh.append(W.line(4.0,4.4,7.0,5.0,color='404040',lw=15875))
sh.append(W.line(12.4,4.4,9.4,5.0,color='404040',lw=15875))
sh.append(W.box(4.2,5.0,8.0,0.95,'ASF ≥ RSF: uzoq muddatli va nolikvid aktivlar\nbarqaror resurslar bilan moliyalashtiriladi',fill=DBLUE,line=DBLUE,sz=20,bold=True,color='FFFFFF'))
sh.append(W.box(0.1,6.15,16.2,0.7,'Gorizont: 1 yil  •  Maqsad: tarkibiy likvidlik riskini va muddatlar nomutanosibligini cheklash',fill='DEEBF7',line=LBLUE,sz=19,italic=True))
W.replace_placeholder(d,'{{FIG:1}}',W.group_run_xml(sh,16.4,7.0,'1-rasm'))
# 2: timeline
sh=[]
Y0=3.0
sh.append(W.line(0.3,Y0,16.1,Y0,color='404040',lw=28575))
ev=[(2010,'Bazel III: likvidlik\nstandartlari e’lon qilindi',DBLUE,True),
    (2014,'Yakuniy standart\n(oktabr)',BLUE,False),
    (2018,'Bazel muddati; Rossiya\n(tizimli banklar) va\nO‘zbekiston – 100%',GREEN,True),
    (2021,'Yevropa Ittifoqi (28 iyun),\nAQSh (1 iyul) – 100%',ORANGE,False),
    (2026,'O‘zbekiston: elementlarni\nBazel III ga to‘liq\nmoslashtirish, monitoring',GOLD,True)]
def X(y): return 0.8+(y-2010)/16*14.6
for yr,t,c,up in ev:
  x=X(yr)
  sh.append(W.box(x-0.35,Y0-0.35,0.7,0.7,'',fill=c,line=c,geom='ellipse'))
  sh.append(W.box(x-0.6,Y0+(0.45 if up else -1.05),1.2,0.55,str(yr),fill=None,line=None,sz=20,bold=True))
  bx=min(max(x-1.9,0.1),12.5)
  if up:
    sh.append(W.box(bx,0.1,3.8,1.6,t,fill='F2F2F2',line=c,sz=17)); sh.append(W.line(x,1.7,x,Y0-0.35,color=c,lw=12700,arrow=False))
  else:
    sh.append(W.box(bx,4.3,3.8,1.4,t,fill='F2F2F2',line=c,sz=17)); sh.append(W.line(x,Y0+0.35,x,4.3,color=c,lw=12700,arrow=False))
W.replace_placeholder(d,'{{FIG:2}}',W.group_run_xml(sh,16.4,5.85,'2-rasm'))
d.save(DST); print('saved')

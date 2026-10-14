"""2.3.7-rasm: ustunlar tafovut zanjiri tartibida, bloklar bo‘yicha rang, vergulli yorliqlar, o‘zbekcha nomlar; izohda 2.3.2-jadval (namuna)."""
import re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W

DST = '/home/user/tarjima/Namuna_2.3.7-rasm.docx'
CAP = '2.3.7-rasm. “O‘zsanoatqurilishbank” ATB ko‘rsatkichlarining milliy va xalqaro yondashuv bo‘yicha nisbiy tafovuti, 2025-yil, %'
CATS = ['Muammoli kreditlar ulushi', 'Kredit zaxiralari', 'Sof foyda', 'Xususiy kapital',
        'CET1 monandligi', 'NSFR me’yori', 'Jami aktivlar']
VALS = [195.7, 44.2, -17.9, -13.5, -13.1, 2.1, -5.1]
BLOCKS = [('Aktivlar sifati', [0, 1], '1F4E79'), ('Daromadlilik', [2], 'C55A11'), ('Kapital', [3, 4], '548235'),
          ('Likvidlik', [5], '7F6000'), ('Balans hajmi', [6], '7F7F7F')]
FN_OLD = '2.3.4-jadval ma’lumotlari'
FN_NEW = '2.3.2- va 2.3.4-jadvallar ma’lumotlari'


def _t(v):
    s = f'{abs(v):.1f}'.replace('.', ',')
    return ('+' if v > 0 else '−') + s


def _lbls(idxs, col):
    one = ''.join(
        f'<c:dLbl><c:idx val="{i}"/><c:tx><c:rich><a:bodyPr/><a:lstStyle/><a:p><a:pPr>{W._rpr(1050, True, col)}</a:pPr>'
        f'<a:r><a:rPr lang="uz-Latn-UZ" sz="1050" b="1"><a:solidFill><a:srgbClr val="{col}"/></a:solidFill>'
        f'<a:latin typeface="Times New Roman"/><a:cs typeface="Times New Roman"/></a:rPr><a:t>{_t(VALS[i])}</a:t></a:r></a:p></c:rich></c:tx>'
        f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr><c:dLblPos val="outEnd"/><c:showLegendKey val="0"/><c:showVal val="1"/>'
        f'<c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>' for i in idxs)
    return ('<c:dLbls>' + one + '<c:showLegendKey val="0"/><c:showVal val="0"/><c:showCatName val="0"/>'
            '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


def chart():
    series = [(b, [VALS[i] if i in idx else None for i in range(len(CATS))], c) for b, idx, c in BLOCKS]
    xml, rows = W.bar_chart(CATS, series, ytitle='%', ymin=-50, ymax=200, major=50, gap=60)
    sers = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)
    for s, (b, idx, c) in zip(sers, BLOCKS):
        ns = re.sub(r'<c:dLbls>.*?</c:dLbls>', _lbls(idx, c), s, flags=re.S)
        xml = xml.replace(s, ns, 1)
    xml = xml.replace('<c:gapWidth val="60"/>', '<c:gapWidth val="60"/><c:overlap val="100"/>', 1)
    return xml, rows


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)

    def par(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, bold=False):
        q = dd.add_paragraph(); q.alignment = align
        for t, col, strike in parts:
            r = q.add_run(t); r.font.size = Pt(size); r.bold = bold; r.font.strike = strike
            if col: r.font.color.rgb = RGBColor.from_string(col)
        return q
    par([('2.3.7-rasm (namuna)', None, False)], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, chart())
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 9.5, '2.3.7-rasm')))
    par([(CAP, None, False)], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    dd.add_paragraph()
    par([('Rasmdagi o‘zgarishlar:', None, False)], size=12, bold=True)
    for t in ['1) ustunlar 2.3.1-rasmdagi tafovut zanjiri tartibida joylashtirildi: aktivlar sifati → daromadlilik → kapital → likvidlik; '
              'zanjirga kirmaydigan jami aktivlar oxirida;',
              '2) har bir blok alohida rangda, legendada blok nomi ko‘rsatildi;',
              '3) yorliqlar vergul va ishora bilan: +195,7; −17,9 va h.k.;',
              '4) toifa nomlari matndagi atamalar bilan bir xil qoldirildi (CET1, NSFR matnda qo‘llangan).']:
        par([(t, None, False)], size=12)
    dd.add_paragraph()
    par([('Rasm izohida (366-izoh):', None, False)], size=12, bold=True)
    par([('Tadqiqot natijasida muallif tomonidan ', None, False), (FN_OLD, '7F7F7F', True), (' ', None, False),
         (FN_NEW, 'FF0000', False), (' asosida ishlab chiqildi.', None, False)], size=12)
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

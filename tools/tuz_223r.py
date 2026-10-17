"""2.2.3-rasm: yangi diagramma (raqamlar ko‘rinadigan, qizil erta ogohlantirish chizig‘i) + ostida sonlar jadvali (ko‘k).

Qizil chiziq diagrammada yozuv bilan, jadvalda rangli belgi bilan ko‘rsatilgani uchun alohida izoh qo‘yilmaydi;
rasm, jadval va sarlavha bir sahifada turishi uchun diagramma balandligi 7,5 sm, jadval qatorlari keepNext.
"""
import copy, re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree
import kamlib as K
import wordart as W
import namuna_223r as N
from pardoz import pardoz

K.RED = '0070C0'
F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
D = K.Doc(F)
cap = D.para('2.2.3-rasm. O‘zbekiston bank tizimida aktivlar sifati')
chp = cap._p.getprevious()
rid = [e.get(qn('r:id')) for e in chp.iter('{http://schemas.openxmlformats.org/drawingml/2006/chart}chart')][0]
part = D.d.part.related_parts[rid]
xrel = [r for r in part.rels.values() if r.reltype.endswith('/package')][0]
xml, rows = N.chart()
part._blob = re.sub(r'(<c:externalData r:id=")[^"]+(")', r'\g<1>' + xrel.rId + r'\g<2>', xml).encode('utf-8')
xrel.target_part._blob = W._xlsx(rows)
for ext in chp.iter('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent'):
    ext.set('cx', str(int(16.0 * 360000))); ext.set('cy', str(int(7.5 * 360000)))

# sonlar jadvali: mavjud jadval xossalari (chegaralar) nusxalanadi
allser = N.SER + N.GR
t = D.d.add_table(rows=1 + len(allser), cols=6)
ref = D.table_after('O‘zbekiston bank tizimining moliyaviy barqarorlik ko‘rsatkichlari')
t._tbl.remove(t._tbl.tblPr)
t._tbl.insert(0, copy.deepcopy(ref._tbl.tblPr))
widths = [Cm(7.4)] + [Cm(1.7)] * 5
for gc, wd in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
    gc.set(qn('w:w'), str(int(wd.twips)))
hdr = ['Ko‘rsatkich'] + N.Y
for i, r in enumerate([hdr] + [[s[0]] + [f'{v:.1f}'.replace('.', ',') for v in s[1]] for s in allser]):
    for j, v in enumerate(r):
        c = t.cell(i, j); c.width = widths[j]; p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and i > 0) else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        if j == 0 and i > 0:
            s = allser[i - 1]
            k = p.add_run('━━ '); k.font.size = Pt(11); k.bold = True
            k.font.color.rgb = RGBColor.from_string(s[2])
        run = p.add_run(v); run.font.size = Pt(12); run.bold = (i == 0); run.font.name = 'Times New Roman'
        if i == 0 or j > 0 or True:
            run.font.color.rgb = RGBColor.from_string('0070C0')   # yangi matn – ko‘k
tr0 = t._tbl.findall(qn('w:tr'))[0]
trpr = tr0.find(qn('w:trPr'))
if trpr is None:
    trpr = etree.SubElement(tr0, qn('w:trPr')); tr0.remove(trpr); tr0.insert(0, trpr)
etree.SubElement(trpr, qn('w:tblHeader'))
chp.addnext(t._tbl)
D.save(F)
print(pardoz(F))

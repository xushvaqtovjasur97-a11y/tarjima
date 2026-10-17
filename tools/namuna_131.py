"""1.3.1-rasm tuzatish namunasi: BAITE olib tashlandi, FIMS = SEER, tasnif manbasi aniqlashtirildi, BCBS WP4 birlamchi manba."""
import copy, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pardoz import pardoz

SRC = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_1.3.1-rasm_yangi.docx'
REP = {'(Xalqaro hisob-kitoblar banki tasnifi)': '(Bazel qo‘mitasi tasnifi asosida, muallif to‘ldirgan)',
       'BAITE, RATE': 'RATE',
       'FIMS, SEER,': 'SEER (avval FIMS),'}
BCBS = 'Sahajwala R., Van den Bergh P. Supervisory risk assessment and early warning systems: Basel Committee on Banking Supervision Working Papers No. 4. – Basel: BIS, December 2000. – P. 5–6 (Table 1)'
OLD_T = 'Xalqaro hisob-kitoblar banki G-10 mamlakatlarida qo‘llanadigan nazorat tizimlarini to‘rt toifaga ajratadi'
NEW_T = 'Bazel bank nazorati qo‘mitasi G-10 ning oltita mamlakatida (AQSh, Buyuk Britaniya, Germaniya, Fransiya, Italiya, Niderlandiya) qo‘llanadigan nazorat tizimlarini to‘rt toifaga ajratadi'
OLD_FN = 'Tadqiqot natijasida muallif tomonidan quyidagi manbalar asosida ishlab chiqildi: Устойчивость'
NEW_FN = 'Tadqiqot natijasida muallif tomonidan quyidagi manbalar asosida ishlab chiqildi: ' + BCBS + '; Устойчивость'


def drawing_el(doc):
    cap = [p for p in doc.paragraphs if p.text.startswith('1.3.1-rasm')][0]
    return cap._p.getprevious()


def fix(el):
    n = 0
    for t in el.iter():
        if t.tag.endswith('}t') and t.text in REP:
            t.text = REP[t.text]; n += 1
    return n


if __name__ == '__main__':
    src = docx.Document(SRC)
    el = copy.deepcopy(drawing_el(src))
    print('almashtirildi:', fix(el))
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, color=None, strike=False):
        q = dd.add_paragraph(); q.alignment = align
        r = q.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.strike = strike
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return q

    par('1.3.1-rasm bo‘yicha tuzatishlar (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    par('1. Rasm (oddiy rangda): sarlavha qutichasida tasnif manbasi aniqlashtirildi (Rossiya banki reytingi va EBA indikatorlari BCBS tasnifiga muallif tomonidan qo‘shilgan); mavjud bo‘lmagan “BAITE” olib tashlandi; FIMS SEER ning avvalgi nomi ekani ko‘rsatildi.')
    p = dd.add_paragraph(); p._p.addnext(el); p._p.getparent().remove(p._p)
    par('1.3.1-rasm. Xorijiy mamlakatlarda banklar moliyaviy barqarorligini baholash tizimlarining tasnifi', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    par('2. Rasmdan oldingi xatboshi:')
    par('Hozir: … ' + OLD_T + ': …', color='7F7F7F', strike=True)
    par('Yangi: … ' + NEW_T + ': …', color='C00000')
    par('3. Rasm manbasi (snoska) – birlamchi manba qo‘shiladi:')
    par('Yangi: ' + NEW_FN.replace('; Устойчивость', '; Устойчивость банковской системы … (qolgani o‘zgarmaydi).'), color='C00000')
    dd.save(DST)
    print(pardoz(DST), DST)

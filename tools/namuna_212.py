"""2.1.2-rasm va 2.1.1-jadval tuzatish namunasi: NSFR – monitoring, standart snoska."""
import copy, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pardoz import pardoz

SRC = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_2.1.2-rasm_yangi.docx'
REP = {'leverej;': 'leverej; LCR va milliy',
       'LCR, NSFR va milliy': 'likvidlik normativlari;',
       'likvidlik normativlari;': 'NSFR – monitoring;'}
OLD_FN = 'Muallif ishlanmasi.'
NEW_FN = 'Tadqiqot natijasida muallif tomonidan ishlab chiqildi.'
T_OLD = ['NSFR', '≥ 100%', '≥ 100%', '2018-yildan']
T_NEW = ['NSFR', '≥ 100% (2025-yil 23-maygacha)', '≥ 100%', '2018-yildan; 2025-yil 23-maydan talab kuchini yo‘qotgan, 2026-yildan monitoring']


def drawing_el(doc):
    cap = [p for p in doc.paragraphs if p.text.startswith('2.1.2-rasm')][0]
    return cap._p.getprevious()


def fix(el):
    n = 0
    for t in el.iter():
        if t.tag.endswith('}t') and t.text in REP:
            t.text = REP[t.text]; n += 1
    return n


if __name__ == '__main__':
    src = docx.Document(SRC)
    el = copy.deepcopy(drawing_el(src))
    print('almashtirildi:', fix(el))
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, color=None, strike=False):
        q = dd.add_paragraph(); q.alignment = align
        r = q.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.strike = strike
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return q

    par('2.1.2-rasm va 2.1.1-jadval bo‘yicha tuzatishlar (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    par('1. Rasm (oddiy rangda): “Prudensial normativlar” qutichasida NSFR monitoring sifatida ko‘rsatildi (2025-yil 23-maydan talab kuchini yo‘qotgan, “yo‘l xaritasi”ga ko‘ra 2026-yildan monitoring).')
    p = dd.add_paragraph(); p._p.addnext(el); p._p.getparent().remove(p._p)
    par('2.1.2-rasm. O‘zbekistonda banklar moliyaviy barqarorligini baholashning amaldagi tizimi', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    par('2. Rasm manbasi (snoska):')
    par('Hozir: ' + OLD_FN, color='7F7F7F', strike=True)
    par('Yangi: ' + NEW_FN, color='C00000')
    par('3. 2.1.1-jadval, NSFR qatori (Ko‘rsatkich | Talab | Bazel III minimumi | Izoh):')
    par('Hozir: ' + ' | '.join(T_OLD), color='7F7F7F', strike=True)
    par('Yangi: ' + ' | '.join(T_NEW), color='C00000')
    dd.save(DST)
    print(pardoz(DST), DST)

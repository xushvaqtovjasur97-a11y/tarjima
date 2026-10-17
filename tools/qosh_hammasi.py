"""Navbatdagi ma’qullangan o‘zgarishlarni asosiy faylga kiritadi (qizil; diagramma va chizmalar oddiy rangda):
1.1.6, 1.3.3, 1.3.4, 1.3.5, 1.3.7-rasmlar, 1.3.3-jadval, adabiyotlar, mundarija va hajm.
1.1.5-rasm alohida – tools/qosh_115.py (oldin ishga tushiriladi)."""
import copy, gc, re, subprocess, sys, os, tempfile
sys.path.insert(0, '/home/user/tarjima/tools')
from lxml import etree
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import kamlib as K
import wordart as W
import namuna_116 as N116
import namuna_133 as N133
import namuna_134 as N134
import namuna_135 as N135
import namuna_137 as N137
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
FNM = lambda s: '[[FN:' + s + ']]'
WNS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
CNS = '{http://schemas.openxmlformats.org/drawingml/2006/chart}chart'
EXT = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent'


def fn_text(D, fid):
    for f in D.fx.findall(qn('w:footnote')):
        if f.get(qn('w:id')) == str(fid):
            return ''.join(x.text or '' for x in f.iter(qn('w:t'))).strip()


def drop_fn(D, ref):
    """Izoh havolasini (run) va izohning o‘zini o‘chiradi."""
    fid = ref.get(qn('w:id'))
    run = ref.getparent(); run.getparent().remove(run)
    for f in D.fx.findall(qn('w:footnote')):
        if f.get(qn('w:id')) == fid:
            D.fx.remove(f)


def set_fn(D, fid, new):
    D.replace(D.fn_para(fid), fn_text(D, fid), new)


def replace_chart(D, cap, xml, rows, cy_cm):
    chp = cap._p.getprevious()
    rid = [e.get(qn('r:id')) for e in chp.iter(CNS)][0]
    part = D.d.part.related_parts[rid]
    xrel = [r for r in part.rels.values() if r.reltype.endswith('/package')][0]
    part._blob = re.sub(r'(<c:externalData r:id=")[^"]+(")', r'\g<1>' + xrel.rId + r'\g<2>', xml).encode('utf-8')
    xrel.target_part._blob = W._xlsx(rows)
    chp.find('.//' + EXT).set('cy', str(int(cy_cm * W.EMU_CM)))


def cap_fid(cap):
    return [r.get(qn('w:id')) for r in cap._p.iter(qn('w:footnoteReference'))][0]


def step(name, fn):
    D = K.Doc(F)
    fn(D)
    D.save(F)
    gc.collect()
    print('OK', name)


# ---------------- 1.1.6-rasm
def s116(D):
    T = N116.texts(FNM)
    p139 = D.para('Bu ko‘rsatkichning amaliy samaradorligini L. Chiaramonte va F. Poli tadqiqoti yaqqol tasdiqlaydi (1.1.6-rasm).')
    cap = D.para('1.1.6-rasm. Z-score ko‘rsatkichiga asoslangan modelning Yevropa banklari')
    p142 = D.para('Mualliflar 12 mamlakatdagi 3125 ta bankning 2001–2011-yillardagi')
    fid = cap_fid(cap)
    old142 = p142.text
    D.replace(p139, 'Bu ko‘rsatkichning amaliy samaradorligini L. Chiaramonte va F. Poli tadqiqoti yaqqol tasdiqlaydi (1.1.6-rasm).', T['s139'], tag='116')
    D.replace(cap, 'Z-score ko‘rsatkichiga asoslangan modelning Yevropa banklari muammolarini bashorat qilish aniqligi, %', T['caption'], tag='116')
    set_fn(D, fid, T['cap_fn'])
    for ref in list(p142._p.iter(qn('w:footnoteReference'))):
        drop_fn(D, ref)
    D.replace(p142, old142.strip(), T['a'], tag='116')
    replace_chart(D, cap, N116.xml, N116.rows, 8.5)


# ---------------- 1.3.3-rasm
def s133(D):
    T = N133.texts(FNM)
    OLD = 'O.I. Lavrushin va I.D. Mamonova keltirgan ma’lumotlarga ko‘ra, 48 306 kuzatuv asosida sinovdan o‘tkazilgan FIMS modeli avvalgi UBSS tizimiga nisbatan bashorat aniqligini sezilarli oshirgan (1.3.3-rasm).'
    pF = D.para(OLD)
    cap = D.para('1.3.3-rasm. AQShda FIMS va UBSS modellarining bashorat aniqligi')
    pA = D.para('Rasm ma’lumotlari FIMS modeli to‘g‘ri bashorat ulushini')
    fid = cap_fid(cap)
    oldA = pA.text
    D.replace(pF, OLD, T['s_fims'], tag='133')
    D.replace(cap, 'AQShda FIMS va UBSS modellarining bashorat aniqligi, %', T['caption'], tag='133')
    set_fn(D, fid, T['cap_fn'])
    for ref in list(pA._p.iter(qn('w:footnoteReference'))):
        drop_fn(D, ref)
    D.replace(pA, oldA.strip(), T['a'], tag='133')
    replace_chart(D, cap, N133.xml, N133.rows, 8.5)


# ---------------- 1.3.4-rasm
def s134(D):
    T = N134.texts(FNM)
    p = D.para('Angliya banki 1997-yildan BAITE tizimini qo‘llab kelmoqda')
    cap = D.para('1.3.4-rasm. Angliya bankining BAITE tizimida risk va nazorat matritsasi')
    pA = D.para('Matritsadan ko‘rinadiki, riski past va ichki nazorati kuchli bank')
    fid = cap_fid(cap)
    txt = p.text
    old = txt[txt.index('Angliya banki 1997-yildan BAITE'):txt.index('(1.3.4-rasm).') + len('(1.3.4-rasm).')]
    lav = [fn_text(D, r.get(qn('w:id'))) for r in p._p.iter(qn('w:footnoteReference'))]
    assert len(lav) == 2 and all('Лаврушин' in x for x in lav), lav
    for ref in list(p._p.iter(qn('w:footnoteReference'))):
        drop_fn(D, ref)
    s_uk = T['s_uk'].replace('3 yilgacha belgilangan.', '3 yilgacha belgilangan' + FNM(lav[0]) + FNM(lav[1]) + '.')
    assert s_uk != T['s_uk']
    D.replace(p, old, s_uk, tag='134')
    D.replace(cap, 'Angliya bankining BAITE tizimida risk va nazorat matritsasi', T['caption'], tag='134')
    set_fn(D, fid, T['cap_fn'])
    oldA = pA.text
    D.replace(pA, oldA.strip(), T['a'], tag='134')
    # chizma: eski matritsa o‘rniga PIF bosqichlari
    drw = cap._p.getprevious()
    for r in [r for r in drw if r.tag == qn('w:r')]:
        drw.remove(r)
    W._counter['docpr'] = 9600; W._counter['shape'] = 9600
    run, h = N134.figure()
    drw.append(etree.fromstring(run))


# ---------------- 1.3.5-rasm
def s135(D):
    xml, part = N135.chart_xml(D.d)
    p = D.para(N135.OLD_S)
    bib = [q for q in D.paras() if re.match(r'^\d+\. ', q.text) and '// Интерфакс. – 2025' in q.text]
    part._blob = xml.encode('utf-8')
    xrel = [r for r in part.rels.values() if r.reltype.endswith('/package')][0]
    xrel.target_part._blob = W._xlsx(N135.rows_from(xml))
    D.replace(p, N135.OLD_S, N135.NEW_S + FNM(N135.PR2024) + '.', tag='135')
    for f in D.fx.findall(qn('w:footnote')):
        for fp in f.findall(qn('w:p')):
            if '// Интерфакс. – 2025' in ''.join(x.text or '' for x in fp.iter(qn('w:t'))):
                D.replace(Paragraph(fp, None), '// Интерфакс. – 2025', '// Интерфакс. – 2023. – 23 ноября', tag='135')
    for q in bib:
        D.replace(q, '// Интерфакс. – 2025', '// Интерфакс. – 2023. – 23 ноября', tag='135')


# ---------------- 1.3.7-rasm
def s137(D):
    T = N137.texts(FNM)
    p = D.para('So‘nggi yillar ma’lumotlari ham vaziyat tubdan o‘zgarmaganini ko‘rsatadi.')
    cap = D.para('1.3.7-rasm. Xitoy tijorat banklarida muammoli kreditlar ulushi')
    OLDA = ('Rasmdan ko‘rinadiki, rasmiy muammoli kreditlar ulushi uch yil davomida 1,5 foiz atrofida barqaror turibdi. Ammo alohida e’tibor talab qiluvchi kreditlar qo‘shilsa, bu ko‘rsatkich ikki yarim barobar, ilmiy baholarga ko‘ra esa uch barobargacha ortadi.')
    pA = D.para(OLDA)
    fid = cap_fid(cap)
    refs = [r for r in p._p.iter(qn('w:footnoteReference')) if 'Supervisory Statistics' in (fn_text(D, r.get(qn('w:id'))) or '')]
    assert len(refs) == 1
    drop_fn(D, refs[0])
    txt = p.text
    old = txt[:txt.index('3,7 foizga yetadi.') + len('3,7 foizga yetadi.')]
    D.replace(p, old, T['s1'], tag='137')
    D.replace(cap, 'Xitoy tijorat banklarida muammoli kreditlar ulushi: rasmiy ma’lumotlar va muqobil baholar, %', T['caption'], tag='137')
    set_fn(D, fid, T['cap_fn'])
    D.replace(pA, OLDA, T['a'], tag='137')
    replace_chart(D, cap, N137.xml, N137.rows, 8.5)


# ---------------- 1.3.3-jadval
def tab133(D):
    t = D.table_after('Xorijiy tajriba va undan O‘zbekiston uchun kelib chiqadigan xulosalar')
    pB = D.para('Buyuk Britaniyada esa risk matritsasi tekshiruv davriyligini belgilaydi')
    R = {r.cells[0].text.strip(): i for i, r in enumerate(t.rows)}
    i = R['AQSh – FIMS']
    D.set_cell(t, i, 0, 'AQSh – FIMS; Yevropa Ittifoqi – EBA modellari', tag='T133')
    D.set_cell(t, i, 1, 'Statistik va mashinali o‘qitish modellari zaif banklarni yuqori aniqlikda bashorat qiladi (AUC 0,91–0,93), chegara xatolar narxiga qarab tanlanadi', tag='T133')
    i = R['Buyuk Britaniya – BAITE, RATE']
    D.set_cell(t, i, 0, 'Buyuk Britaniya – RATE, PRA', tag='T133')
    D.set_cell(t, i, 1, 'Nazorat intensivligi bankning tizimga ta’siri toifasi va faol aralashuv tizimining bosqichiga bog‘langan', tag='T133')
    D.replace(t.rows[R['Yevropa Ittifoqi – EBA KRI']].cells[2].paragraphs[0], ', kredit riskiga ortiqcha urg‘u bermaslik', '', tag='T133')
    D.insert(t.rows[R['Germaniya – BAKIS']].cells[2].paragraphs[0], '', ', kredit riskiga ortiqcha urg‘u bermaslik', tag='T133')
    D.insert(t.rows[R['Rossiya']].cells[1].paragraphs[0], '', '; 2025-yildan tizimli banklar uchun Bazel LCR o‘rniga milliy normativ', tag='T133')
    D.insert(t.rows[R['Rossiya']].cells[2].paragraphs[0], '', '; xalqaro standartlarni milliy bozor imkoniyatlariga moslashtirish', tag='T133')
    D.insert(t.rows[R['Xitoy']].cells[1].paragraphs[0], '', '; 2024-yildan banklar hajmiga qarab proporsional talablar', tag='T133')
    D.insert(t.rows[R['Xitoy']].cells[2].paragraphs[0], '', '; kichik banklar uchun soddalashtirilgan talablar', tag='T133')
    D.replace(pB, 'Buyuk Britaniyada esa risk matritsasi tekshiruv davriyligini belgilaydi', 'Buyuk Britaniyada esa faol aralashuv tizimining bosqichi majburiy nazorat choralarini belgilaydi', tag='T133')


# ---------------- adabiyotlar
def bib(D):
    entries = [p for p in D.paras() if re.match(r'^\d+\. ', p.text)]
    heads = [p for p in D.paras() if re.match(r'^(I|II|III|IV|V|VI)\. ', p.text.strip())]

    def section(start):
        h = [p for p in heads if p.text.strip().startswith(start)][-1]
        allp = D.paras(); i0 = [k for k, q in enumerate(allp) if q._p is h._p][0]
        out = []
        for q in allp[i0 + 1:]:
            if re.match(r'^(I|II|III|IV|V|VI)\. ', q.text.strip()) or q.text.strip() == 'ILOVALAR':
                break
            if re.match(r'^\d+\. ', q.text):
                out.append(q)
        return out

    def add_sorted(start, text):
        sec = section(start)
        lat = [q for q in sec if re.match(r'^\d+\. [A-Za-zČ]', q.text)]
        key = text.lower()
        prev = None
        for q in lat:
            if re.sub(r'^\d+\. ', '', q.text).lower() > key:
                break
            prev = q
        D.new_para_after(prev or lat[0], '0. ' + text, tag='BIB')
        gc.collect()

    add_sorted('II.', N134.PRA)
    add_sorted('IV.', N133.SRC)
    a = [p for p in entries if 'Банк России устанавливает порядок выхода из послабления' in p.text][-1]
    D.new_para_after(a, '0. ' + N135.PR2024 + '.', tag='BIB')


def renumber(D):
    start = False; n = 0
    for p in D.paras():
        if p.text.strip() == 'FOYDALANILGAN ADABIYOTLAR RO‘YXATI':
            start = True; continue
        if start and p.text.strip() == 'ILOVALAR':
            break
        if start:
            m = re.match(r'^(\d+)\. ', p.text)
            if m:
                n += 1
                r0 = D._runs(p)[0].find(qn('w:t'))
                if r0.text.startswith(m.group(1) + '.'):
                    r0.text = str(n) + r0.text[len(m.group(1)):]
                elif r0.text == m.group(1):
                    r0.text = str(n)
    hp = D.para('Dissertatsiyaning tuzilishi va hajmi.')
    D.replace(hp, re.search(r'ro‘yxati \d+ ta manbani', hp.text).group(0), f'ro‘yxati {n} ta manbani', tag='BIB')
    print('adabiyotlar:', n)


if __name__ == '__main__':
    for nm, f in (('1.1.6', s116), ('1.3.3', s133), ('1.3.4', s134), ('1.3.5', s135), ('1.3.7', s137), ('1.3.3-jadval', tab133), ('adabiyot', bib), ('raqamlash', renumber)):
        step(nm, f)
    print('pardoz:', pardoz(F))

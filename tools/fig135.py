import sys
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import wordart as W
BLUE,ORANGE,GREEN,DBLUE='2A78D6','EB6834','70AD47','1F4E79'
d=docx.Document()
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
Y=list(range(2015,2027))
data=[('Bazel qo‘mitasi',[60,70,80,90]+[100]*8,BLUE,'circle'),
      ('Rossiya',[None,70,80,90,100,100,100,None,None,100,100,100],ORANGE,'square'),
      ('O‘zbekiston',[None,80,90]+[100]*9,GREEN,'triangle')]
for k,(nm,ys,col,mk) in enumerate(data):
  t=d.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER; t.paragraph_format.space_before=Pt(10)
  rr=t.add_run(('a','b','c')[k]+') '+nm); rr.bold=True; rr.font.name='Times New Roman'; rr.font.size=Pt(12)
  p=d.add_paragraph('{{FIG%d}}'%k); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  show=[i for i,v in enumerate(ys) if v is not None]
  xml,rows=W.scatter_chart([(nm,Y,ys,col,28575,None,mk,('t',show))],
    xtitle=None,ytitle='%',xmin=2014,xmax=2027,xmajor=1,ymin=50,ymax=115,ymajor=10,xfmt='0',legend=False)
  rid=W.add_chart(d,(xml,rows))
  W.replace_placeholder(d,'{{FIG%d}}'%k,W.chart_run_xml(rid,16.0,5.2,'1.3.5-rasm-'+str(k)))
c=d.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=c.add_run('1.3.5-rasm. Likvidlikni qoplash koeffitsiyentiga qo‘yiladigan talablar: islohotning boshlanishi va davomi, %')
r.bold=True; r.font.name='Times New Roman'; r.font.size=Pt(14)
n=c.add_paragraph if False else None
d.save('/home/user/tarjima/paragraflar/rasmlar/1.3.5-rasm.docx'); print('ok')

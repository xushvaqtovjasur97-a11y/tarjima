"""2.3.5-rasm: yashirin kredit riski koeffitsiyenti – chiziqli diagramma, yorliqlar ustma-ust tushmaydi (namuna)."""
import re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_2.3.5-rasm.docx'
Y = ['2021', '2022', '2023', '2024', '2025']
SER = [('“O‘zsanoatqurilishbank”', [1.65, 2.61, 2.14, 2.08, 2.96], '2A78D6'),
       ('“Mikrokreditbank”', [2.95, 2.81, 2.77, 3.30, 2.80], 'EB6834'),
       ('“Aloqabank”', [2.46, 4.90, 4.01, 4.01, 2.83], '70AD47'),
       ('Muvofiqlik darajasi (1,0)', [1.0] * 5, '7F7F7F', 'dash')]
# har bir nuqta yorlig‘ining joyi
POS = [['b', 'b', 'b', 'b', 't'],
       ['t', 't', 'b', 't', 'b'],
       ['b', 't', 't', 't', 'r']]


def dl(col, poss):
    pts = ''.join(f'<c:dLbl><c:idx val="{i}"/><c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
                  f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr>{W._rpr(1000, True, col)}</a:pPr><a:endParaRPr lang="uz-Latn-UZ"/></a:p></c:txPr>'
                  f'<c:dLblPos val="{p}"/><c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/>'
                  f'<c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbl>' for i, p in enumerate(poss))
    return (f'<c:dLbls>{pts}<c:numFmt formatCode="0.00" sourceLinked="0"/><c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
            f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr>{W._rpr(1000, True, col)}</a:pPr><a:endParaRPr lang="uz-Latn-UZ"/></a:p></c:txPr>'
            '<c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/><c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


def chart():
    xml, rows = W.line_table_chart(Y, SER, ytitle='Koeffitsiyent', ymin=0, ymax=5.5, major=1, fmt='0.0', w=28575, dtable=False)
    xml = xml.replace('<c:tickLblPos val="none"/>', '<c:tickLblPos val="low"/>', 1)
    sers = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)
    for s, sr, p in zip(sers, SER, POS):
        xml = xml.replace(s, s.replace('<c:cat>', dl(sr[2], p) + '<c:cat>', 1), 1)
    s3 = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)[3]
    xml = xml.replace(s3, re.sub(r'<c:marker>.*?</c:marker>', '<c:marker><c:symbol val="none"/></c:marker>', s3, flags=re.S), 1)
    xml = xml.replace('<c:size val="5"/>', '<c:size val="6"/>')
    xml = xml.replace('<c:numFmt formatCode="0" sourceLinked="0"/>', '<c:numFmt formatCode="0.0" sourceLinked="0"/>', 1)
    xml = xml.replace('</c:plotArea>', '</c:plotArea>' + W._legend('b'), 1)
    return xml, rows


OLD = 'Uning ikkala bank bo‘yicha dinamikasi 2.3.5-rasmda keltirildi.'
NEW = 'Uning uch bank bo‘yicha dinamikasi 2.3.5-rasmda keltirildi.'

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, color=None, strike=False):
        q = dd.add_paragraph(); q.alignment = align
        r = q.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.strike = strike
        if color:
            r.font.color.rgb = RGBColor.from_string(color)

    par('2.3.5-rasm (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, chart())
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.5, '2.3.5-rasm')))
    par('2.3.5-rasm. “O‘zsanoatqurilishbank” ATB, “Mikrokreditbank” ATB va “Aloqabank” ATB da yashirin kredit riski koeffitsiyenti',
        bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    par('Rasmdan oldingi xatboshida:')
    par('Hozir: ' + OLD, color='7F7F7F', strike=True)
    par('Yangi: ' + NEW, color='FF0000')
    dd.save(DST)
    print(pardoz(DST), DST)

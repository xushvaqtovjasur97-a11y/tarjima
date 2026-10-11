import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/3.3-paragraf.md'; DST='/home/user/tarjima/paragraflar/3.3-paragraf.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p33_raw.docx'],check=True)
d=docx.Document(S+'/p33_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^3\.3\.\d-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^3\.3\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^3\.3\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'
PURPLE='7030A0'

def setw(t,ws):
  for row in t.rows:
    for k,c in enumerate(row.cells): c.width=Cm(ws[k])
  for k,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(ws[k]*567)))
  t.autofit=False
for t in d.tables:
  h=[c.text.strip() for c in t.rows[0].cells]; n=len(h)
  if h[0]=='Gipoteza' and n==4: setw(t,[1.6,7.0,4.6,3.2])
  elif h[0]=='Gipoteza': setw(t,[3.6,3.4,2.6,2.6,4.2])
  elif n==7: setw(t,[4.4,2.0,2.0,2.0,2.0,2.0,2.0])
  elif n==5: setw(t,[3.6,3.2,3.2,3.2,3.2])
  elif n==4: setw(t,[4.0,4.1,4.1,4.2])
  elif n==3: setw(t,[6.0,5.2,5.2])
  for row in t.rows[1:]:
    for c in row.cells[1:]:
      if h[0]!='Gipoteza':
        for p in c.paragraphs: p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
# "Izoh:" paragraflari – 12 pt, xatboshisiz
for p in d.paragraphs:
  if p.text.startswith('Izoh:'):
    p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(6)
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.font.size=Pt(12)
  if p.text.startswith('MBII') and '=' in p.text:
    p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0)
    for r in p.runs: r.italic=True; r.font.size=Pt(13)
import json
O=json.load(open('/home/user/tarjima/manbalar/hisob_kitob/model33.json'))['guruh']
Y=['2021','2022','2023','2024','2025']
ser=[('Davlat banklari',[2021,2022,2023,2024,2025],[O[y+'_1']['MBII'] for y in Y],BLUE,28575,None,'circle',('b',[0,1,2,3,4])),
     ('Boshqa banklar',[2021,2022,2023,2024,2025],[O[y+'_0']['MBII'] for y in Y],ORANGE,28575,None,'square',('t',[0,1,2,3,4]))]
rid=W.add_chart(d,W.scatter_chart(ser,'Yillar','MBII, ball',2020.5,2025.5,1,2.5,4.0,0.5,xfmt='0',yfmt='0.00',legend=True))
W.replace_placeholder(d,'{{FIG:1}}',W.chart_run_xml(rid,16.0,8.0,'3.3.1-rasm'))
BN=['Kapital','Aktivlar sifati','Likvidlik','Daromadlilik']; BL=['bK','bA','bL','bD']
rid=W.add_chart(d,W.bar_chart(BN,[('Davlat banklari',[O['2025_1'][b] for b in BL],BLUE),('Boshqa banklar',[O['2025_0'][b] for b in BL],ORANGE)],ytitle='Ball',ymax=4.5,ymin=0,major=1,label_fmt='0.00',gap=60))
W.replace_placeholder(d,'{{FIG:2}}',W.chart_run_xml(rid,16.0,8.0,'3.3.2-rasm'))
YY=['2022','2023','2024','2025']
rid=W.add_chart(d,W.bar_chart(YY,[('Davlat banklari',[O[y+'_1']['lnZ'] for y in YY],BLUE),('Boshqa banklar',[O[y+'_0']['lnZ'] for y in YY],ORANGE)],ytitle='ln(Z-score)',ymax=3.5,ymin=0,major=0.5,label_fmt='0.00',gap=70))
W.replace_placeholder(d,'{{FIG:3}}',W.chart_run_xml(rid,16.0,8.0,'3.3.3-rasm'))
d.save(DST); print('saved')

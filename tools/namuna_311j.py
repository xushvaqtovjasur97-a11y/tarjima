"""3.1.1-jadval: 2- va 5-qatorlardagi dalillar tuzatilgan (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_3.1.1-jadval.docx'
FIX = {1: ('2.3-§: yashirin kredit riski koeffitsiyenti 1,65–3,30', '2.3-§: yashirin kredit riski koeffitsiyenti 1,65–4,90'),
       4: ('2.1-§: normativlar faqat bajarilgan yoki bajarilmagani bo‘yicha nazorat qilinadi',
           '2.1-§: erta ogohlantirish va tezkor tuzatuvchi choralar tizimi mavjud emas, chora normativ buzilgandan keyingina ko‘riladi')}

src = docx.Document('/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx')
cap = [p for p in src.paragraphs if p.text.startswith('Amaldagi baholash tizimi muammolari')][0]
tbl = docx.table.Table(cap._p.getnext(), src)
rows = [[c.text for c in r.cells] for r in tbl.rows]

dd = docx.Document()
sec = dd.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(0)

def par(text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    q = dd.add_paragraph(); q.alignment = align
    r = q.add_run(text); r.bold = bold

par('3.1.1-jadval', bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
par(cap.text, bold=True)
t = dd.add_table(rows=len(rows), cols=3); t.style = 'Table Grid'
for i, row in enumerate(rows):
    for j, txt in enumerate(row):
        c = t.cell(i, j); q = c.paragraphs[0]
        if j == 1 and i in FIX:
            old, new = FIX[i]
            r = q.add_run(old); r.font.strike = True; r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F); r.font.size = Pt(11)
            q2 = c.add_paragraph(); r = q2.add_run(new); r.font.color.rgb = RGBColor(0xFF, 0, 0); r.font.size = Pt(11)
        else:
            r = q.add_run(txt); r.font.size = Pt(11); r.bold = (i == 0)
            if i == 0: q.alignment = WD_ALIGN_PARAGRAPH.CENTER
for j, w in enumerate([5.0, 5.5, 6.0]):
    for c in t.columns[j].cells: c.width = Cm(w)
dd.save(DST)
print(pardoz(DST), DST)

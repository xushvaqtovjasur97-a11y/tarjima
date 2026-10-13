"""1.2.5-rasm o‘rniga “standart × blok” matritsa-jadval (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_1.2.5-matritsa.docx'
BLOK = ['Kapital yetarliligi', 'Aktivlar sifati', 'Daromadlilik', 'Likvidlik', 'Bozor (valyuta) riskiga sezgirlik']
# ● – asosiy mezon manbai, ○ – bilvosita ta’sir, – – qamramaydi
M = [
    ('Bazel III', ['● CET1 ≥ 4,5%, jami ≥ 8% (+2,5% bufer), leverej ≥ 3%', '–', '–', '● LCR ≥ 100%, NSFR ≥ 100%', '○ bozor riski uchun kapital talabi']),
    ('MHXS 9', ['○ kutilgan zarar zaxiralari orqali', '● kutilgan kredit zararlari, 3 bosqichli tasnif', '○ zaxira xarajatlari orqali', '–', '–']),
    ('XVJ moliyaviy barqarorlik ko‘rsatkichlari', ['● regulyativ kapital / RWA', '● muammoli kreditlar / jami kreditlar', '● ROA, ROE', '● likvid aktivlar / jami aktivlar', '● sof ochiq valyuta pozitsiyasi / kapital']),
    ('CAMELS', ['● C – kapital', '● A – aktivlar sifati', '● E – daromadlilik', '● L – likvidlik', '● S – bozor riskiga sezgirlik']),
    ('Reyting agentliklari', ['● kapitallashuv', '● aktivlar sifati', '● daromadlilik', '● moliyalashtirish va likvidlik', '● xorijiy valyutadagi likvidlik']),
]


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)
    tcPr.append(sh)


def put(cell, text, bold=False, size=11, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = cell.paragraphs[0]; p.alignment = align
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    q.add_run('1.2.5-jadval').italic = True
    c = dd.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.add_run('Xalqaro standartlarning moliyaviy barqarorlikni integral baholash bloklari bilan bog‘liqligi').bold = True

    t = dd.add_table(rows=1 + len(M) + 1, cols=6); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    W_ = [Cm(3.0)] + [Cm(2.7)] * 5
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), W_):
        gc.set(qn('w:w'), str(int(w.twips)))
    for row in t.rows:
        for i, cl in enumerate(row.cells):
            cl.width = W_[i]
    put(t.cell(0, 0), 'Standart / blok', bold=True); shade(t.cell(0, 0), 'D9E2F3')
    for j, b in enumerate(BLOK, 1):
        put(t.cell(0, j), b, bold=True, size=10); shade(t.cell(0, j), 'D9E2F3')
    for i, (s, cells) in enumerate(M, 1):
        put(t.cell(i, 0), s, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT); shade(t.cell(i, 0), 'F2F2F2')
        for j, v in enumerate(cells, 1):
            put(t.cell(i, j), v, size=9.5)
    last = len(M) + 1
    put(t.cell(last, 0), 'Integral bahoga hissasi', bold=True, align=WD_ALIGN_PARAGRAPH.LEFT); shade(t.cell(last, 0), 'D9E2F3')
    roles = ['Bazel III mezonlari', 'MHXS 9 bo‘yicha baholash', 'XVJ va CAMELS ko‘rsatkichlari', 'Bazel III mezonlari', 'XVJ ko‘rsatkichi, dollarlashuvni hisobga olish']
    for j, v in enumerate(roles, 1):
        put(t.cell(last, j), v, bold=True, size=9.5); shade(t.cell(last, j), 'D9E2F3')

    n = dd.add_paragraph(); n.paragraph_format.line_spacing = 1.0
    r = n.add_run('Izoh: ● – blok uchun asosiy mezon yoki ko‘rsatkich manbai; ○ – bilvosita ta’sir; – – blokni qamramaydi.'); r.font.size = Pt(11); r.italic = True
    f = dd.add_paragraph(); f.paragraph_format.line_spacing = 1.0
    r = f.add_run('Manba: Tadqiqot natijasida muallif tomonidan ishlab chiqildi.'); r.font.size = Pt(11)
    dd.save(DST)
    print(pardoz(DST), DST)

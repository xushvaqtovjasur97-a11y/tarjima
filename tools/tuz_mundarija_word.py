"""Mundarija sahifa raqamlari va Kirishdagi hajm – Word sahifalashiga moslangan render (render_uno.py, izoh ajratgichi 4,1+4,1 mm) bo‘yicha.
Foydalanish: python3 tuz_mundarija_word.py [docx]"""
import sys, subprocess, tempfile, os, re
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.oxml.ns import qn
import kamlib as K
F = sys.argv[1] if len(sys.argv) > 1 else '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'

def render(path):
    pdf = os.path.join(tempfile.mkdtemp(), 'm.pdf')
    subprocess.run(['python3', '/home/user/tarjima/tools/sahifalash/render_uno.py', path, pdf, '--fnline', '4.1', '--fntext', '4.1'],
                   capture_output=True, timeout=900)
    n = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout).group(1))
    return [' '.join(subprocess.run(['pdftotext', '-f', str(p), '-l', str(p), pdf, '-'], capture_output=True, text=True).stdout.split())
            for p in range(1, n + 1)]

def mundarija(path, pages):
    d = docx.Document(path)
    start = [i for i, p in enumerate(d.paragraphs) if p.text.strip() == 'MUNDARIJA'][0]
    cur, changed = 2, 0
    for p in d.paragraphs[start + 1:]:
        if '\t' not in p.text: break
        title, pg = p.text.rsplit('\t', 1)
        key = ' '.join(title.split()[:6])
        hit = next((i + 1 for i in range(cur, len(pages)) if key in pages[i]), None)
        if hit is None: continue
        cur = hit - 1
        if str(hit) != pg.strip():
            t = [x for x in p._p.iter(qn('w:t'))][-1]
            if t.text.strip() == pg.strip(): t.text = str(hit); changed += 1
    d.save(path)
    return changed

pages = render(F)
bib = next(i + 1 for i in range(5, len(pages)) if pages[i].startswith('FOYDALANILGAN ADABIYOTLAR'))
vol = bib - 4
D = K.Doc(F); K.RED = '000000'
p = D.para('Dissertatsiyaning asosiy matni')
m = re.search(r'asosiy matni (\d+) betni', p.text)
if m.group(1) != str(vol):
    D.replace(p, m.group(1) + ' betni', f'{vol} betni')
D.save(F)
print('adabiyotlar', bib, 'hajm', vol, 'mundarija:', mundarija(F, pages))
pages = render(F)
print('2-bosqich mundarija:', mundarija(F, pages))

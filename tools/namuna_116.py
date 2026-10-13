"""1.1.6-rasmning yangi varianti (namuna): buxgalteriya ko‘rsatkichlariga asoslangan modellarning
AQSh banklari muvaffaqiyatsizligini bashorat qilish aniqligi (AUC), 1959–2024-yillar.
Manba: Correia S., Luck S., Verner E. Failing Banks (QJE, 2026; arXiv:2506.06082, Table 1, Panel B)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_1.1.6-rasm_yangi.docx'

# Table 1, Panel B (1959–2024), 1 yillik ufq: (1) faqat to‘lovga qobiliyatsizlik, (2) faqat noasosiy moliyalashtirish,
# (3) ikkalasi va ularning o‘zaro ta’siri, (4) + o‘sish va umumiy sharoitlar
CATS = ['Faqat sof foyda /\naktivlar', 'Faqat noasosiy\nmoliyalashtirish', 'Ikkala ko‘rsatkich\nbirgalikda', 'To‘liq model']
INS = [0.949, 0.846, 0.954, 0.953]
OOS = [0.942, 0.792, 0.946, 0.945]
SER = [('Model baholangan ma’lumotlarda', INS, 'EB6834'), ('Keyingi davrlarda (tanlanmadan tashqari)', OOS, '2A78D6')]
xml, rows = W.bar_chart(CATS, SER, ytitle='AUC', ymax=1.0, ymin=0.5, major=0.1, fmt='0.0', label_fmt='0.00', gap=80)

SRC = 'Correia S., Luck S., Verner E. Failing Banks // Quarterly Journal of Economics. – 2026. – Vol. 141, No. 1. – P. 147–204; ishchi varianti: arXiv:2506.06082. – 2025.'


def texts(fn):
    s139 = 'Bu ko‘rsatkich tarkibiy qismlarining amaliy samaradorligini S. Korreya, S. Lak va E. Verner tadqiqoti yaqqol tasdiqlaydi (1.1.6-rasm).'
    caption = 'Buxgalteriya ko‘rsatkichlariga asoslangan modellarning AQSh banklari faoliyati to‘xtatilishini bashorat qilish aniqligi (AUC), 1959–2024-yillar'
    cap_fn = 'Tadqiqot natijasida muallif tomonidan quyidagi manba ma’lumotlari asosida ishlab chiqildi: ' + SRC[:-1] + ' – P. 30. Izoh: AUC – ROC egri chizig‘i ostidagi maydon (0,5 – tasodifiy bashorat, 1,0 – mukammal bashorat); bashorat ufqi – bir yil.'
    a = ('Mualliflar AQShning 1959–2024-yillardagi 23 mingdan ortiq banki ma’lumotlari asosida faqat sof foydaning aktivlarga nisbati, ya’ni Z-score ning rentabellik tarkibiy qismi banklar faoliyatining bir yil ichida to‘xtatilishini 0,95 AUC bilan bashorat qilishini, bu aniqlik model baholanishida foydalanilmagan keyingi davrlarda ham 0,94 darajasida saqlanishini ko‘rsatgan. Noasosiy moliyalashtirish ulushining o‘zi kamroq aniqlik beradi (0,85), ikkala ko‘rsatkich birgalikda esa eng yuqori natijaga erishadi (0,95). Mualliflar hisob-kitobiga ko‘ra, noto‘g‘ri signallar ulushi 10 foiz bo‘lganda faoliyati to‘xtatilgan banklarning 92 foizini oldindan aniqlash mumkin'
         + fn(SRC[:-1] + ' – P. 1, 30–31.')
         + '. Tadqiqot xulosasiga ko‘ra, banklar faoliyatining to‘xtatilishi ochiq moliyaviy hisobotlardagi oddiy buxgalteriya ko‘rsatkichlari yordamida yuqori aniqlikda bashorat qilinadi. Bu xulosa aksiyalar bozori rivojlanmagan O‘zbekiston sharoitida buxgalteriya ma’lumotlariga asoslangan integral baholash usullarining afzalligini asoslaydi.')
    return dict(s139=s139, caption=caption, cap_fn=cap_fn, a=a)


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=14, color=None):
        q = dd.add_paragraph(); q.alignment = align
        q.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
        r = q.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)

    FN = []

    def fn(t):
        FN.append(t); return f'[{len(FN)}]'

    FN.append(texts(lambda x: '')['cap_fn'])
    T = texts(fn)
    par('1.1.6-rasmning yangi varianti (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    par('Ma’qul bo‘lsa: 1.1-§ dagi “Bu ko‘rsatkichning amaliy samaradorligini L. Chiaramonte va F. Poli tadqiqoti yaqqol tasdiqlaydi (1.1.6-rasm).” jumlasi quyidagi birinchi jumla bilan, 1.1.6-rasm va undan keyingi “Mualliflar 12 mamlakatdagi …” xatboshi esa quyidagi rasm va tahlil bilan almashtiriladi.',
        italic=True, size=12, indent=False, color='595959')
    par('… ' + T['s139'])
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, (xml, rows))
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.5, '1.1.6-rasm')))
    par('1.1.6-rasm. ' + T['caption'] + '[1]', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    par(T['a'])
    par('Sahifa osti snoskalari', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    for i, t in enumerate(FN, 1):
        par(f'[{i}] {t}', size=11, indent=False)
    dd.save(DST)
    print(pardoz(DST), DST)

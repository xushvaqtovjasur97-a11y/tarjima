"""4.1.1-rasm: ROC egri chiziqlari – to‘liq empirik ROC (pog‘onali), MBII-8 chegaralari aynan egri chiziq ustida,
har bir chegara alohida belgi bilan legendada (yozuvlar bir-birini qoplamaydi) – namuna."""
import json, re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W

DST = '/home/user/tarjima/Namuna_4.1.1-rasm.docx'
O = json.load(open('/home/user/tarjima/manbalar/hisob_kitob/ews41.json'))
BLUE, ORANGE, GREEN, GREY = '2A78D6', 'EB6834', '70AD47', '808080'
CAP = '4.1.1-rasm. Integral indeks, muammoli kreditlar ulushi va Z-score bo‘yicha ROC egri chiziqlari'
W_CM, H_CM = 16.0, 9.5
PL = (0.085, 0.03, 0.895, 0.80)        # chizish maydoni (x, y, w, h) – diagramma ulushida
LG = (0.50, 0.465, 0.475, 0.37)        # legenda: chizish maydonining bo‘sh o‘ng-pastki qismida (II tur xato > 47 %, ulush < 45 %)


def corners(pts):
    """Empirik ROC nuqtalari (ulush) → foiz; bir chiziqdagi oraliq nuqtalar tashlanadi."""
    p = []
    for fp, tp in pts:
        q = (round(100 * fp, 2), round(100 * tp, 2))
        if not p or q != p[-1]:
            p.append(q)
    if p[-1] != (100.0, 100.0):
        p.append((100.0, 100.0))
    out = [p[0]]
    for a, b in zip(p[1:], p[2:] + [None]):
        if b is not None and ((out[-1][0] == a[0] == b[0]) or (out[-1][1] == a[1] == b[1])):
            continue
        out.append(a)
    return out


def ser(name, pts, color, w, dash=None, marker=None):
    return (name, [x for x, _ in pts], [y for _, y in pts], color, w, dash, marker)


def chart():
    R = O['ROC']
    sg = {round(r['th'], 2): r for r in O['signal']}
    mb = corners(R['MBII'])
    S = [ser('MBII-8 (AUC = 0,86)', mb, BLUE, 28575),
         ser('Muammoli kreditlar ulushi (AUC = 0,86)', corners(R['Muammoli kreditlar ulushi']), GREEN, 22225, 'dash'),
         ser('ln(Z-score) (AUC = 0,42)', corners(R['ln(Z-score)']), ORANGE, 28575),
         ser('Tasodifiy bashorat', [(0, 0), (100, 100)], GREY, 9525, 'sysDash')]
    # LibreOffice: keyingi seriya uzunroq bo‘lsa oldingilari chizilmaydi – chiziqlar oxirgi nuqta bilan tenglashtiriladi
    n = max(len(s[1]) for s in S)
    S = [(a, xs + [xs[-1]] * (n - len(xs)), ys + [ys[-1]] * (n - len(ys)), *rest) for a, xs, ys, *rest in S]
    TH = [(3.25, 'square', '7030A0', 'MBII-8 = 3,25 (monitoring)'),
          (2.85, 'triangle', '000000', 'MBII-8 = 2,85 (optimal)'),
          (2.75, 'diamond', 'C00000', 'MBII-8 = 2,75 (nazorat choralari)')]
    for th, mk, col, nm in TH:
        r = sg[th]
        x = round(100 * r['FP'] / (r['FP'] + r['TN']), 2)
        y = round(100 * r['TP'] / (r['TP'] + r['FN']), 2)
        full = [(round(100 * a, 2), round(100 * b, 2)) for a, b in R['MBII']]
        assert (x, y) in full, (th, x, y)          # chegara nuqtasi MBII-8 egri chizig‘ining aynan ustida
        S.append(ser(nm, [(x, y)], col, 0, None, mk))
    xml, rows = W.scatter_chart(S, 'Noto‘g‘ri signallar ulushi (II tur xato), %', 'To‘g‘ri aniqlangan hodisalar, %',
                                0, 100, 20, 0, 100, 20, xfmt='0', yfmt='0', legend=True)
    # chegara belgilari: kattaroq, oq hoshiyali (chiziq ustida aniq ko‘rinishi uchun)
    xml = re.sub(r'<c:size val="7"/><c:spPr><a:solidFill><a:srgbClr val="(\w+)"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>',
                 r'<c:size val="9"/><c:spPr><a:solidFill><a:srgbClr val="\1"/></a:solidFill>'
                 r'<a:ln w="9525"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:ln></c:spPr>', xml)
    # chizish maydoni qo‘lda (o‘q nomlari to‘liq ko‘rinadi); legenda egri chiziqlardan bo‘sh o‘ng-pastki burchakda
    xml = xml.replace('<c:legend><c:legendPos val="b"/><c:overlay val="0"/>' + W._txpr(1050) + '</c:legend>',
                      '<c:legend><c:legendPos val="r"/><c:layout><c:manualLayout><c:xMode val="edge"/><c:yMode val="edge"/>'
                      f'<c:x val="{LG[0]}"/><c:y val="{LG[1]}"/><c:w val="{LG[2]}"/><c:h val="{LG[3]}"/></c:manualLayout></c:layout>'
                      '<c:overlay val="1"/><c:spPr><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>'
                      '<a:ln w="6350"><a:solidFill><a:srgbClr val="BFBFBF"/></a:solidFill></a:ln></c:spPr>' + W._txpr(1000) + '</c:legend>', 1)
    assert 'legendPos val="r"' in xml
    xml = xml.replace('<c:plotArea><c:layout/>',
                      '<c:plotArea><c:layout><c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
                      f'<c:x val="{PL[0]}"/><c:y val="{PL[1]}"/><c:w val="{PL[2]}"/><c:h val="{PL[3]}"/></c:manualLayout></c:layout>', 1)
    return xml, rows


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.5
    h = dd.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run('4.1.1-rasm – tuzatilgan ko‘rinish (namuna)'); r.bold = True
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, chart())
    q._p.append(etree.fromstring(W.chart_run_xml(rid, W_CM, H_CM, '4.1.1-rasm')))
    c = dd.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; rr = c.add_run(CAP); rr.bold = True
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

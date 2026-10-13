"""2.2-§: uchta bank bo‘yicha alohida jadvallar (2.2.2–2.2.4) va tahlil – namuna."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_2.2.2-2.2.4-jadvallar.docx'
RED, BLUE = 'FF0000', '0070C0'
ROWS = ['Jami aktivlar, trln so‘m', 'Regulyativ kapital monandligi, %', 'Asosiy kapital (CET1) monandligi, %',
        'Muammoli kreditlar ulushi, %', 'Substandart kreditlar ulushi, %', 'ROA, %', 'ROE, %', 'LCR, %', 'NSFR, %',
        'Kreditlar / depozitlar, %']
SQB = [[56.5, 64.3, 74.6, 87.6, 102.4], [15.8, 15.3, 16.1, 15.6, 17.0], [11.9, 12.1, 11.0, 10.6, 10.2], [3.8, 2.8, 2.2, 2.8, 2.5],
       [6.2, 5.3, 11.0, 14.8, 18.4], [1.9, 1.0, 1.3, 1.6, 2.0], [13.9, 7.6, 11.0, 13.9, 17.2], [164.5, 200.1, 177.3, 148.4, 222.9],
       [115.0, 110.4, 104.2, 113.1, 122.2], [342.2, 323.0, 414.0, 315.7, 206.5]]
MKB = [[13.5, 16.8, 19.1, 23.5, 30.7], [13.8, 20.7, 18.6, 15.4, 18.5], [13.2, 20.0, 17.6, 13.0, 16.4], [5.9, 4.8, 6.0, 4.6, 4.4],
       [9.1, 19.3, 24.5, 21.3, 9.0], [0.3, 0.1, 0.2, -8.0, 0.6], [1.9, 0.7, 1.1, -45.5, 3.5], [165.2, 156.4, 131.7, 169.0, 231.9],
       [117.1, 114.3, 108.3, 113.1, 113.0], [343.0, 246.2, 227.0, 222.4, 208.5]]
ALQ = [[12.6, 16.1, 16.0, 23.3, 31.5], [13.2, 12.5, 16.4, 13.9, 14.1], [11.4, 9.4, 11.5, 11.8, 10.1], [4.1, 2.0, 1.5, 1.9, 2.5],
       [8.0, 1.1, 11.8, 10.2, 6.4], [1.0, 1.6, 3.8, 1.0, 0.9], [6.6, 13.2, 27.4, 6.6, 7.3], [110.9, 226.8, 103.1, 144.4, 142.0],
       [108.0, 100.9, 102.8, 106.6, 103.0], [114.1, 119.0, 111.6, 93.7, 121.1]]

FORM = 'ning Markaziy bankka taqdim etilgan prudensial hisobot shakllari ({0}C, {0}L, {0}O): bank balansi, daromadlar, regulyativ kapital hisob-kitobi, aktivlarni tasniflash va prudensial me’yorlar jadvallari, 2020–2025-yillar {1}-dekabr holatiga.'
SRC = 'Tadqiqot natijasida muallif tomonidan quyidagi manba ma’lumotlari asosida hisoblandi: '
IZOH = 'Izoh: muammoli kreditlar – qoniqarsiz, shubhali va umidsiz deb tasniflangan kreditlar; ROA va ROE sof foydaning aktivlar va kapitalning o‘rtacha yillik qiymatiga nisbati sifatida hisoblandi.'

INTRO = ('Tizim darajasidagi agregat ko‘rsatkichlar alohida banklar o‘rtasidagi farqlarni yashiradi. Shu sababli tahlilni uchta davlat banki – '
         'yirik korporativ va infratuzilma loyihalarini moliyalashtiruvchi tizimli ahamiyatga molik “O‘zsanoatqurilishbank” ATB, '
         'davlat dasturlari doirasida mikrokreditlash bilan shug‘ullanuvchi “Mikrokreditbank” ATB va tizimli ahamiyatga molik banklar qatoriga kirmaydigan '
         '“Aloqabank” ATB misolida chuqurlashtirish maqsadga muvofiq. Avval har bir bankning 2021–2025-yillardagi asosiy ko‘rsatkichlari alohida '
         '(2.2.2–2.2.4-jadvallar), so‘ngra ular o‘zaro qiyosiy tahlil qilinadi.')
A_SQB = ('“O‘zsanoatqurilishbank” ATB ning aktivlari 2021–2025-yillarda 56,5 trln so‘mdan 102,4 trln so‘mga, ya’ni 1,8 barobar o‘sgan. '
         'Jadvaldan ko‘rinadiki, bankning asosiy muammosi kapital va aktivlar sifati o‘rtasidagi nomutanosiblikdir: regulyativ kapital monandligi '
         '15,3–17,0 foiz oralig‘ida barqaror bo‘lsa-da, asosiy kapital monandligi 12,1 foizdan 10,2 foizgacha izchil pasaygan, substandart kreditlar ulushi esa '
         '6,2 foizdan 18,4 foizgacha oshgan. Daromadlilik 2022-yildagi pasayishdan keyin tiklanib, ROE 2025-yilda 17,2 foizga yetgan, likvidlik koeffitsiyentlari '
         'esa Bazel III talablaridan yuqori. Biroq kreditlarning depozitlarga nisbati 206,5–414,0 foiz bo‘lib, bank kreditlarni asosan qarz mablag‘lari hisobiga moliyalashtiradi.')
A_MKB = ('“Mikrokreditbank” ATB ning aktivlari shu davrda 13,5 trln so‘mdan 30,7 trln so‘mga, ya’ni 2,3 barobar o‘sgan. Bankning o‘ziga xos xususiyati – '
         'daromadlilikning doimiy pastligi: ROE 2021–2023-yillarda 0,7–1,9 foizni tashkil etgan, 2024-yilda esa bank yirik zarar ko‘rib, ROE −45,5 foizga tushgan. '
         'Shunga qaramay, regulyativ kapital monandligi 13,8–20,7 foiz oralig‘ida saqlangan, bu esa kapital bankning o‘z foydasi emas, balki tashqi manbalar hisobiga '
         'to‘ldirilganidan dalolat beradi. Aktivlar sifati ham zaif: rasmiy muammoli kreditlar ulushi 4,4–6,0 foiz, substandart kreditlar ulushi 2023-yilda 24,5 foizga '
         'yetgan va faqat 2025-yilda 9,0 foizgacha kamaygan.')
A_ALQ = ('“Aloqabank” ATB tahlilga natijalarning umumiyligini tekshirish uchun kiritildi. Bankning aktivlari 12,6 trln so‘mdan 31,5 trln so‘mga, ya’ni 2,5 barobar '
         'o‘sgan. Uning yuqoridagi ikki bankdan asosiy farqi moliyalashtirish tuzilmasida: kreditlarning depozitlarga nisbati 93,7–121,1 foiz bo‘lib, bank kreditlarni '
         'asosan mijozlar depozitlari hisobiga moliyalashtiradi. Biroq kapital va likvidlik bo‘yicha zaxira uch bank ichida eng kichik: regulyativ kapital monandligi '
         '2023-yildan tashqari barcha yillarda 12,5–14,1 foiz bo‘lib, 13 foizlik minimum atrofida bo‘lgan (2022-yilda hisoblangan qiymat minimumdan 0,5 foiz punkt past), '
         'LCR 2023-yilda 103,1 foizgacha, NSFR esa 2022-yilda 100,9 foizgacha tushib, Bazel III minimumiga yaqinlashgan. Rasmiy muammoli kreditlar ulushi 1,5–4,1 foiz '
         'bilan past, ammo substandart kreditlar ulushi 2023–2024-yillarda 10,2–11,8 foizga ko‘tarilgan. Daromadlilik 2023-yilda eng yuqori bo‘lgan (ROE 27,4 foiz), '
         'keyingi yillarda esa 6,6–7,3 foizgacha pasaygan.')
QIYOS = ('Uch jadvalni solishtirish banklarning moliyaviy barqarorlik profili turlicha bo‘lsa-da, ularda umumiy qonuniyatlar ham borligini ko‘rsatadi. '
         'Kapital bo‘yicha uchala bankda ham regulyativ kapital monandligi minimumga yaqin yoki beqaror: “O‘zsanoatqurilishbank” ATB da zaxira 2,3–4,0 foiz punktni, '
         '“Aloqabank” ATB da 2023-yildan tashqari 1,1 foiz punktdan oshmagan, “Mikrokreditbank” ATB da esa ko‘rsatkich keskin tebranadi: 2021-yilda 13,8 foiz bilan '
         'minimum chegarasida bo‘lgan ko‘rsatkich 2022-yilda 20,7 foizga ko‘tarilgan, 2024-yilda 15,4 foizgacha tushib, 2025-yilda yana 18,5 foizga yetgan (2.2.5-rasm).')
NEXT = [
    '2.2.5-rasm (regulyativ kapital monandligi), 2.2.6-rasm (muammoli va substandart kreditlar) va 2.2.8-rasm (ROE) diagrammalariga “Aloqabank” ATB qatori qo‘shiladi (diagrammalar oddiy rangda); 2.2.7-rasmda u allaqachon bor.',
    'Keyingi xatboshilardagi “ikki bank”, “ikkala bank” iboralari “uch bank”, “uchala bank” ga o‘zgartiriladi va har bir xatboshiga “Aloqabank” ATB bo‘yicha bir-ikki jumla qo‘shiladi (masalan, likvidlik xatboshisiga: “Aloqabank” ATB da kreditlarning depozitlarga nisbati 100 foiz atrofida, bu esa bank depozit bazasiga tayanishini ko‘rsatadi).',
    'Hozirgi 2.2.3-jadval (umumlashtiruvchi) 2.2.5-jadval bo‘ladi va unga “Aloqabank” ATB ustuni qo‘shiladi; matndagi havolalar yangilanadi.',
    '2.2-§ boshidagi “Tahlil ikki darajada amalga oshirildi …” xatboshisida “tizimli ahamiyatga ega ikki davlat banki – “O‘zsanoatqurilishbank” ATB va “Mikrokreditbank” ATB” deyilgan. “Mikrokreditbank” ATB 2026-yil uchun belgilangan yetti tizimli bank ro‘yxatida yo‘q, shuning uchun bu jumla “uchta davlat banki – …” shaklida tuzatiladi va “Aloqabank” ATB qo‘shiladi.',
]

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=14, color=None, ls=None):
        q = dd.add_paragraph(); q.alignment = align
        q.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
        if ls:
            q.paragraph_format.line_spacing = ls
        r = q.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return q

    FN = []

    def table(num, title, data, color, src):
        par(num + '-jadval', italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False, color=color)
        FN.append(src)
        par(title + f' [{len(FN)}]', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, color=color, ls=1.0)
        t = dd.add_table(rows=len(ROWS) + 1, cols=6); t.style = 'Table Grid'
        t.autofit = False
        W_ = [Cm(6.0)] + [Cm(2.1)] * 5
        for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), W_):
            gc.set(qn('w:w'), str(int(w.twips)))
        for row in t.rows:
            for cc, w in zip(row.cells, W_):
                cc.width = w
        hdr = ['Ko‘rsatkich', '2021', '2022', '2023', '2024', '2025']
        for j, h in enumerate(hdr):
            c = t.cell(0, j); r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(12)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraphs[0].paragraph_format.line_spacing = 1.0
        for i, (nm, vals) in enumerate(zip(ROWS, data), 1):
            c = t.cell(i, 0); r = c.paragraphs[0].add_run(nm); r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(color)
            c.paragraphs[0].paragraph_format.line_spacing = 1.0
            for j, v in enumerate(vals, 1):
                c = t.cell(i, j); s = f'{v:.1f}'.replace('.', ',').replace('-', '−')
                r = c.paragraphs[0].add_run(s); r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(color)
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraphs[0].paragraph_format.line_spacing = 1.0
        par(IZOH, italic=True, size=12, indent=False, color=color, ls=1.0)

    par('2.2-§: uch bank bo‘yicha alohida jadvallar va tahlil (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    par('Qizil – yangi yoki o‘zgartirilgan matn; ko‘k – “Aloqabank” ATB bo‘yicha qo‘shimchalar. Hozirgi bitta 2.2.2-jadval uchta jadvalga ajratiladi.', italic=True, size=12, indent=False, color='595959')
    p = par(INTRO, color=RED)
    table('2.2.2', '“O‘zsanoatqurilishbank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari', SQB, RED, SRC + '“O‘zsanoatqurilishbank” ATB ' + FORM.format('crs003', 30))
    par(A_SQB, color=RED)
    table('2.2.3', '“Mikrokreditbank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari', MKB, RED, SRC + '“Mikrokreditbank” ATB ' + FORM.format('crs005', 30))
    par(A_MKB, color=RED)
    table('2.2.4', '“Aloqabank” ATB ning moliyaviy barqarorlik ko‘rsatkichlari', ALQ, BLUE, SRC + '“Aloqabank” ATB ' + FORM.format('crs031', 31))
    par(A_ALQ, color=BLUE)
    par(QIYOS, color=RED)
    par('Keyingi qismdagi tegishli o‘zgarishlar (diss’ga kiritishda bajariladi):', bold=True, size=12, indent=False)
    for i, s in enumerate(NEXT, 1):
        par(f'{i}. {s}', size=12, indent=False, ls=1.15)
    par('Sahifa osti snoskalari', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12)
    for i, s in enumerate(FN, 1):
        par(f'[{i}] {s}', size=11, indent=False, ls=1.0)
    dd.save(DST)
    print(pardoz(DST), DST)

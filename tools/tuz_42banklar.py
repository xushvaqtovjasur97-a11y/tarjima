"""4.2-§: tijorat banklari darajasidagi prognoz (ko‘k).
1) yangi 4.2.4-rasm – “O‘zsanoatqurilishbank”, “Mikrokreditbank”, “Aloqabank” ATB: muammoli kreditlar ulushi, haqiqiy va ssenariylar (4.2.3-rasmdan keyin);
2) yangi 4.2.5-jadval – bank darajasidagi model (2-model) bilan tekshiruv (4.2.4-jadval tahlilidan keyin).
Raqamlar: rasmlar 4.2.4→4.2.5, 4.2.5→4.2.6; jadvallar 4.2.5→4.2.6, 4.2.6→4.2.7. Diagramma va hisob: tools/namuna_42banklar.py, manbalar/hisob_kitob/bank42.py."""
import copy, sys
sys.path.insert(0, '/home/user/tarjima/tools')
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from lxml import etree
import kamlib as K
import wordart as W
import namuna_42banklar as N
from pardoz import pardoz

B = N.B
K.RED = '0070C0'
F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
W._counter['chart'] = 320; W._counter['docpr'] = 9800
D = K.Doc(F)
FN = 'Tadqiqotlar jarayonida muallif tomonidan tuzilgan.'
c = N.c


def tx(el):
    return ''.join(x.text or '' for x in el.iter(qn('w:t')))


# 1) qayta raqamlash (kamayish tartibida)
n = 0
for old, new in (('4.2.5-rasm', '4.2.6-rasm'), ('4.2.4-rasm', '4.2.5-rasm'), ('4.2.6-jadval', '4.2.7-jadval'), ('4.2.5-jadval', '4.2.6-jadval')):
    for p in D.d.paragraphs:
        for _ in range(tx(p._p).count(old)):
            D.replace(p, old, new); n += 1
    for f in D.fx.findall(qn('w:footnote')):
        for pe in f.findall(qn('w:p')):
            for _ in range(tx(pe).count(old)):
                D.replace(Paragraph(pe, None), old, new); n += 1
assert n == 9, n

# 2) yangi 4.2.4-rasm
p714 = D.para('Rasmdan ko‘rinib turibdiki, makroiqtisodiy shok muammoli kreditlar ulushida darhol emas')
last = 'Bu natijalar asosida 25 ta bank uchun integral indeks va erta ogohlantirish zonalari prognoz qilindi (4.2.4-jadval).'
assert tx(p714._p).count(last) == 1
D.replace(p714, last, 'Tizim darajasidagi prognoz har bir tijorat bankiga qo‘llanilganda ssenariylarning ta’siri banklar kesimida turlicha '
          'namoyon bo‘ladi. Ikkinchi bobda batafsil tahlil qilingan uch bank bo‘yicha natijalar 4.2.4-rasmda keltirilgan.')
cap423 = D.para('4.2.3-rasm. Bank tizimida muammoli kreditlar ulushining haqiqiy dinamikasi')
chart_ppr = cap423._p.getprevious().find(qn('w:pPr'))
iz_src = D.para('Izoh: qavs ichida standart xatolar (1-modelda')
anchor = p714
for k, b in enumerate(B.BANKS):
    st = D.new_para_after(cap423, f'{"abv"[k]}) {N.NM[b]}')
    for r in st.runs:
        r.font.size = Pt(12)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.line_spacing = 1.0; st.paragraph_format.space_before = Pt(0); st.paragraph_format.space_after = Pt(0)
    anchor._p.addnext(st._p); anchor = st
    cp = D.d.add_paragraph()
    rid = W.add_chart(D.d, N.bank_chart(b, legend=(k == 2)))
    cp._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 6.6 if k == 2 else 5.0, f'4.2.4-rasm {"abv"[k]}')))
    cp._p.insert(0, copy.deepcopy(chart_ppr))
    cp.paragraph_format.keep_with_next = True
    anchor._p.addnext(cp._p); anchor = cp
iz = D.new_para_after(iz_src, 'Izoh: haqiqiy qiymatlar – choraklik, milliy hisobot; prognoz – tizim modeli (4.2.2-jadval, 1-model) bo‘yicha og‘ishning '
                      'bankning 2025-yil IV choragidagi muammoli kreditlar ulushiga proporsional taqsimlanishi.')
anchor._p.addnext(iz._p)
cap = D.new_para_after(cap423, '4.2.4-rasm. Tanlangan tijorat banklarida muammoli kreditlar ulushining haqiqiy dinamikasi (2021–2025-yillar) '
                       'va ssenariylar bo‘yicha prognozi (2026–2028-yillar), %[[FN:' + FN + ']]')
iz._p.addnext(cap._p)
pk = {b: {s: B.PEAK[('P', s, b)] for s in N.ORD} for b in B.BANKS}
s25 = {b: B.HIST[b][-1][1] for b in B.BANKS}
sq, mk, al = B.BANKS
an = D.new_para_after(p714,
    f'Rasmdan ko‘rinib turibdiki, “O‘zsanoatqurilishbank” va “Aloqabank” ATBlarda muammoli kreditlar ulushi 2025-yil oxirida deyarli bir xil '
    f'({c(s25[sq])} va {c(s25[al])} foiz) bo‘lib, keskin salbiy ssenariyda 2027-yil oxirida {c(pk[sq]["Keskin salbiy"])}–{c(pk[al]["Keskin salbiy"])} '
    'foizgacha oshadi, ya’ni 5 foizlik erta ogohlantirish chegarasidan past qoladi. “Mikrokreditbank” ATB da esa ulush 2025-yilning o‘zida '
    f'{c(s25[mk])} foizni tashkil etib, salbiy ssenariyda {c(pk[mk]["Salbiy"])} foizga, keskin salbiy ssenariyda {c(pk[mk]["Keskin salbiy"])} '
    'foizga yetadi va chegaradan sezilarli oshadi. Asosiy va ijobiy ssenariylarda esa uch bankning birortasida ham ulush 5 foizdan oshmaydi. '
    'Demak, bir xil makroiqtisodiy shok kredit portfeli sifati past bo‘lgan bankda ancha kuchli ta’sir ko‘rsatadi. ' + last)
cap._p.addnext(an._p)

# 3) yangi 4.2.5-jadval: bank darajasidagi model bilan tekshiruv
p723 = D.para('Jadvaldan ko‘rinib turibdiki, asosiy va ijobiy ssenariylarda bank tizimining moliyaviy barqarorligi')   # 4.2.4-jadval tahlili
intro = D.new_para_after(p723,
    'Prognoz natijalari tanlangan usulga bog‘liq emasligini tekshirish uchun muammoli kreditlar ulushi har bir bank uchun bevosita '
    '4.2.2-jadvaldagi 2-model (25 ta bank paneli) asosida ham prognoz qilindi. Bu modelda YaIM o‘sishining ta’siri o‘sha chorakning o‘zida '
    'namoyon bo‘ladi va barcha banklarda bir xil foiz punktga teng bo‘ladi, asosiy usulda esa tizim bo‘yicha og‘ish to‘rt chorakdan keyin '
    'namoyon bo‘lib, banklarga ularning ulushiga proporsional taqsimlanadi. Natijalar 4.2.5-jadvalda keltirilgan.')
lab424 = [q for q in D.d.paragraphs if q.text.strip() == '4.2.4-jadval'][0]
tit424 = D.para('Tanlangan 25 ta bank bo‘yicha moliyaviy barqarorlik ko‘rsatkichlarining ssenariylar bo‘yicha prognozi')
tb424 = D.table_after('Tanlangan 25 ta bank bo‘yicha moliyaviy barqarorlik ko‘rsatkichlarining ssenariylar bo‘yicha prognozi')
lab = D.new_para_after(lab424, '4.2.5-jadval'); intro._p.addnext(lab._p)
tit = D.new_para_after(tit424, 'Bank darajasidagi prognozning ikki usul bo‘yicha natijalari: salbiy va keskin salbiy ssenariylar[[FN:' + FN + ']]')
lab._p.addnext(tit._p)
T, PK = B.T, B.PEAK
rows = [['Ssenariy, yil', 'Muammoli kreditlar ulushi, %', 'Qizil zonadagi banklar soni', 'Qizil zonadagi banklar aktivlari ulushi, %']]
for s in ('Salbiy', 'Keskin salbiy'):
    for y in (2026, 2027, 2028):
        p_, b_ = T.loc[('P', s, y)], T.loc[('B', s, y)]
        rows.append([f'{s}, {y}', f'{c(p_.S3_w)} / {c(b_.S3_w)}', f'{int(p_.qizil)} / {int(b_.qizil)}', f'{c(p_.qizil_aktiv, 1)} / {c(b_.qizil_aktiv, 1)}'])
rows.append(['Eng yuqori ulush (keskin salbiy), %',
             '; '.join(f'{x}: {c(PK[("P", "Keskin salbiy", b)])} / {c(PK[("B", "Keskin salbiy", b)])}'
                       for x, b in zip(('“O‘zsanoatqurilishbank”', '“Mikrokreditbank”', '“Aloqabank”'), B.BANKS)), '', ''])
t = D.d.add_table(rows=len(rows), cols=4)
t._tbl.remove(t._tbl.tblPr); t._tbl.insert(0, copy.deepcopy(tb424._tbl.tblPr))
for gc, wv in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), (2700, 2200, 2000, 2400)):
    gc.set(qn('w:w'), str(wv))
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        p = t.cell(i, j).paragraphs[0]
        run = p.add_run(v); run.font.size = Pt(12); run.font.name = 'Times New Roman'; run.bold = (i == 0)
        run.font.color.rgb = RGBColor.from_string('0070C0')
tr0 = t._tbl.findall(qn('w:tr'))[0]
trpr = tr0.find(qn('w:trPr'))
if trpr is None:
    trpr = etree.Element(qn('w:trPr')); tr0.insert(0, trpr)
etree.SubElement(trpr, qn('w:tblHeader'))
lr = t.rows[-1].cells; m = lr[1].merge(lr[3])
for p in m.paragraphs[1:]:
    p._p.getparent().remove(p._p)
tit._p.addnext(t._tbl)
izt = D.new_para_after(D.para('Izoh: muammoli kreditlar ulushi va ROA aktivlar bo‘yicha tortilgan o‘rtacha; MBII – 8 ko‘rsatkichli panel indeksi (MBII-8); aktivlar ulushi'),
                       'Izoh: har bir katakda – asosiy usul / bank modeli bo‘yicha natija; asosiy usul – tizim modeli bo‘yicha og‘ishning banklarga '
                       'proporsional taqsimlanishi (4.2.4-jadval), bank modeli – 4.2.2-jadvaldagi 2-model; muammoli kreditlar ulushi – 25 ta bank bo‘yicha '
                       'aktivlar bilan tortilgan o‘rtacha; asosiy va ijobiy ssenariylarda ikkala usulda ham qizil zonada bank yo‘q.')
t._tbl.addnext(izt._p)
r27 = T.loc[('P', 'Keskin salbiy', 2027)]
an2 = D.new_para_after(p723,
    'Jadvaldan ko‘rinib turibdiki, ikki usul bir xil asosiy natija beradi. Keskin salbiy ssenariyda qizil zonaga ikkala usulda ham aynan bir xil '
    f'to‘rt bank – “Agrobank”, “Asaka bank”, “Mikrokreditbank” va “Turonbank” ATBlar tushadi, ularning aktivlari ulushi esa {c(r27.qizil_aktiv, 1)} '
    'foizni tashkil etadi. Salbiy ssenariyda ham ikkala usulda “Mikrokreditbank” va “Turonbank” ATBlar qizil zonaga tushadi. Usullar o‘rtasidagi '
    'farq asosan vaqtda namoyon bo‘ladi. Bank modelida shok o‘sha yilning o‘zida (2026-yil), asosiy usulda esa bir yildan keyin (2027-yil) eng '
    'yuqori ta’sir ko‘rsatadi. Demak, erta ogohlantirish zonalari bo‘yicha xulosalar prognoz '
    'usuliga bog‘liq emas, ya’ni barqaror hisoblanadi.')
izt._p.addnext(an2._p)
D.save(F)
print(n, pardoz(F))

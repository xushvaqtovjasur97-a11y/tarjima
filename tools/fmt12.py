import docx,sys
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
src,dst=sys.argv[1],sys.argv[2]
d=docx.Document(src)
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6)
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
  else:
    pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
    t=p.text.strip()
    if 'jadval.' in t[:25] and t[:1].isdigit():
      pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6)
    if t.startswith('Manba:'): pf.first_line_indent=Cm(0); pf.space_after=Pt(6)
    if t.startswith('Z = Σ'): pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0)
    for r in p.runs: fmt_run(r,14)
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    for c in row.cells:
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,12)
          if i==0: r.bold=True
d.save(dst)

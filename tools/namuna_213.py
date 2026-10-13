"""2.1.3/2.1.4-rasm va snoskalar bo‘yicha kiritilgan tuzatishlar namunasi (asosiy fayldan olinadi)."""
import sys, re, subprocess, io
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W
import namuna_135 as N135
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_2.1.3-2.1.4_snoskalar.docx'
src = docx.Document(F)
fx = etree.fromstring(src.part.package.part_related_by if False else [r.target_part for r in src.part.rels.values() if r.reltype.endswith('/footnotes')][0].blob)
FT = {f.get(qn('w:id')): ''.join(x.text or '' for x in f.iter(qn('w:t'))).strip() for f in fx.findall(qn('w:footnote'))}


def where(fid):
    for p in src.paragraphs:
        if any(r.get(qn('w:id')) == fid for r in p._p.iter(qn('w:footnoteReference'))):
            return p.text.strip()
    for t in src.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if any(r.get(qn('w:id')) == fid for r in p._p.iter(qn('w:footnoteReference'))):
                        return 'Jadval: ' + p.text.strip()
    return '?'


CH = [  # (yangi matn boshi, eski matn)
    ('Tadqiqot natijasida muallif tomonidan quyidagi manbalar asosida ishlab chiqildi: O‘zbekiston Respublikasi Markaziy banki. Statistik byulleten', 'Muallif tomonidan quyidagi manbalar asosida tuzilgan: …'),
    ('Tadqiqot natijasida muallif tomonidan quyidagi manbalar asosida ishlab chiqildi: The Central Bank of Uzbekistan. Financial Stability Report for H1', 'Muallif tomonidan quyidagi manbalar asosida tuzilgan: …'),
    ('Tadqiqot natijasida muallif tomonidan quyidagi manbalar asosida ishlab chiqildi: 3697-son', 'Muallif tomonidan quyidagi manbalar asosida tuzilgan: …'),
    ('Tadqiqot natijasida muallif tomonidan quyidagi manbalar asosida ishlab chiqildi: IMF Country Report No. 25/227', 'Muallif tomonidan quyidagi manbalar asosida tuzilgan: …'),
    ('Tadqiqot natijasida muallif tomonidan hisoblandi (25 ta bank paneli', 'Muallif hisob-kitobi: manbalar/hisob_kitob/robust_kam.py va robust_kam.json (25 ta bank paneli, …).'),
    ('Tadqiqot natijasida muallif tomonidan “Aloqabank” ATB ning prudensial hisobot shakllari', '… asosida hisoblandi (manbalar/hisob_kitob/aloqa23.py, aloqa_mbii.json).'),
]

if __name__ == '__main__':
    cap = [p for p in src.paragraphs if p.text.startswith('2.1.4-rasm')][0]
    rid = [c.get(qn('r:id')) for c in cap._p.getprevious().iter('{http://schemas.openxmlformats.org/drawingml/2006/chart}chart')][0]
    xml = src.part.related_parts[rid].blob.decode('utf-8')
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(4)

    def par(text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, color=None, strike=False, italic=False):
        q = dd.add_paragraph(); q.alignment = align
        r = q.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.strike = strike
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return q

    par('2.1.3–2.1.4-rasmlar va snoskalar bo‘yicha kiritilgan tuzatishlar', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    par('1. 2.1.4-rasm: “Regulyativ minimum” ustuni 2026-yildan amaldagi talab (buferlar bilan) bo‘yicha – CET1 8,5% (oldin 8%), jami kapital 12% (oldin 13%).')
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid2 = W.add_chart(dd, (re.sub(r'(<c:externalData r:id=")[^"]+(")', r'\g<1>rId1\g<2>', xml), N135.rows_from(xml)))
    q._p.append(etree.fromstring(W.chart_run_xml(rid2, 16.0, 8.0, '2.1.4-rasm')))
    par(cap.text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
    par('2. Snoskalar standart ifodaga keltirildi va ulardan ichki fayl nomlari olib tashlandi:')
    for new, old in CH:
        ids = [k for k, v in FT.items() if v.startswith(new)]
        if 'Statistik byulleten' in new:
            ids = [k for k in ids if where(k).startswith('2.1.3-rasm')]
        for fid in ids:
            par('Joyi: ' + where(fid)[:110] + ' …', italic=True, size=10, color='595959')
            par('Oldin: ' + old, size=11, color='7F7F7F', strike=True)
            par('Hozir: ' + FT[fid], size=11, color='C00000')
    dd.save(DST)
    print(pardoz(DST), DST)

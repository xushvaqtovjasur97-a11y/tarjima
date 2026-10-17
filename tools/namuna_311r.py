"""3.1.1-rasm: 6-bosqichda vaznlar blok harflari bilan (namuna)."""
import copy, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pardoz import pardoz

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_3.1.1-rasm.docx'
WPS = '{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}wsp'
OLD = ['0,25; 0,30; 0,20;', '0,15; 0,10']
NEW = ['K 0,25; A 0,30;', 'L 0,20; D 0,15;', 'V 0,10']


def tuzat(p_el):
    """Chizma paragrafidagi barcha (asosiy va fallback) 6-qutichalarni yangilaydi."""
    n = 0
    for tx in p_el.iter(qn('w:txbxContent')):
        ps = tx.findall(qn('w:p'))
        txt = [''.join(t.text or '' for t in p.iter(qn('w:t'))) for p in ps]
        if txt[:1] != ['6. Vaznlash:'] or txt[1:] != OLD:
            continue
        for p in ps[2:]:
            tx.remove(p)
        tmpl = ps[1]
        tx.remove(tmpl)
        for s in NEW:
            q = copy.deepcopy(tmpl)
            ts = list(q.iter(qn('w:t')))
            ts[0].text = s
            for t in ts[1:]:
                t.text = ''
            tx.append(q)
        n += 1
    return n


if __name__ == '__main__':
    src = docx.Document(F)
    cap = [p for p in src.paragraphs if p.text.startswith('3.1.1-rasm.')][0]
    fig = copy.deepcopy(cap._p.getprevious())
    print('qutichalar:', tuzat(fig))
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)
    def par(text, bold=False, color=None, size=14, align=WD_ALIGN_PARAGRAPH.CENTER):
        q = dd.add_paragraph(); q.alignment = align
        r = q.add_run(text); r.bold = bold; r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
    par('3.1.1-rasm (namuna)', bold=True)
    dd.element.body.insert(len(dd.element.body) - 1, fig)
    par('3.1.1-rasm. Moliyaviy barqarorlikning integral indeksini hisoblash algoritmi', bold=True)
    par('6-bosqich: hozir – “0,25; 0,30; 0,20; 0,15; 0,10”; yangi – “K 0,25; A 0,30; L 0,20; D 0,15; V 0,10” '
        '(K – kapital yetarliligi, A – aktivlar sifati, L – likvidlik, D – daromadlilik, V – valyuta riskiga sezgirlik; '
        'harflar 3.1.2-jadvaldagi ko‘rsatkichlar belgilariga mos).', color='FF0000', size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    dd.save(DST)
    print(pardoz(DST), DST)

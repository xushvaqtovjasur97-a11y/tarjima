"""Adabiyotlar ro‘yxatiga 11 ta MHXS hisoboti; 3.2.2-jadval izohi umumiy shaklda (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

DST = '/home/user/tarjima/Namuna_adabiyotlar_11_hisobot.docx'
# (qaysi yozuvdan keyin, yangi yozuv)
NEW = [
    ('Консолидированная финансовая отчетность за год, закончившийся 31 декабря 2021 года, и аудиторское заключение. – Ташкент, 2022.',
     'Акционерный коммерческий банк “Узбекский промышленно-строительный банк” и его дочерние организации. Консолидированная финансовая отчетность за год, закончившийся 31 декабря 2022 года, и аудиторское заключение независимого аудитора. – Ташкент, 2023.'),
    ('за годы, закончившиеся 31 декабря 2023, 2022 и 2021 гг. – Ташкент, 2024.',
     'Joint Stock Commercial Bank “Uzbek Industrial and Construction Bank” and its subsidiaries. Consolidated Financial Statements and Independent Auditor’s Report for the year ended 31 December 2024. – Tashkent, 2025.'),
    ('Финансовая отчетность за год, закончившийся 31 декабря 2021 года, и заключение независимых аудиторов. – Ташкент, 2022.',
     'Акционерно-коммерческий банк “Микрокредитбанк”. Финансовая отчетность за год, закончившийся 31 декабря 2022 года, и заключение независимых аудиторов. – Ташкент, 2023.'),
    ('Финансовая отчетность за год, закончившийся 31 декабря 2023 года, и заключение независимых аудиторов. – Ташкент, 2024.',
     'Акционерно-коммерческий банк “Микрокредитбанк”. Консолидированная финансовая отчетность за год, закончившийся 31 декабря 2024 года, и заключение независимых аудиторов. – Ташкент, 2025.'),
    ('Joint-Stock Commercial Bank “Business Development Bank”. Consolidated Financial Statements for the year ended 31 December 2025',
     'Joint-Stock Commercial Bank “Qishloq Qurilish Bank”. Consolidated Financial Statements for the year ended 31 December 2022 and Independent Auditor’s Report. – Tashkent: Deloitte & Touche, 2023.'),
    ('NEXT',
     'Joint-Stock Commercial Bank “Business Development Bank”. Consolidated Financial Statements for the year ended 31 December 2023 and Independent Auditor’s Report. – Tashkent: Deloitte & Touche, 2024.'),
    ('NEXT',
     'Joint-Stock Commercial Bank “Business Development Bank”. Consolidated Financial Statements for the year ended 31 December 2024 and Independent Auditor’s Report. – Tashkent: Deloitte & Touche, 2025.'),
    ('independent auditor’s report for the year ended 31 December 2022. – Tashkent: Ernst & Young, 2023.',
     'Joint-Stock Commercial “Xalq Bank”. Consolidated financial statements and independent auditor’s report for the year ended 31 December 2023. – Tashkent: Ernst & Young, 2024.'),
    ('NEXT',
     'Joint-Stock Commercial “Xalq Bank”. Consolidated financial statements and independent auditor’s report for the year ended 31 December 2024. – Tashkent: Ernst & Young, 2025.'),
    ('“Ipoteka Bank” and its subsidiaries. Consolidated Financial Statements and Independent Auditor’s Report for the year ended 31 December 2025',
     'Joint Stock Commercial Mortgage Bank “Ipoteka Bank” and its subsidiaries. Consolidated Financial Statements and Independent Auditors’ Report for the year ended 31 December 2022. – Tashkent: KPMG Audit, 2023.'),
    ('Independent Auditors’ Report for the year ended 31 December 2023. – Tashkent: Ernst & Young, 2024.',
     'Joint Stock Commercial Mortgage Bank “Ipoteka-bank” and its subsidiaries. Consolidated Financial Statements and Independent Auditors’ Report for the year ended 31 December 2024. – Tashkent: Ernst & Young, 2025.'),
]
FN_NEW = ('Tadqiqot natijasida muallif tomonidan “O‘zsanoatqurilishbank”, “Mikrokreditbank”, “Biznesni rivojlantirish banki”, '
          '“Xalq banki” va “Ipoteka-bank” ATBlarining 2021–2025-yillardagi MHXS bo‘yicha moliyaviy hisobotlari, “SNS” reyting '
          'agentligi bazasi va Markaziy bank statistikasi asosida hisoblandi.')

if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(3)

    def par(t, col=None, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12):
        q = dd.add_paragraph(); q.alignment = align
        r = q.add_run(t); r.bold = bold; r.font.size = Pt(size)
        if col: r.font.color.rgb = RGBColor.from_string(col)
    par('Adabiyotlar ro‘yxati va jadval izohi (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    par('1) 3.2.2-jadval nomidagi izoh – uzun ro‘yxat o‘rniga umumiy shakl:', bold=True)
    par(FN_NEW, 'FF0000', size=11)
    par('2) Adabiyotlar ro‘yxatining VI bo‘limiga (har biri o‘z bankining yozuvlari yonida, yillar tartibida) qo‘shiladigan 11 ta hisobot:', bold=True)
    for i, (_, t) in enumerate(NEW, 1):
        par(f'{i}) {t}', 'FF0000', size=11)
    par('Izoh: ro‘yxat raqamlari qayta tartiblanadi (197 → 208 manba), Kirishdagi manbalar soni ham 208 ga o‘zgaradi. '
        '“O‘zsanoatqurilishbank” ATB ning 2023-yil hisoboti ro‘yxatda (181-raqam) borligi sababli qo‘shilmaydi.', size=11)
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)

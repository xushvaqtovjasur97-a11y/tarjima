"""1.3.7-rasmning yangi varianti (namuna): Xitoy tijorat banklarida muammoli va alohida e’tibor talab qiluvchi
kreditlar ulushi, 2021–2025-yillar oxiri (NFRA/CBIRC), ilmiy baholar (Charoenwong va boshq., 2025).
Alohida e’tibor talab qiluvchi kreditlar ulushi = SML / (normal kreditlar + muammoli kreditlar), trln yuan."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from lxml import etree
import wordart as W
from pardoz import pardoz

DST = '/home/user/tarjima/Namuna_1.3.7-rasm_yangi.docx'
Y = ['2021', '2022', '2023', '2024', '2025']
NPLR = [1.73, 1.63, 1.59, 1.50, 1.50]                     # rasmiy, %
NPL = [2.8, 3.0, 3.2, 3.3, 3.5]                            # trln yuan
PERF = [162.0, 179.7, 199.3, 214.7, 230.2]                 # normal (ishlaydigan) kreditlar, trln yuan
SML = [3.8, 4.1, 4.5, 4.8, 5.1]                            # alohida e’tibor talab qiluvchi, trln yuan
SMLR = [round(100 * s / (p + n), 2) for s, p, n in zip(SML, PERF, NPL)]
LO, HI = round(NPLR[-1] * 2, 1), round(NPLR[-1] * 3, 1)    # rasmiy ko‘rsatkichdan 2–3 barobar
print(SMLR, LO, HI)

CATS = [f'{y}-y. oxiri' for y in Y] + ['Ilmiy baho:\nquyi chegara', 'Ilmiy baho:\nyuqori chegara']
SER = [('Muammoli kreditlar (rasmiy)', NPLR + [None, None], '2A78D6', 'FFFFFF'),
       ('Alohida e’tibor talab qiluvchi kreditlar', SMLR + [None, None], 'FFC000', '000000'),
       ('Muqobil baholar (2025-yil)', [None] * 5 + [LO, HI], 'EB6834', 'FFFFFF')]
xml, rows = W.bar_chart(CATS, SER, ytitle='Kreditlardagi ulushi, %', stacked=True, ymax=5, ymin=0, major=1, label_fmt='0.00', gap=60)

NFRA = ('China Banking and Insurance Regulatory Commission, National Financial Regulatory Administration. Supervisory Statistics of the Banking and Insurance Sectors (Major regulatory indicators of the banking and insurance sectors): 2021–2025-yillar IV choraklari. – https://www.nfra.gov.cn')
CHA = 'Charoenwong B., Miao M., Ruan T. Nonperforming Loan Disposals Without Resolution // Management Science. – 2025. – Vol. 71, No. 1. – P. 898–916.'
IMF = 'International Monetary Fund. People’s Republic of China: 2025 Article IV Consultation – Press Release; Staff Report; and Statement by the Executive Director. IMF Country Report No. 26/44. – Washington, DC: IMF, 2026.'


def texts(fn):
    s1 = ('So‘nggi yillar ma’lumotlari ham vaziyat tubdan o‘zgarmaganini ko‘rsatadi. Milliy moliyaviy nazorat boshqarmasi ma’lumotlariga ko‘ra, tijorat banklarida muammoli kreditlar ulushi 2021-yil oxirida 1,73 foiz, 2022-yil oxirida 1,63 foiz, 2023-yil oxirida 1,59 foiz, 2024-yil va 2025-yil oxirida esa 1,50 foizni tashkil etgan'
          + fn(NFRA)
          + '. Biroq ularning yonida “alohida e’tibor talab qiluvchi” kreditlar ham bor bo‘lib, ularning ulushi shu yillarda 2,2–2,3 foiz atrofida saqlanib kelmoqda; ikkalasi qo‘shilganda 2025-yil oxirida muammoli aktivlar ulushi 3,7 foizga yetadi.')
    caption = 'Xitoy tijorat banklarida muammoli kreditlar ulushi: rasmiy ma’lumotlar va muqobil baholar, 2021–2025-yillar, %'
    cap_fn = ('Tadqiqot natijasida muallif tomonidan quyidagi manbalar ma’lumotlari asosida ishlab chiqildi: ' + NFRA + '; ' + CHA[:-1]
              + '. Izoh: alohida e’tibor talab qiluvchi kreditlar ulushi muallif tomonidan ushbu kreditlar qoldig‘ining jami kreditlarga (ishlaydigan va muammoli kreditlar yig‘indisiga) nisbati sifatida hisoblangan; muqobil baholar – 2025-yil oxiridagi rasmiy ko‘rsatkichning ikki va uch barobari.')
    a = ('Rasmdan ko‘rinadiki, rasmiy muammoli kreditlar ulushi besh yil davomida 1,73 foizdan 1,50 foizgacha sekin pasaygan, alohida e’tibor talab qiluvchi kreditlar ulushi esa deyarli o‘zgarmagan. Ular qo‘shilsa, bu ko‘rsatkich ikki yarim barobar, ilmiy baholarga ko‘ra esa uch barobargacha ortadi.')
    return dict(s1=s1, caption=caption, cap_fn=cap_fn, a=a)


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)

    def par(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, size=14, color=None):
        q = dd.add_paragraph(); q.alignment = align
        q.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
        r = q.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return q

    FN = []

    def fn(t):
        FN.append(t); return f'[{len(FN)}]'

    T = texts(fn)
    par('1.3.7-rasmning yangi varianti (namuna)', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    par('Ma’qul bo‘lsa: 1.3-§ dagi “So‘nggi yillar ma’lumotlari ham …” xatboshining birinchi qismi (“… 3,7 foizga yetadi.” gacha) quyidagi matn bilan almashtiriladi; xatboshining davomi (Charoenwong va boshq., XVJ) o‘zgarmaydi. 1.3.7-rasm yangilanadi, rasmdan keyingi xatboshining birinchi ikki jumlasi quyidagi jumlalar bilan almashtiriladi. XVJ 2016-yilgi “risk ostidagi kreditlar” (15,5 foiz) ustuni tahlil davridan tashqari bo‘lgani uchun olib tashlanadi.',
        italic=True, size=12, indent=False, color='595959')
    par(T['s1'])
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, (xml, rows))
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.5, '1.3.7-rasm')))
    FN.append(T['cap_fn']); par('1.3.7-rasm. ' + T['caption'] + f'[{len(FN)}]', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    par(T['a'] + ' …')
    par('Sahifa osti snoskalari', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    for i, t in enumerate(FN, 1):
        par(f'[{i}] {t}', size=11, indent=False)
    dd.save(DST)
    print(pardoz(DST), DST)

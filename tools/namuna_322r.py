"""3.2.2-rasm: yorliqsiz chiziqlar, sinf chegaralari qizil, ostida sonlar jadvali (namuna)."""
import re, sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree
import wordart as W
import namuna_313r as B

DST = '/home/user/tarjima/Namuna_3.2.2-rasm.docx'
Y = ['2021', '2022', '2023', '2024', '2025']
SER = [('“Xalq banki”', [2.66, 2.71, 2.96, 3.20, 3.38], '70AD47', None),
       ('“Biznesni rivojlantirish banki”', [2.51, 2.25, 2.03, 2.58, 3.23], 'FFC000', None),
       ('“Ipoteka-bank”', [3.32, 3.37, 2.69, 3.13, 3.11], '7030A0', None),
       ('“O‘zsanoatqurilishbank”', [2.92, 2.81, 2.91, 2.72, 2.81], '2A78D6', None),
       ('“Mikrokreditbank”', [2.18, 2.31, 2.22, 2.25, 2.47], 'EB6834', None),
       ('“Aloqabank” (tanlanmadan tashqari)', [3.01, 2.97, 3.17, 3.10, 2.78], '00A3A3', 'dash')]
GR = [('I va II sinflar chegarasi', [3.25] * 5, 'C00000', 'sysDash'), ('II va III sinflar chegarasi', [2.5] * 5, 'C00000', 'sysDash')]
IZOH = 'Izoh: qizil uzuq chiziqlar – I va II sinflar (3,25 ball) hamda II va III sinflar (2,50 ball) chegaralari.'


def chart():
    xml, rows = W.line_table_chart(Y, SER + GR, ytitle='ball', ymin=1.75, ymax=3.75, major=0.25, fmt='0.00', w=28575, dtable=False,
                                   layout=(0.11, 0.04, 0.80, 0.84))
    xml = xml.replace('<c:tickLblPos val="none"/>', '<c:tickLblPos val="low"/>', 1)
    xml = xml.replace('<c:numFmt formatCode="0" sourceLinked="0"/>', '<c:numFmt formatCode="0.00" sourceLinked="0"/>', 1)
    sers = re.findall(r'<c:ser>.*?</c:ser>', xml, re.S)
    for s, (t, ix, ps) in zip(sers[6:], [('I va II sinflar chegarasi (3,25)', 2, 't'), ('II va III sinflar chegarasi (2,50)', 2, 't')]):
        ns = re.sub(r'<c:marker>.*?</c:marker>', '<c:marker><c:symbol val="none"/></c:marker>', s, flags=re.S)
        lb = B._lastlbl(t, ix, ps)
        if t.startswith('II va III'):
            lb = lb.replace('<c:tx>', '<c:layout><c:manualLayout><c:x val="-0.06"/><c:y val="0"/></c:manualLayout></c:layout><c:tx>', 1)
        if t.startswith('I va II'):
            lb = lb.replace('<c:tx>', '<c:layout><c:manualLayout><c:x val="0"/><c:y val="-0.10"/></c:manualLayout></c:layout><c:tx>', 1)
        ns = ns.replace('<c:cat>', lb + '<c:cat>', 1).replace('w="28575"', 'w="19050"')
        xml = xml.replace(s, ns, 1)
    return xml, rows


def table(doc, red=None):
    t = doc.add_table(rows=1 + len(SER), cols=6)
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(6.5)] + [Cm(1.8)] * 5
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    for gc, wd in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(wd / 635)))
    data = [['Bank'] + Y] + [[s[0]] + [f'{v:.2f}'.replace('.', ',') for v in s[1]] for s in SER]
    for i, r in enumerate(data):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.width = widths[j]; p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format; pf.space_before = pf.space_after = 0; pf.first_line_indent = Cm(0); pf.line_spacing = 1.0
            if j == 0 and i > 0:
                s = SER[i - 1]
                k = p.add_run('━━ ' if s[3] is None else '╍╍ '); k.font.size = Pt(10); k.bold = True
                k.font.color.rgb = RGBColor.from_string(s[2]); k.font.name = 'Times New Roman'
            run = p.add_run(v); run.font.size = Pt(11); run.bold = (i == 0); run.font.name = 'Times New Roman'
            if red: run.font.color.rgb = RGBColor.from_string(red)
    return t


if __name__ == '__main__':
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)
    h = dd.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = h.add_run('3.2.2-rasm (namuna)'); r.bold = True
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rid = W.add_chart(dd, chart())
    q._p.append(etree.fromstring(W.chart_run_xml(rid, 16.0, 8.0, '3.2.2-rasm')))
    table(dd, 'FF0000')
    p = dd.add_paragraph(); r = p.add_run(IZOH); r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string('FF0000')
    p = dd.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('3.2.2-rasm. Tanlangan banklarda moliyaviy barqarorlikning integral indeksi, MHXS hisobotlari asosida, 2021–2025-yillar'); r.bold = True
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST))

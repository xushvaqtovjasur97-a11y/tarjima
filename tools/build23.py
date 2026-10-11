import sys, subprocess
sys.path.insert(0,'/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import wordart as W

S='/home/user/tarjima/tools'
SRC='/home/user/tarjima/paragraflar/2.3-paragraf.md'; DST='/home/user/tarjima/paragraflar/2.3-paragraf.docx'
subprocess.run(['pandoc',SRC,'-o',S+'/p23_raw.docx'],check=True)
d=docx.Document(S+'/p23_raw.docx')
for s in d.sections:
  s.left_margin=Cm(3); s.right_margin=Cm(1.5); s.top_margin=Cm(2); s.bottom_margin=Cm(2)
def fmt_run(r,size):
  r.font.name='Times New Roman'; r.font.size=Pt(size)
  rPr=r._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
  if rf is None: rf=OxmlElement('w:rFonts'); rPr.append(rf)
  for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Times New Roman')
  r.font.color.rgb=None
for st in d.styles:
  try: st.font.name='Times New Roman'; st.font.color.rgb=None
  except Exception: pass
fs=d.styles['Footnote Text']; fs.font.size=Pt(10)
fs.paragraph_format.line_spacing=1.0; fs.paragraph_format.first_line_indent=Cm(0)
fs.paragraph_format.space_after=Pt(0); fs.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
prev_j=False
for p in d.paragraphs:
  pf=p.paragraph_format; name=p.style.name; t=p.text.strip()
  pj=prev_j; prev_j=bool(__import__('re').match(r'^2\.3\.\d-jadval$',t))
  pf.line_spacing=1.5; pf.space_before=Pt(0); pf.space_after=Pt(0)
  if name.startswith('Heading'):
    lvl=name[-1]
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl=='2' else WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent=Cm(0 if lvl=='2' else 1.25); pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: fmt_run(r,14); r.bold=True; r.italic=False
    continue
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  pf.first_line_indent=Cm(0) if ('List' in name or 'Compact' in name) else Cm(1.25)
  for r in p.runs:
    if r.style is not None and 'Footnote' in r.style.name: continue
    fmt_run(r,14)
  import re
  if re.match(r'^2\.3\.\d-rasm\.',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.space_after=Pt(12); pf.line_spacing=1.0
    for r in p.runs:
      if r.style is None or 'Footnote' not in r.style.name: r.bold=True
  if re.match(r'^2\.3\.\d-jadval$',t):
    pf.alignment=WD_ALIGN_PARAGRAPH.RIGHT; pf.first_line_indent=Cm(0); pf.space_before=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.italic=True
  if pj:
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0); pf.space_after=Pt(6); pf.keep_with_next=True; pf.line_spacing=1.0
for t in d.tables:
  tblPr=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
  for e in ('top','left','bottom','right','insideH','insideV'):
    x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); b.append(x)
  tblPr.append(b)
  for i,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr()
    cs=OxmlElement('w:cantSplit'); trPr.append(cs)
    if i==0:
      hd=OxmlElement('w:tblHeader'); trPr.append(hd)
    for c in row.cells:
      if i==0:
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),'DEEBF7'); tcPr.append(sh)
      for p in c.paragraphs:
        pf=p.paragraph_format; pf.line_spacing=1.0; pf.first_line_indent=Cm(0)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
          fmt_run(r,11)
          if i==0: r.bold=True

BLUE,ORANGE,GREEN,GOLD,GREY,DBLUE,LBLUE='2A78D6','EB6834','70AD47','FFC000','A5A5A5','1F4E79','9DC3E6'

for t in d.tables:
  n=len(t.columns)
  if n==6:
    first=6.0 if t.rows[0].cells[1].text.strip()=='2021' else 3.6
    rest=(16.4-first)/5
    for row in t.rows:
      for k,c in enumerate(row.cells): c.width=Cm(first if k==0 else rest)
    grid=t._tbl.tblGrid
    for k,gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int((first if k==0 else rest)*567)))
    t.autofit=False

def setw(t,ws):
  for row in t.rows:
    for k,c in enumerate(row.cells): c.width=Cm(ws[k])
  for k,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(ws[k]*567)))
  t.autofit=False
for t in d.tables:
  n=len(t.columns)
  if n==4: setw(t,[3.0,4.4,4.6,4.4])
  if n==5: setw(t,[2.6,5.0,2.8,2.8,3.2])
Y=['2021','2022','2023','2024','2025']
# 2.3.1 scheme
sh=[]
steps=[('Kreditlarni\ntasniflash','Toifa (milliy) /\nbosqich (MHXS 9)',DBLUE),('Zaxiralar','Qat’iy stavkalar /\nkutilgan zararlar',BLUE),
       ('Sof foyda','Zaxira xarajati,\ndastlabki tan olish',ORANGE),('Xususiy\nkapital','Taqsimlanmagan\nfoyda',GREEN),('Kapital\nmonandligi va\nNSFR','Regulyativ\nko‘rsatkichlar','7030A0')]
bw,gap=2.8,0.6
for i,(h,dt,c) in enumerate(steps):
  x=0.1+i*(bw+gap)
  sh.append(W.box(x,0.3,bw,1.5,h,fill=c,line=c,sz=22,bold=True,color='FFFFFF'))
  sh.append(W.box(x,2.1,bw,1.3,dt,fill='F2F2F2',line='A6A6A6',sz=20))
  if i<4: sh.append(W.line(x+bw+0.02,1.05,x+bw+gap-0.02,1.05,color='404040',lw=19050))
sh.append(W.line(0.3,4.0,16.6,4.0,color='C00000',lw=28575))
sh.append(W.box(0.1,4.2,4.2,0.9,'Tafovut: 2–3 barobar',fill='FCE4D6',line='C00000',sz=20,bold=True))
sh.append(W.box(6.2,4.2,4.2,0.9,'Tafovut: 10–20 foiz',fill='FCE4D6',line='C00000',sz=20,bold=True))
sh.append(W.box(12.3,4.2,4.3,0.9,'Tafovut: 2–3 foiz punkt',fill='FCE4D6',line='C00000',sz=20,bold=True))
sh.append(W.box(0.1,5.3,16.5,0.6,'Ta’sir zanjir bo‘ylab kuchsizlanib boradi: aktivlar sifati → daromadlilik → kapital → likvidlik',fill=None,line=None,sz=20,italic=True))
W.replace_placeholder(d,'{{FIG:1}}',W.group_run_xml(sh,16.7,6.0,'2.3.1-rasm'))
rid=W.add_chart(d,W.bar_chart(Y,[('Milliy hisobot',[998.2,588.9,911.7,1303.2,1880.1],BLUE),('MHXS',[857.0,633.7,856.2,1115.0,1543.8],ORANGE)],
  ytitle='mlrd so‘m',ymax=2000,major=500,label_fmt='#,##0.0',gap=80))
W.replace_placeholder(d,'{{FIG:2}}',W.chart_run_xml(rid,16.0,8.0,'2.3.2-rasm'))
rid=W.add_chart(d,W.bar_chart(Y,[('Milliy hisobot',[3.6,2.7,2.9,3.4,4.5],BLUE),('MHXS',[4.5,5.6,4.0,4.7,6.0],ORANGE)],
  ytitle='Jami kreditlarga nisbatan, %',ymax=7,major=1,label_fmt='0.0',gap=80))
W.replace_placeholder(d,'{{FIG:3}}',W.chart_run_xml(rid,16.0,8.0,'2.3.3-rasm'))
rid=W.add_chart(d,W.bar_chart(['2024','2025'],[('Standart',[1013.3,746.4],GREEN),('Substandart',[577.8,2290.9],GOLD),('Qoniqarsiz',[1326.1,243.9],ORANGE),
  ('Shubhali',[628.2,642.1],'C00000'),('Umidsiz',[540.2,1386.9],'7F7F7F')],stacked=True,ytitle='mlrd so‘m',ymax=6000,major=1000,label_fmt='#,##0',gap=70))
W.replace_placeholder(d,'{{FIG:4}}',W.chart_run_xml(rid,16.0,8.0,'2.3.4-rasm'))
rid=W.add_chart(d,W.scatter_chart([
  ('«O‘zsanoatqurilishbank»',[2021,2022,2023,2024,2025],[1.65,2.61,2.14,2.08,2.96],BLUE,28575,None,'circle',('t',[0,1,2,3,4])),
  ('«Mikrokreditbank»',[2021,2022,2023,2024,2025],[2.95,2.81,2.77,3.30,2.80],ORANGE,28575,None,'square',('t',[0,1,2,3,4])),
  ('Muvofiqlik darajasi (1,0)',[2021,2025],[1,1],GREY,19050,'dash',None,None)],
  'Yillar','Koeffitsiyent',2020.5,2025.5,1,0,4,0.5,xfmt='0',yfmt='0.0',legend=True))
W.replace_placeholder(d,'{{FIG:5}}',W.chart_run_xml(rid,16.0,8.0,'2.3.5-rasm'))
rid=W.add_chart(d,W.bar_chart(Y,[('Milliy hisobot',[38.4,21.3,41.1,-1703.8,160.8],GREEN),('MHXS',[5.1,10.4,-471.2,-961.3,42.1],ORANGE)],
  ytitle='mlrd so‘m',ymax=500,ymin=-2000,major=500,label_fmt='#,##0.0',gap=80))
W.replace_placeholder(d,'{{FIG:6}}',W.chart_run_xml(rid,16.0,8.0,'2.3.6-rasm'))
rid=W.add_chart(d,W.bar_chart(['Muammoli kreditlar ulushi','Kredit zaxiralari','NSFR','Jami aktivlar','CET1 monandligi','Xususiy kapital','Sof foyda'],
  [('Nisbiy tafovut',[195.7,44.2,2.1,-5.1,-13.1,-13.5,-17.9],DBLUE)],ytitle='%',ymax=200,ymin=-50,major=50,label_fmt='0.0',gap=60,legend=False))
W.replace_placeholder(d,'{{FIG:7}}',W.chart_run_xml(rid,16.0,8.0,'2.3.7-rasm'))
d.save(DST); print('saved')

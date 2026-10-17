"""3.1.4-jadval: “Tavsifi” ustuni ballar orqali (namuna)."""
import sys
sys.path.insert(0, '/home/user/tarjima/tools')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

F = '/home/user/tarjima/Dissertatsiya_Maxmudov_O.T_asosiy.docx'
DST = '/home/user/tarjima/Namuna_3.1.4-jadval.docx'
YANGI = {1: 'Ko‘rsatkichlarning aksariyati 3–4 ball darajasida, alohida zaif ko‘rsatkichlar boshqa bloklar hisobiga qoplanadi',
         2: 'Ayrim bloklarda zaiflik, ko‘rsatkichlar asosan me’yoriy talablar darajasida (2–3 ball)',
         3: 'Bir nechta blokda ko‘rsatkichlar me’yorga yaqin yoki undan past (1–2 ball)',
         4: 'Ko‘rsatkichlarning ko‘pchiligi minimal xalqaro talablardan past (1 ball atrofida)'}

if __name__ == '__main__':
    src = docx.Document(F)
    cap = [p for p in src.paragraphs if p.text.startswith('Moliyaviy barqarorlik sinflari va ularga mos choralar')][0]
    rows = [[c.text for c in r.cells] for r in docx.table.Table(cap._p.getnext(), src).rows]
    dd = docx.Document()
    sec = dd.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    st = dd.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(0)

    def run(q, t, size=14, bold=False, red=False, strike=False):
        r = q.add_run(t); r.font.size = Pt(size); r.bold = bold; r.font.strike = strike
        if red: r.font.color.rgb = RGBColor(0xFF, 0, 0)
        if strike: r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.RIGHT; run(q, '3.1.4-jadval', bold=True)
    q = dd.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; run(q, cap.text, bold=True)
    t = dd.add_table(rows=len(rows), cols=4); t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            q = t.cell(i, j).paragraphs[0]
            if i == 0 or j == 1: q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if j == 2 and i in YANGI:
                run(q, txt, 11, strike=True)
                run(t.cell(i, j).add_paragraph(), YANGI[i], 11, red=True)
            else:
                run(q, txt, 11, bold=(i == 0))
    dd.save(DST)
    from pardoz import pardoz
    print(pardoz(DST), DST)
